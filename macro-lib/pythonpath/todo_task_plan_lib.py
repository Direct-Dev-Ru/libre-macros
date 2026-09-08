# -*- coding: utf-8 -*-
"""План выполнения задачи: хранение, рабочие дни, пересчёт (todo_task_plan)."""
from __future__ import print_function, unicode_literals
MACRO_VERSION = "3.10.711"
import re

import todo_task_cfg as _tcfg
import todo_task_plan_cfg as _pcfg

try:
    unicode
except NameError:
    unicode = str


def _log(msg):
    try:
        from todo_task_settings_lib import get_console_log_enabled

        if not get_console_log_enabled():
            return
    except Exception:
        pass
    try:
        print("[todo_task/plan] %s" % msg)
    except Exception:
        pass


def _import_lib():
    import todo_task_lib as _lib

    return _lib


def _cell_text(cell):
    return _import_lib()._cell_text(cell)


def _set_cell_text(cell, text):
    _import_lib()._set_cell_text(cell, text)


def _clear_cell_value(cell):
    _import_lib()._clear_cell_value(cell)


def _get_sheet_by_name(doc, name):
    return _import_lib()._get_sheet_by_name(doc, name)


def _hide_sheet(sh):
    return _import_lib()._hide_sheet(sh)


def _used_last_row(sheet, min_cols=1):
    return _import_lib()._used_last_row(sheet, min_cols=min_cols)


def _new_guid():
    return _import_lib()._new_guid()


def _now_timestamp_text():
    return _import_lib()._now_timestamp_text()


def parse_due_parts(text):
    return _import_lib().parse_due_parts(text)


def _read_due_from_cell(cell):
    return _import_lib()._read_due_from_cell(cell)


def set_cell_due_date(doc, cell, day, month, year):
    _import_lib().set_cell_due_date(doc, cell, day, month, year)


def _norm_header(text):
    return _import_lib()._norm_header(text)


def _parts_to_date(parts):
    if not parts:
        return None
    try:
        from datetime import date

        return date(int(parts[2]), int(parts[1]), int(parts[0]))
    except Exception:
        return None


def _date_to_parts(d):
    if d is None:
        return None
    return int(d.day), int(d.month), int(d.year)


def format_plan_date_display(doc, parts):
    if not parts:
        return u""
    try:
        from datetime import date

        d = date(int(parts[2]), int(parts[1]), int(parts[0]))
        return u"%s %s" % (
            format_plan_date_field(parts),
            plan_date_weekday_label(parts),
        )
    except Exception:
        return format_plan_date_field(parts)


_PLAN_WEEKDAY_NAMES = (
    u"Пн",
    u"Вт",
    u"Ср",
    u"Чт",
    u"Пт",
    u"Сб",
    u"Вс",
)


def normalize_plan_date_input(text):
    """Нормализация ввода: 01012026, 01б012026, 01.01.26, запятая/«ю» → точка."""
    s = unicode(text or u"").strip()
    if not s:
        return u""
    s = re.sub(
        r"\s+[Пп][нвсчт][гтрб]?\s*$",
        u"",
        s,
        flags=re.UNICODE,
    )
    digits = re.sub(r"\D", u"", s)
    if len(digits) == 8:
        try:
            d = int(digits[0:2])
            mo = int(digits[2:4])
            y = int(digits[4:8])
            return u"%02d.%02d.%04d" % (d, mo, y)
        except (TypeError, ValueError):
            pass
    if len(digits) == 6:
        try:
            from datetime import date

            d = int(digits[0:2])
            mo = int(digits[2:4])
            y2 = int(digits[4:6])
            y = (date.today().year // 100) * 100 + y2
            return u"%02d.%02d.%04d" % (d, mo, y)
        except (TypeError, ValueError):
            pass
    s = s.replace(u",", u".").replace(u"ю", u".").replace(u"Ю", u".")
    s = re.sub(r"[^\d.\-/]+", u".", s, flags=re.UNICODE)
    s = re.sub(r"[.\-/]+", u".", s)
    s = s.strip(u".")
    m2 = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2})$", s)
    if m2:
        from datetime import date

        d = int(m2.group(1))
        mo = int(m2.group(2))
        y2 = int(m2.group(3))
        y = (date.today().year // 100) * 100 + y2
        s = u"%02d.%02d.%04d" % (d, mo, y)
    return s


def parse_plan_date_parts(text):
    """(day, month, year) из поля визарда плана или None."""
    s = normalize_plan_date_input(text)
    if not s:
        return None
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", s)
    if not m:
        m = re.match(r"^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", s)
    if not m:
        return None
    d = int(m.group(1))
    mo = int(m.group(2))
    y = int(m.group(3))
    if not (1 <= d <= 31 and 1 <= mo <= 12 and 1900 <= y <= 2100):
        return None
    try:
        from datetime import date

        date(y, mo, d)
    except ValueError:
        return None
    return d, mo, y


def format_plan_date_field(parts):
    if not parts:
        return u""
    return u"%02d.%02d.%04d" % (int(parts[0]), int(parts[1]), int(parts[2]))


def plan_date_weekday_label(parts):
    if not parts:
        return u""
    try:
        from datetime import date

        d = date(int(parts[2]), int(parts[1]), int(parts[0]))
        return _PLAN_WEEKDAY_NAMES[d.weekday()]
    except Exception:
        return u""


def _write_plan_date_cell(doc, cell, parts):
    if cell is None:
        return
    if not parts:
        _clear_cell_value(cell)
        return
    set_cell_due_date(doc, cell, parts[0], parts[1], parts[2])


def _read_plan_date_cell(cell):
    parts = _read_due_from_cell(cell)
    if parts:
        return parts
    return parse_due_parts(_cell_text(cell))


def _guid_key(guid):
    g = unicode(guid or u"").strip()
    if not g:
        return u""
    if hasattr(g, "casefold"):
        return g.casefold()
    return g.lower()


def plan_sheet_name_for_guid(guid):
    """Имя листа плана для задачи: __План_задач_<guid> (до 31 символа LO/Excel)."""
    g = unicode(guid or u"").strip()
    prefix = unicode(getattr(_pcfg, "PLAN_SHEET_PREFIX", u"__План_задач_"))
    max_len = int(getattr(_pcfg, "PLAN_SHEET_NAME_MAX", 31))
    if not g:
        return unicode(_pcfg.PLAN_SHEET_NAME)
    safe = re.sub(r"[^\w\-]", u"_", g, flags=re.UNICODE)
    safe = safe.replace(u"-", u"")
    room = max_len - len(prefix)
    if room < 4:
        room = 4
    return (prefix + safe[:room])[:max_len]


_LAST_PLAN_IO_ERROR = u""


def get_last_plan_io_error():
    return unicode(_LAST_PLAN_IO_ERROR or u"")


def _set_plan_io_error(msg):
    global _LAST_PLAN_IO_ERROR
    _LAST_PLAN_IO_ERROR = unicode(msg or u"")


def _task_book_path_and_stem(doc):
    try:
        from todo_task_settings_lib import get_doc_book_path
    except Exception:
        return u"", u""
    path = unicode(get_doc_book_path(doc) or u"").strip()
    if not path:
        return u"", u""
    import os

    base = os.path.basename(path)
    stem, _ext = os.path.splitext(base)
    stem = unicode(stem or u"").strip()
    if not stem:
        return path, u""
    # Имя каталога/файла — как у книги задач, без опасных символов ОС.
    stem = re.sub(r'[\\/:*?"<>|]+', u"_", stem).strip(u" ._")
    return path, stem


def plans_store_dir_path(doc):
    """Каталог Plans_for_<stem> рядом с книгой задач."""
    import os

    path, stem = _task_book_path_and_stem(doc)
    if not path or not stem:
        return u""
    prefix = unicode(getattr(_pcfg, "PLAN_STORE_DIR_PREFIX", u"Plans_for_"))
    return os.path.join(os.path.dirname(path), prefix + stem)


def plans_workbook_path(doc):
    """Путь к Plans_for_<stem>.xlsx в каталоге планов."""
    import os

    d = plans_store_dir_path(doc)
    if not d:
        return u""
    _path, stem = _task_book_path_and_stem(doc)
    if not stem:
        return u""
    prefix = unicode(getattr(_pcfg, "PLAN_STORE_FILE_PREFIX", u"Plans_for_"))
    ext = unicode(getattr(_pcfg, "PLAN_STORE_FILE_EXT", u".xlsx"))
    return os.path.join(d, prefix + stem + ext)


def plans_workbook_path_for_book(book_path):
    """Путь Plans_for_<stem>.xlsx для указанной книги задач."""
    import os

    path, stem = _task_book_path_and_stem_from_path(book_path)
    if not path or not stem:
        return u""
    prefix_dir = unicode(getattr(_pcfg, "PLAN_STORE_DIR_PREFIX", u"Plans_for_"))
    prefix_file = unicode(getattr(_pcfg, "PLAN_STORE_FILE_PREFIX", u"Plans_for_"))
    ext = unicode(getattr(_pcfg, "PLAN_STORE_FILE_EXT", u".xlsx"))
    d = os.path.join(os.path.dirname(path), prefix_dir + stem)
    return os.path.join(d, prefix_file + stem + ext)


def _task_book_path_and_stem_from_path(book_path):
    import os

    path = unicode(book_path or u"").strip()
    if not path:
        return u"", u""
    try:
        from todo_task_settings_lib import normalize_book_path

        path = normalize_book_path(path)
    except Exception:
        try:
            path = os.path.abspath(os.path.normpath(path))
        except Exception:
            pass
    base = os.path.basename(path)
    stem, _ext = os.path.splitext(base)
    stem = unicode(stem or u"").strip()
    if not stem:
        return path, u""
    stem = re.sub(r'[\\/:*?"<>|]+', u"_", stem).strip(u" ._")
    return path, stem


def _sheet_name_key(name):
    n = unicode(name or u"").strip()
    if not n:
        return u""
    return n.casefold() if hasattr(n, "casefold") else n.lower()


def _doc_sheet_by_name(doc, name):
    nm = unicode(name or u"").strip()
    if not nm or doc is None:
        return None
    try:
        sheets = doc.Sheets
        if sheets.hasByName(nm):
            return sheets.getByName(nm)
    except Exception:
        pass
    key = _sheet_name_key(nm)
    try:
        sheets = doc.Sheets
        n = int(sheets.getCount())
        i = 0
        while i < n:
            sh = sheets.getByIndex(i)
            sn = unicode(sh.getName())
            if _sheet_name_key(sn) == key:
                return sh
            i = i + 1
    except Exception:
        pass
    return None


def _iter_doc_inline_plan_task_sheet_names(doc):
    """Имена листов __План_задач_<id> в книге (без общего legacy __План_задач)."""
    out = []
    if doc is None:
        return out
    legacy = unicode(_pcfg.PLAN_SHEET_NAME)
    prefix = unicode(getattr(_pcfg, "PLAN_SHEET_PREFIX", u"__План_задач_"))
    try:
        sheets = doc.Sheets
        n = int(sheets.getCount())
        i = 0
        while i < n:
            sn = unicode(sheets.getByIndex(i).getName())
            if sn != legacy and sn.startswith(prefix):
                out.append(sn)
            i = i + 1
    except Exception as exc:
        _log("iter inline plan sheets: %s" % exc)
    return out


def _guid_from_inline_plan_sheet(doc, sheet_name, guids_hint=None):
    """GUID задачи для листа __План_задач_* (свод, имя листа, данные)."""
    sn = unicode(sheet_name or u"").strip()
    prefix = unicode(getattr(_pcfg, "PLAN_SHEET_PREFIX", u"__План_задач_"))
    guids_hint = guids_hint or []
    gi = 0
    while gi < len(guids_hint):
        g = unicode(guids_hint[gi] or u"").strip()
        gi = gi + 1
        if not g:
            continue
        psn = plan_sheet_name_for_guid(g)
        if psn == sn or _sheet_name_key(psn) == _sheet_name_key(sn):
            return g
    if sn.startswith(prefix):
        suffix = sn[len(prefix) :]
        sk = suffix.casefold() if hasattr(suffix, "casefold") else suffix.lower()
        gi = 0
        while gi < len(guids_hint):
            g = unicode(guids_hint[gi] or u"").strip()
            gi = gi + 1
            gk = _guid_key(g)
            if not gk:
                continue
            safe = re.sub(r"[^\w\-]", u"_", g, flags=re.UNICODE).replace(u"-", u"")
            room = int(getattr(_pcfg, "PLAN_SHEET_NAME_MAX", 31)) - len(prefix)
            if room < 4:
                room = 4
            short = safe[:room]
            short_k = short.casefold() if hasattr(short, "casefold") else short.lower()
            if sk == short_k or gk.startswith(sk) or sk in gk:
                return g
    sh = _doc_sheet_by_name(doc, sn)
    if sh is not None:
        try:
            cols = _build_plan_col_map(sh)
            if cols is not None:
                cols = _ensure_missing_plan_columns(sh, cols)
                rows = _load_rows_from_sheet(sh, cols, doc)
                ri = 0
                while ri < len(rows):
                    rg = unicode(rows[ri].get(u"guid") or u"").strip()
                    if rg:
                        return rg
                    ri = ri + 1
        except Exception as exc:
            _log("guid from plan sheet %s: %s" % (sn, exc))
    return u""


def _collect_inline_plan_sheets_for_merge(doc, guids=None):
    """Пары (имя_листа, guid) — все __План_задач_* в книге руководителя."""
    guids = [
        unicode(g or u"").strip() for g in (guids or []) if unicode(g or u"").strip()
    ]
    names = _iter_doc_inline_plan_task_sheet_names(doc)
    out = []
    ni = 0
    while ni < len(names):
        sn = names[ni]
        ni = ni + 1
        g = _guid_from_inline_plan_sheet(doc, sn, guids_hint=guids)
        out.append((sn, g))
    return out


def _remove_doc_plan_sheet_by_name(doc, sheet_name):
    sn = unicode(sheet_name or u"").strip()
    if not sn or doc is None:
        return
    try:
        if doc.Sheets.hasByName(sn):
            doc.Sheets.removeByName(sn)
    except Exception as exc:
        _log("remove plan sheet %s: %s" % (sn, exc))


def _doc_has_inline_plan(doc, guid):
    """План задачи на листах внутри книги задач (не только во внешнем Plans_for_…)."""
    gk = _guid_key(guid)
    if not gk or doc is None:
        return False
    name = plan_sheet_name_for_guid(guid)
    if _doc_sheet_by_name(doc, name) is not None:
        return True
    names = _iter_doc_inline_plan_task_sheet_names(doc)
    gi = 0
    while gi < len(names):
        g = _guid_from_inline_plan_sheet(doc, names[gi], guids_hint=[guid])
        if _guid_key(g) == gk:
            return True
        gi = gi + 1
    try:
        return bool(_load_legacy_plan_by_guid(doc, guid))
    except Exception:
        return False


def _is_plan_task_sheet_name(name):
    n = unicode(name or u"")
    if not n:
        return False
    if n == unicode(_pcfg.PLAN_SHEET_NAME):
        return True
    return n.startswith(unicode(getattr(_pcfg, "PLAN_SHEET_PREFIX", u"__План_задач_")))


def preview_plan_transfer_for_merge(manager_doc, guids, employee_task_books=None):
    """Сухой прогон: сколько листов __План_задач_* в книге руководителя переносится."""
    guids = [unicode(g or u"").strip() for g in (guids or []) if unicode(g or u"").strip()]
    items = _collect_inline_plan_sheets_for_merge(manager_doc, guids)
    found = len(items)
    covered = set()
    ii = 0
    while ii < len(items):
        gk = _guid_key(items[ii][1])
        if gk:
            covered.add(gk)
        ii = ii + 1
    missing = 0
    gi = 0
    while gi < len(guids):
        gk = _guid_key(guids[gi])
        if gk and gk not in covered and not _doc_has_inline_plan(manager_doc, guids[gi]):
            missing = missing + 1
        gi = gi + 1
    return {
        u"inline_sheets_found": found,
        u"guids_missing_plan": missing,
        u"no_inline_sheets": found == 0,
        u"guids_total": len(guids),
    }


def transfer_employee_plans_on_merge(manager_doc, guids, employee_task_books=None):
    """
    Перенести все листы __План_задач_* из книги руководителя во внешний Plans_for_….xlsx
    (файл создаётся при отсутствии). employee_task_books не используется.
    """
    guids = [unicode(g or u"").strip() for g in (guids or []) if unicode(g or u"").strip()]
    result = {
        u"moved": [],
        u"missing": [],
        u"errors": [],
        u"lines": [],
        u"inline_found": 0,
    }
    items = _collect_inline_plan_sheets_for_merge(manager_doc, guids)
    if not items:
        return result
    dest_path = plans_workbook_path(manager_doc)
    if not dest_path:
        result[u"errors"].append(
            unicode(
                getattr(
                    _pcfg,
                    "PLAN_STORE_DOC_UNSAVED",
                    u"Сохраните книгу задач на диск.",
                )
            )
        )
        return result
    ii = 0
    while ii < len(items):
        sn, guid = items[ii]
        ii = ii + 1
        sh = _doc_sheet_by_name(manager_doc, sn)
        if sh is None:
            continue
        try:
            cols = _build_plan_col_map(sh)
            if cols is None:
                result[u"errors"].append(u"%s: нет заголовков плана" % sn)
                continue
            cols = _ensure_missing_plan_columns(sh, cols)
            rows = _load_rows_from_sheet(sh, cols, manager_doc)
        except Exception as exc:
            result[u"errors"].append(u"%s: %s" % (sn, exc))
            continue
        if not rows:
            result[u"missing"].append(guid or sn)
            continue
        if not guid:
            guid = unicode(rows[0].get(u"guid") or u"").strip()
        if not guid:
            result[u"errors"].append(u"%s: не удалось определить GUID задачи" % sn)
            continue
        ri = 0
        while ri < len(rows):
            if not rows[ri].get(u"guid"):
                rows[ri][u"guid"] = guid
            ri = ri + 1
        try:
            if save_plan_by_guid(manager_doc, guid, rows):
                _remove_doc_plan_sheet_by_name(manager_doc, sn)
                result[u"moved"].append(guid)
                result[u"lines"].append(
                    u"→ %s (%s строк) → %s"
                    % (sn, len(rows), dest_path)
                )
            else:
                err = get_last_plan_io_error()
                result[u"errors"].append(
                    u"%s: %s" % (sn, err or u"не удалось сохранить")
                )
        except Exception as exc:
            result[u"errors"].append(u"%s: %s" % (sn, exc))
    result[u"inline_found"] = len(result[u"moved"])
    return result


def format_merge_plan_preview_line(preview):
    preview = preview or {}
    if preview.get(u"no_inline_sheets"):
        return unicode(
            getattr(_tcfg, "MERGE_PLAN_PREVIEW_NO_SOURCES", u"")
        )
    return unicode(getattr(_tcfg, "MERGE_PLAN_PREVIEW_FOUND", u"")) % (
        int(preview.get(u"inline_sheets_found") or 0),
        int(preview.get(u"guids_missing_plan") or 0),
    )


def _import_openpyxl():
    import openpyxl_bundled  # noqa: F401
    from openpyxl import Workbook, load_workbook

    return Workbook, load_workbook


def _xlsx_cell_text(val):
    if val is None:
        return u""
    try:
        from datetime import date, datetime

        if isinstance(val, datetime):
            return format_plan_date_field((val.day, val.month, val.year))
        if isinstance(val, date):
            return format_plan_date_field((val.day, val.month, val.year))
    except Exception:
        pass
    return unicode(val).strip()


def _xlsx_read_date_parts(val):
    if val is None or val == u"":
        return None
    try:
        from datetime import date, datetime

        if isinstance(val, datetime):
            return int(val.day), int(val.month), int(val.year)
        if isinstance(val, date):
            return int(val.day), int(val.month), int(val.year)
    except Exception:
        pass
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        try:
            from openpyxl.utils.datetime import from_excel

            d = from_excel(val)
            return int(d.day), int(d.month), int(d.year)
        except Exception:
            pass
    return parse_due_parts(_xlsx_cell_text(val))


def _xlsx_write_date(cell, parts):
    if not parts:
        cell.value = None
        return
    try:
        from datetime import date

        cell.value = date(int(parts[2]), int(parts[1]), int(parts[0]))
        cell.number_format = u"DD.MM.YYYY"
    except Exception:
        cell.value = format_plan_date_field(parts)


def _xlsx_build_col_map(ws):
    hdr_row = int(_pcfg.PLAN_HEADER_ROW) + 1  # openpyxl 1-based
    titles = _plan_header_titles()
    norm_want = {}
    i = 0
    while i < len(titles):
        norm_want[_norm_header(titles[i])] = titles[i]
        i = i + 1
    found = {}
    c = 1
    while c <= 40:
        t = _xlsx_cell_text(ws.cell(row=hdr_row, column=c).value)
        if t:
            nk = _norm_header(t)
            if nk in norm_want:
                found[norm_want[nk]] = c
        c = c + 1
    if len(found) < 3:
        return None
    key_to_col = {}
    ci = 0
    while ci < len(_pcfg.PLAN_COLUMN_SPECS):
        kid, hdr = _pcfg.PLAN_COLUMN_SPECS[ci]
        if hdr in found:
            key_to_col[kid] = found[hdr]
        ci = ci + 1
    if u"guid" not in key_to_col or u"item_no" not in key_to_col:
        return None
    key_to_col[u"_hdr_row"] = hdr_row
    return key_to_col


def _xlsx_write_headers(ws):
    hdr_row = int(_pcfg.PLAN_HEADER_ROW) + 1
    i = 0
    while i < len(_pcfg.PLAN_COLUMN_SPECS):
        _kid, hdr = _pcfg.PLAN_COLUMN_SPECS[i]
        ws.cell(row=hdr_row, column=i + 1, value=unicode(hdr))
        i = i + 1


def _xlsx_ensure_missing_columns(ws, cols):
    if cols is None:
        return cols
    hdr_row = int(cols.get(u"_hdr_row", int(_pcfg.PLAN_HEADER_ROW) + 1))
    max_c = 0
    for k, v in cols.items():
        if unicode(k).startswith(u"_"):
            continue
        try:
            max_c = max(max_c, int(v))
        except Exception:
            pass
    ci = 0
    while ci < len(_pcfg.PLAN_COLUMN_SPECS):
        kid, hdr = _pcfg.PLAN_COLUMN_SPECS[ci]
        if kid not in cols:
            max_c = max_c + 1
            ws.cell(row=hdr_row, column=max_c, value=unicode(hdr))
            cols[kid] = max_c
        ci = ci + 1
    return cols


def _xlsx_row_get(ws, cols, row, key):
    if key not in cols:
        return u""
    return _xlsx_cell_text(ws.cell(row=row, column=cols[key]).value)


def _xlsx_row_set(ws, cols, row, key, value):
    if key not in cols:
        return
    cell = ws.cell(row=row, column=cols[key])
    if value is None or value == u"":
        cell.value = None
        return
    if key in (u"item_no", u"parent_item"):
        s = normalize_item_no_for_store(value)
        cell.value = s if s else None
        if s:
            cell.number_format = u"@"
        return
    if key == u"duration_wd":
        try:
            cell.value = int(value)
            return
        except Exception:
            pass
    if key == u"sort_key":
        try:
            cell.value = float(value)
            return
        except Exception:
            pass
    cell.value = unicode(value)


def _xlsx_row_to_dict(ws, cols, row, doc=None):
    start_p = None
    end_p = None
    if u"date_start" in cols:
        start_p = _xlsx_read_date_parts(
            ws.cell(row=row, column=cols[u"date_start"]).value
        )
    if u"date_end" in cols:
        end_p = _xlsx_read_date_parts(
            ws.cell(row=row, column=cols[u"date_end"]).value
        )
    dur = _xlsx_row_get(ws, cols, row, u"duration_wd")
    try:
        dur_i = int(float(dur)) if dur else 0
    except (TypeError, ValueError):
        dur_i = 0
    ino = normalize_item_no_from_store(_xlsx_row_get(ws, cols, row, u"item_no"))
    sk_f = item_no_to_sort_key(ino)
    summ = _xlsx_row_get(ws, cols, row, u"is_summary")
    return {
        u"guid": _xlsx_row_get(ws, cols, row, u"guid"),
        u"item_no": ino,
        u"is_summary": summ == _pcfg.PLAN_SUMMARY_YES,
        u"date_start_parts": start_p,
        u"date_end_parts": end_p,
        u"date_start": format_plan_date_display(doc, start_p),
        u"date_end": format_plan_date_display(doc, end_p),
        u"start_no_snap": _bool_from_plan_flag(
            _xlsx_row_get(ws, cols, row, u"start_no_snap")
        )
        if u"start_no_snap" in cols
        else False,
        u"end_no_snap": _bool_from_plan_flag(
            _xlsx_row_get(ws, cols, row, u"end_no_snap")
        )
        if u"end_no_snap" in cols
        else False,
        u"is_done": _bool_from_plan_flag(_xlsx_row_get(ws, cols, row, u"is_done"))
        if u"is_done" in cols
        else False,
        u"duration_wd": dur_i,
        u"assignees": _xlsx_row_get(ws, cols, row, u"assignees"),
        u"stage_name": _xlsx_row_get(ws, cols, row, u"stage_name")
        if u"stage_name" in cols
        else u"",
        u"result_doc": _xlsx_row_get(ws, cols, row, u"result_doc"),
        u"criticality": _xlsx_row_get(ws, cols, row, u"criticality"),
        u"sort_key": sk_f,
        u"parent_item": normalize_item_no_from_store(
            _xlsx_row_get(ws, cols, row, u"parent_item")
        ),
        u"row_id": _xlsx_row_get(ws, cols, row, u"row_id"),
        u"created_at": _xlsx_row_get(ws, cols, row, u"created_at"),
        u"modified_at": _xlsx_row_get(ws, cols, row, u"modified_at"),
    }


def _xlsx_dict_to_row_write(ws, cols, row, item, is_new=False):
    now = _now_timestamp_text()
    if is_new or not item.get(u"created_at"):
        item[u"created_at"] = now
    item[u"modified_at"] = now
    if not item.get(u"row_id"):
        item[u"row_id"] = _new_guid()
    item_no = normalize_item_no_from_store(item.get(u"item_no"))
    if not item.get(u"parent_item") and item_no:
        item[u"parent_item"] = parent_of_item_no(item_no)
    if not item.get(u"sort_key"):
        item[u"sort_key"] = item_no_to_sort_key(item_no)
    parent_item = normalize_item_no_from_store(item.get(u"parent_item"))
    _xlsx_row_set(ws, cols, row, u"guid", item.get(u"guid"))
    _xlsx_row_set(ws, cols, row, u"item_no", item_no)
    _xlsx_row_set(
        ws,
        cols,
        row,
        u"is_summary",
        _pcfg.PLAN_SUMMARY_YES if item.get(u"is_summary") else u"",
    )
    if u"date_start" in cols:
        _xlsx_write_date(
            ws.cell(row=row, column=cols[u"date_start"]),
            item.get(u"date_start_parts"),
        )
    if u"date_end" in cols:
        _xlsx_write_date(
            ws.cell(row=row, column=cols[u"date_end"]),
            item.get(u"date_end_parts"),
        )
    if u"start_no_snap" in cols:
        _xlsx_row_set(
            ws,
            cols,
            row,
            u"start_no_snap",
            _plan_flag_from_bool(item.get(u"start_no_snap")),
        )
    if u"end_no_snap" in cols:
        _xlsx_row_set(
            ws,
            cols,
            row,
            u"end_no_snap",
            _plan_flag_from_bool(item.get(u"end_no_snap")),
        )
    if u"is_done" in cols:
        _xlsx_row_set(
            ws, cols, row, u"is_done", _plan_flag_from_bool(item.get(u"is_done"))
        )
    _xlsx_row_set(ws, cols, row, u"duration_wd", item.get(u"duration_wd"))
    _xlsx_row_set(ws, cols, row, u"assignees", item.get(u"assignees"))
    if u"stage_name" in cols:
        _xlsx_row_set(ws, cols, row, u"stage_name", item.get(u"stage_name"))
    _xlsx_row_set(ws, cols, row, u"result_doc", item.get(u"result_doc"))
    _xlsx_row_set(ws, cols, row, u"criticality", item.get(u"criticality"))
    _xlsx_row_set(ws, cols, row, u"sort_key", item.get(u"sort_key"))
    _xlsx_row_set(ws, cols, row, u"parent_item", parent_item)
    _xlsx_row_set(ws, cols, row, u"row_id", item.get(u"row_id"))
    _xlsx_row_set(ws, cols, row, u"created_at", item.get(u"created_at"))
    _xlsx_row_set(ws, cols, row, u"modified_at", item.get(u"modified_at"))


def _xlsx_load_rows_from_ws(ws, cols, doc, guid_filter=None):
    hdr = int(cols.get(u"_hdr_row", int(_pcfg.PLAN_HEADER_ROW) + 1))
    last = int(ws.max_row or hdr)
    out = []
    r = hdr + 1
    while r <= last:
        item_no = _xlsx_row_get(ws, cols, r, u"item_no")
        guid_v = _xlsx_row_get(ws, cols, r, u"guid")
        if not item_no and not guid_v:
            r = r + 1
            continue
        if guid_filter:
            if _guid_key(guid_v) != guid_filter:
                r = r + 1
                continue
        out.append(_xlsx_row_to_dict(ws, cols, r, doc=doc))
        r = r + 1
    return out


def _xlsx_clear_data_rows(ws, cols):
    hdr = int(cols.get(u"_hdr_row", int(_pcfg.PLAN_HEADER_ROW) + 1))
    last = int(ws.max_row or hdr)
    if last <= hdr:
        return
    # openpyxl: удалить снизу вверх
    r = last
    while r > hdr:
        ws.delete_rows(r, 1)
        r = r - 1


def _plans_xlsx_open(doc, create=False):
    """Открыть книгу планов. create=True — создать каталог/файл при отсутствии."""
    import os

    _set_plan_io_error(u"")
    path = plans_workbook_path(doc)
    if not path:
        _set_plan_io_error(
            unicode(
                getattr(
                    _pcfg,
                    "PLAN_STORE_DOC_UNSAVED",
                    u"Сохраните книгу задач на диск.",
                )
            )
        )
        return None, u""
    try:
        Workbook, load_workbook = _import_openpyxl()
    except Exception as exc:
        _set_plan_io_error(
            unicode(getattr(_pcfg, "PLAN_STORE_IO_ERROR", u"%s")) % exc
        )
        return None, path
    if not os.path.isfile(path):
        if not create:
            return None, path
        try:
            d = os.path.dirname(path)
            if d and not os.path.isdir(d):
                os.makedirs(d)
            wb = Workbook()
            ws0 = wb.active
            placeholder = unicode(
                getattr(_pcfg, "PLAN_STORE_PLACEHOLDER_SHEET", u"_index")
            )
            try:
                ws0.title = placeholder[:31]
            except Exception:
                pass
            wb.save(path)
        except Exception as exc:
            _set_plan_io_error(
                unicode(getattr(_pcfg, "PLAN_STORE_IO_ERROR", u"%s")) % exc
            )
            return None, path
    try:
        wb = load_workbook(path)
    except Exception as exc:
        _set_plan_io_error(
            unicode(getattr(_pcfg, "PLAN_STORE_IO_ERROR", u"%s")) % exc
        )
        return None, path
    return wb, path


def _plans_xlsx_save(wb, path):
    try:
        wb.save(path)
        return True
    except Exception as exc:
        _set_plan_io_error(
            unicode(getattr(_pcfg, "PLAN_STORE_IO_ERROR", u"%s")) % exc
        )
        _log("plans xlsx save: %s" % exc)
        return False


def _xlsx_get_or_create_plan_ws(wb, guid, create=True):
    name = plan_sheet_name_for_guid(guid)
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        if not create:
            return None, None
        ws = wb.create_sheet(title=name[:31])
        _xlsx_write_headers(ws)
    cols = _xlsx_build_col_map(ws)
    if cols is None:
        _xlsx_write_headers(ws)
        cols = _xlsx_build_col_map(ws)
    cols = _xlsx_ensure_missing_columns(ws, cols)
    return ws, cols


def _xlsx_ensure_archive_ws(wb, plan_cols):
    name = unicode(_pcfg.PLAN_ARCHIVE_SHEET_NAME)
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(title=name[:31])
    hdr_row = int(_pcfg.PLAN_HEADER_ROW) + 1
    if not _xlsx_cell_text(ws.cell(row=hdr_row, column=1).value):
        ci = 0
        while ci < len(_pcfg.PLAN_COLUMN_SPECS):
            kid, hdr = _pcfg.PLAN_COLUMN_SPECS[ci]
            if kid in plan_cols:
                ws.cell(row=hdr_row, column=plan_cols[kid], value=unicode(hdr))
            ci = ci + 1
        max_c = 0
        for k, v in plan_cols.items():
            if unicode(k).startswith(u"_"):
                continue
            try:
                max_c = max(max_c, int(v))
            except Exception:
                pass
        ws.cell(row=hdr_row, column=max_c + 1, value=unicode(_pcfg.PLAN_ARCHIVE_COL_AT))
        ws.cell(
            row=hdr_row, column=max_c + 2, value=unicode(_pcfg.PLAN_ARCHIVE_COL_SHEET)
        )
        ws.cell(
            row=hdr_row, column=max_c + 3, value=unicode(_pcfg.PLAN_ARCHIVE_COL_TASK)
        )
        return ws, max_c + 1
    # найти мета-колонки по заголовкам
    max_c = 0
    for k, v in plan_cols.items():
        if unicode(k).startswith(u"_"):
            continue
        try:
            max_c = max(max_c, int(v))
        except Exception:
            pass
    c = 1
    at_c = max_c + 1
    while c <= max_c + 10:
        t = _xlsx_cell_text(ws.cell(row=hdr_row, column=c).value)
        if t == unicode(_pcfg.PLAN_ARCHIVE_COL_AT):
            at_c = c
            break
        c = c + 1
    return ws, at_c


def _bool_from_plan_flag(text):
    return unicode(text or u"").strip() == _pcfg.PLAN_SUMMARY_YES


def _plan_flag_from_bool(val):
    return _pcfg.PLAN_SUMMARY_YES if val else u""


def normalize_item_no_from_store(item_no):
    """Убрать текстовый префикс ' при чтении номера пункта из ячейки."""
    s = unicode(item_no or u"").strip()
    if s.startswith(u"'"):
        s = s[1:].strip()
    return s


def normalize_item_no_for_store(item_no):
    """Записать номер пункта как текст: '1, '1.1, '1.2 …"""
    s = normalize_item_no_from_store(item_no)
    if not s:
        return u""
    return u"'" + s


_ITEM_NO_SORT_BASE = 10000.0


def item_no_sort_parts(item_no):
    """Части номера пункта как целые: «3.10» → (3, 10)."""
    s = normalize_item_no_from_store(item_no)
    if not s:
        return ()
    parts = s.split(u".")
    out = []
    i = 0
    while i < len(parts):
        p = unicode(parts[i] or u"").strip()
        if not p:
            i = i + 1
            continue
        try:
            out.append(int(p))
        except (TypeError, ValueError):
            out.append(0)
        i = i + 1
    return tuple(out)


def item_no_to_sort_key(item_no):
    """Числовой ключ для сортировки иерархии 1 < 1.2 < 1.10."""
    parts = item_no_sort_parts(item_no)
    if not parts:
        return 0.0
    key = 0.0
    i = 0
    while i < len(parts):
        key = key * _ITEM_NO_SORT_BASE + float(parts[i])
        i = i + 1
    return key


def plan_row_order_key(row):
    """Ключ сортировки строки плана (по item_no, иначе sort_key)."""
    parts = item_no_sort_parts((row or {}).get(u"item_no"))
    if parts:
        return parts
    try:
        return (float((row or {}).get(u"sort_key") or 0.0),)
    except (TypeError, ValueError):
        return (0.0,)


def sort_plan_rows(rows):
    return sorted(list(rows or []), key=plan_row_order_key)


def parent_of_item_no(item_no):
    s = normalize_item_no_from_store(item_no)
    if u"." not in s:
        return u""
    return s.rsplit(u".", 1)[0]


def _plan_header_titles():
    return [hdr for _k, hdr in _pcfg.PLAN_COLUMN_SPECS]


def _build_plan_col_map(sheet):
    hdr_row = int(_pcfg.PLAN_HEADER_ROW)
    titles = _plan_header_titles()
    norm_want = {}
    i = 0
    while i < len(titles):
        norm_want[_norm_header(titles[i])] = titles[i]
        i = i + 1
    found = {}
    c = 0
    while c < 40:
        t = _cell_text(sheet.getCellByPosition(c, hdr_row)).strip()
        if t:
            nk = _norm_header(t)
            if nk in norm_want:
                found[norm_want[nk]] = c
        c = c + 1
    if len(found) < 3:
        return None
    key_to_col = {}
    ci = 0
    while ci < len(_pcfg.PLAN_COLUMN_SPECS):
        kid, hdr = _pcfg.PLAN_COLUMN_SPECS[ci]
        if hdr in found:
            key_to_col[kid] = found[hdr]
        ci = ci + 1
    if u"guid" not in key_to_col or u"item_no" not in key_to_col:
        return None
    key_to_col[u"_hdr_row"] = hdr_row
    return key_to_col


def _write_plan_headers(sheet):
    hdr_row = int(_pcfg.PLAN_HEADER_ROW)
    i = 0
    while i < len(_pcfg.PLAN_COLUMN_SPECS):
        _kid, hdr = _pcfg.PLAN_COLUMN_SPECS[i]
        _set_cell_text(sheet.getCellByPosition(i, hdr_row), unicode(hdr))
        i = i + 1


def _ensure_missing_plan_columns(sheet, cols):
    if cols is None:
        return cols
    hdr_row = int(_pcfg.PLAN_HEADER_ROW)
    max_c = _table_max_from_cols(cols)
    ci = 0
    while ci < len(_pcfg.PLAN_COLUMN_SPECS):
        kid, hdr = _pcfg.PLAN_COLUMN_SPECS[ci]
        if kid not in cols:
            max_c = max_c + 1
            _set_cell_text(sheet.getCellByPosition(max_c, hdr_row), unicode(hdr))
            cols[kid] = max_c
        ci = ci + 1
    return cols


def _ensure_plan_sheet(doc, guid=None):
    if guid:
        name = plan_sheet_name_for_guid(guid)
    else:
        name = unicode(_pcfg.PLAN_SHEET_NAME)
    sheets = doc.Sheets
    if sheets.hasByName(name):
        sh = sheets.getByName(name)
    else:
        pos = int(sheets.getCount())
        sheets.insertNewByName(name, pos)
        sh = sheets.getByName(name)
        _write_plan_headers(sh)
    _hide_sheet(sh)
    cols = _build_plan_col_map(sh)
    if cols is None:
        _write_plan_headers(sh)
        cols = _build_plan_col_map(sh)
    cols = _ensure_missing_plan_columns(sh, cols)
    return sh, cols


def _ensure_legacy_plan_sheet(doc):
    return _ensure_plan_sheet(doc, guid=None)


def _ensure_plan_archive_sheet(doc, plan_cols):
    name = unicode(_pcfg.PLAN_ARCHIVE_SHEET_NAME)
    sheets = doc.Sheets
    if sheets.hasByName(name):
        sh = sheets.getByName(name)
    else:
        pos = int(sheets.getCount())
        sheets.insertNewByName(name, pos)
        sh = sheets.getByName(name)
    _hide_sheet(sh)
    hdr_row = int(_pcfg.PLAN_HEADER_ROW)
    if not _cell_text(sh.getCellByPosition(plan_cols.get(u"guid", 0), hdr_row)):
        ci = 0
        while ci < len(_pcfg.PLAN_COLUMN_SPECS):
            kid, hdr = _pcfg.PLAN_COLUMN_SPECS[ci]
            if kid in plan_cols:
                _set_cell_text(
                    sh.getCellByPosition(plan_cols[kid], hdr_row), unicode(hdr)
                )
            ci = ci + 1
        max_c = _table_max_from_cols(plan_cols)
        _set_cell_text(
            sh.getCellByPosition(max_c + 1, hdr_row),
            unicode(_pcfg.PLAN_ARCHIVE_COL_AT),
        )
        _set_cell_text(
            sh.getCellByPosition(max_c + 2, hdr_row),
            unicode(_pcfg.PLAN_ARCHIVE_COL_SHEET),
        )
        _set_cell_text(
            sh.getCellByPosition(max_c + 3, hdr_row),
            unicode(_pcfg.PLAN_ARCHIVE_COL_TASK),
        )
    max_c = _table_max_from_cols(plan_cols)
    return sh, max_c + 1


def _row_get(sheet, cols, row, key):
    if key not in cols:
        return u""
    return _cell_text(sheet.getCellByPosition(cols[key], row))


def _row_set(sheet, cols, row, key, value):
    if key not in cols:
        return
    if key in (u"item_no", u"parent_item"):
        value = normalize_item_no_for_store(value)
    _set_cell_text(sheet.getCellByPosition(cols[key], row), unicode(value or u""))


def _row_to_dict(sheet, cols, row, doc=None):
    start_p = _read_plan_date_cell(
        sheet.getCellByPosition(cols[u"date_start"], row)
        if u"date_start" in cols
        else None
    )
    end_p = _read_plan_date_cell(
        sheet.getCellByPosition(cols[u"date_end"], row)
        if u"date_end" in cols
        else None
    )
    dur = _row_get(sheet, cols, row, u"duration_wd")
    try:
        dur_i = int(float(dur)) if dur else 0
    except (TypeError, ValueError):
        dur_i = 0
    ino = normalize_item_no_from_store(_row_get(sheet, cols, row, u"item_no"))
    sk_f = item_no_to_sort_key(ino)
    summ = _row_get(sheet, cols, row, u"is_summary")
    return {
        u"guid": _row_get(sheet, cols, row, u"guid"),
        u"item_no": ino,
        u"is_summary": summ == _pcfg.PLAN_SUMMARY_YES,
        u"date_start_parts": start_p,
        u"date_end_parts": end_p,
        u"date_start": format_plan_date_display(doc, start_p),
        u"date_end": format_plan_date_display(doc, end_p),
        u"start_no_snap": _bool_from_plan_flag(
            _row_get(sheet, cols, row, u"start_no_snap")
        )
        if u"start_no_snap" in cols
        else False,
        u"end_no_snap": _bool_from_plan_flag(
            _row_get(sheet, cols, row, u"end_no_snap")
        )
        if u"end_no_snap" in cols
        else False,
        u"is_done": _bool_from_plan_flag(_row_get(sheet, cols, row, u"is_done"))
        if u"is_done" in cols
        else False,
        u"duration_wd": dur_i,
        u"assignees": _row_get(sheet, cols, row, u"assignees"),
        u"stage_name": _row_get(sheet, cols, row, u"stage_name")
        if u"stage_name" in cols
        else u"",
        u"result_doc": _row_get(sheet, cols, row, u"result_doc"),
        u"criticality": _row_get(sheet, cols, row, u"criticality"),
        u"sort_key": sk_f,
        u"parent_item": normalize_item_no_from_store(
            _row_get(sheet, cols, row, u"parent_item")
        ),
        u"row_id": _row_get(sheet, cols, row, u"row_id"),
        u"created_at": _row_get(sheet, cols, row, u"created_at"),
        u"modified_at": _row_get(sheet, cols, row, u"modified_at"),
    }


def _dict_to_row_write(sheet, cols, row, item, doc=None, is_new=False):
    now = _now_timestamp_text()
    if is_new or not item.get(u"created_at"):
        item[u"created_at"] = now
    item[u"modified_at"] = now
    if not item.get(u"row_id"):
        item[u"row_id"] = _new_guid()
    item_no = normalize_item_no_from_store(item.get(u"item_no"))
    if not item.get(u"parent_item") and item_no:
        item[u"parent_item"] = parent_of_item_no(item_no)
    if not item.get(u"sort_key"):
        item[u"sort_key"] = item_no_to_sort_key(item_no)
    parent_item = normalize_item_no_from_store(item.get(u"parent_item"))
    _row_set(sheet, cols, row, u"guid", item.get(u"guid"))
    _row_set(sheet, cols, row, u"item_no", item_no)
    _row_set(
        sheet,
        cols,
        row,
        u"is_summary",
        _pcfg.PLAN_SUMMARY_YES if item.get(u"is_summary") else u"",
    )
    if u"date_start" in cols:
        _write_plan_date_cell(
            doc,
            sheet.getCellByPosition(cols[u"date_start"], row),
            item.get(u"date_start_parts"),
        )
    if u"date_end" in cols:
        _write_plan_date_cell(
            doc,
            sheet.getCellByPosition(cols[u"date_end"], row),
            item.get(u"date_end_parts"),
        )
    if u"start_no_snap" in cols:
        _row_set(
            sheet,
            cols,
            row,
            u"start_no_snap",
            _plan_flag_from_bool(item.get(u"start_no_snap")),
        )
    if u"end_no_snap" in cols:
        _row_set(
            sheet,
            cols,
            row,
            u"end_no_snap",
            _plan_flag_from_bool(item.get(u"end_no_snap")),
        )
    if u"is_done" in cols:
        _row_set(
            sheet,
            cols,
            row,
            u"is_done",
            _plan_flag_from_bool(item.get(u"is_done")),
        )
    _row_set(sheet, cols, row, u"duration_wd", item.get(u"duration_wd"))
    _row_set(sheet, cols, row, u"assignees", item.get(u"assignees"))
    if u"stage_name" in cols:
        _row_set(sheet, cols, row, u"stage_name", item.get(u"stage_name"))
    _row_set(sheet, cols, row, u"result_doc", item.get(u"result_doc"))
    _row_set(sheet, cols, row, u"criticality", item.get(u"criticality"))
    _row_set(sheet, cols, row, u"sort_key", item.get(u"sort_key"))
    _row_set(sheet, cols, row, u"parent_item", parent_item)
    _row_set(sheet, cols, row, u"row_id", item.get(u"row_id"))
    _row_set(sheet, cols, row, u"created_at", item.get(u"created_at"))
    _row_set(sheet, cols, row, u"modified_at", item.get(u"modified_at"))


def load_holidays(doc):
    """Множество date — нерабочие дни из справочника «Праздники»."""
    out = set()
    try:
        from datetime import date

        ref_sheet = _get_sheet_by_name(
            doc, getattr(_tcfg, "REF_SHEET_NAME", u"__Справочники_задачи")
        )
        if ref_sheet is None:
            return out
        dicts = _import_lib().discover_ref_dictionaries(ref_sheet)
        target = _norm_header(_pcfg.REF_HOLIDAYS_TITLE)
        ref_dict = None
        di = 0
        while di < len(dicts):
            d = dicts[di]
            if _norm_header(d.get(u"title") or u"") == target or (
                d.get(u"titles")
                and _norm_header(d[u"titles"][0]) == target
            ):
                ref_dict = d
                break
            di = di + 1
        if ref_dict is None:
            return out
        items = _import_lib().read_ref_dictionary_items(ref_sheet, ref_dict)
        ii = 0
        while ii < len(items):
            vals = items[ii]
            ii = ii + 1
            vi = 0
            while vi < len(vals):
                raw = unicode(vals[vi] or u"").strip()
                vi = vi + 1
                if not raw:
                    continue
                parts = parse_due_parts(raw)
                if parts:
                    try:
                        out.add(date(parts[2], parts[1], parts[0]))
                    except Exception:
                        pass
        return out
    except Exception as exc:
        _log("load_holidays: %s" % exc)
        return out


def holidays_ref_info(doc):
    """(есть_блок_«Праздники», число_дней)."""
    try:
        ref_sheet = _get_sheet_by_name(
            doc, getattr(_tcfg, "REF_SHEET_NAME", u"__Справочники_задачи")
        )
        if ref_sheet is None:
            return False, 0
        dicts = _import_lib().discover_ref_dictionaries(ref_sheet)
        target = _norm_header(_pcfg.REF_HOLIDAYS_TITLE)
        has_block = False
        di = 0
        while di < len(dicts):
            d = dicts[di]
            if _norm_header(d.get(u"title") or u"") == target or (
                d.get(u"titles")
                and _norm_header(d[u"titles"][0]) == target
            ):
                has_block = True
                break
            di = di + 1
        return has_block, len(load_holidays(doc))
    except Exception:
        return False, len(load_holidays(doc))


def load_criticality_choices(doc):
    defaults = list(_pcfg.REF_CRITICALITY_DEFAULTS)
    try:
        ref_sheet = _get_sheet_by_name(
            doc, getattr(_tcfg, "REF_SHEET_NAME", u"__Справочники_задачи")
        )
        if ref_sheet is None:
            return defaults
        dicts = _import_lib().discover_ref_dictionaries(ref_sheet)
        target = _norm_header(_pcfg.REF_CRITICALITY_TITLE)
        ref_dict = None
        di = 0
        while di < len(dicts):
            d = dicts[di]
            tit = d.get(u"titles") or []
            if tit and _norm_header(tit[0]) == target:
                ref_dict = d
                break
            di = di + 1
        if ref_dict is None:
            return defaults
        items = _import_lib().read_ref_dictionary_items(ref_sheet, ref_dict)
        out = []
        ii = 0
        while ii < len(items):
            vals = items[ii]
            ii = ii + 1
            if vals and unicode(vals[0] or u"").strip():
                out.append(unicode(vals[0]).strip())
        return out if out else defaults
    except Exception:
        return defaults


def ensure_plan_ref_blocks(doc):
    """Создать блоки «Праздники» и «Критичность плана» на листе справочников."""
    try:
        ref_sheet = _get_sheet_by_name(
            doc, getattr(_tcfg, "REF_SHEET_NAME", u"__Справочники_задачи")
        )
        if ref_sheet is None:
            return False
        hdr = int(getattr(_tcfg, "REF_LIST_HEADER_ROW", 0))
        dicts = _import_lib().discover_ref_dictionaries(ref_sheet)
        have = set()
        di = 0
        while di < len(dicts):
            tit = dicts[di].get(u"titles") or []
            if tit:
                have.add(_norm_header(tit[0]))
            di = di + 1
        last_col = 0
        try:
            cursor = ref_sheet.createCursor()
            cursor.gotoEndOfUsedArea(False)
            last_col = int(cursor.getRangeAddress().EndColumn)
        except Exception:
            pass
        col = last_col + 2
        if _norm_header(_pcfg.REF_HOLIDAYS_TITLE) not in have:
            _set_cell_text(
                ref_sheet.getCellByPosition(col, hdr),
                _pcfg.REF_HOLIDAYS_TITLE,
            )
            col = col + 2
        if _norm_header(_pcfg.REF_CRITICALITY_TITLE) not in have:
            _set_cell_text(
                ref_sheet.getCellByPosition(col, hdr),
                _pcfg.REF_CRITICALITY_TITLE,
            )
            ci = 0
            data_start = int(getattr(_tcfg, "REF_LIST_DATA_START_ROW", 1))
            while ci < len(_pcfg.REF_CRITICALITY_DEFAULTS):
                _set_cell_text(
                    ref_sheet.getCellByPosition(
                        col, data_start + ci
                    ),
                    _pcfg.REF_CRITICALITY_DEFAULTS[ci],
                )
                ci = ci + 1
        return True
    except Exception as exc:
        _log("ensure_plan_ref_blocks: %s" % exc)
        return False


def is_workday(d, holidays=None):
    if d is None:
        return False
    if d.weekday() >= 5:
        return False
    if holidays and d in holidays:
        return False
    return True


def snap_workday_left(d, holidays=None):
    if d is None:
        return None
    try:
        from datetime import timedelta

        guard = 0
        while not is_workday(d, holidays) and guard < 366:
            d = d - timedelta(days=1)
            guard = guard + 1
        return d
    except Exception:
        return d


def next_workday_after(d, holidays=None):
    if d is None:
        return None
    try:
        from datetime import timedelta

        nd = d + timedelta(days=1)
        guard = 0
        while not is_workday(nd, holidays) and guard < 366:
            nd = nd + timedelta(days=1)
            guard = guard + 1
        return nd
    except Exception:
        return d


def add_workdays_inclusive(start, n, holidays=None, snap_start=True):
    """start + (n-1) рабочих дней, интервал включительно."""
    if start is None or n is None:
        return None
    try:
        n = int(n)
    except (TypeError, ValueError):
        n = 1
    if n < 1:
        n = 1
    d = snap_workday_left(start, holidays) if snap_start else start
    if n == 1:
        return d
    try:
        from datetime import timedelta

        count = 1
        guard = 0
        while count < n and guard < 3660:
            d = d + timedelta(days=1)
            if is_workday(d, holidays):
                count = count + 1
            guard = guard + 1
        return d
    except Exception:
        return start


def workdays_between(start, end, holidays=None):
    if start is None or end is None:
        return 0
    if start > end:
        return 0
    try:
        from datetime import timedelta

        n = 0
        d = start
        guard = 0
        while d <= end and guard < 3660:
            if is_workday(d, holidays):
                n = n + 1
            d = d + timedelta(days=1)
            guard = guard + 1
        return n
    except Exception:
        return 0


def _children_items(rows, parent_no):
    out = []
    p = unicode(parent_no or u"").strip()
    i = 0
    while i < len(rows):
        r = rows[i]
        i = i + 1
        if r.get(u"is_summary"):
            continue
        par = unicode(r.get(u"parent_item") or u"").strip()
        ino = unicode(r.get(u"item_no") or u"").strip()
        if p:
            if par == p:
                out.append(r)
        elif par == u"" and ino and u"." not in ino:
            out.append(r)
    return out


def _summary_children(rows, parent_no):
    out = []
    p = unicode(parent_no or u"").strip()
    i = 0
    while i < len(rows):
        r = rows[i]
        i = i + 1
        par = unicode(r.get(u"parent_item") or u"").strip()
        if par == p:
            out.append(r)
    return out


def _apply_row_date_snap(row, holidays):
    if row.get(u"is_summary"):
        return
    if not row.get(u"start_no_snap"):
        sp = row.get(u"date_start_parts")
        ds = _parts_to_date(sp)
        if ds:
            row[u"date_start_parts"] = _date_to_parts(snap_workday_left(ds, holidays))
    if not row.get(u"end_no_snap"):
        ep = row.get(u"date_end_parts")
        de = _parts_to_date(ep)
        if de:
            row[u"date_end_parts"] = _date_to_parts(snap_workday_left(de, holidays))


def compute_row_end_from_duration(row, holidays=None):
    """Дата окончания по началу и длительности с учётом флагов «Не корр.»."""
    if holidays is None:
        holidays = set()
    ds = _parts_to_date(row.get(u"date_start_parts"))
    dur = row.get(u"duration_wd")
    try:
        dur = int(dur) if dur else 0
    except (TypeError, ValueError):
        dur = 0
    if not ds or dur < 1:
        return None
    snap_start = not row.get(u"start_no_snap")
    de = add_workdays_inclusive(ds, dur, holidays, snap_start=snap_start)
    if de and not row.get(u"end_no_snap"):
        de = snap_workday_left(de, holidays)
    return _date_to_parts(de) if de else None


def plan_row_display_name(row):
    """Название этапа для списка (stage_name, иначе legacy result_doc)."""
    stage = unicode(row.get(u"stage_name") or u"").strip()
    if not stage:
        stage = unicode(row.get(u"result_doc") or u"").strip()
    return stage


def confirm_early_plan_done(doc, row):
    """True — можно отметить выполненным; False — отмена досрочного выполнения."""
    ep = _parts_to_date(row.get(u"date_end_parts"))
    if not ep:
        return True
    try:
        from datetime import date

        if date.today() >= ep:
            return True
    except Exception:
        return True
    try:
        from todo_task_dialog_lib import show_todo_task_confirm

        disp = format_plan_date_field(row.get(u"date_end_parts"))
        msg = unicode(getattr(_pcfg, "PLAN_DONE_EARLY_CONFIRM_MSG", u"")) % disp
        title = unicode(getattr(_pcfg, "PLAN_DONE_EARLY_CONFIRM_TITLE", u"Выполнение"))
        return bool(show_todo_task_confirm(doc=doc, text=msg, title=title))
    except Exception:
        return True


def plan_row_status_code(row, holidays=None):
    """Код срочности: done | overdue | soon | normal."""
    if row.get(u"is_done"):
        return unicode(getattr(_pcfg, "PLAN_ROW_STATUS_DONE", u"done"))
    ep = _parts_to_date(row.get(u"date_end_parts"))
    if not ep:
        return unicode(getattr(_pcfg, "PLAN_ROW_STATUS_NORMAL", u"normal"))
    try:
        from datetime import date

        today = date.today()
        if ep < today:
            return unicode(getattr(_pcfg, "PLAN_ROW_STATUS_OVERDUE", u"overdue"))
        wd_left = workdays_between(today, ep, holidays)
        soon_lim = int(getattr(_pcfg, "PLAN_ROW_SOON_WORKDAYS", 2))
        if wd_left <= soon_lim:
            return unicode(getattr(_pcfg, "PLAN_ROW_STATUS_SOON", u"soon"))
        return unicode(getattr(_pcfg, "PLAN_ROW_STATUS_NORMAL", u"normal"))
    except Exception:
        return unicode(getattr(_pcfg, "PLAN_ROW_STATUS_NORMAL", u"normal"))


def _cascade_scope_parent(row):
    """Родитель для каскада подэтапов: дети итогового — под ним, иначе — общий родитель."""
    ino = unicode(row.get(u"item_no") or u"").strip()
    par = unicode(row.get(u"parent_item") or u"").strip()
    if not par:
        par = parent_of_item_no(ino)
    if row.get(u"is_summary"):
        return ino
    return par if par else ino


def _rollup_plan_summaries(rows, holidays):
    summaries = []
    i = 0
    while i < len(rows):
        if rows[i].get(u"is_summary"):
            summaries.append(rows[i])
        i = i + 1
    summaries.sort(
        key=lambda r: unicode(r.get(u"item_no") or u"").count(u"."),
        reverse=True,
    )
    si = 0
    while si < len(summaries):
        srow = summaries[si]
        si = si + 1
        kids = _summary_children(rows, srow.get(u"item_no"))
        starts = []
        ends = []
        ki = 0
        while ki < len(kids):
            k = kids[ki]
            ki = ki + 1
            d1 = _parts_to_date(k.get(u"date_start_parts"))
            d2 = _parts_to_date(k.get(u"date_end_parts"))
            if d1:
                starts.append(d1)
            if d2:
                ends.append(d2)
        if starts and ends:
            mn = min(starts)
            mx = max(ends)
            srow[u"date_start_parts"] = _date_to_parts(mn)
            srow[u"date_end_parts"] = _date_to_parts(mx)
            srow[u"duration_wd"] = workdays_between(mn, mx, holidays)
    return rows


def _resolve_cascade_row_index(rows, row_index, row_id=None):
    rows = sort_plan_rows(rows)
    if row_id:
        ri = 0
        while ri < len(rows):
            if rows[ri].get(u"row_id") == row_id:
                return rows, ri
            ri = ri + 1
    return rows, int(row_index)


def _cascade_siblings_after_row(rows, row_index, holidays):
    """Цепочка подэтапов того же родителя ниже row_index."""
    if row_index < 0 or row_index >= len(rows):
        return rows
    row = rows[row_index]
    if row.get(u"is_summary"):
        return rows
    scope = _cascade_scope_parent(row)
    cur_sort = plan_row_order_key(row)
    prev = row
    i = row_index + 1
    while i < len(rows):
        nxt = rows[i]
        if nxt.get(u"is_summary"):
            i = i + 1
            continue
        nxt_par = unicode(nxt.get(u"parent_item") or u"").strip()
        if not nxt_par:
            nxt_par = parent_of_item_no(nxt.get(u"item_no"))
        if nxt_par != scope:
            break
        if plan_row_order_key(nxt) <= cur_sort:
            i = i + 1
            continue
        if nxt.get(u"start_no_snap"):
            break
        ep = _parts_to_date(prev.get(u"date_end_parts"))
        if not ep:
            break
        ns = next_workday_after(ep, holidays)
        nxt[u"date_start_parts"] = _date_to_parts(ns)
        if not nxt.get(u"start_no_snap"):
            ds2 = _parts_to_date(nxt.get(u"date_start_parts"))
            if ds2:
                nxt[u"date_start_parts"] = _date_to_parts(
                    snap_workday_left(ds2, holidays)
                )
        ndur = nxt.get(u"duration_wd")
        try:
            ndur = int(ndur) if ndur else int(_pcfg.PLAN_DEFAULT_DURATION)
        except (TypeError, ValueError):
            ndur = int(_pcfg.PLAN_DEFAULT_DURATION)
        ns2 = _parts_to_date(nxt.get(u"date_start_parts"))
        if ns2 and ndur >= 1:
            nxt[u"duration_wd"] = ndur
            eparts = compute_row_end_from_duration(nxt, holidays)
            if eparts:
                nxt[u"date_end_parts"] = eparts
                nxt[u"duration_wd"] = workdays_between(
                    ns2, _parts_to_date(eparts), holidays
                )
        prev = nxt
        i = i + 1
    return rows


def cascade_plan_from_row(rows, row_index, holidays=None, row_id=None):
    """После смены длительности: конец строки и цепочка подэтапов того же родителя вниз."""
    if holidays is None:
        holidays = set()
    if not rows:
        return rows
    rows, row_index = _resolve_cascade_row_index(rows, row_index, row_id=row_id)
    if row_index < 0 or row_index >= len(rows):
        return rows
    row = rows[row_index]
    if row.get(u"is_summary"):
        return rows

    eparts = compute_row_end_from_duration(row, holidays)
    if eparts:
        row[u"date_end_parts"] = eparts
        ds = _parts_to_date(row.get(u"date_start_parts"))
        de = _parts_to_date(eparts)
        if ds and de:
            row[u"duration_wd"] = workdays_between(ds, de, holidays)

    rows = _cascade_siblings_after_row(rows, row_index, holidays)
    return _rollup_plan_summaries(rows, holidays)


def cascade_plan_from_end_change(rows, row_index, holidays=None, row_id=None):
    """После смены даты окончания: пересчёт длительности и цепочка подэтапов вниз."""
    if holidays is None:
        holidays = set()
    if not rows:
        return rows
    rows, row_index = _resolve_cascade_row_index(rows, row_index, row_id=row_id)
    if row_index < 0 or row_index >= len(rows):
        return rows
    row = rows[row_index]
    if row.get(u"is_summary"):
        return rows
    ds = _parts_to_date(row.get(u"date_start_parts"))
    de = _parts_to_date(row.get(u"date_end_parts"))
    if ds and de and de >= ds:
        row[u"duration_wd"] = workdays_between(ds, de, holidays)
    elif ds and de and de < ds:
        row[u"duration_wd"] = 0
    rows = _cascade_siblings_after_row(rows, row_index, holidays)
    return _rollup_plan_summaries(rows, holidays)


def recalc_plan_rows(rows, holidays=None, auto_chain=None):
    """Пересчитать даты, длительность и итоговые пункты в памяти."""
    if auto_chain is None:
        auto_chain = bool(_pcfg.PLAN_AUTO_CHAIN_DATES)
    if not rows:
        return rows
    rows = sort_plan_rows(rows)

    i = 0
    while i < len(rows):
        r = rows[i]
        if r.get(u"is_summary"):
            i = i + 1
            continue
        _apply_row_date_snap(r, holidays)
        sp = r.get(u"date_start_parts")
        ep = r.get(u"date_end_parts")
        ds = _parts_to_date(sp)
        de = _parts_to_date(ep)
        dur = r.get(u"duration_wd")
        try:
            dur = int(dur) if dur else 0
        except (TypeError, ValueError):
            dur = 0
        if ds and dur >= 1 and not de:
            eparts = compute_row_end_from_duration(r, holidays)
            if eparts:
                r[u"date_end_parts"] = eparts
        elif ds and de:
            if de >= ds:
                r[u"duration_wd"] = workdays_between(ds, de, holidays)
            else:
                r[u"duration_wd"] = 0
        elif ds and dur >= 1:
            eparts = compute_row_end_from_duration(r, holidays)
            if eparts:
                r[u"date_end_parts"] = eparts
        i = i + 1

    if auto_chain:
        i = 0
        while i < len(rows) - 1:
            cur = rows[i]
            nxt = rows[i + 1]
            if cur.get(u"is_summary") or nxt.get(u"is_summary"):
                i = i + 1
                continue
            ep = _parts_to_date(cur.get(u"date_end_parts"))
            if ep:
                if nxt.get(u"start_no_snap"):
                    i = i + 1
                    continue
                ns = next_workday_after(ep, holidays)
                if ns:
                    if not nxt.get(u"start_no_snap"):
                        ns = snap_workday_left(ns, holidays)
                    nxt[u"date_start_parts"] = _date_to_parts(ns)
                    nd = _parts_to_date(nxt.get(u"date_end_parts"))
                    dur = nxt.get(u"duration_wd")
                    try:
                        dur = int(dur) if dur else int(_pcfg.PLAN_DEFAULT_DURATION)
                    except (TypeError, ValueError):
                        dur = int(_pcfg.PLAN_DEFAULT_DURATION)
                    if dur >= 1:
                        nxt[u"duration_wd"] = dur
                        eparts = compute_row_end_from_duration(nxt, holidays)
                        if eparts:
                            nxt[u"date_end_parts"] = eparts
                            ns3 = _parts_to_date(nxt.get(u"date_start_parts"))
                            ne = _parts_to_date(eparts)
                            if ns3 and ne:
                                nxt[u"duration_wd"] = workdays_between(
                                    ns3, ne, holidays
                                )
            i = i + 1

    return _rollup_plan_summaries(rows, holidays)


def _load_rows_from_sheet(sh, cols, doc, guid_filter=None):
    hdr = int(cols.get(u"_hdr_row", _pcfg.PLAN_HEADER_ROW))
    last = _used_last_row(sh, min_cols=_table_max_from_cols(cols))
    out = []
    r = hdr + 1
    while r <= last:
        if guid_filter:
            rg = _guid_key(_row_get(sh, cols, r, u"guid"))
            if rg != guid_filter:
                r = r + 1
                continue
        out.append(_row_to_dict(sh, cols, r, doc=doc))
        r = r + 1
    return out


def _load_legacy_plan_by_guid(doc, guid):
    gk = _guid_key(guid)
    if not gk:
        return []
    legacy = unicode(_pcfg.PLAN_SHEET_NAME)
    if not doc.Sheets.hasByName(legacy):
        return []
    sh = doc.Sheets.getByName(legacy)
    cols = _build_plan_col_map(sh)
    if cols is None:
        return []
    cols = _ensure_missing_plan_columns(sh, cols)
    return _load_rows_from_sheet(sh, cols, doc, guid_filter=gk)


def _load_uno_plan_by_guid(doc, guid):
    """Чтение legacy-плана со скрытых листов книги задач (миграция)."""
    gk = _guid_key(guid)
    if not gk or doc is None:
        return []
    name = plan_sheet_name_for_guid(guid)
    try:
        if doc.Sheets.hasByName(name):
            sh = doc.Sheets.getByName(name)
            cols = _build_plan_col_map(sh)
            if cols is not None:
                cols = _ensure_missing_plan_columns(sh, cols)
                out = _load_rows_from_sheet(sh, cols, doc)
                if out:
                    return out
    except Exception as exc:
        _log("uno plan load: %s" % exc)
    return _load_legacy_plan_by_guid(doc, guid)


def load_plan_by_guid(doc, guid):
    """Загрузить план задачи из внешнего .xlsx (Plans_for_…)."""
    gk = _guid_key(guid)
    if not gk:
        return []
    out = []
    wb, path = _plans_xlsx_open(doc, create=False)
    if wb is not None:
        try:
            ws, cols = _xlsx_get_or_create_plan_ws(wb, guid, create=False)
            if ws is not None and cols is not None:
                out = _xlsx_load_rows_from_ws(ws, cols, doc)
        finally:
            try:
                wb.close()
            except Exception:
                pass
    if not out:
        out = _load_uno_plan_by_guid(doc, guid)
    return sort_plan_rows(out)


def count_plan_items(doc, guid):
    gk = _guid_key(guid)
    if not gk:
        return 0
    import os

    path = plans_workbook_path(doc)
    if path and os.path.isfile(path):
        wb, _p = _plans_xlsx_open(doc, create=False)
        if wb is not None:
            try:
                ws, cols = _xlsx_get_or_create_plan_ws(wb, guid, create=False)
                if ws is not None and cols is not None:
                    return len(_xlsx_load_rows_from_ws(ws, cols, doc))
            finally:
                try:
                    wb.close()
                except Exception:
                    pass
    return len(_load_uno_plan_by_guid(doc, guid))


def _table_max_from_cols(cols):
    m = 0
    for k, v in cols.items():
        if unicode(k).startswith(u"_"):
            continue
        try:
            m = max(m, int(v))
        except Exception:
            pass
    return m


def _clear_plan_data_rows(sh, cols):
    hdr = int(cols.get(u"_hdr_row", _pcfg.PLAN_HEADER_ROW))
    last = _used_last_row(sh, min_cols=_table_max_from_cols(cols))
    if last <= hdr:
        return True
    try:
        sh.getRows().removeByIndex(hdr + 1, last - hdr)
        return True
    except Exception as exc:
        _log("clear plan sheet: %s" % exc)
        return False


def _delete_legacy_plan_rows(doc, guid):
    gk = _guid_key(guid)
    if not gk:
        return True
    legacy = unicode(_pcfg.PLAN_SHEET_NAME)
    if not doc.Sheets.hasByName(legacy):
        return True
    sh = doc.Sheets.getByName(legacy)
    cols = _build_plan_col_map(sh)
    if cols is None:
        return True
    hdr = int(cols.get(u"_hdr_row", _pcfg.PLAN_HEADER_ROW))
    last = _used_last_row(sh, min_cols=_table_max_from_cols(cols))
    to_del = []
    r = hdr + 1
    while r <= last:
        if _guid_key(_row_get(sh, cols, r, u"guid")) == gk:
            to_del.append(int(r))
        r = r + 1
    to_del.sort(reverse=True)
    i = 0
    while i < len(to_del):
        try:
            sh.getRows().removeByIndex(to_del[i], 1)
        except Exception as exc:
            _log("delete legacy plan row: %s" % exc)
            return False
        i = i + 1
    return True


def _delete_plan_rows_on_sheet(doc, guid):
    """Удалить legacy-листы плана в книге задач (после переноса во внешний файл)."""
    gk = _guid_key(guid)
    if not gk:
        return True
    name = plan_sheet_name_for_guid(guid)
    sheets = doc.Sheets
    if sheets.hasByName(name):
        sh = sheets.getByName(name)
        cols = _build_plan_col_map(sh)
        if cols and not _clear_plan_data_rows(sh, cols):
            return False
        try:
            sheets.removeByName(name)
        except Exception as exc:
            _log("remove plan sheet: %s" % exc)
    return _delete_legacy_plan_rows(doc, guid)


def save_plan_by_guid(doc, guid, rows):
    """Сохранить план задачи во внешний .xlsx (Plans_for_…)."""
    gk = _guid_key(guid)
    if not gk:
        _set_plan_io_error(u"Пустой GUID задачи.")
        return False
    rows = recalc_plan_rows(list(rows or []), holidays=load_holidays(doc))
    wb, path = _plans_xlsx_open(doc, create=True)
    if wb is None:
        return False
    try:
        ws, cols = _xlsx_get_or_create_plan_ws(wb, guid, create=True)
        if ws is None or cols is None:
            _set_plan_io_error(u"Не удалось создать лист плана.")
            return False
        _xlsx_clear_data_rows(ws, cols)
        hdr = int(cols.get(u"_hdr_row", int(_pcfg.PLAN_HEADER_ROW) + 1))
        dest = hdr + 1
        i = 0
        while i < len(rows):
            item = dict(rows[i])
            item[u"guid"] = unicode(guid)
            _xlsx_dict_to_row_write(ws, cols, dest, item, is_new=True)
            dest = dest + 1
            i = i + 1
        # убрать служебный _index, если есть реальные листы планов
        placeholder = unicode(
            getattr(_pcfg, "PLAN_STORE_PLACEHOLDER_SHEET", u"_index")
        )
        if placeholder in wb.sheetnames and len(wb.sheetnames) > 1:
            try:
                del wb[placeholder]
            except Exception:
                pass
        if not _plans_xlsx_save(wb, path):
            return False
    finally:
        try:
            wb.close()
        except Exception:
            pass
    # После успешной записи во внешний файл — убрать legacy-листы в книге задач.
    try:
        _delete_plan_rows_on_sheet(doc, guid)
    except Exception as exc:
        _log("cleanup uno plan after save: %s" % exc)
    _log("plan saved: %s sheet=%s rows=%s" % (path, plan_sheet_name_for_guid(guid), len(rows)))
    return True


def delete_plan_by_guid(doc, guid):
    """Удалить лист плана задачи из внешнего .xlsx."""
    gk = _guid_key(guid)
    if not gk:
        return True
    name = plan_sheet_name_for_guid(guid)
    wb, path = _plans_xlsx_open(doc, create=False)
    if wb is not None:
        try:
            if name in wb.sheetnames:
                try:
                    del wb[name]
                except Exception as exc:
                    _log("remove plan xlsx sheet: %s" % exc)
                    return False
                if not wb.sheetnames:
                    placeholder = unicode(
                        getattr(_pcfg, "PLAN_STORE_PLACEHOLDER_SHEET", u"_index")
                    )
                    wb.create_sheet(title=placeholder[:31])
                if not _plans_xlsx_save(wb, path):
                    return False
        finally:
            try:
                wb.close()
            except Exception:
                pass
    try:
        _delete_plan_rows_on_sheet(doc, guid)
    except Exception as exc:
        _log("cleanup uno plan after delete: %s" % exc)
    return True


def archive_plan_by_guid(doc, guid, source_sheet_name=u"", task_name=u""):
    """Архивировать план задачи на лист __Архив_план во внешнем .xlsx."""
    gk = _guid_key(guid)
    if not gk:
        return True
    rows = load_plan_by_guid(doc, guid)
    if not rows:
        return True
    wb, path = _plans_xlsx_open(doc, create=True)
    if wb is None:
        return False
    try:
        ws_plan, cols = _xlsx_get_or_create_plan_ws(wb, guid, create=True)
        if cols is None:
            return False
        arch, arch_meta_base = _xlsx_ensure_archive_ws(wb, cols)
        hdr = int(_pcfg.PLAN_HEADER_ROW) + 1
        dest = int(arch.max_row or hdr) + 1
        if dest <= hdr:
            dest = hdr + 1
        stamp = _now_timestamp_text()
        i = 0
        while i < len(rows):
            item = dict(rows[i])
            item[u"guid"] = unicode(guid)
            _xlsx_dict_to_row_write(arch, cols, dest, item, is_new=False)
            arch.cell(row=dest, column=arch_meta_base + 0).value = stamp
            arch.cell(row=dest, column=arch_meta_base + 1).value = unicode(
                source_sheet_name or plan_sheet_name_for_guid(guid)
            )
            arch.cell(row=dest, column=arch_meta_base + 2).value = unicode(
                task_name or u""
            )
            dest = dest + 1
            i = i + 1
        if not _plans_xlsx_save(wb, path):
            return False
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return True


def plan_max_end_parts(rows):
    mx = None
    i = 0
    while i < len(rows or []):
        ep = _parts_to_date((rows[i] or {}).get(u"date_end_parts"))
        if ep and (mx is None or ep > mx):
            mx = ep
        i = i + 1
    return _date_to_parts(mx) if mx else None


def apply_plan_save_to_task(doc, sheet, cols, row, rows):
    """После сохранения плана: предложить срок; обновить метки редакции задачи.

    Возвращает dict: synced_due, due_parts, due, comment (актуальные значения с листа).
    """
    lib = _import_lib()
    out = {
        u"synced_due": False,
        u"due_parts": None,
        u"due": u"",
        u"comment": None,
    }
    parts = plan_max_end_parts(rows)
    if parts and cols and u"due" in cols and sheet is not None and row is not None:
        disp = format_plan_date_display(doc, parts)
        from todo_task_dialog_lib import show_todo_task_confirm

        due_cell = sheet.getCellByPosition(cols[u"due"], int(row))
        old_parts = lib._read_due_from_cell(due_cell)
        if show_todo_task_confirm(
            doc=doc,
            text=unicode(_pcfg.PLAN_SYNC_DUE_MSG) % disp,
            title=unicode(_pcfg.PLAN_SYNC_DUE_TITLE),
        ):
            lib.set_cell_due_date(doc, due_cell, parts[0], parts[1], parts[2])
            out[u"synced_due"] = True
            out[u"due_parts"] = parts
            out[u"due"] = u"%02d.%02d.%04d" % (
                int(parts[0]),
                int(parts[1]),
                int(parts[2]),
            )
            if old_parts != parts and u"comment" in cols:
                cmt_cell = sheet.getCellByPosition(cols[u"comment"], int(row))
                old_cmt = lib._cell_text(cmt_cell)
                note = unicode(
                    getattr(
                        _pcfg,
                        u"PLAN_SYNC_DUE_COMMENT",
                        u"Дата окончания исполнения задачи изменена согласно плану выполнения.",
                    )
                )
                merged = lib._accumulate_comment_text(old_cmt, note, None, None)
                lib._set_cell_text(cmt_cell, merged)
                try:
                    cmt_cell.IsTextWrapped = True
                except Exception:
                    pass
                out[u"comment"] = merged
    if sheet is not None and cols is not None and row is not None:
        lib._stamp_task_timestamps(doc, sheet, cols, int(row), is_new=False)
        lib._stamp_task_owner(sheet, cols, int(row), is_new=False, doc=doc)
        try:
            lib._autofit_task_rows(sheet, [int(row)])
        except Exception:
            pass
        try:
            fresh = lib.read_task_fields(sheet, cols, int(row))
            if fresh.get(u"due_parts"):
                out[u"due_parts"] = fresh.get(u"due_parts")
                out[u"due"] = fresh.get(u"due") or out.get(u"due") or u""
            if fresh.get(u"comment") is not None:
                out[u"comment"] = fresh.get(u"comment")
        except Exception:
            pass
    return out


def default_plan_rows(doc=None, guid=u""):
    """Один итоговый пункт «1» с датами по умолчанию."""
    holidays = load_holidays(doc) if doc else set()
    parts = _import_lib().default_due_parts()
    ds = _parts_to_date(parts)
    if ds:
        ds = snap_workday_left(ds, holidays)
        parts = _date_to_parts(ds)
    dur = int(_pcfg.PLAN_DEFAULT_DURATION)
    de = add_workdays_inclusive(ds, dur, holidays) if ds else None
    return [
        {
            u"guid": unicode(guid or u""),
            u"item_no": u"1",
            u"is_summary": True,
            u"parent_item": u"",
            u"sort_key": item_no_to_sort_key(u"1"),
            u"date_start_parts": parts,
            u"date_end_parts": _date_to_parts(de),
            u"duration_wd": dur,
            u"assignees": u"",
            u"stage_name": u"",
            u"result_doc": u"",
            u"criticality": u"",
            u"row_id": _new_guid(),
        }
    ]


def plan_root_item_no(item_no):
    """Корневой номер пункта: «1.2» → «1», «2» → «2»."""
    s = normalize_item_no_from_store(item_no)
    if not s:
        return u""
    if u"." in s:
        return s.split(u".", 1)[0]
    return s


def _plan_delete_targets(rows, selected_ino):
    """Номера пунктов к удалению: выбранный + прямые/косвенные потомки."""
    sel = unicode(selected_ino or u"").strip()
    if not sel:
        return set()
    kill = {sel}
    changed = True
    while changed:
        changed = False
        i = 0
        while i < len(rows):
            ino = unicode(rows[i].get(u"item_no") or u"").strip()
            if not ino or ino in kill:
                i = i + 1
                continue
            par = unicode(rows[i].get(u"parent_item") or u"").strip()
            if not par:
                par = parent_of_item_no(ino)
            if par in kill:
                kill.add(ino)
                changed = True
            elif ino.startswith(sel + u"."):
                kill.add(ino)
                changed = True
            i = i + 1
    return kill


def next_root_item_no(rows):
    mx = 0
    i = 0
    while i < len(rows):
        ino = normalize_item_no_from_store(rows[i].get(u"item_no"))
        if ino and u"." not in ino:
            try:
                mx = max(mx, int(ino))
            except ValueError:
                pass
        i = i + 1
    return unicode(mx + 1)


def next_child_item_no(rows, parent_no):
    p = normalize_item_no_from_store(parent_no)
    mx = 0
    i = 0
    while i < len(rows):
        ino = normalize_item_no_from_store(rows[i].get(u"item_no"))
        if ino.startswith(p + u"."):
            tail = ino[len(p) + 1 :]
            if tail and u"." not in tail:
                try:
                    mx = max(mx, int(tail))
                except ValueError:
                    pass
        i = i + 1
    return u"%s.%s" % (p, mx + 1)


def _safe_export_filename(task_name, guid=u""):
    import re

    base = unicode(task_name or u"").strip()
    if not base:
        base = unicode(guid or u"").strip()[:8] or u"задача"
    base = re.sub(r'[\\/:*?"<>|\r\n]+', u"_", base)
    base = base.strip(u"._ ") or u"задача"
    if len(base) > 48:
        base = base[:48].rstrip(u"._ ")
    try:
        from datetime import date

        stamp = date.today().strftime("%Y%m%d")
    except Exception:
        stamp = u"export"
    return u"%s_%s_%s.xlsx" % (
        unicode(_pcfg.PLAN_EXPORT_FILE_PREFIX),
        base,
        stamp,
    )


def _item_no_depth(item_no):
    s = normalize_item_no_from_store(item_no)
    if not s:
        return 0
    return max(0, s.count(u"."))


def _pick_plan_export_path(suggested_name=u"plan.xlsx"):
    try:
        import os

        import uno

        from todo_task_dialog_lib import _uno_context

        ctx = _uno_context()
        if ctx is None:
            return None
        sm = ctx.getServiceManager()
        picker = sm.createInstanceWithContext(
            "com.sun.star.ui.dialogs.FilePicker", ctx
        )
        try:
            from com.sun.star.ui.dialogs.TemplateDescription import FILESAVE_SIMPLE

            picker.initialize((FILESAVE_SIMPLE,))
        except Exception:
            pass
        try:
            picker.setTitle(unicode(_pcfg.PLAN_EXPORT_TITLE))
        except Exception:
            pass
        try:
            picker.appendFilter(u"Excel", u"*.xlsx")
            picker.setCurrentFilter(u"Excel")
        except Exception:
            pass
        try:
            home = os.path.expanduser(u"~")
            if os.path.isdir(home):
                picker.setDisplayDirectory(uno.systemPathToFileUrl(home))
        except Exception:
            pass
        try:
            picker.setDefaultName(unicode(suggested_name or u"plan.xlsx"))
        except Exception:
            pass
        if int(picker.execute()) != 1:
            return None
        urls = picker.getFiles()
        if not urls or len(urls) < 1:
            return None
        path = os.path.abspath(uno.fileUrlToSystemPath(urls[0]))
        if not path.lower().endswith(u".xlsx"):
            path = path + u".xlsx"
        return path
    except Exception as exc:
        _log("export picker: %s" % exc)
        return None


def _open_exported_workbook(path):
    try:
        import os

        import uno

        from todo_task_dialog_lib import _desktop

        desk = _desktop()
        if desk is None:
            return False
        url = uno.systemPathToFileUrl(os.path.abspath(unicode(path)))
        desk.loadComponentFromURL(url, u"_blank", 0, ())
        return True
    except Exception as exc:
        _log("open export: %s" % exc)
        return False


def _xlsx_cell_display_len(val):
    if val is None:
        return 0
    try:
        from datetime import date, datetime

        if isinstance(val, datetime):
            return 10
        if isinstance(val, date):
            return 10
    except Exception:
        pass
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return max(3, len(unicode(val)))
    s = unicode(val).replace(u"\r\n", u"\n")
    parts = s.split(u"\n")
    mx = 0
    i = 0
    while i < len(parts):
        mx = max(mx, len(parts[i]))
        i = i + 1
    return mx


def _xlsx_autofit_worksheet(ws, hdr_row, data_row, ncol, min_w=8, max_w=55):
    from openpyxl.utils import get_column_letter

    ci = 1
    while ci <= int(ncol):
        letter = get_column_letter(ci)
        mx = float(min_w)
        r = 1
        while r <= ws.max_row:
            ln = _xlsx_cell_display_len(ws.cell(row=r, column=ci).value)
            if ln > mx:
                mx = float(ln)
            r = r + 1
        try:
            ws.column_dimensions[letter].width = min(float(max_w), mx * 1.08 + 2.0)
        except Exception:
            pass
        ci = ci + 1
    _xlsx_autofit_row_heights(ws, hdr_row, data_row, ncol)


def _xlsx_autofit_row_heights(ws, hdr_row, data_row, ncol):
    """Автоподбор высоты строк данных по переносам и ширине колонок."""
    from openpyxl.utils import get_column_letter

    try:
        ws.row_dimensions[int(hdr_row)].height = 22.0
    except Exception:
        pass
    r = int(data_row)
    while r <= ws.max_row:
        mx_lines = 1
        ci = 1
        while ci <= int(ncol):
            cell = ws.cell(row=r, column=ci)
            val = cell.value
            txt = unicode(val or u"")
            if u"\n" in txt:
                mx_lines = max(mx_lines, txt.count(u"\n") + 1)
            ln = _xlsx_cell_display_len(val)
            try:
                col_w = float(ws.column_dimensions[get_column_letter(ci)].width or 10)
            except Exception:
                col_w = 10.0
            if col_w > 0 and ln > col_w * 0.9:
                est = int(ln / max(1.0, col_w * 0.95)) + 1
                if est > mx_lines:
                    mx_lines = est
            ci = ci + 1
        try:
            ws.row_dimensions[r].height = max(18.0, 15.0 * float(mx_lines))
        except Exception:
            pass
        r = r + 1


def _plan_export_date_value(row, parts_key):
    """datetime.date для openpyxl (число Excel + формат), иначе None."""
    return _parts_to_date(row.get(parts_key))


def _plan_export_duration_value(row):
    dv = row.get(u"duration_wd")
    if dv in (None, u""):
        return None
    try:
        n = float(dv)
        if n == int(n):
            return int(n)
        return n
    except (TypeError, ValueError):
        return None


def _plan_export_cell_value(row, key, doc=None):
    if key == u"item_no":
        depth = _item_no_depth(row.get(u"item_no"))
        return (u"  " * depth) + unicode(row.get(u"item_no") or u"")
    if key == u"is_summary":
        return _pcfg.PLAN_SUMMARY_YES if row.get(u"is_summary") else u""
    if key == u"is_done":
        return _pcfg.PLAN_SUMMARY_YES if row.get(u"is_done") else u""
    if key in (u"start_no_snap", u"end_no_snap"):
        return _pcfg.PLAN_SUMMARY_YES if row.get(key) else u""
    if key == u"date_start":
        return _plan_export_date_value(row, u"date_start_parts")
    if key == u"date_end":
        return _plan_export_date_value(row, u"date_end_parts")
    if key == u"duration_wd":
        return _plan_export_duration_value(row)
    return unicode(row.get(key) or u"")


def _plan_export_row_fill(row, holidays=None):
    """PatternFill: done=зелёный, overdue=красный, soon=оранжевый; иначе None."""
    from openpyxl.styles import PatternFill

    code = plan_row_status_code(row, holidays=holidays)
    done = unicode(getattr(_pcfg, "PLAN_ROW_STATUS_DONE", u"done"))
    overdue = unicode(getattr(_pcfg, "PLAN_ROW_STATUS_OVERDUE", u"overdue"))
    soon = unicode(getattr(_pcfg, "PLAN_ROW_STATUS_SOON", u"soon"))
    if code == done:
        rgb = unicode(getattr(_pcfg, "PLAN_EXPORT_FILL_DONE", u"C8E6C9"))
    elif code == overdue:
        rgb = unicode(getattr(_pcfg, "PLAN_EXPORT_FILL_OVERDUE", u"FFCDD2"))
    elif code == soon:
        rgb = unicode(getattr(_pcfg, "PLAN_EXPORT_FILL_SOON", u"FFE0B2"))
    else:
        return None
    rgb = unicode(rgb or u"").strip()
    if not rgb:
        return None
    return PatternFill(fill_type="solid", fgColor=rgb)


def export_plan_to_xlsx(doc, guid, task_name=u"", dest_path=None, open_after=True):
    """
    Выгрузить план задачи в новую книгу .xlsx (openpyxl).
    Возвращает (ok, message_or_path).
    """
    rows = load_plan_by_guid(doc, guid)
    if not rows:
        return False, unicode(_pcfg.PLAN_EXPORT_NO_PLAN)
    fname = _safe_export_filename(task_name, guid=guid)
    if not dest_path:
        dest_path = _pick_plan_export_path(fname)
        if not dest_path:
            return False, u""
    dest_path = unicode(dest_path)
    try:
        import openpyxl_bundled  # noqa: F401
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        return False, u"openpyxl недоступен: %s" % exc

    cols = list(getattr(_pcfg, "PLAN_EXPORT_COLUMNS", ()) or ())
    if not cols:
        cols = [
            (u"item_no", _pcfg.PLAN_COL_ITEM, 10),
            (u"stage_name", _pcfg.PLAN_COL_STAGE, 28),
            (u"is_done", _pcfg.PLAN_COL_DONE, 10),
            (u"date_start", _pcfg.PLAN_COL_START, 14),
            (u"date_end", _pcfg.PLAN_COL_END, 14),
        ]
    ncol = len(cols)
    font_name = unicode(getattr(_pcfg, "PLAN_EXPORT_FONT_NAME", u"PT Sans"))
    title_size = int(getattr(_pcfg, "PLAN_EXPORT_FONT_TITLE_SIZE", 12))
    hdr_size = int(getattr(_pcfg, "PLAN_EXPORT_FONT_HDR_SIZE", 12))
    body_size = int(getattr(_pcfg, "PLAN_EXPORT_FONT_BODY_SIZE", 11))
    date_fmt = unicode(getattr(_pcfg, "PLAN_EXPORT_DATE_FORMAT", u"DD.MM.YYYY"))
    dur_fmt = unicode(getattr(_pcfg, "PLAN_EXPORT_DURATION_FORMAT", u"0"))
    wb = Workbook()
    ws = wb.active
    try:
        ws.title = unicode(_pcfg.PLAN_EXPORT_SHEET)[:31]
    except Exception:
        pass

    title_fill = PatternFill(fill_type="solid", fgColor="0D47A1")
    title_font = Font(name=font_name, bold=True, size=title_size, color="FFFFFF")
    hdr_fill = PatternFill(fill_type="solid", fgColor="1565C0")
    hdr_font = Font(name=font_name, bold=True, size=hdr_size, color="FFFFFF")
    body_font = Font(name=font_name, size=body_size)
    sum_font = Font(name=font_name, bold=True, size=body_size)
    thin = Side(style="thin", color="B0B0B0")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    title = u"%s: %s" % (
        unicode(_pcfg.PLAN_EXPORT_DOC_TITLE),
        unicode(task_name or guid or u"").strip() or u"—",
    )
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=max(1, ncol)
    )
    tc = ws.cell(row=1, column=1, value=title)
    tc.font = title_font
    tc.fill = title_fill
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    hdr_row = 3
    ci = 0
    while ci < ncol:
        key, label, width = cols[ci]
        col_i = ci + 1
        hc = ws.cell(row=hdr_row, column=col_i, value=unicode(label))
        hc.font = hdr_font
        hc.fill = hdr_fill
        hc.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        hc.border = cell_border
        try:
            ws.column_dimensions[get_column_letter(col_i)].width = float(width)
        except Exception:
            pass
        ci = ci + 1
    ws.row_dimensions[hdr_row].height = 22

    holidays = None
    try:
        holidays = load_holidays(doc)
    except Exception:
        holidays = None

    data_row = hdr_row + 1
    ri = 0
    while ri < len(rows):
        row = rows[ri]
        excel_r = data_row + ri
        is_sum = bool(row.get(u"is_summary"))
        try:
            row_fill = _plan_export_row_fill(row, holidays=holidays)
        except Exception:
            row_fill = None
        ci = 0
        while ci < ncol:
            key = cols[ci][0]
            col_i = ci + 1
            val = _plan_export_cell_value(row, key, doc=doc)
            cell = ws.cell(row=excel_r, column=col_i, value=val)
            cell.border = cell_border
            cell.font = sum_font if is_sum else body_font
            if row_fill is not None:
                cell.fill = row_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if key == u"item_no":
                depth = _item_no_depth(row.get(u"item_no"))
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                    indent=min(depth, 6),
                    wrap_text=True,
                )
            elif key in (u"date_start", u"date_end"):
                if val is not None:
                    cell.number_format = date_fmt
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif key == u"duration_wd":
                if val is not None:
                    cell.number_format = dur_fmt
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            elif key in (u"is_done", u"start_no_snap", u"end_no_snap", u"criticality"):
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            ci = ci + 1
        ri = ri + 1

    try:
        ws.freeze_panes = ws.cell(row=data_row, column=1)
    except Exception:
        pass
    if rows:
        try:
            last_col = get_column_letter(ncol)
            last_row = data_row + len(rows) - 1
            ws.auto_filter.ref = u"A%d:%s%d" % (hdr_row, last_col, last_row)
        except Exception:
            pass

    try:
        _xlsx_autofit_worksheet(ws, hdr_row, data_row, ncol)
    except Exception as exc:
        _log("export autofit: %s" % exc)
    try:
        _xlsx_autofit_row_heights(ws, hdr_row, data_row, ncol)
    except Exception as exc:
        _log("export row autofit: %s" % exc)

    try:
        wb.save(dest_path)
    except Exception as exc:
        return False, u"Не удалось сохранить файл: %s" % exc
    if open_after:
        _open_exported_workbook(dest_path)
    return True, dest_path


def remove_plan_with_archive( doc, guid, task_name=u"", source_sheet_name=u"", skip_confirm=False ):
    """Архивировать и удалить план задачи. Возвращает (ok, message)."""
    if count_plan_items(doc, guid) <= 0:
        return False, unicode(_pcfg.PLAN_DELETE_NO_PLAN)
    if not skip_confirm:
        from todo_task_dialog_lib import show_todo_task_confirm

        msg = unicode(_pcfg.PLAN_DELETE_CONFIRM) % (
            unicode(task_name or guid or u"—").strip(),
            unicode(_pcfg.PLAN_ARCHIVE_SHEET_NAME),
        )
        if not show_todo_task_confirm(doc=doc, text=msg, title=unicode(_pcfg.PLAN_BTN_DELETE)):
            return False, u""
    if not archive_plan_by_guid(
        doc,
        guid,
        source_sheet_name=source_sheet_name,
        task_name=task_name,
    ):
        return False, u"Не удалось архивировать план."
    if not delete_plan_by_guid(doc, guid):
        return False, u"Не удалось удалить план."
    return True, unicode(_pcfg.PLAN_DELETE_DONE)
