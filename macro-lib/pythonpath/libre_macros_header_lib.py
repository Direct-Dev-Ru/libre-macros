# -*- coding: utf-8 -*-
"""
Сборка итоговых имён столбцов из многоэтажного заголовка (в памяти).

Расширенный режим «Строка_Заголовков»: JSON с диапазоном строк и разделителем.
Исходные файлы не меняются.
"""
from __future__ import print_function, unicode_literals

MACRO_VERSION = "3.10.689"
import re
import unicodedata

try:
    unicode
except NameError:
    unicode = str
try:
    unichr
except NameError:
    unichr = chr

# Excel OOXML: управляющие символы в sharedStrings как _x000D_ / _x000A_ / …
_EXCEL_XML_ESC_RE = re.compile(r"_x([0-9A-Fa-f]{4})_")

HEADER_FN_KEY = u"строка_заголовков"
DEFAULT_SEPARATOR = u" "
EMPTY_RUN_STOP = 2
SCAN_MAX_COLS = 500

_ROWS_RE = re.compile(
    r"^\s*(\d+)\s*[-:–—]\s*(\d+)\s*$",
    re.UNICODE,
)


def col_letters(col_index):
    n = int(col_index) + 1
    letters = u""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = unichr(65 + rem) + letters
    return letters


def synthetic_header(col_index):
    """Пустой итоговый заголовок → Колонка_A (как в простом режиме без шапки)."""
    return u"Колонка_%s" % col_letters(col_index)


def _unescape_excel_xml_char(match):
    try:
        return unichr(int(match.group(1), 16))
    except Exception:
        return match.group(0)


def clean_header_name(text):
    """
    Очистка имени столбца: переносы (LF/CR/CRLF/NEL/LS/PS), табы, нечитаемые
    пробелы и control → пробел или удаление.

    Windows: CRLF и Unicode line/paragraph separators; Excel XML _x000D_/_x000A_.
    NBSP/тонкие пробелы → обычный пробел; zero-width / soft hyphen / BOM — убрать;
    подряд идущие пробелы схлопываются. Края — strip.
    """
    s = unicode(text or u"")
    if s == u"":
        return u""
    if u"_x" in s:
        try:
            s = _EXCEL_XML_ESC_RE.sub(_unescape_excel_xml_char, s)
        except Exception:
            pass
    try:
        s = u" ".join(s.splitlines())
    except Exception:
        pass
    out = []
    i = 0
    while i < len(s):
        ch = s[i]
        o = ord(ch)
        try:
            cat = unicodedata.category(ch)
        except Exception:
            cat = u""
        if ch.isspace() or cat in (u"Zs", u"Zl", u"Zp"):
            out.append(u" ")
        elif cat == u"Cf" or o < 32 or o == 0x7F or (0x80 <= o <= 0x9F):
            pass
        else:
            out.append(ch)
        i = i + 1
    s = u"".join(out)
    while u"  " in s:
        s = s.replace(u"  ", u" ")
    return s.strip()


def cell_to_header_text(value):
    if value is None:
        return u""
    if isinstance(value, bool):
        return clean_header_name(unicode(value))
    if isinstance(value, float):
        if value == int(value):
            return clean_header_name(unicode(int(value)))
        return clean_header_name(unicode(value))
    return clean_header_name(unicode(value))


def parse_header_rows_range(text):
    """
    «2-5» / «2:5» / «2» → (from_1based, to_1based) включительно.

    Невалидно → None.
    """
    raw = unicode(text or u"").strip()
    if raw == u"":
        return None
    m = _ROWS_RE.match(raw)
    if m:
        a = int(m.group(1))
        b = int(m.group(2))
        if a <= 0 or b <= 0:
            return None
        if a > b:
            a, b = b, a
        return (a, b)
    try:
        n = int(raw)
    except (TypeError, ValueError):
        return None
    if n <= 0:
        return None
    return (n, n)


def format_header_rows_range(row_from_1, row_to_1):
    a = int(row_from_1)
    b = int(row_to_1)
    if a > b:
        a, b = b, a
    if a == b:
        return unicode(a)
    return u"%d-%d" % (a, b)


def normalize_header_separator(raw):
    """Подпись визарда / код → строка-разделитель."""
    if raw is None:
        return DEFAULT_SEPARATOR
    if not isinstance(raw, (unicode, str)):
        raw = unicode(raw)
    s = raw
    key = unicode(s).strip().casefold()
    aliases = {
        u"пробел": u" ",
        u"space": u" ",
        u" ": u" ",
        u"подчёркивание _": u"_",
        u"подчеркивание _": u"_",
        u"подчёркивание": u"_",
        u"подчеркивание": u"_",
        u"_": u"_",
        u"дефис -": u"-",
        u"дефис": u"-",
        u"-": u"-",
        u"слэш /": u"/",
        u"слеш /": u"/",
        u"слэш": u"/",
        u"слеш": u"/",
        u"/": u"/",
        u"без разделителя": u"",
        u"нет": u"",
        u"empty": u"",
        u"(пусто)": u"",
    }
    if key in aliases:
        return aliases[key]
    # Не strip: пользователь может задать несколько пробелов.
    return s


def spec_from_block(block):
    """dict JSON-блока → spec или None."""
    if not isinstance(block, dict):
        return None
    rows_raw = block.get("rows")
    if rows_raw in (None, u""):
        a = block.get("row_from")
        b = block.get("row_to")
        if a in (None, u"") and b in (None, u""):
            a = block.get("from")
            b = block.get("to")
        if a not in (None, u"") or b not in (None, u""):
            try:
                fa = int(a if a not in (None, u"") else b)
                fb = int(b if b not in (None, u"") else a)
            except (TypeError, ValueError):
                return None
            if fa <= 0 or fb <= 0:
                return None
            if fa > fb:
                fa, fb = fb, fa
            rows_pair = (fa, fb)
        else:
            return None
    else:
        rows_pair = parse_header_rows_range(rows_raw)
        if rows_pair is None:
            return None
    sep = normalize_header_separator(
        block.get("separator") if "separator" in block else block.get("sep")
    )
    return {
        u"row0_from": rows_pair[0] - 1,
        u"row0_to": rows_pair[1] - 1,
        u"row_from_1": rows_pair[0],
        u"row_to_1": rows_pair[1],
        u"separator": sep,
        u"rows_disp": format_header_rows_range(rows_pair[0], rows_pair[1]),
    }


def parse_header_spec_from_text(raw_text):
    """
    JSON колонки C (или целиком ячейка) → spec.

    Не JSON / пусто / без диапазона → None (простой режим).
    """
    text = unicode(raw_text or u"").strip()
    if text == u"":
        return None
    if text[:1] not in (u"[", u"{"):
        return None
    try:
        from libre_macros_param_codec import param_decode

        blocks = param_decode(HEADER_FN_KEY, text)
    except Exception:
        blocks = []
    if not blocks:
        return None
    for block in blocks:
        spec = spec_from_block(block)
        if spec is not None:
            return spec
    return None


def join_vertical_parts(parts, separator=DEFAULT_SEPARATOR):
    """Склеить уровни заголовка; пустые и подряд идущие дубли пропускаются."""
    sep = separator if separator is not None else DEFAULT_SEPARATOR
    out = []
    prev = None
    for raw in parts or ():
        t = cell_to_header_text(raw)
        if t == u"":
            continue
        if prev is not None and t == prev:
            continue
        out.append(t)
        prev = t
    if not out:
        return u""
    return sep.join(out)


def fill_grid_from_merges(grid, merges):
    """
    Заполнить пропуски объединённых ячеек значением верхнего левого угла.

    Только **горизонтально** (в строке r0): C2:D2 с «2024» → и C2, и D2.
    Вертикальное объединение **не** протягивается вниз — значение остаётся
    только в верхней строке диапазона (r0).

    merges: iterable (c0, r0, c1, r1) в координатах grid (0-based, включительно).
    """
    if not grid or not merges:
        return grid
    n_rows = len(grid)
    for m in merges:
        try:
            c0, r0, c1, r1 = int(m[0]), int(m[1]), int(m[2]), int(m[3])
        except (TypeError, ValueError, IndexError):
            continue
        if r0 < 0:
            r0 = 0
        if c0 < 0:
            c0 = 0
        if r0 >= n_rows:
            continue
        row0 = grid[r0]
        if row0 is None:
            continue
        val = row0[c0] if c0 < len(row0) else None
        text = cell_to_header_text(val)
        if text == u"":
            continue
        cc = c0
        while cc <= c1:
            while len(row0) <= cc:
                row0.append(u"")
            if cell_to_header_text(row0[cc]) == u"":
                row0[cc] = val
            cc += 1
    return grid


def _row_width(grid):
    w = 0
    for row in grid or ():
        if row is None:
            continue
        n = len(row)
        if n > w:
            w = n
    return w


def scan_last_header_col(grid, start_col, col_count=None, empty_run_limit=EMPTY_RUN_STOP):
    """Последний значащий столбец шапки после заполнения объединений."""
    start_col = int(start_col or 0)
    if start_col < 0:
        start_col = 0
    width = _row_width(grid)
    if col_count not in (None, u""):
        try:
            n = int(col_count)
        except (TypeError, ValueError):
            n = 0
        if n > 0:
            last = start_col + n - 1
            if last >= width:
                last = width - 1
            return last
    last = start_col - 1
    empty_run = 0
    c = start_col
    limit = min(width, start_col + SCAN_MAX_COLS)
    n_rows = len(grid or ())
    while c < limit:
        nonempty = False
        ri = 0
        while ri < n_rows:
            row = grid[ri]
            if row is not None and c < len(row) and cell_to_header_text(row[c]) != u"":
                nonempty = True
                break
            ri += 1
        if nonempty:
            last = c
            empty_run = 0
        else:
            empty_run += 1
            if empty_run >= int(empty_run_limit) and last >= start_col:
                break
        c += 1
    return last


def disambiguate_headers(names):
    """
    Повторяющиеся итоговые имена → Имя_1, Имя_2, … в порядке появления.
    Уникальные имена не трогаются.
    """
    counts = {}
    for name in names:
        key = unicode(name or u"")
        counts[key] = counts.get(key, 0) + 1
    seen = {}
    out = []
    for name in names:
        key = unicode(name or u"")
        if counts.get(key, 0) <= 1:
            out.append(key)
            continue
        n = seen.get(key, 0) + 1
        seen[key] = n
        out.append(u"%s_%d" % (key, n))
    return out


def flatten_header_names(grid, start_col=0, end_col=None, separator=DEFAULT_SEPARATOR, merges=None):
    """
    grid — строки диапазона заголовка (уже вырезанные), 0-based столбцы листа.

    Возвращает list[(col_index, name)].
    """
    if not grid:
        return []
    fill_grid_from_merges(grid, merges or ())
    start_col = int(start_col or 0)
    if start_col < 0:
        start_col = 0
    if end_col is None:
        end_col = scan_last_header_col(grid, start_col)
    try:
        end_col = int(end_col)
    except (TypeError, ValueError):
        end_col = start_col - 1
    if end_col < start_col:
        return []
    n_rows = len(grid)
    raw_names = []
    cols = []
    c = start_col
    while c <= end_col:
        parts = []
        ri = 0
        while ri < n_rows:
            row = grid[ri]
            if row is not None and c < len(row):
                parts.append(row[c])
            else:
                parts.append(u"")
            ri += 1
        name = join_vertical_parts(parts, separator)
        if name == u"":
            name = synthetic_header(c)
        raw_names.append(name)
        cols.append(c)
        c += 1
    uniq = disambiguate_headers(raw_names)
    out = []
    i = 0
    while i < len(cols):
        out.append((cols[i], uniq[i]))
        i += 1
    return out


def titles_map_from_pairs(pairs):
    out = {}
    for item in pairs or ():
        try:
            out[int(item[0])] = unicode(item[1] or u"")
        except (TypeError, ValueError, IndexError):
            continue
    return out


def clip_grid_rows(full_grid, row0_from, row0_to):
    """Вырезать строки заголовка из полной матрицы листа (дополнить пустыми)."""
    a = int(row0_from)
    b = int(row0_to)
    if a < 0:
        a = 0
    if b < a:
        b = a
    out = []
    r = a
    n = len(full_grid or ())
    while r <= b:
        if r < n and full_grid[r] is not None:
            out.append(list(full_grid[r]))
        else:
            out.append([])
        r += 1
    return out


def shift_merges_to_header_grid(merges, row0_from):
    """Перенести merge-прямоугольники листа в координаты вырезанного grid."""
    base = int(row0_from)
    out = []
    for m in merges or ():
        try:
            c0, r0, c1, r1 = int(m[0]), int(m[1]), int(m[2]), int(m[3])
        except (TypeError, ValueError, IndexError):
            continue
        out.append((c0, r0 - base, c1, r1 - base))
    return out


def flatten_from_sheet_grid(full_grid, spec, start_col=0, col_count=None, merges=None):
    """full_grid — строки листа 0-based; spec из spec_from_block."""
    if not spec:
        return []
    row0_from = int(spec.get(u"row0_from", 0))
    row0_to = int(spec.get(u"row0_to", row0_from))
    sep = spec.get(u"separator")
    if sep is None:
        sep = DEFAULT_SEPARATOR
    fill_grid_from_merges(full_grid, merges or ())
    header_grid = clip_grid_rows(full_grid, row0_from, row0_to)
    end_col = scan_last_header_col(header_grid, start_col, col_count)
    return flatten_header_names(
        header_grid, start_col=start_col, end_col=end_col, separator=sep, merges=None
    )


def uno_used_end_col(sheet, fallback=80):
    try:
        cursor = sheet.createCursor()
        cursor.gotoEndOfUsedArea(False)
        addr = cursor.getRangeAddress()
        return int(addr.EndColumn)
    except Exception:
        return int(fallback)


def uno_header_data_array(sheet, start_col, end_col, row0_from, row0_to):
    rng = sheet.getCellRangeByPosition(
        int(start_col), int(row0_from), int(end_col), int(row0_to)
    )
    data = rng.getDataArray()
    grid = []
    for row in data or ():
        grid.append(list(row))
    return grid


def uno_merges_in_rect(sheet, c0, r0, c1, r1):
    """Объединения, пересекающие прямоугольник (координаты листа)."""
    seen = set()
    merges = []
    r = int(r0)
    r1 = int(r1)
    c0 = int(c0)
    c1 = int(c1)
    while r <= r1:
        c = c0
        while c <= c1:
            if (c, r) in seen:
                c += 1
                continue
            try:
                cell = sheet.getCellByPosition(c, r)
                merged = bool(cell.IsMerged)
            except Exception:
                merged = False
            if merged:
                try:
                    cursor = sheet.createCursorByRange(cell)
                    cursor.collapseToMergedArea()
                    addr = cursor.getRangeAddress()
                    sc = int(addr.StartColumn)
                    sr = int(addr.StartRow)
                    ec = int(addr.EndColumn)
                    er = int(addr.EndRow)
                except Exception:
                    c += 1
                    continue
                rr = sr
                while rr <= er:
                    cc = sc
                    while cc <= ec:
                        seen.add((cc, rr))
                        cc += 1
                    rr += 1
                merges.append((sc, sr, ec, er))
            c += 1
        r += 1
    return merges


def flatten_headers_uno(sheet, spec, start_col=0, col_count=None):
    """UNO-лист → list[(col, name)]. Исходный лист не меняется."""
    if sheet is None or not spec:
        return []
    start_col = int(start_col or 0)
    row0_from = int(spec.get(u"row0_from", 0))
    row0_to = int(spec.get(u"row0_to", row0_from))
    if col_count not in (None, u""):
        try:
            n = int(col_count)
        except (TypeError, ValueError):
            n = 0
        if n > 0:
            end_col = start_col + n - 1
        else:
            end_col = uno_used_end_col(sheet)
    else:
        end_col = uno_used_end_col(sheet)
    if end_col < start_col:
        return []
    rel = None
    try:
        rel = uno_header_data_array(sheet, start_col, end_col, row0_from, row0_to)
    except Exception:
        rel = None
    if rel is None:
        rel = []
        r = row0_from
        while r <= row0_to:
            row = []
            c = start_col
            while c <= end_col:
                try:
                    row.append(sheet.getCellByPosition(c, r).String)
                except Exception:
                    row.append(u"")
                c += 1
            rel.append(row)
            r += 1
    full = []
    i = 0
    while i < row0_from:
        full.append([])
        i += 1
    for row in rel:
        full.append(([u""] * start_col) + list(row))
    merges = uno_merges_in_rect(sheet, start_col, row0_from, end_col, row0_to)
    return flatten_from_sheet_grid(
        full, spec, start_col=start_col, col_count=col_count, merges=merges
    )


def xlsx_merges_and_grid(ws, row0_from, row0_to, start_col, end_col):
    """openpyxl worksheet (не read_only) → (grid_full_rows_slice, merges_sheet)."""
    grid = []
    r = int(row0_from)
    while r <= int(row0_to):
        row = []
        c = 0
        while c <= int(end_col):
            try:
                val = ws.cell(row=r + 1, column=c + 1).value
            except Exception:
                val = None
            row.append(val)
            c += 1
        grid.append(row)
        r += 1
    merges = []
    try:
        ranges = getattr(ws, u"merged_cells", None)
        if ranges is not None:
            for mrange in ranges.ranges:
                try:
                    min_col = int(mrange.min_col) - 1
                    min_row = int(mrange.min_row) - 1
                    max_col = int(mrange.max_col) - 1
                    max_row = int(mrange.max_row) - 1
                except Exception:
                    continue
                if max_row < int(row0_from) or min_row > int(row0_to):
                    continue
                if max_col < int(start_col) or min_col > int(end_col):
                    continue
                merges.append((min_col, min_row, max_col, max_row))
    except Exception:
        pass
    full = []
    i = 0
    while i < int(row0_from):
        full.append([])
        i += 1
    for row in grid:
        full.append(row)
    return full, merges


def flatten_headers_xlsx(path, sheet_index, spec, start_col=0, col_count=None):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    if not spec:
        return []
    start_col = int(start_col or 0)
    row0_from = int(spec.get(u"row0_from", 0))
    row0_to = int(spec.get(u"row0_to", row0_from))
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        used_col = int(getattr(ws, u"max_column", 0) or 0) - 1
        if used_col < start_col:
            used_col = start_col
        if col_count not in (None, u""):
            try:
                n = int(col_count)
            except (TypeError, ValueError):
                n = 0
            if n > 0:
                end_col = start_col + n - 1
            else:
                end_col = used_col
        else:
            end_col = min(used_col, start_col + SCAN_MAX_COLS - 1)
        full, merges = xlsx_merges_and_grid(ws, row0_from, row0_to, start_col, end_col)
        return flatten_from_sheet_grid(
            full, spec, start_col=start_col, col_count=col_count, merges=merges
        )
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _ods_local_name(elem):
    qn = getattr(elem, u"qname", None)
    if isinstance(qn, (tuple, list)) and len(qn) >= 2:
        return unicode(qn[1] or u"")
    tag = getattr(elem, u"tagName", None) or getattr(elem, u"tag", None)
    s = unicode(tag or u"")
    if u":" in s:
        return s.split(u":")[-1]
    return s


def _ods_int_attr(elem, *names):
    for name in names:
        try:
            raw = elem.getAttribute(name)
        except Exception:
            raw = None
        if raw not in (None, u""):
            try:
                n = int(raw)
                if n >= 1:
                    return n
            except (TypeError, ValueError):
                pass
    return 1


def ods_header_grid_and_merges(table_rows, row0_from, row0_to, start_col, end_col_cap):
    """
    Строки ODS → (full_grid, merges) с учётом column/row span и covered-cell.
    """
    try:
        import libre_macros_direct_lib as direct

        cell_value = direct._ods_cell_value
    except Exception:
        def cell_value(cell):
            try:
                return unicode(cell)
            except Exception:
                return u""

    full = []
    merges = []
    n_rows = len(table_rows or ())
    i = 0
    while i < int(row0_from):
        full.append([])
        i += 1
    r = int(row0_from)
    cap = int(end_col_cap)
    while r <= int(row0_to):
        row_vals = []
        c = 0
        if r < n_rows:
            row_elem = table_rows[r]
            children = list(getattr(row_elem, u"childNodes", None) or ())
            if not children:
                try:
                    from odf.table import TableCell

                    children = list(row_elem.getElementsByType(TableCell))
                except Exception:
                    children = []
            for child in children:
                lname = _ods_local_name(child).replace(u"_", u"-")
                if lname not in (u"table-cell", u"covered-table-cell"):
                    continue
                repeat = _ods_int_attr(
                    child, u"numbercolumnsrepeated", u"number-columns-repeated"
                )
                if lname == u"covered-table-cell":
                    c += repeat
                    continue
                cspan = _ods_int_attr(
                    child, u"numbercolumnsspanned", u"number-columns-spanned"
                )
                rspan = _ods_int_attr(
                    child, u"numberrowsspanned", u"number-rows-spanned"
                )
                try:
                    val = cell_value(child)
                except Exception:
                    val = u""
                ri = 0
                while ri < repeat:
                    cidx = c + ri
                    while len(row_vals) <= cidx:
                        row_vals.append(u"")
                    if cell_to_header_text(row_vals[cidx]) == u"":
                        row_vals[cidx] = val
                    if cspan > 1 or rspan > 1:
                        merges.append(
                            (cidx, r, cidx + cspan - 1, r + rspan - 1)
                        )
                    ri += 1
                c += repeat
        full.append(row_vals)
        r += 1
        if cap >= 0 and c > cap + 50:
            pass
    return full, merges


def flatten_headers_ods(path, sheet_index, spec, start_col=0, col_count=None):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    if not spec:
        return []
    start_col = int(start_col or 0)
    row0_from = int(spec.get(u"row0_from", 0))
    row0_to = int(spec.get(u"row0_to", row0_from))
    doc = load(path)
    tables = doc.spreadsheet.getElementsByType(Table)
    if int(sheet_index) < 0 or int(sheet_index) >= len(tables):
        return []
    rows = tables[int(sheet_index)].getElementsByType(TableRow)
    if col_count not in (None, u""):
        try:
            n = int(col_count)
        except (TypeError, ValueError):
            n = 0
        end_cap = start_col + n - 1 if n > 0 else start_col + SCAN_MAX_COLS - 1
    else:
        end_cap = start_col + SCAN_MAX_COLS - 1
    full, merges = ods_header_grid_and_merges(
        rows, row0_from, row0_to, start_col, end_cap
    )
    return flatten_from_sheet_grid(
        full, spec, start_col=start_col, col_count=col_count, merges=merges
    )


def flatten_headers_path(path, src_kind, sheet_index, spec, start_col=0, col_count=None):
    kind = unicode(src_kind or u"").strip().lower()
    if kind in (u"xlsx", u"xlsm", u"xls"):
        return flatten_headers_xlsx(
            path, sheet_index, spec, start_col=start_col, col_count=col_count
        )
    return flatten_headers_ods(
        path, sheet_index, spec, start_col=start_col, col_count=col_count
    )


def apply_flattened_names_uno(sheet, spec, pairs, unmerge=True):
    """
    Записать итоговые имена в нижнюю строку диапазона на листе результата.
    Источник не должен попадать сюда.
    """
    if sheet is None or not spec or not pairs:
        return
    row0_from = int(spec.get(u"row0_from", 0))
    row0_to = int(spec.get(u"row0_to", row0_from))
    if unmerge:
        merges = uno_merges_in_rect(
            sheet,
            0,
            row0_from,
            max(p[0] for p in pairs) if pairs else 0,
            row0_to,
        )
        # снимаем с конца, чтобы индексы не плыли
        i = len(merges) - 1
        while i >= 0:
            c0, r0, c1, r1 = merges[i]
            try:
                rng = sheet.getCellRangeByPosition(c0, r0, c1, r1)
                if rng.IsMerged:
                    rng.merge(False)
            except Exception:
                pass
            i -= 1
    for col, name in pairs:
        try:
            sheet.getCellByPosition(int(col), int(row0_to)).String = unicode(name or u"")
        except Exception:
            pass
