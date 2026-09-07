# -*- coding: utf-8 -*-
"""
Прямой перенос данных через работу с файлами (.ods/.xlsx) без UNO/буфера.

Ограничения первой версии:
  - поддерживается режим «На один лист» и «На разные листы»
  - переносим значения с типами (числа/даты/время) при возможности
  - источники: .xlsx и .ods
  - цель: .ods
"""

from __future__ import print_function, unicode_literals
MACRO_VERSION = "3.10.691"
import calendar
import os
import re
import datetime
import time

try:
    unicode  # type: ignore[name-defined]
except NameError:  # pragma: no cover (py3)
    unicode = str

# Отладка прямого режима. Включается:
#  - автоматически, если включён LIBRE_MACROS_DEBUG из libre_macros_lib (если доступен),
#  - или если DIRECT_DEBUG=True (меняется installer/enable_debug_flags.py и сборкой бандла).
DIRECT_DEBUG = False


def _direct_debug_enabled():
    if DIRECT_DEBUG:
        return True
    try:
        from libre_macros_lib import LIBRE_MACROS_DEBUG

        return bool(LIBRE_MACROS_DEBUG)
    except Exception:
        return False


_direct_yield_counter = 0
_direct_t0 = None
_direct_external_ui = None
_direct_journal_buffer = []
_direct_xml_pp_formula_done = set()
_direct_xml_pp_as_values_pending = []
_direct_xml_pp_freeze_pending = []
# [{name, sheet, extra}] — шаги Постобработка_xml, которые остаются на UNO после reopen
# (autofit / заголовок_плюс_высота / перенос_и_авто_высота — надёжнее в Calc).
_direct_xml_pp_uno_pending = []
# Уникальные имена листов, на которых file-visual уже применён в файле.
_direct_xml_pp_file_visual_sheets = []
# Столбцы заголовка, оформленные format_столбцы (include_header/header_only) в файле.
# sheet_name -> set(0-based col indices)
_direct_xml_pp_format_header_cols = {}
# Листы, созданные в xml (сводная_таблица / копировать_лист) — в UNO-пайплайн после reopen.
_direct_xml_pp_created_sheets = []
_direct_options = {
    "convert_to_numbers": True,
    "datetime_formats": "",
    "default_date_format": "DD.MM.YYYY",
    "number_format": "# ##0,00",
}
_DIRECT_CONSOLE_LOG_HOOK = None


def direct_set_console_log_hook(callback):
    global _DIRECT_CONSOLE_LOG_HOOK
    _DIRECT_CONSOLE_LOG_HOOK = callback


def _direct_console_log(line):
    hook = globals().get("_DIRECT_CONSOLE_LOG_HOOK")
    if hook is not None:
        try:
            hook(line)
        except Exception:
            pass


# Визуальные шаги Постобработка_xml, выполняемые в файле (openpyxl / odfpy).
_DIRECT_XML_FILE_VISUAL_KEYS = frozenset(
    (
        u"сетка",
        u"тонкая_сетка",
        u"толстая_сетка",
        u"перенос",
        u"высота_строки",
        u"левое_выравнивание",
        u"отступ",
        u"шрифт",
        u"зебра_диапазон",
        u"формат_деньги",
        u"формат_даты",
        u"формат_столбцы",
        u"градиент",
        u"ширина_столбцов",
        u"раскрасить_блоки",
    )
)

# Не отключают встроенное UNO-оформление после reopen (шрифт, заголовок, числа).
# Сетка — только границы; формат_* — только целевые столбцы.
_DIRECT_XML_PRESERVE_UNO_DECORATION_KEYS = frozenset(
    (
        u"сетка",
        u"тонкая_сетка",
        u"толстая_сетка",
        u"формат_деньги",
        u"формат_даты",
        u"формат_столбцы",
    )
)
# Обратная совместимость имени (то же множество).
_DIRECT_XML_BORDER_ONLY_VISUAL_KEYS = _DIRECT_XML_PRESERVE_UNO_DECORATION_KEYS

# format_* в ODS/xlsx до reopen; после встроенного UNO-оформления — повтор через UNO.
_DIRECT_XML_FORMAT_UNO_REAPPLY_KEYS = frozenset(
    (
        u"формат_деньги",
        u"формат_даты",
        u"формат_столбцы",
    )
)

# У fn_key поле sheets — список источников, не фильтр контекстного листа.
_DIRECT_PP_SHEETS_IS_SOURCE_LIST_KEYS = frozenset()

# Структурные шаги Постобработка_xml в файле (без UNO).
_DIRECT_XML_FILE_STRUCTURAL_KEYS = frozenset(
    (
        u"удалить_столбцы",
        u"конкатенация_столбцов",
        u"замена_значений",
        u"текстовые_операции",
        u"заполнение_вниз",
        u"заполнить_вверх",
        u"заполнение_вниз_вычислить",
        u"сортировка",
        u"переименовать_лист",
        u"переименовать_столбцы",
        u"переставить_столбцы",
        u"копировать_лист",
        u"сводная_таблица",
    )
)

# Как UNO lm_pp_pivot_post_values_format / collapse headers
_DIRECT_PP_PIVOT_HEADER_DELIM = u"\\"
_DIRECT_PP_PIVOT_HEADER_EMPTY = u"[пусто]"
_DIRECT_PP_PIVOT_TOTAL_LABEL = u"Итого"
_DIRECT_PP_PIVOT_NUM_FMT = u"#,##0.00"
_DIRECT_PP_PIVOT_MAX_COL_WIDTH = 50.0  # мм → openpyxl width ≈ mm/1.8

# Шаги, которые остаются в очереди UNO после reopen.
_DIRECT_XML_UNO_PILOT_KEYS = frozenset(
    (
        u"заголовок_плюс_высота",
        u"перенос_и_авто_высота",
        u"авто_высота",
        u"авто_ширина",
        u"автофильтр",
    )
)

# Фильтр sheets (JSON), как в param_codec._SHEET_LIST_FNS
_DIRECT_XML_SHEET_LIST_KEYS = frozenset(
    (
        u"тонкая_сетка",
        u"толстая_сетка",
        u"авто_высота",
        u"авто_ширина",
        u"левое_выравнивание",
        u"закрепить_заголовок",
        u"автофильтр",
    )
)


def direct_bind_ui(hooks=None):
    """Подключить колбэки UI из collect_workbooks (прогресс, pulse, журнал)."""
    global _direct_external_ui
    _direct_external_ui = hooks


def direct_bind_options(opts=None):
    """
    Опции direct-фазы (только xml_excel_ods).
    convert_to_numbers — преобразовывать текст, похожий на число/дату (по умолчанию True).
    datetime_formats — строка со списком допустимых форматов (через запятую).
    default_date_format / number_format — целевые форматы ячеек после преобразования.
    font_name / font_size — шрифт по умолчанию (Постобработка_xml «шрифт» и т.п.).
    locale_language / locale_country — локаль книги или LibreOffice (CharLocale / Setup/L10N).
    """
    global _direct_options, _DIRECT_PP_DEFAULT_FONT_NAME, _DIRECT_PP_DEFAULT_FONT_SIZE
    if opts is None:
        _direct_options = {
            "convert_to_numbers": True,
            "datetime_formats": "",
            "default_date_format": "DD.MM.YYYY",
            "number_format": "# ##0,00",
            "locale_language": "",
            "locale_country": "",
            "font_name": "PT Sans",
            "font_size": 11.0,
        }
        return
    merged = dict(_direct_options or {})
    merged.update(opts or {})
    _direct_options = merged
    fn = unicode(merged.get("font_name") or u"").strip()
    if fn:
        _DIRECT_PP_DEFAULT_FONT_NAME = fn
    try:
        raw_sz = merged.get("font_size")
        if raw_sz is not None and unicode(raw_sz).strip() != u"":
            sz = float(unicode(raw_sz).replace(u",", u"."))
            if sz >= 1.0:
                _DIRECT_PP_DEFAULT_FONT_SIZE = float(sz)
    except (TypeError, ValueError):
        pass


def _direct_wants_convert_to_numbers():
    return bool((_direct_options or {}).get("convert_to_numbers", True))


def _direct_convert_auto_probe():
    """Автоопределение столбцов (≥3 совпадений); false — только явные convert_to_numbers."""
    opts = _direct_options or {}
    if "convert_auto_probe" in opts:
        return bool(opts.get("convert_auto_probe"))
    return _direct_wants_convert_to_numbers()


def _direct_default_date_format():
    raw = unicode(
        (_direct_options or {}).get("default_date_format", u"DD.MM.YYYY") or u"DD.MM.YYYY"
    ).strip()
    return raw or u"DD.MM.YYYY"


def _direct_translate_cyrillic_date_format(fmt):
    try:
        from libre_macros_lib import _lm_pp_translate_cyrillic_date_format

        return unicode(_lm_pp_translate_cyrillic_date_format(fmt) or u"")
    except Exception:
        return unicode(fmt or u"")


def _direct_normalize_calc_date_format(fmt):
    """Excel-совместимые токены → формат Calc/ODS (DDD → NNN и т.д.)."""
    s = _direct_translate_cyrillic_date_format(fmt).strip()
    if s == u"":
        return u"DD.MM.YYYY"
    s = re.sub(r"DDDD", "NNNN", s, flags=re.IGNORECASE)
    s = re.sub(r"DDD", "NNN", s, flags=re.IGNORECASE)
    return s


def _direct_normalize_excel_date_format(fmt):
    """Формат для xlsx/openpyxl (NNN → DDD; NNN в Excel — это минуты)."""
    s = _direct_translate_cyrillic_date_format(fmt).strip()
    if s == u"":
        return u"DD.MM.YYYY"
    s = re.sub(r"NNNN", "DDDD", s, flags=re.IGNORECASE)
    s = re.sub(r"NNN", "DDD", s, flags=re.IGNORECASE)
    return s


def _direct_default_datetime_format():
    fmt = _direct_default_date_format()
    low = fmt.casefold()
    if u"hh" in low or u"чч" in low:
        return fmt
    return u"%s HH:MM" % fmt


class DirectServiceDatetime(datetime.datetime):
    """#ДатаВремяФайла — отдельный формат отображения (не default_date_format)."""


def _direct_service_datetime_format(dt):
    """Формат ячейки #ДатаВремяФайла после записи datetime."""
    if isinstance(dt, DirectServiceDatetime):
        if (
            dt.hour == 0
            and dt.minute == 0
            and dt.second == 0
            and dt.microsecond == 0
        ):
            return u"DD.MM.YYYY"
        return u"DD.MM.YYYY HH:MM:SS"
    return None


def _direct_date_format_for_parsed(parsed, for_ods=False):
    """Формат отображения после преобразования текста → date/datetime."""
    svc_fmt = _direct_service_datetime_format(parsed)
    if svc_fmt:
        normalize = (
            _direct_normalize_calc_date_format
            if for_ods
            else _direct_normalize_excel_date_format
        )
        return normalize(svc_fmt)
    normalize = (
        _direct_normalize_calc_date_format
        if for_ods
        else _direct_normalize_excel_date_format
    )
    if isinstance(parsed, datetime.datetime):
        if (
            parsed.hour == 0
            and parsed.minute == 0
            and parsed.second == 0
            and parsed.microsecond == 0
        ):
            return normalize(_direct_default_date_format())
        fmt = _direct_default_date_format()
        low = fmt.casefold()
        if u"hh" in low or u"чч" in low:
            return normalize(fmt)
        return normalize(u"%s HH:MM" % fmt)
    return normalize(_direct_default_date_format())


def _direct_default_number_format():
    return unicode((_direct_options or {}).get("number_format", u"") or u"").strip()


def _direct_xlsx_number_format():
    """Формат числа для xlsx (OOXML: десятичная точка, не запятая Calc)."""
    fmt = _direct_default_number_format()
    if fmt == u"":
        return u""
    out = fmt
    m = re.search(r",(\d+)$", out)
    if m:
        out = out[: m.start()] + u"." + m.group(1)
    return _direct_pp_excel_number_format(out)


def _direct_xlsx_apply_converted_cell_style(cell, parsed):
    """number_format openpyxl после coerce (даты/числа)."""
    if cell is None or parsed is None:
        return
    try:
        svc_fmt = _direct_service_datetime_format(parsed)
        if svc_fmt:
            cell.number_format = _direct_normalize_excel_date_format(svc_fmt)
            return
        if not _direct_wants_convert_to_numbers():
            return
        if isinstance(parsed, (datetime.date, datetime.datetime)):
            fmt = _direct_date_format_for_parsed(parsed, for_ods=False)
            if fmt:
                cell.number_format = fmt
            return
        if isinstance(parsed, (int, float)):
            num_fmt = _direct_xlsx_number_format()
            if num_fmt and num_fmt != u"General":
                cell.number_format = num_fmt
    except Exception:
        pass


def _direct_xlsx_write_row_coerced(ws, values):
    """Запись строки в xlsx с форматами из глобальных настроек."""
    row_num = (ws.max_row or 0) + 1
    ci = 1
    for val in values or []:
        cell = ws.cell(row=row_num, column=ci, value=val)
        _direct_xlsx_apply_converted_cell_style(cell, val)
        ci += 1


def _direct_ods_apply_converted_cell_style(cache, cell, parsed):
    if cell is None or parsed is None:
        return
    try:
        svc_fmt = _direct_service_datetime_format(parsed)
        if svc_fmt:
            sn = _direct_ods_cell_style_name(
                cache, {u"data_style": u"custom:" + svc_fmt}
            )
            cell.setAttribute(u"stylename", sn)
            return
        if isinstance(parsed, (datetime.date, datetime.datetime)):
            fmt = _direct_date_format_for_parsed(parsed, for_ods=True)
            sn = _direct_ods_cell_style_name(
                cache, {u"data_style": u"custom:" + fmt}
            )
            cell.setAttribute(u"stylename", sn)
            return
        if isinstance(parsed, (int, float)):
            num_fmt = _direct_default_number_format()
            if num_fmt != u"":
                sn = _direct_ods_cell_style_name(
                    cache, {u"data_style": u"custom:" + num_fmt}
                )
                cell.setAttribute(u"stylename", sn)
    except Exception:
        pass


_DT_FORMAT_ALIASES = {
    # даты (латиница)
    "DD.MM.YYYY": "DD.MM.YYYY",
    "DD.MM.YY": "DD.MM.YY",
    "D.M.YYYY": "DD.MM.YYYY",
    "YYYY-MM-DD": "YYYY-MM-DD",
    "YY-MM-DD": "YY-MM-DD",
    "DD/MM/YYYY": "DD/MM/YYYY",
    "DD/MM/YY": "DD/MM/YY",
    "MM/DD/YYYY": "MM/DD/YYYY",
    "MM/DD/YY": "MM/DD/YY",
    "DD.MMM.YYYY": "DD.MMM.YYYY",
    "DD.MMM.YY": "DD.MMM.YY",
    "D MMMM YYYY": "D MMMM YYYY",
    "DD MMM YYYY": "DD MMM YYYY",
    # даты (кириллица — аналоги токенов Calc)
    "ДД.ММ.ГГГГ": "DD.MM.YYYY",
    "ДД.ММ.ГГ": "DD.MM.YY",
    "Д.М.ГГГГ": "DD.MM.YYYY",
    "ГГГГ-ММ-ДД": "YYYY-MM-DD",
    "ГГ-ММ-ДД": "YY-MM-DD",
    "ДД/ММ/ГГГГ": "DD/MM/YYYY",
    "ДД/ММ/ГГ": "DD/MM/YY",
    "ММ/ДД/ГГГГ": "MM/DD/YYYY",
    "ММ/ДД/ГГ": "MM/DD/YY",
    "ДД.МММ.ГГГГ": "DD.MMM.YYYY",
    "Д ММММ ГГГГ": "D MMMM YYYY",
    "ДД МММ ГГГГ": "DD MMM YYYY",
    # дата+время (только DMY; YMD datetime по умолчанию НЕ включаем)
    "DD.MM.YYYY HH:MM": "DD.MM.YYYY HH:MM",
    "DD.MM.YYYY HH:MM:SS": "DD.MM.YYYY HH:MM:SS",
    "DD.MM.YYYYTHH:MM": "DD.MM.YYYYTHH:MM",
    "DD.MM.YYYYTHH:MM:SS": "DD.MM.YYYYTHH:MM:SS",
    "DD.MM.YY HH:MM": "DD.MM.YY HH:MM",
    "DD.MM.YY HH:MM:SS": "DD.MM.YY HH:MM:SS",
    "DD.MM.YYTHH:MM": "DD.MM.YYTHH:MM",
    "DD.MM.YYTHH:MM:SS": "DD.MM.YYTHH:MM:SS",
    "ДД.ММ.ГГГГ ЧЧ:ММ": "DD.MM.YYYY HH:MM",
    "ДД.ММ.ГГГГ ЧЧ:ММ:СС": "DD.MM.YYYY HH:MM:SS",
    # опционально (через флаг): ISO
    "YYYY-MM-DD HH:MM": "YYYY-MM-DD HH:MM",
    "YYYY-MM-DD HH:MM:SS": "YYYY-MM-DD HH:MM:SS",
    "YYYY-MM-DDTHH:MM": "YYYY-MM-DDTHH:MM",
    "YYYY-MM-DDTHH:MM:SS": "YYYY-MM-DDTHH:MM:SS",
    "YY-MM-DD HH:MM": "YY-MM-DD HH:MM",
    "YY-MM-DD HH:MM:SS": "YY-MM-DD HH:MM:SS",
    "YY-MM-DDTHH:MM": "YY-MM-DDTHH:MM",
    "YY-MM-DDTHH:MM:SS": "YY-MM-DDTHH:MM:SS",
    "ГГГГ-ММ-ДД ЧЧ:ММ": "YYYY-MM-DD HH:MM",
    "ГГГГ-ММ-ДД ЧЧ:ММ:СС": "YYYY-MM-DD HH:MM:SS",
}

_DIRECT_COPY_DT_COL_MATCH_MIN = 3
_DIRECT_COPY_DT_ROW_BATCH = 200
_DIRECT_COPY_COL_PROBE_ROWS = 60
_RE_GROUPED_NUM_BODY = re.compile(r"^\d+\.?\d*$")

_DIRECT_DT_FORMATS_BASE = frozenset(
    {
        "DD.MM.YYYY",
        "DD.MM.YY",
        "YYYY-MM-DD",
        "YY-MM-DD",
        "DD.MM.YYYY HH:MM",
        "DD.MM.YYYY HH:MM:SS",
        "DD.MM.YYYYTHH:MM",
        "DD.MM.YYYYTHH:MM:SS",
        "DD.MM.YY HH:MM",
        "DD.MM.YY HH:MM:SS",
        "DD.MM.YYTHH:MM",
        "DD.MM.YYTHH:MM:SS",
        "YYYY-MM-DD HH:MM",
        "YYYY-MM-DD HH:MM:SS",
        "YYYY-MM-DDTHH:MM",
        "YYYY-MM-DDTHH:MM:SS",
        "YY-MM-DD HH:MM",
        "YY-MM-DD HH:MM:SS",
        "YY-MM-DDTHH:MM",
        "YY-MM-DDTHH:MM:SS",
    }
)


def _direct_options_locale():
    opts = _direct_options or {}
    lang = str(opts.get("locale_language", "") or "").strip().lower()
    country = str(opts.get("locale_country", "") or "").strip().upper()
    return lang, country


def _direct_locale_extra_datetime_formats(lang, country):
    """
    Дополнительные форматы по локали LibreOffice / книги.
    ru — слэш DMY и текстовые месяцы; en-US — MM/DD.
    """
    extras = {
        "DD/MM/YYYY",
        "DD/MM/YY",
        "DD.MMM.YYYY",
        "DD.MMM.YY",
        "D MMMM YYYY",
        "DD MMM YYYY",
    }
    if lang.startswith("en") and country in ("US", "USA"):
        extras.update({"MM/DD/YYYY", "MM/DD/YY"})
    elif lang.startswith("ru") or country in ("RU", "BY", "KZ", "UA"):
        pass
    return extras


_MONTH_NAME_FMT_TOKENS = frozenset(
    {
        "MMMM",
        "MMM",
        u"ММММ",
        u"МММ",
    }
)


def _direct_canonicalize_datetime_format_token(token):
    """
    Один элемент MERGE_XML_DATETIME_FORMATS → канонический ключ или None.

    Классические алиасы (DD.MM.YYYY, …) и шаблоны month-only
    (FIRST_DAY.MMMM.YEAR, 01.ММММ.2027, …).
    """
    t = unicode(token or u"").strip()
    if t == u"":
        return None
    up = t.upper()
    key = _DT_FORMAT_ALIASES.get(up)
    if key is None:
        key = _DT_FORMAT_ALIASES.get(t)
    if key is not None:
        return key
    construct = _direct_parse_month_construct_template(up)
    if construct is not None:
        return construct[0]
    # уже нормализованный классический ключ или неизвестный токен как есть
    return up


def _direct_datetime_formats_from_raw(raw):
    """Строка форматов (`,`/`;`) → set канонических ключей. Пустая → пустой set."""
    out = set()
    for part in re.split(r"[,;]", unicode(raw or u"")):
        key = _direct_canonicalize_datetime_format_token(part)
        if key is not None:
            out.add(key)
    return out


def _direct_allowed_datetime_formats():
    """
    Разрешённые форматы распознавания даты/времени в тексте.

    Если флаг пуст — базовый набор (точка/дефис, ISO) плюс дополнения по локали
    (слэш DMY, DD.MMM.YYYY с рус./англ. месяцами, для en-US — MM/DD).
    """
    raw = str((_direct_options or {}).get("datetime_formats", "") or "").strip()
    if raw == "":
        lang, country = _direct_options_locale()
        return set(_DIRECT_DT_FORMATS_BASE) | _direct_locale_extra_datetime_formats(
            lang, country
        )
    return _direct_datetime_formats_from_raw(raw)


def direct_datetime_formats_from_raw(raw):
    """Публичная обёртка: строка форматов → set (для PP convert / parse_formats)."""
    return _direct_datetime_formats_from_raw(raw)


def direct_take_journal_buffer():
    """Снять буфер строк журнала после direct-фазы (для записи после reopen)."""
    global _direct_journal_buffer
    buf = list(_direct_journal_buffer)
    _direct_journal_buffer = []
    return buf


def direct_take_xml_as_values_pending():
    """Столбцы Постобработка_xml с as_values — материализовать после reopen."""
    global _direct_xml_pp_as_values_pending
    buf = list(_direct_xml_pp_as_values_pending)
    _direct_xml_pp_as_values_pending = []
    return buf


def direct_take_xml_freeze_pending():
    """
    Листы для «закрепить_заголовок» из Постобработка_xml.

    Freeze panes через openpyxl/odfpy Calc часто не подхватывает после reopen,
    поэтому шаг только ставит в очередь — реальное закрепление делает UNO.
    """
    global _direct_xml_pp_freeze_pending
    buf = list(_direct_xml_pp_freeze_pending)
    _direct_xml_pp_freeze_pending = []
    return buf


def direct_take_xml_header_height_pending():
    """Устарело: используйте direct_take_xml_uno_pending. Оставлено для совместимости."""
    items = direct_take_xml_uno_pending()
    out = []
    for it in items:
        if (it or {}).get(u"name") == u"заголовок_плюс_высота":
            out.append(
                {
                    u"sheet": (it or {}).get(u"sheet") or u"",
                    u"extra": (it or {}).get(u"extra") or u"",
                }
            )
    return out


def direct_take_xml_created_sheets():
    """
    Листы, созданные шагами Постобработка_xml (сводная / копия).

    После reopen их нужно включить в UNO-оформление и пайплайн
    Постобработка_Диапазон / Строка / финал (как новые вкладки результата).
    """
    global _direct_xml_pp_created_sheets
    buf = list(_direct_xml_pp_created_sheets)
    _direct_xml_pp_created_sheets = []
    return buf


def _direct_pp_track_created_sheet(sheet_name, header_row=0, source=""):
    """Запомнить новый лист результата для UNO-пайплайна после reopen."""
    global _direct_xml_pp_created_sheets
    key = unicode(sheet_name or u"").strip()
    if key == u"":
        return
    for it in _direct_xml_pp_created_sheets:
        if _merge_identity_key((it or {}).get(u"name")) == _merge_identity_key(key):
            return
    _direct_xml_pp_created_sheets.append(
        {
            u"name": key,
            u"header_row": int(header_row or 0),
            u"source": unicode(source or u"").strip(),
        }
    )


def direct_take_xml_uno_pending():
    """
    Очередь пилотных шагов Постобработка_xml для UNO после reopen.

    Элементы: {"name": fn_key, "sheet": имя, "extra": JSON колонки C}.
    """
    global _direct_xml_pp_uno_pending
    buf = list(_direct_xml_pp_uno_pending)
    _direct_xml_pp_uno_pending = []
    return buf


def direct_take_xml_file_visual_sheets():
    """
    Листы, на которых file-visual Постобработка_xml уже применён в файле.

    Нужны collect_workbooks, чтобы не затирать стили встроенным оформлением.
    """
    global _direct_xml_pp_file_visual_sheets
    buf = list(_direct_xml_pp_file_visual_sheets)
    _direct_xml_pp_file_visual_sheets = []
    return buf


def direct_take_xml_format_header_cols():
    """
    Столбцы заголовка, уже оформленные format_столбцы в файле.

    Возвращает dict: имя листа -> set(0-based индексы столбцов).
    """
    global _direct_xml_pp_format_header_cols
    buf = {
        k: set(v) for k, v in (_direct_xml_pp_format_header_cols or {}).items()
    }
    _direct_xml_pp_format_header_cols = {}
    return buf


def direct_xml_pp_applied_formulas():
    """Пары (лист, столбец), для которых применить_формулу уже выполнена в файле."""
    return frozenset(_direct_xml_pp_formula_done)


def direct_xml_pp_applied_steps():
    """Устаревшее: имена функций. Оставлено для совместимости."""
    if not _direct_xml_pp_formula_done:
        return set()
    return {u"применить_формулу"}


def _direct_journal(source, sheet, column, rows_n, status, note=""):
    """Строка журнала «Сбор_книг_лог» (буфер + live-колбэк из collect_workbooks)."""
    global _direct_journal_buffer
    try:
        entry = (
            unicode(source or u""),
            unicode(sheet or u""),
            unicode(column or u""),
            rows_n,
            unicode(status or u""),
            unicode(note or u"")[:500],
        )
    except Exception:
        entry = (u"", u"", u"", u"", u"ошибка", u"журнал")
    _direct_journal_buffer.append(entry)
    ui = _direct_external_ui
    if ui:
        fn = ui.get("journal")
        if fn is not None:
            try:
                fn(*entry)
            except Exception:
                pass
    if _direct_debug_enabled():
        try:
            print(
                "[direct][log] %s | %s | %s | %s | %s | %s"
                % entry
            )
        except Exception:
            pass
    hook = globals().get("_DIRECT_CONSOLE_LOG_HOOK")
    if hook is not None:
        try:
            line = u"[direct][log] %s | %s | %s | %s | %s | %s" % entry
            _direct_console_log(line)
        except Exception:
            pass


def _direct_ui_notify_source(src_path, sheet_name=u"", row_hint=None):
    ui = _direct_external_ui
    if not ui:
        return
    fn = ui.get("source")
    if fn is None:
        return
    try:
        fn(src_path, sheet_name=sheet_name, row_hint=row_hint)
    except Exception:
        pass


def _direct_now():
    try:
        return time.time()
    except Exception:
        return 0.0


def _direct_since_start():
    global _direct_t0
    if _direct_t0 is None:
        _direct_t0 = _direct_now()
        return 0.0
    return max(0.0, _direct_now() - float(_direct_t0 or 0.0))


def _direct_ui_yield(force=False):
    """
    Отдать события UI, чтобы Calc не выглядел «зависшим».
    Безопасно: если UNO недоступен (запуск вне LO) — просто ничего не делает.
    """
    global _direct_yield_counter
    _direct_yield_counter += 1
    if not (force or (_direct_yield_counter % 20 == 0)):
        return
    ui = _direct_external_ui
    if ui:
        pulse = ui.get("pulse")
        if pulse is not None:
            try:
                pulse(force=force)
            except Exception:
                pass
    try:
        # UNO доступен только внутри LibreOffice; вне LO просто пропускаем.
        import uno  # type: ignore

        ctx = uno.getComponentContext()
        sm = ctx.ServiceManager
        toolkit = sm.createInstanceWithContext("com.sun.star.awt.Toolkit", ctx)
        if toolkit is None:
            return
        if hasattr(toolkit, "processEventsToIdle"):
            toolkit.processEventsToIdle(200)
        elif hasattr(toolkit, "processEventsToIdleWithMode"):
            toolkit.processEventsToIdleWithMode(200, 0)
    except Exception:
        return


def _log(msg):
    try:
        line = "[direct] %s" % msg
        hook = globals().get("_DIRECT_CONSOLE_LOG_HOOK")
        if _direct_debug_enabled() or hook is not None:
            print(line)
            _direct_console_log(line)
    except Exception:
        pass


def _direct_trace(where, msg=u""):
    """Подробный тайминг в консоль при MERGE_DEBUG / LIBRE_MACROS_DEBUG."""
    try:
        hook = globals().get("_DIRECT_CONSOLE_LOG_HOOK")
        if not _direct_debug_enabled() and hook is None:
            return
        line = u"[direct +%.1fs] %s" % (_direct_since_start(), unicode(where or u""))
        if msg not in (None, u""):
            line += u": " + unicode(msg)
        print(line)
        _direct_console_log(line)
    except Exception:
        pass


_DIRECT_PARAM_SHEET_PREFIX = u"collect_params"
_DIRECT_PARAM_SHEET_PREFIX_LEGACY = u"Параметры_Объединения"
_DIRECT_SERVICE_SHEET_NAMES = frozenset(
    (
        u"Справка_макроса",
        u"Сбор_книг_лог",
        u"Индексы_Источников",
    )
)


def _direct_is_service_sheet_name(name):
    sn = unicode(name or u"").strip()
    if sn == u"":
        return True
    try:
        from libre_macros_global_settings_lib import is_param_sheet_name as _gs_is_param

        if _gs_is_param(sn):
            return True
    except Exception:
        for pref in (_DIRECT_PARAM_SHEET_PREFIX, _DIRECT_PARAM_SHEET_PREFIX_LEGACY):
            if sn.startswith(pref):
                return True
    key = _merge_identity_key(sn)
    for svc in _DIRECT_SERVICE_SHEET_NAMES:
        if key == _merge_identity_key(svc):
            return True
    return False


def _direct_collect_file_note(blk, sheet_specs, col_specs):
    bits = []
    if blk:
        try:
            bits.append(u"столбец %d" % (int(blk.get("start_col", 0)) + 1))
        except Exception:
            pass
        rc = blk.get("row_count")
        if rc not in (None, u""):
            bits.append(u"строк=%s" % rc)
        cc = blk.get("col_count")
        if cc not in (None, u""):
            bits.append(u"кол=%s" % cc)
    if sheet_specs:
        bits.append(
            u"листы: %s"
            % u",".join(unicode(s) for s in sheet_specs[:6])
        )
    if col_specs:
        bits.append(
            u"столбцы: %s" % u",".join(unicode(s) for s in col_specs[:6])
        )
    return u" | ".join(bits)


def _direct_journal_source_label(item, path_display):
    if item.get("is_current_book"):
        return u"(текущая книга)"
    src = item.get("src_path") or u""
    return _format_source_path(src, path_display if path_display != "none" else "short")


def _direct_service_path_value(item, path_mode):
    if path_mode == "none":
        return u""
    sn = unicode(item.get("src_sheet_name") or u"").strip()
    if item.get("is_current_book"):
        if path_mode == "long":
            p = unicode(item.get("src_path") or u"").strip()
            if sn:
                return u"%s#%s" % (p, sn) if p else sn
            return p
        if sn:
            return u"#%s" % sn
        return u"(текущая книга)"
    return _service_path_value(item.get("src_path"), path_mode, item.get("src_sheet_name"))


def _direct_service_datetime_value(item, datetime_mode):
    if datetime_mode == "none" or item.get("is_current_book"):
        return u""
    return _service_datetime_value(item.get("src_path"), datetime_mode)


def _dbg(msg):
    """Сокращённый debug-лог (синоним _log)."""
    _log(msg)


def _is_xlsx(path):
    return str(path or "").strip().lower().endswith(".xlsx")

def _is_xlsm(path):
    return str(path or "").strip().lower().endswith(".xlsm")

def _is_xls(path):
    return str(path or "").strip().lower().endswith(".xls")


def _is_ods(path):
    return str(path or "").strip().lower().endswith(".ods")


def _is_xml(path):
    try:
        from libre_macros_xml_source_lib import is_xml_source_path

        return is_xml_source_path(path)
    except Exception:
        return str(path or "").strip().lower().endswith(".xml")


def _is_csv(path):
    try:
        from libre_macros_csv_source_lib import is_csv_source_path

        return is_csv_source_path(path)
    except Exception:
        return str(path or "").strip().lower().endswith(".csv")


def _direct_resolve_read_path(path, xml_temp_paths=None, sheet_specs=None, source_extra=None):
    """Локальный путь; .xml/.csv → временный .xlsx для openpyxl."""
    p = unicode(path or u"").strip()
    if p == u"":
        return p
    try:
        from libre_macros_bundle_paths import is_remote_url, normalize_source_path

        if is_remote_url(p):
            raise ValueError(u"URL-источник (%s): режим xml_excel_ods работает только с локальными путями; для smb:// используйте UNO" % p)
        p = normalize_source_path(p)
    except ValueError:
        raise
    except Exception:
        if u"://" in p and not p.lower().startswith(u"file://"):
            raise ValueError(u"URL-источник (%s): режим xml_excel_ods работает только с локальными путями; для smb:// используйте UNO" % p)
        if p.lower().startswith(u"file://"):
            try:
                from libre_macros_bundle_paths import normalize_source_path

                p = normalize_source_path(p)
            except Exception:
                pass
    ap = os.path.abspath(p)
    if _is_xml(ap):
        from libre_macros_xml_source_lib import xml_source_to_temp_xlsx

        tmp = xml_source_to_temp_xlsx(ap, sheet_specs=sheet_specs)
        if not tmp:
            raise ValueError(u"не удалось разобрать XML-источник: %s" % ap)
        if xml_temp_paths is not None:
            xml_temp_paths.append(tmp)
        _log("xml→xlsx: %s → %s" % (ap, tmp))
        return tmp
    if _is_csv(ap):
        from libre_macros_csv_source_lib import (
            csv_source_to_temp_xlsx,
            source_extra_delimiter,
            source_extra_encoding,
        )

        delim = source_extra_delimiter(source_extra)
        encoding = source_extra_encoding(source_extra)
        tmp = csv_source_to_temp_xlsx(ap, delimiter=delim, encoding=encoding)
        if not tmp:
            raise ValueError(u"не удалось разобрать CSV-источник: %s" % ap)
        if xml_temp_paths is not None:
            xml_temp_paths.append(tmp)
        _log("csv→xlsx: %s → %s (delimiter=%r encoding=%r)" % (ap, tmp, delim, encoding))
        _direct_journal(
            os.path.basename(ap),
            u"",
            u"",
            0,
            u"csv→xlsx",
            u"временный файл: %s" % tmp,
        )
        return tmp
    return ap


def _direct_do_empty_sheets(settings):
    """MERGE_DO_EMPTY_SHEETS из settings.merge_flags (умолч. True)."""
    try:
        flags = (settings or {}).get("merge_flags") or {}
        if "MERGE_DO_EMPTY_SHEETS" in flags:
            return bool(flags.get("MERGE_DO_EMPTY_SHEETS"))
    except Exception:
        pass
    return True


def _direct_empty_sheet_message():
    try:
        from libre_macros_collect_cfg import MERGE_EMPTY_SHEET_MESSAGE

        return unicode(MERGE_EMPTY_SHEET_MESSAGE)
    except Exception:
        return u"данных для переноса не обнаружено"


def _direct_cleanup_xml_temp_paths(settings):
    keep = bool(DIRECT_DEBUG)
    try:
        flags = (settings or {}).get("merge_flags") or {}
        if flags.get("MERGE_DEBUG"):
            keep = True
    except Exception:
        pass
    paths = (settings or {}).get("_direct_xml_temp_paths") or []
    while paths:
        tp = paths.pop()
        if keep:
            _log("temp keep (debug): %s" % tp)
            _direct_journal(u"", u"", u"", 0, u"temp_keep", u"временный файл не удалён: %s" % tp)
            continue
        try:
            if tp and os.path.isfile(tp):
                os.remove(tp)
        except Exception:
            pass

def _direct_patch_odf_style_families():
    """
    LO/AO пишут style:family=smart-table|pivot-table — stock odfpy падает на load().
    Патч словаря attrconverters (ссылка на cnv_family закэширована при import).
    """
    try:
        import odf_bundled  # noqa: F401
        from odf.namespaces import STYLENS
        import odf.attrconverters as ac
    except Exception:
        return
    allowed = (
        "text",
        "paragraph",
        "section",
        "ruby",
        "table",
        "table-column",
        "table-row",
        "table-cell",
        "graphic",
        "presentation",
        "drawing-page",
        "chart",
        "smart-table",
        "pivot-table",
    )

    def cnv_family(attribute, arg, element):
        if str(arg) not in allowed:
            raise ValueError("'%s' not allowed" % str(arg))
        return str(arg)

    try:
        ac.cnv_family = cnv_family
        key = ((STYLENS, u"family"), None)
        if hasattr(ac, "attrconverters") and key in ac.attrconverters:
            ac.attrconverters[key] = cnv_family
    except Exception:
        pass


def _direct_ods_load(path):
    """load ODS с патчем style:family для LO/AO smart-table/pivot-table."""
    import odf_bundled  # noqa: F401
    from odf.opendocument import load

    _direct_patch_odf_style_families()
    return load(path)


def _is_xlsx_target(path):
    # Приёмник прямого режима: только .ods или .xlsx
    return _is_xlsx(path)


def _merge_identity_key(text):
    if text is None:
        return ""
    return unicode(text).strip().casefold()


def _text_match(text, pattern):
    import fnmatch

    pattern = unicode(pattern or u"").strip()
    text = unicode(text or u"").strip()
    if _merge_identity_key(pattern) == u"все" or pattern == u"*":
        return True
    if u"*" in pattern or u"?" in pattern:
        return fnmatch.fnmatchcase(text.casefold(), pattern.casefold())
    return _merge_identity_key(text) == _merge_identity_key(pattern)


def _sheet_spec_select_token(spec):
    """Левая часть «a->b» или весь токен; !hidden → ''."""
    try:
        st = unicode(spec or u"").strip()
    except Exception:
        st = str(spec or "").strip()
    if st == u"":
        return u""
    if unicode(st).casefold() in (u"!hidden", u"ignore_hidden", u"игнорировать_скрытые"):
        return u""
    if u"->" in st:
        st = unicode(st.split(u"->", 1)[0] or u"").strip()
    return st


def _classify_sheet_specs(specs):
    if not specs:
        return "ALL"
    cleaned = []
    for s in specs:
        try:
            st = unicode(s or u"").strip()
        except Exception:
            st = str(s or "").strip()
        if st == u"":
            continue
        if unicode(st).casefold() in (u"!hidden", u"ignore_hidden", u"игнорировать_скрытые"):
            continue
        cleaned.append(s)
    specs = cleaned
    if not specs:
        return "ALL"
    for s in specs:
        sel = _sheet_spec_select_token(s)
        if sel.strip().lower() in (u"все", u"*"):
            return "ALL"
    has_num = False
    has_name = False
    for s in specs:
        sel = _sheet_spec_select_token(s)
        if sel == u"":
            continue
        if sel.strip().isdigit():
            has_num = True
        else:
            has_name = True
    if has_num and has_name:
        return "MIXED"
    if has_num:
        return "INDEX"
    return "NAME"


def _xlsx_sheet_names(path):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [ws.title for ws in wb.worksheets]
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _ods_sheet_names(path):
    import odf_bundled  # noqa: F401
    from odf.table import Table

    doc = _direct_ods_load(path)
    names = []
    for t in doc.spreadsheet.getElementsByType(Table):
        names.append(unicode(t.getAttribute("name") or u""))
    return names


def _direct_unique_slot_target_name(base_name, used_keys):
    """Спр + конфликт → Спр_2 (как _merge_unique_slot_target_name)."""
    try:
        base = unicode(base_name or u"").strip()
    except Exception:
        base = str(base_name or "").strip()
    if base == u"":
        base = u"Лист"
    if len(base) > 31:
        base = base[:31]
    candidate = base
    suffix = 2
    while _merge_identity_key(candidate) in used_keys:
        tail = u"_%d" % suffix
        head_len = 31 - len(tail)
        if head_len < 1:
            head_len = 1
        candidate = base[:head_len] + tail
        suffix = suffix + 1
        if suffix > 999:
            break
    used_keys.add(_merge_identity_key(candidate))
    return candidate


def _iter_sheet_slots_file(source_sheet_names, sheet_specs, sheet_kind):
    IGNORE_HIDDEN_TOKENS = set([u"!hidden", u"ignore_hidden", u"игнорировать_скрытые"])
    ARROW_SEP = u"->"
    rename_entries = []
    for spec in sheet_specs or []:
        try:
            s = unicode(spec or u"").strip()
        except Exception:
            s = str(spec or "").strip()
        if s == u"":
            continue
        if unicode(s).casefold() in IGNORE_HIDDEN_TOKENS:
            continue
        if ARROW_SEP in s:
            left, right = s.split(ARROW_SEP, 1)
            left = unicode(left or u"").strip()
            right = unicode(right or u"").strip()
            if left != u"":
                rename_entries.append((left, right if right != u"" else None))
            continue
        rename_entries.append((s, None))

    if sheet_kind == "ALL":
        for i, n in enumerate(source_sheet_names):
            yield {"target_key": n, "source_name": n, "source_index": i}
        return
    if sheet_kind == "INDEX":
        used_tgt = set()
        for spec, tgt in rename_entries:
            try:
                idx = int(unicode(spec).strip()) - 1
            except Exception:
                continue
            if idx >= 0 and idx < len(source_sheet_names):
                n = source_sheet_names[idx]
                if tgt is not None and unicode(tgt).strip() != u"":
                    desired = unicode(tgt).strip()
                else:
                    desired = n
                tkey = _direct_unique_slot_target_name(desired, used_tgt)
                yield {
                    "target_key": tkey,
                    "source_name": n,
                    "source_index": idx,
                }
        return
    picked = []
    used_tgt = set()
    for n in source_sheet_names:
        for match_spec, tgt in rename_entries:
            if _text_match(n, match_spec):
                desired = tgt if tgt is not None and unicode(tgt).strip() != u"" else n
                target_name = _direct_unique_slot_target_name(desired, used_tgt)
                picked.append((n, target_name))
                break
    for n, target_name in picked:
        yield {"target_key": target_name, "source_name": n, "source_index": -1}


_DIRECT_MERGE_ONE_SHEET = u"Merge"
_DIRECT_HEADER_SCAN_LIMIT = 500
_DIRECT_SKIP_HEADER_SCAN_LIMIT = 300
# Fallback, если библиотека не даёт значащий диапазон (как Ctrl+End в LO).
_DIRECT_EMPTY_ROWS_END = 20
_DIRECT_EMPTY_COLS_END = 4
_DIRECT_SCAN_MAX_COL = 2000
_DIRECT_SCAN_MAX_ROW = 500000


def _col_label(col_index):
    n = int(col_index) + 1
    letters = u""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _col_letters_to_index(letters):
    letters = unicode(letters or u"").strip().upper()
    n = 0
    i = 0
    while i < len(letters):
        n = n * 26 + (ord(letters[i]) - ord("A") + 1)
        i = i + 1
    return n - 1


def _col_index_to_letters(col_index):
    """0-based индекс столбца → буквы Excel/Calc (A, B, ..., Z, AA...)."""
    try:
        n = int(col_index) + 1
    except Exception:
        n = 1
    if n <= 0:
        n = 1
    letters = u""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters


def _is_col_letters(spec):
    s = unicode(spec or u"").strip().upper()
    # Calc/Excel: максимум XFD (3 буквы). Длиннее — заголовок (GUID, NAME…).
    if s == u"" or len(s) > 3:
        return False
    i = 0
    while i < len(s):
        ch = s[i]
        if ch < u"A" or ch > u"Z":
            return False
        i = i + 1
    return True


def _parse_col_spec(spec):
    spec = unicode(spec or u"").strip()
    if spec == u"":
        return -1
    if _is_col_letters(spec):
        return _col_letters_to_index(spec)
    return int(spec) - 1


def _split_filter_specs(text):
    parts = []
    normalized = unicode(text or u"").replace(u";", u",")
    for chunk in normalized.split(u","):
        st = chunk.strip()
        if st != u"":
            parts.append(st)
    return parts


def _classify_column_specs(specs):
    if not specs:
        return "ALL"
    for s in specs:
        if unicode(s).strip().lower() in (u"все", u"*"):
            return "ALL"
    has_num = False
    has_letter = False
    has_name = False
    for s in specs:
        st = unicode(s).strip()
        if st.isdigit():
            has_num = True
        elif _is_col_letters(st):
            has_letter = True
        else:
            has_name = True
    kinds = int(has_num) + int(has_letter) + int(has_name)
    if kinds > 1:
        return "MIXED"
    if has_num:
        return "INDEX"
    if has_letter:
        return "LETTER"
    if has_name:
        return "NAME"
    return "ALL"


def _direct_clean_header_text(raw):
    t = unicode(raw or u"").strip()
    try:
        from libre_macros_header_lib import clean_header_name

        t = unicode(clean_header_name(t) or u"").strip()
    except Exception:
        pass
    return t


def _header_name_at(raw_text, col_index):
    t = _direct_clean_header_text(raw_text)
    if t != u"":
        return t
    return u"Колонка_%s" % _col_label(col_index)


def _is_service_header(hname):
    if not hname:
        return False
    h = _merge_identity_key(hname)
    return h in (_merge_identity_key(u"#Путь"), _merge_identity_key(u"#ДатаВремяФайла"))


def _direct_parse_skip_spec(spec_text, source_sheet_name=None):
    """
    Разбор «Пропуск_строк_источника».

    Legacy: «1;A;ФИО» или формула.
    JSON: [{"v":1,"fn":"пропуск_строк_источника","sheet":"Заказы","columns":["1"]}]
      sheet — имя листа *источника* (не результата); пусто = все листы файла.
    """
    text = unicode(spec_text or u"").strip()
    if text == u"":
        return {"mode": "none"}
    # JSON (визард)
    if text[:1] in (u"[", u"{"):
        try:
            from libre_macros_param_codec import param_decode

            blocks = param_decode(u"пропуск_строк_источника", text)
        except Exception:
            blocks = []
        if blocks:
            src = unicode(source_sheet_name or u"").strip()
            fallback = None
            for block in blocks:
                if not isinstance(block, dict):
                    continue
                sheet_ref = unicode(block.get("sheet") or u"").strip()
                if sheet_ref == u"":
                    fallback = block
                    continue
                if src == u"":
                    continue
                try:
                    from libre_macros_lib import _lm_pp_parse_sheet_names_from_spec

                    names = _lm_pp_parse_sheet_names_from_spec(sheet_ref)
                except Exception:
                    names = [sheet_ref]
                if _direct_pp_sheet_name_matches(src, names):
                    return _direct_skip_block_to_parsed(block)
            if fallback is not None:
                return _direct_skip_block_to_parsed(fallback)
            # JSON есть, но ни один блок не для этого листа источника
            return {"mode": "none"}
    if re.search(r"R(?:\[[+-]?\d+\]|\d+)?[CcСс]", text):
        return {"mode": "formula", "formula": text}
    if re.search(r"\[[^\]]+\]", text):
        return {"mode": "formula", "formula": text}
    if re.search(r"\{[^}]*\brow_id\b[^}]*\}", text) or re.search(
        r"\$\{[^}]*\brow_id\b[^}]*\}", text
    ):
        return {"mode": "formula", "formula": text}
    if re.search(r"[A-Za-zА-Яа-яЁё_][A-Za-zА-Яа-яЁё0-9_]*\{[^}]*\brow_id\b[^}]*\}", text):
        return {"mode": "formula", "formula": text}
    if re.search(r"[A-Za-z]+\{", text):
        return {"mode": "formula", "formula": text}
    if re.search(r"[A-Za-z]+\d+", text) and re.search(r"[><=!]", text):
        return {"mode": "formula", "formula": text}
    if re.search(r"[A-Za-z]+\d+\s*[\+\-\*/]", text):
        return {"mode": "formula", "formula": text}
    parts = _split_filter_specs(text)
    if len(parts) == 0:
        return {"mode": "none"}
    return {"mode": "columns", "parts": parts}


def _direct_skip_block_to_parsed(block):
    """JSON-блок пропуск_строк_источника → parsed dict."""
    mode = unicode((block or {}).get("mode") or u"columns").strip().casefold()
    if mode == u"formula":
        fo = unicode((block or {}).get("formula") or u"").strip()
        if fo == u"":
            return {"mode": "none"}
        return {"mode": "formula", "formula": fo}
    if mode in (u"none", u""):
        return {"mode": "none"}
    cols = (block or {}).get("columns") or []
    parts = [unicode(c).strip() for c in cols if unicode(c).strip()]
    if not parts:
        return {"mode": "none"}
    return {"mode": "columns", "parts": parts}


def _direct_resolve_skip_indices(header_pairs, parts):
    indices = []
    for part in parts or []:
        st = unicode(part).strip()
        if st == u"":
            continue
        if st.isdigit():
            idx = int(st) - 1
        elif _is_col_letters(st):
            idx = _col_letters_to_index(st)
        else:
            idx = -1
            for c, h in header_pairs:
                if _text_match(h, st):
                    idx = c
                    break
        if idx >= 0 and idx not in indices:
            indices.append(idx)
    return indices


def _direct_row_skip_by_empty(row_map, col_indices):
    if not col_indices:
        return False
    for ci in col_indices:
        v = row_map.get(ci)
        if v is not None and unicode(v).strip() != u"":
            return False
    return True


def _direct_value_empty(v):
    if v is None:
        return True
    if isinstance(v, (str, unicode)):
        return unicode(v).strip() == u""
    return False


def _direct_effective_limits(used_end_col, used_end_row, hdr_row0, start_col0, row_count, col_count):
    """
    Границы блока с учётом значащей области и параметров Число_Строк / Число_Столбцов.

    Возвращает (last_col, last_data_row, first_data_row) — 0-based, включительно.
    Если данных нет: last_data_row < first_data_row или last_col < start_col0.
    """
    hdr_row0 = int(hdr_row0)
    start_col0 = int(start_col0)
    first_data_row = hdr_row0 + 1
    used_end_col = int(used_end_col)
    used_end_row = int(used_end_row)

    if used_end_row < first_data_row:
        last_data_row = first_data_row - 1
    else:
        available_rows = used_end_row - first_data_row + 1
        if row_count is not None:
            n_rows = min(int(row_count), available_rows)
        else:
            n_rows = available_rows
        last_data_row = (
            first_data_row + n_rows - 1 if n_rows > 0 else first_data_row - 1
        )

    if used_end_col < start_col0:
        last_col = start_col0 - 1
    else:
        available_cols = used_end_col - start_col0 + 1
        if col_count is not None:
            n_cols = min(int(col_count), available_cols)
        else:
            n_cols = available_cols
        last_col = start_col0 + n_cols - 1 if n_cols > 0 else start_col0 - 1

    return last_col, last_data_row, first_data_row


def _direct_xlsx_bounds_from_dimension(ws):
    """openpyxl: calculate_dimension / max_row — (end_col, end_row) 0-based или None."""
    try:
        from openpyxl.utils.cell import range_boundaries

        dim = ws.calculate_dimension()
        if dim:
            _c0, _r0, max_c, max_r = range_boundaries(dim)
            if max_c is not None and max_r is not None:
                return int(max_c) - 1, int(max_r) - 1
    except Exception:
        pass
    try:
        mr = getattr(ws, "max_row", None)
        mc = getattr(ws, "max_column", None)
        if mr and mc:
            return int(mc) - 1, int(mr) - 1
    except Exception:
        pass
    return None


def _direct_scan_bounds_xlsx(ws, hdr_row0, start_col0, max_row_cap=None, max_col_cap=None):
    """
    Значащая область по сканированию: 4 пустых столбца в строке заголовков,
    20 пустых строк данных в блоке столбцов.
    max_*_cap — верхняя граница скана из dimension (не доверяем dimension как итог).
    """
    hdr_row0 = int(hdr_row0)
    start_col0 = int(start_col0)
    header_row_1 = hdr_row0 + 1
    data_row_1 = hdr_row0 + 2
    max_col_scan = min(
        start_col0 + _DIRECT_SCAN_MAX_COL,
        start_col0 + _DIRECT_HEADER_SCAN_LIMIT - 1,
    )
    if max_col_cap is not None:
        max_col_scan = min(max_col_scan, int(max_col_cap))

    end_col = start_col0 - 1
    empty_cols = 0
    try:
        header_vals = next(
            ws.iter_rows(
                min_row=header_row_1,
                max_row=header_row_1,
                min_col=start_col0 + 1,
                max_col=max_col_scan + 1,
                values_only=True,
            ),
            (),
        )
    except Exception:
        header_vals = ()
    ci = 0
    while ci < len(header_vals):
        col_idx = start_col0 + ci
        if _direct_value_empty(header_vals[ci]):
            empty_cols += 1
            if empty_cols >= _DIRECT_EMPTY_COLS_END:
                break
        else:
            end_col = col_idx
            empty_cols = 0
        ci += 1
    if end_col < start_col0:
        return start_col0 - 1, hdr_row0

    end_row = hdr_row0
    empty_rows = 0
    row_scan_limit = min(hdr_row0 + _DIRECT_SCAN_MAX_ROW, hdr_row0 + 200000)
    if max_row_cap is not None:
        row_scan_limit = min(row_scan_limit, int(max_row_cap))
    rr = hdr_row0 + 1
    try:
        for row_vals in ws.iter_rows(
            min_row=data_row_1,
            max_row=row_scan_limit + 1,
            min_col=start_col0 + 1,
            max_col=end_col + 1,
            values_only=True,
        ):
            if rr % 200 == 0:
                _direct_ui_yield()
            non_empty = False
            for v in row_vals:
                if not _direct_value_empty(v):
                    non_empty = True
                    break
            if non_empty:
                end_row = rr
                empty_rows = 0
            else:
                empty_rows += 1
                if empty_rows >= _DIRECT_EMPTY_ROWS_END:
                    break
            rr += 1
    except Exception:
        pass
    if end_row < hdr_row0 + 1:
        end_row = hdr_row0
    return end_col, end_row


def _direct_resolve_last_row_scan_col_xlsx(ws, hdr_row0, start_col0, end_col, spec):
    spec = unicode(spec or u"").strip()
    if spec == u"" or end_col < start_col0:
        return None
    if spec.isdigit():
        try:
            idx = int(spec) - 1
        except (TypeError, ValueError):
            idx = None
        if idx is not None and start_col0 <= idx <= end_col:
            return idx
    if _is_col_letters(spec):
        try:
            idx = _col_letters_to_index(spec)
        except Exception:
            idx = None
        if idx is not None and start_col0 <= idx <= end_col:
            return idx
    try:
        header_vals = next(
            ws.iter_rows(
                min_row=int(hdr_row0) + 1,
                max_row=int(hdr_row0) + 1,
                min_col=int(start_col0) + 1,
                max_col=int(end_col) + 1,
                values_only=True,
            ),
            (),
        )
    except Exception:
        header_vals = ()
    ci = 0
    while ci < len(header_vals):
        col_idx = int(start_col0) + ci
        name = _header_name_at(
            unicode(header_vals[ci]).strip() if header_vals[ci] is not None else u"",
            col_idx,
        )
        if _text_match(name, spec):
            return col_idx
        ci += 1
    return None


def _direct_refine_last_row_xlsx(ws, hdr_row0, first_data_row0, start_col0, end_col, end_row, spec):
    scan_col = _direct_resolve_last_row_scan_col_xlsx(
        ws, hdr_row0, start_col0, end_col, spec
    )
    if scan_col is None:
        return end_row
    first_data_row0 = int(first_data_row0)
    end_row = int(end_row)
    lo = first_data_row0
    hi = end_row
    best = first_data_row0 - 1
    while lo <= hi:
        mid = (lo + hi) // 2
        try:
            mid_val = ws.cell(row=mid + 1, column=scan_col + 1).value
        except Exception:
            mid_val = u""
        filled = not _direct_value_empty(mid_val)
        _log(
            "last_row_column[xlsx]: cell row=%d col=%d -> %s | range=%d..%d"
            % (
                mid + 1,
                scan_col + 1,
                "filled" if filled else "empty",
                lo + 1,
                hi + 1,
            )
        )
        if filled:
            best = mid
            _log(
                "last_row_column[xlsx]: go down -> next range=%d..%d"
                % (mid + 2, hi + 1)
            )
            lo = mid + 1
        else:
            _log(
                "last_row_column[xlsx]: go up -> next range=%d..%d"
                % (lo + 1, mid)
            )
            hi = mid - 1
    _log(
        "last_row_column[xlsx]: result row=%d"
        % (best + 1 if best >= first_data_row0 else 0)
    )
    return best


def _ods_row_col_values(row_elem):
    """(col_index, value) по строке ODS с учётом numbercolumnsrepeated."""
    from odf.table import TableCell

    col = 0
    for cell in row_elem.getElementsByType(TableCell):
        try:
            repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        except Exception:
            repeat = 1
        if repeat < 1:
            repeat = 1
        val = _ods_cell_value(cell)
        ri = 0
        while ri < repeat:
            yield col, val
            col += 1
            ri += 1


def _direct_scan_bounds_ods(rows, hdr_row0, start_col0):
    hdr_row0 = int(hdr_row0)
    start_col0 = int(start_col0)
    if hdr_row0 < 0 or hdr_row0 >= len(rows):
        return start_col0 - 1, hdr_row0

    end_col = start_col0 - 1
    empty_cols = 0
    for col_idx, val in _ods_row_col_values(rows[hdr_row0]):
        if col_idx < start_col0:
            continue
        if col_idx > start_col0 + _DIRECT_SCAN_MAX_COL:
            break
        if _direct_value_empty(val):
            empty_cols += 1
            if empty_cols >= _DIRECT_EMPTY_COLS_END:
                break
        else:
            end_col = col_idx
            empty_cols = 0
    if end_col < start_col0:
        return start_col0 - 1, hdr_row0

    end_row = hdr_row0
    empty_rows = 0
    rr = hdr_row0 + 1
    limit = min(len(rows), hdr_row0 + _DIRECT_SCAN_MAX_ROW)
    while rr < limit:
        if rr % 200 == 0:
            _direct_ui_yield()
        non_empty = False
        if rr < len(rows):
            for col_idx, val in _ods_row_col_values(rows[rr]):
                if col_idx < start_col0 or col_idx > end_col:
                    continue
                if not _direct_value_empty(val):
                    non_empty = True
                    break
        if non_empty:
            end_row = rr
            empty_rows = 0
        else:
            empty_rows += 1
            if empty_rows >= _DIRECT_EMPTY_ROWS_END:
                break
        rr += 1
    if end_row < hdr_row0 + 1:
        end_row = hdr_row0
    return end_col, end_row


def _direct_resolve_last_row_scan_col_ods(rows, hdr_row0, start_col0, end_col, spec):
    spec = unicode(spec or u"").strip()
    if spec == u"" or hdr_row0 < 0 or hdr_row0 >= len(rows):
        return None
    if spec.isdigit():
        try:
            idx = int(spec) - 1
        except (TypeError, ValueError):
            idx = None
        if idx is not None and start_col0 <= idx <= end_col:
            return idx
    if _is_col_letters(spec):
        try:
            idx = _col_letters_to_index(spec)
        except Exception:
            idx = None
        if idx is not None and start_col0 <= idx <= end_col:
            return idx
    for col_idx, val in _ods_row_col_values(rows[int(hdr_row0)]):
        if col_idx < start_col0 or col_idx > end_col:
            continue
        name = _header_name_at(unicode(val).strip() if val is not None else u"", col_idx)
        if _text_match(name, spec):
            return col_idx
    return None


def _direct_refine_last_row_ods(rows, hdr_row0, first_data_row0, start_col0, end_col, end_row, spec):
    scan_col = _direct_resolve_last_row_scan_col_ods(
        rows, hdr_row0, start_col0, end_col, spec
    )
    if scan_col is None:
        return end_row
    first_data_row0 = int(first_data_row0)
    end_row = int(end_row)
    lo = first_data_row0
    hi = end_row
    best = first_data_row0 - 1
    while lo <= hi:
        mid = (lo + hi) // 2
        row_map = {}
        if mid < len(rows):
            row_map = dict(_ods_row_col_values(rows[mid]))
        mid_val = row_map.get(scan_col, u"")
        filled = not _direct_value_empty(mid_val)
        _log(
            "last_row_column[ods]: cell row=%d col=%d -> %s | range=%d..%d"
            % (
                mid + 1,
                scan_col + 1,
                "filled" if filled else "empty",
                lo + 1,
                hi + 1,
            )
        )
        if filled:
            best = mid
            _log(
                "last_row_column[ods]: go down -> next range=%d..%d"
                % (mid + 2, hi + 1)
            )
            lo = mid + 1
        else:
            _log(
                "last_row_column[ods]: go up -> next range=%d..%d"
                % (lo + 1, mid)
            )
            hi = mid - 1
    _log(
        "last_row_column[ods]: result row=%d"
        % (best + 1 if best >= first_data_row0 else 0)
    )
    return best


def _direct_resolve_sheet_used_bounds(read_path, src_kind, sheet_index, hdr_row0, start_col0, last_row_col_spec=u""):
    """
    Значащий диапазон листа (end_col, end_row) 0-based, включительно.
    Сначала метаданные библиотеки; иначе скан 4 пустых столбца / 20 пустых строк.
    """
    hdr_row0 = int(hdr_row0)
    start_col0 = int(start_col0)
    if src_kind == "xlsx":
        import openpyxl_bundled  # noqa: F401
        from openpyxl import load_workbook

        wb = load_workbook(read_path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[int(sheet_index)]
            dim_bounds = _direct_xlsx_bounds_from_dimension(ws)
            max_row_cap = None
            max_col_cap = None
            if dim_bounds is not None:
                ec_d, er_d = dim_bounds
                if ec_d >= start_col0 and er_d >= hdr_row0:
                    max_col_cap = ec_d
                    max_row_cap = er_d
            ec, er = _direct_scan_bounds_xlsx(
                ws, hdr_row0, start_col0, max_row_cap, max_col_cap
            )
            if dim_bounds is not None:
                _log(
                    "bounds(xlsx): sheet=%d scan col=%d row=%d (dim cap col=%d row=%d)"
                    % (
                        int(sheet_index),
                        ec,
                        er,
                        dim_bounds[0],
                        dim_bounds[1],
                    )
                )
            else:
                _log(
                    "bounds(xlsx): sheet=%d scan col=%d row=%d"
                    % (int(sheet_index), ec, er)
                )
            spec = unicode(last_row_col_spec or u"").strip()
            if spec != u"" and er >= hdr_row0 + 1:
                er = _direct_refine_last_row_xlsx(
                    ws, hdr_row0, hdr_row0 + 1, start_col0, ec, er, spec
                )
            return ec, er
        finally:
            try:
                wb.close()
            except Exception:
                pass

    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    doc = load(read_path)
    sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
    rows = sheet.getElementsByType(TableRow)
    ec, er = _direct_scan_bounds_ods(rows, hdr_row0, start_col0)
    spec = unicode(last_row_col_spec or u"").strip()
    if spec != u"" and er >= hdr_row0 + 1:
        er = _direct_refine_last_row_ods(
            rows, hdr_row0, hdr_row0 + 1, start_col0, ec, er, spec
        )
    _log(
        "bounds(ods): sheet=%d scan col=%d row=%d" % (int(sheet_index), ec, er)
    )
    return ec, er


def _direct_read_header_pairs_xlsx(path, sheet_index, hdr_row0, start_col0, last_col0):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    last_col0 = int(last_col0)
    start_col0 = int(start_col0)
    if last_col0 < start_col0:
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        header_row_1 = int(hdr_row0) + 1
        c0 = start_col0 + 1
        c1_cap = last_col0 + 1
        try:
            header_vals = next(
                ws.iter_rows(
                    min_row=header_row_1,
                    max_row=header_row_1,
                    min_col=c0,
                    max_col=c1_cap,
                    values_only=True,
                ),
                (),
            )
        except Exception:
            header_vals = ()
        pairs = []
        ci = 0
        while ci <= last_col0 - start_col0:
            col_idx = start_col0 + ci
            t = u""
            if ci < len(header_vals):
                v = header_vals[ci]
                t = unicode(v).strip() if v is not None else u""
            pairs.append((col_idx, _header_name_at(t, col_idx)))
            ci += 1
        return pairs
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_read_header_pairs_ods(path, sheet_index, hdr_row0, start_col0, last_col0):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    last_col0 = int(last_col0)
    start_col0 = int(start_col0)
    if last_col0 < start_col0:
        return []
    doc = load(path)
    sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
    rows = sheet.getElementsByType(TableRow)
    hdr_r = int(hdr_row0)
    if hdr_r < 0 or hdr_r >= len(rows):
        return []
    hdr_by_col = {}
    for col_idx, val in _ods_row_col_values(rows[hdr_r]):
        if start_col0 <= col_idx <= last_col0:
            hdr_by_col[col_idx] = unicode(val).strip() if not _direct_value_empty(val) else u""
    pairs = []
    col_idx = start_col0
    while col_idx <= last_col0:
        t = hdr_by_col.get(col_idx, u"")
        pairs.append((col_idx, _header_name_at(t, col_idx)))
        col_idx += 1
    return pairs


def _direct_read_header_pairs(read_path, src_kind, sheet_index, hdr_row0, start_col0, last_col0):
    if src_kind == "xlsx":
        return _direct_read_header_pairs_xlsx(
            read_path, sheet_index, hdr_row0, start_col0, last_col0
        )
    return _direct_read_header_pairs_ods(
        read_path, sheet_index, hdr_row0, start_col0, last_col0
    )


def _direct_build_column_list(header_pairs, col_specs, col_kind, start_col0, col_count):
    if not header_pairs:
        return []
    result = []
    seen = []

    if col_kind == "ALL":
        for c, h in header_pairs:
            if _is_service_header(h):
                continue
            result.append((c, h))
        return result

    for spec in col_specs:
        st = unicode(spec).strip()
        if st == u"":
            continue
        if col_kind == "INDEX":
            c = int(st) - 1
        elif col_kind == "LETTER":
            c = _col_letters_to_index(st)
        else:
            c = -1
            for hc, h in header_pairs:
                if _text_match(h, st):
                    c = hc
                    break
            if c < 0:
                continue
        if c in seen:
            continue
        h = u""
        for hc, ht in header_pairs:
            if hc == c:
                h = ht
                break
        if h == u"":
            h = _header_name_at(u"", c)
        if _is_service_header(h):
            continue
        seen.append(c)
        result.append((c, h))
    return result


def _direct_read_cell_xlsx(path, sheet_index, col0, row0):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        cell = ws.cell(row=int(row0) + 1, column=int(col0) + 1)
        return cell.value
    except Exception:
        return None
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_read_cell_ods(path, sheet_index, col0, row0):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell

    try:
        doc = load(path)
        sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
        rows = sheet.getElementsByType(TableRow)
        if row0 < 0 or row0 >= len(rows):
            return None
        cells = rows[int(row0)].getElementsByType(TableCell)
        if col0 < 0 or col0 >= len(cells):
            return None
        return _ods_cell_value(cells[int(col0)])
    except Exception:
        return None


def _direct_read_cell(read_path, src_kind, sheet_index, col0, row0):
    if src_kind == "xlsx":
        return _direct_read_cell_xlsx(read_path, sheet_index, col0, row0)
    return _direct_read_cell_ods(read_path, sheet_index, col0, row0)


def _direct_coerce_cell_value(v):
    if isinstance(v, (bytes, str, unicode)):
        if not _direct_wants_convert_to_numbers():
            return unicode(v) if v is not None else u""
        return _coerce_text_value(v)
    return v


def _direct_prefetch_cell_values(read_path, src_kind, sheet_index, specs):
    out = {}
    if not specs:
        return out
    for spec in specs:
        if spec is None:
            continue
        col_name = unicode(spec.get("column") or u"").strip()
        if col_name == u"":
            continue
        if col_name in out:
            continue
        v = _direct_read_cell(
            read_path,
            src_kind,
            sheet_index,
            int(spec.get("src_col", 0)),
            int(spec.get("row", 0)),
        )
        out[col_name] = _direct_coerce_cell_value(v)
    return out


def _direct_prefetch_add_to_variables_map( fi, fpath, is_current_book, path_display, sheet_name, read_path, src_kind, sheet_index ):
    """
    Заполнить runtime-карту из «Добавить_в_карту_переменных» при direct/xml чтении.
    Глобальные переменные уже в карте после merge_run_context_begin.
    """
    try:
        import libre_macros_collect_cfg as _cw_cfg

        var_all = getattr(_cw_cfg, "_MERGE_ADD_TO_VARIABLES_PER_FILE", None) or []
        if fi >= len(var_all) or not var_all[fi]:
            return
        specs = var_all[fi]
    except Exception:
        return
    try:
        import collect_workbooks as cw
    except Exception:
        try:
            import importlib

            cw = importlib.import_module("collect_workbooks")
        except Exception as err:
            _log("add_to_variables_map direct: import collect_workbooks: %s" % err)
            return
    try:
        log_fpath = fpath if not is_current_book else (cw.get_current_book_path() or u"")
        file_key = cw.cell_to_column_file_key(log_fpath, path_display)
    except Exception:
        file_key = unicode(fpath or u"")
    for spec in specs:
        if spec is None:
            continue
        try:
            if not cw._source_variables_file_matches(spec.get("file"), file_key):
                continue
            if not cw._source_variables_sheet_matches(spec.get("sheet"), sheet_name):
                continue
        except Exception:
            continue
        try:
            src_col = int(spec.get("src_col", -1))
            row_idx = int(spec.get("row", -1))
        except Exception:
            continue
        if src_col < 0 or row_idx < 0:
            continue
        raw = _direct_read_cell(read_path, src_kind, sheet_index, src_col, row_idx)
        try:
            snap = cw._snapshot_from_python_value(raw)
            text_val = u"" if snap is None else unicode(snap.get(u"text") or u"")
            token = cw._source_variables_type_token_from_value(raw, snap)
        except Exception:
            text_val = unicode(raw if raw is not None else u"")
            token = u"text"
        try:
            cw._merge_source_variables_put(
                spec.get("name"),
                sheet_name,
                file_key,
                (text_val, token),
                report=None,
                cell_ref=spec.get("cell") or u"",
            )
        except Exception as err:
            _log("add_to_variables_map direct put: %s" % err)


def _direct_cell_specs_per_file(cell_to_column_plan, n_files):
    per_file = []
    fi = 0
    while fi < n_files:
        per_file.append([])
        fi += 1
    if not cell_to_column_plan:
        return per_file
    for entry in cell_to_column_plan:
        specs = (entry or {}).get("per_file") or []
        fi = 0
        while fi < n_files:
            if fi < len(specs) and specs[fi] is not None:
                per_file[fi].append(specs[fi])
            fi += 1
    return per_file


def _direct_cell_spec_sheet_filter_ref(spec, role):
    """role: target — лист результата; source — лист источника."""
    if spec is None:
        return u""
    if role == u"source":
        return unicode(
            (spec or {}).get("source_sheet")
            or (spec or {}).get("src_sheet")
            or u""
        ).strip()
    return unicode(
        (spec or {}).get("sheet_name") or (spec or {}).get("sheet") or u""
    ).strip()


def _direct_cell_sheet_name_candidates(sheet_name):
    sn = unicode(sheet_name or u"").strip()
    out = []
    if sn == u"":
        return out
    out.append(sn)
    m = re.match(r"^(\d+)_(.+)$", sn)
    if m is not None:
        base = unicode(m.group(2) or u"").strip()
        if base != u"" and base not in out:
            out.append(base)
    return out


def _direct_cell_spec_applies_to_sheet(spec, sheet_name, role=u"target"):
    """Фильтр Ячейка_В_Столбец по имени листа (wildcards *, список через запятую)."""
    if spec is None:
        return False
    ref = _direct_cell_spec_sheet_filter_ref(spec, role)
    if ref == u"":
        return True
    candidates = _direct_cell_sheet_name_candidates(sheet_name)
    if not candidates:
        return role != u"source"
    for part in unicode(ref).replace(u";", u",").split(u","):
        p = part.strip()
        if p == u"":
            continue
        for cand in candidates:
            if _text_match(cand, p):
                return True
        m_pat = re.match(r"^(\d+)_(.+)$", p)
        if m_pat is not None:
            base_pat = unicode(m_pat.group(2) or u"").strip()
            if base_pat != u"":
                for cand in candidates:
                    if _text_match(cand, base_pat):
                        return True
    return False


def _direct_cell_spec_applies_to_target_sheet(spec, sheet_name):
    return _direct_cell_spec_applies_to_sheet(spec, sheet_name, u"target")


def _direct_cell_spec_applies_to_source_sheet(spec, sheet_name):
    return _direct_cell_spec_applies_to_sheet(spec, sheet_name, u"source")


def _direct_copy_prefetch_cell_values_from_rows(rows, specs):
    """Снимок ячеек из уже прочитанных строк источника (до обрезки)."""
    out = {}
    if not specs or not rows:
        return out
    for spec in specs:
        if spec is None:
            continue
        col_name = unicode(spec.get("column") or u"").strip()
        if col_name == u"" or col_name in out:
            continue
        try:
            c0 = int(spec.get("src_col", 0))
            r0 = int(spec.get("row", 0))
        except Exception:
            continue
        if r0 < 0 or c0 < 0 or r0 >= len(rows):
            continue
        row = rows[r0] or []
        if c0 >= len(row):
            continue
        out[col_name] = _direct_coerce_cell_value(row[c0])
    return out


def _direct_copy_inject_cell_to_column(rows, cell_values):
    """Добавить столбцы Ячейка_В_Столбец в матрицу строк (заголовок уже в rows[0])."""
    if not rows or not cell_values:
        return rows
    rows = [list(r or []) for r in rows]
    max_w = 0
    for r in rows:
        if len(r) > max_w:
            max_w = len(r)
    for r in rows:
        while len(r) < max_w:
            r.append(u"")
    for col_name, val in cell_values.items():
        cn = unicode(col_name or u"").strip()
        if cn == u"":
            continue
        header = rows[0]
        idx = -1
        hi = 0
        while hi < len(header):
            if _merge_identity_key(unicode(header[hi] or u"")) == _merge_identity_key(cn):
                idx = hi
                break
            hi += 1
        if idx < 0:
            idx = len(header)
            rows[0].append(cn)
            ri = 1
            while ri < len(rows):
                while len(rows[ri]) < idx:
                    rows[ri].append(u"")
                rows[ri].append(val)
                ri += 1
        else:
            ri = 1
            while ri < len(rows):
                while len(rows[ri]) <= idx:
                    rows[ri].append(u"")
                rows[ri][idx] = val
                ri += 1
    return rows

def _direct_iter_filtered_rows( read_path, src_kind, sheet_index, first_data_row, last_data_row, columns, skip_parsed):
    col_indices = [c[0] for c in columns]
    if not col_indices:
        return
    first_data_row = int(first_data_row)
    last_data_row = int(last_data_row)
    if last_data_row < first_data_row:
        return
    skip_indices = (skip_parsed or {}).get("col_indices") or []
    min_col = min(col_indices + skip_indices) if skip_indices else min(col_indices)
    max_col = max(col_indices + skip_indices) if skip_indices else max(col_indices)

    if src_kind == "xlsx":
        import openpyxl_bundled  # noqa: F401
        from openpyxl import load_workbook

        wb = load_workbook(read_path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[int(sheet_index)]
            rr = first_data_row
            for row_vals in ws.iter_rows(
                min_row=first_data_row + 1,
                max_row=last_data_row + 1,
                min_col=min_col + 1,
                max_col=max_col + 1,
                values_only=True,
            ):
                if rr % 200 == 0:
                    _direct_ui_yield()
                row_map = {}
                off = 0
                while off < len(row_vals):
                    row_map[min_col + off] = row_vals[off]
                    off += 1
                if (skip_parsed or {}).get("mode") == "columns":
                    if _direct_row_skip_by_empty(row_map, skip_indices):
                        rr += 1
                        continue
                out_row = []
                for ci in col_indices:
                    out_row.append(_direct_coerce_cell_value(row_map.get(ci)))
                yield out_row
                rr += 1
        finally:
            try:
                wb.close()
            except Exception:
                pass
        return

    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    doc = load(read_path)
    sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
    rows = sheet.getElementsByType(TableRow)
    rr = first_data_row
    while rr <= last_data_row:
        if rr % 200 == 0:
            _direct_ui_yield()
        row_map = {}
        if rr < len(rows):
            for col_idx, val in _ods_row_col_values(rows[rr]):
                if min_col <= col_idx <= max_col:
                    row_map[col_idx] = val
        if (skip_parsed or {}).get("mode") == "columns":
            if _direct_row_skip_by_empty(row_map, skip_indices):
                rr += 1
                continue
        out_row = []
        for cidx in col_indices:
            out_row.append(_direct_coerce_cell_value(row_map.get(cidx)))
        yield out_row
        rr += 1


def _direct_resolve_delete_top_count(sheet_name, mapping, skipped_top_by_sheet=None):
    """
    N для удаления верхних строк.

    Удаление_верхних_строк выполняется ДО Постобработка_xml (до переименовать_лист).
    Поэтому в JSON часто пишут итоговое имя «Заказы», а в файле ещё «1_Заказы» —
    сопоставляем оба варианта (как у фильтров sheet в постобработке).
    Шаблоны ``*``/``?``: ``*Отчет*`` ↔ ``1_Отчет_Q1``.

    skipped_top_by_sheet — для one/many: сколько строк над шапкой уже не
    переносилось (hdr_row 0-based); вычитается из N.
    """
    if not mapping or sheet_name is None:
        return None
    name = unicode(sheet_name).strip()
    if name == u"":
        return None
    raw_n = None
    if name in mapping:
        raw_n = mapping[name]
    name_key = _merge_identity_key(name)
    m_sheet = re.match(r"^(\d+)_(.+)$", name)
    sheet_base = unicode(m_sheet.group(2)).strip() if m_sheet else u""
    sheet_base_key = _merge_identity_key(sheet_base) if sheet_base else u""

    if raw_n is None:
        for k, v in mapping.items():
            mk = unicode(k or u"").strip()
            if mk == u"":
                continue
            if _merge_identity_key(mk) == name_key:
                raw_n = v
                break
            # лист «1_Заказы», в карте «Заказы»
            if sheet_base_key and _merge_identity_key(mk) == sheet_base_key:
                raw_n = v
                break
            # лист «Заказы», в карте «1_Заказы»
            m_map = re.match(r"^(\d+)_(.+)$", mk)
            if m_map is not None:
                map_base = unicode(m_map.group(2)).strip()
                if map_base and _merge_identity_key(map_base) == name_key:
                    raw_n = v
                    break
            if u"*" in mk or u"?" in mk:
                if _text_match(name, mk):
                    raw_n = v
                    break
                if sheet_base and _text_match(sheet_base, mk):
                    raw_n = v
                    break
    if raw_n is None:
        return None
    try:
        n = int(raw_n)
    except Exception:
        return None
    skipped = 0
    if skipped_top_by_sheet:
        try:
            skipped = int(skipped_top_by_sheet.get(name) or 0)
        except Exception:
            skipped = 0
        if skipped <= 0 and sheet_base:
            try:
                skipped = int(skipped_top_by_sheet.get(sheet_base) or 0)
            except Exception:
                skipped = 0
    n_eff = n - int(skipped or 0)
    if n_eff <= 0:
        return None
    return n_eff


def _direct_ods_delete_top_rows(path, delete_map, skipped_top_by_sheet=None):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    doc = load(path)
    changed = False
    for table in doc.spreadsheet.getElementsByType(Table):
        name = unicode(table.getAttribute("name") or u"")
        n = _direct_resolve_delete_top_count(name, delete_map, skipped_top_by_sheet)
        if not n or int(n) <= 0:
            continue
        removed = 0
        while removed < int(n):
            row_nodes = table.getElementsByType(TableRow)
            if not row_nodes:
                break
            try:
                table.removeChild(row_nodes[0])
            except Exception:
                break
            removed += 1
        if removed > 0:
            _log("delete_top_rows: лист «%s» — удалено %d" % (name, removed))
            changed = True
    if changed:
        doc.save(path)


def _direct_xlsx_delete_top_rows(path, delete_map, skipped_top_by_sheet=None):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        changed = False
        for sheet_name in wb.sheetnames:
            n = _direct_resolve_delete_top_count(
                sheet_name, delete_map, skipped_top_by_sheet
            )
            if not n or int(n) <= 0:
                continue
            ws = wb[sheet_name]
            max_r = ws.max_row or 0
            if max_r <= 0:
                continue
            to_del = min(int(n), int(max_r))
            if to_del > 0:
                ws.delete_rows(1, to_del)
                _log(
                    "delete_top_rows: лист «%s» — удалено %d"
                    % (unicode(sheet_name), to_del)
                )
                changed = True
        if changed:
            wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_resolve_copy_trim_spec(sheet_name, trim_map):
    """Спецификация обрезки для листа (как _direct_resolve_delete_top_count)."""
    if not trim_map or sheet_name is None:
        return None
    name = unicode(sheet_name).strip()
    if name == u"":
        return None
    if name in trim_map:
        return trim_map[name]
    name_key = _merge_identity_key(name)
    m_sheet = re.match(r"^(\d+)_(.+)$", name)
    sheet_base = unicode(m_sheet.group(2)).strip() if m_sheet else u""
    sheet_base_key = _merge_identity_key(sheet_base) if sheet_base else u""

    for k, v in trim_map.items():
        mk = unicode(k or u"").strip()
        if mk == u"":
            continue
        if _merge_identity_key(mk) == name_key:
            return v
        if sheet_base_key and _merge_identity_key(mk) == sheet_base_key:
            return v
        m_map = re.match(r"^(\d+)_(.+)$", mk)
        if m_map is not None:
            map_base = unicode(m_map.group(2)).strip()
            if map_base and _merge_identity_key(map_base) == name_key:
                return v
    return None


def _direct_ods_copy_sheets_row_trim(path, trim_map):
    _direct_ods_copy_sheets_trim_and_convert(
        path, trim_map, [], hdr_row0=0, do_convert=False
    )


def _direct_ods_copy_sheets_trim_and_convert( path, trim_map, sheet_names, hdr_row0=0, do_convert=True ):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    trim_map = trim_map or {}
    sheet_names = sheet_names or []
    target_keys = _direct_copy_sheet_target_keys(sheet_names)
    want_convert = bool(do_convert and _direct_wants_convert_to_numbers())
    doc = load(path)
    style_cache = _direct_ods_style_cache_new(doc) if want_convert else None
    changed = False
    for table in doc.spreadsheet.getElementsByType(Table):
        name = unicode(table.getAttribute("name") or u"").strip()
        spec = _direct_resolve_copy_trim_spec(name, trim_map)
        in_convert = want_convert and _direct_copy_sheet_name_matches(
            name, sheet_names, target_keys
        )
        if not spec and not in_convert:
            continue
        if spec:
            top, after = _direct_ods_trim_table_rows(table, spec)
            if top > 0 or after > 0:
                _log(
                    "copy_sheets_row_trim: лист «%s» — сверху %d, после заголовка %d"
                    % (name, top, after)
                )
                changed = True
        if in_convert:
            rows = table.getElementsByType(TableRow)
            if _direct_ods_convert_table_rows(name, rows, int(hdr_row0), style_cache):
                changed = True
    if changed:
        doc.save(path)


def _direct_xlsx_copy_sheets_row_trim(path, trim_map):
    _direct_xlsx_copy_sheets_trim_and_convert(
        path, trim_map, [], hdr_row0=0, do_convert=False
    )


def _direct_xlsx_convert_ws(sheet_name, ws, hdr_row0):
    """Конвертация текстовых дат/чисел на листе xlsx (даты, затем числа)."""
    end_col, end_row = _direct_scan_bounds_xlsx(ws, hdr_row0, 0)
    if end_col < 0 or end_row <= hdr_row0:
        return False
    max_col_1 = end_col + 1
    max_row_1 = end_row + 1
    data_r0_1 = int(hdr_row0) + 2

    def _row_map(rr):
        row1 = int(rr) + 1
        if row1 > max_row_1:
            return {}
        out = {}
        try:
            row_vals = next(
                ws.iter_rows(
                    min_row=row1,
                    max_row=row1,
                    min_col=1,
                    max_col=max_col_1,
                    values_only=True,
                ),
                (),
            )
        except Exception:
            return {}
        ci = 0
        while ci < len(row_vals):
            if row_vals[ci] is not None:
                out[ci] = row_vals[ci]
            ci += 1
        return out

    date_cols, num_cols = _direct_build_convert_column_sets_probe(
        _row_map, hdr_row0, end_col, end_row
    )
    if not date_cols and not num_cols:
        return False
    nm = unicode(sheet_name or u"")
    if date_cols:
        _log(
            "copy_sheets_convert_dt: лист «%s» — столбцов %d (%s)"
            % (nm, len(date_cols), sorted(date_cols))
        )
    if num_cols:
        _log(
            "copy_sheets_convert_num: лист «%s» — столбцов %d (%s)"
            % (nm, len(num_cols), sorted(num_cols))
        )
    converted_dt = 0
    converted_num = 0
    batch_start = data_r0_1
    if date_cols:
        while batch_start <= max_row_1:
            batch_end = min(
                max_row_1,
                batch_start + _DIRECT_COPY_DT_ROW_BATCH - 1,
            )
            for row_cells in ws.iter_rows(
                min_row=batch_start,
                max_row=batch_end,
                min_col=1,
                max_col=max_col_1,
            ):
                for col in date_cols:
                    if col >= len(row_cells):
                        continue
                    cell = row_cells[col]
                    v = cell.value
                    if not _direct_is_plain_string_for_datetime(v):
                        continue
                    parsed = _direct_try_parse_datetime_text(v)
                    if parsed is None:
                        continue
                    cell.value = parsed
                    _direct_xlsx_apply_converted_cell_style(cell, parsed)
                    converted_dt += 1
            _direct_ui_yield()
            batch_start = batch_end + 1
        _log(
            "copy_sheets_convert_dt: лист «%s» — ячеек %d"
            % (nm, int(converted_dt))
        )
    batch_start = data_r0_1
    if num_cols:
        while batch_start <= max_row_1:
            batch_end = min(
                max_row_1,
                batch_start + _DIRECT_COPY_DT_ROW_BATCH - 1,
            )
            for row_cells in ws.iter_rows(
                min_row=batch_start,
                max_row=batch_end,
                min_col=1,
                max_col=max_col_1,
            ):
                for col in num_cols:
                    if col >= len(row_cells):
                        continue
                    cell = row_cells[col]
                    v = cell.value
                    if not _direct_is_plain_string_for_datetime(v):
                        continue
                    parsed = _direct_try_parse_number_text(v)
                    if parsed is None:
                        continue
                    cell.value = parsed
                    _direct_xlsx_apply_converted_cell_style(cell, parsed)
                    converted_num += 1
            _direct_ui_yield()
            batch_start = batch_end + 1
        _log(
            "copy_sheets_convert_num: лист «%s» — ячеек %d"
            % (nm, int(converted_num))
        )
    return True


def _direct_xlsx_copy_sheets_trim_and_convert( path, trim_map, sheet_names, hdr_row0=0, do_convert=True ):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    trim_map = trim_map or {}
    sheet_names = sheet_names or []
    target_keys = _direct_copy_sheet_target_keys(sheet_names)
    want_convert = bool(do_convert and _direct_wants_convert_to_numbers())
    wb = load_workbook(path)
    try:
        changed = False
        for sheet_name in wb.sheetnames:
            nm = unicode(sheet_name or u"").strip()
            spec = _direct_resolve_copy_trim_spec(nm, trim_map)
            in_convert = want_convert and _direct_copy_sheet_name_matches(
                nm, sheet_names, target_keys
            )
            if not spec and not in_convert:
                continue
            ws = wb[sheet_name]
            if spec:
                try:
                    top = max(0, int(spec.get("top") or 0))
                except Exception:
                    top = 0
                try:
                    after = max(0, int(spec.get("after_header") or 0))
                except Exception:
                    after = 0
                if top > 0:
                    ws.delete_rows(1, top)
                if after > 0:
                    ws.delete_rows(2, after)
                if top > 0 or after > 0:
                    _log(
                        "copy_sheets_row_trim: лист «%s» — сверху %d, после заголовка %d"
                        % (nm, top, after)
                    )
                    changed = True
            if in_convert:
                if _direct_xlsx_convert_ws(nm, ws, int(hdr_row0)):
                    changed = True
        if changed:
            wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def direct_apply_copy_sheets_row_trim(path, trim_map):
    """
    Обрезка скопированных листов в файле результата (xml_excel_ods, режим копирования).

    trim_map: имя листа → {"top": N, "after_header": M}.
    """
    direct_apply_copy_sheets_row_trim_and_convert(
        path, trim_map, [], header_row=0, convert=False
    )


def direct_apply_copy_sheets_row_trim_and_convert( path, trim_map, sheet_names, header_row=0, convert=True ):
    """
    Один проход по файлу: обрезка + (опционально) текст → число/дата.

    Для «Копирование листов» + xml_excel_ods: после UNO-копии листов один load/save
    вместо отдельных проходов trim и convert. convert=False — только обрезка.
    """
    trim_map = trim_map or {}
    names = []
    seen = set()
    for nm in list(trim_map.keys()) + list(sheet_names or []):
        t = unicode(nm or u"").strip()
        if t == u"" or t in seen:
            continue
        seen.add(t)
        names.append(t)
    want_convert = bool(convert and _direct_wants_convert_to_numbers())
    if not trim_map and not (want_convert and names):
        return
    hdr = int(header_row or 0)
    if _is_ods(path):
        _direct_ods_copy_sheets_trim_and_convert(
            path, trim_map, names, hdr_row0=hdr, do_convert=want_convert
        )
    elif _is_xlsx_target(path):
        _direct_xlsx_copy_sheets_trim_and_convert(
            path, trim_map, names, hdr_row0=hdr, do_convert=want_convert
        )


def _direct_hdr_row_after_copy_trim(blk, trim_spec=None):
    """0-based строка заголовков на листе после автообрезки top."""
    blk = blk or {}
    if blk.get("no_header"):
        return -1
    hr, _last = _direct_block_header_first_last_1(blk)
    base = max(0, int(hr) - 1)
    if not trim_spec:
        return base
    top = int((trim_spec or {}).get("top") or 0)
    return max(0, base - top)


def _direct_apply_col_rename_to_header_row(row, rename_items, start_col=0, col_count=None):
    """Переименовать ячейки строки заголовков (list) по правилам «Столбцы»."""
    if not row or not rename_items:
        return row, 0
    try:
        from libre_macros_param_codec import columns_map_apply_renames
    except Exception:
        return row, 0
    cols = []
    c = int(start_col or 0)
    if col_count is not None and int(col_count) > 0:
        limit = c + int(col_count)
    else:
        limit = max(len(row), c + 1)
        limit = min(limit, c + 500)
    while c < limit:
        val = u""
        if c < len(row):
            val = unicode(row[c] or u"").strip()
        if val == u"" and col_count is None:
            break
        cols.append((c, val))
        c += 1
    if not cols:
        return row, 0
    new_cols = columns_map_apply_renames(cols, rename_items)
    changed = False
    out = list(row)
    for col_idx, new_name in new_cols:
        old = u""
        if col_idx < len(out):
            old = unicode(out[col_idx] or u"").strip()
        nn = unicode(new_name or u"").strip()
        if nn != u"" and nn != old:
            while len(out) <= col_idx:
                out.append(u"")
            out[col_idx] = nn
            changed = True
    return (out, 1 if changed else 0)


def _direct_copy_apply_col_rename_to_rows(rows, blk, rename_items):
    """Переименование заголовков в отфильтрованной матрице (direct copy ODS)."""
    if not rows or not rename_items or not blk or blk.get("no_header"):
        return rows
    start_col = int((blk or {}).get("start_col", 0) or 0)
    col_count = (blk or {}).get("col_count")
    new_hdr, n = _direct_apply_col_rename_to_header_row(
        rows[0], rename_items, start_col=start_col, col_count=col_count
    )
    if n <= 0:
        return rows
    out = list(rows)
    out[0] = new_hdr
    return out


def _direct_xlsx_copy_sheets_column_rename(path, settings, trim_map=None):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    block_by_name = (settings or {}).get("copy_sheets_block_by_name") or {}
    file_index_by_name = (settings or {}).get("copy_sheets_file_index_by_name") or {}
    per_file = (settings or {}).get("per_file_col_rename_items") or []
    sheet_names = (settings or {}).get("direct_result_sheet_names") or list(block_by_name.keys())
    trim_map = trim_map or {}
    if not sheet_names or not per_file:
        return
    wb = load_workbook(path)
    changed = False
    for nm in sheet_names:
        if nm not in wb.sheetnames:
            continue
        fi = file_index_by_name.get(nm)
        rename_items = []
        if fi is not None and fi < len(per_file):
            rename_items = per_file[fi] or []
        if not rename_items:
            continue
        blk = block_by_name.get(nm) or {}
        hdr_row = _direct_hdr_row_after_copy_trim(blk, trim_map.get(nm))
        if hdr_row < 0:
            continue
        ws = wb[nm]
        row_vals = []
        try:
            row_vals = list(
                next(
                    ws.iter_rows(
                        min_row=int(hdr_row) + 1,
                        max_row=int(hdr_row) + 1,
                        values_only=True,
                    ),
                    (),
                )
            )
        except Exception:
            row_vals = []
        new_row, n = _direct_apply_col_rename_to_header_row(
            row_vals,
            rename_items,
            start_col=int(blk.get("start_col", 0) or 0),
            col_count=blk.get("col_count"),
        )
        if n <= 0:
            continue
        ci = 0
        while ci < len(new_row):
            ws.cell(row=int(hdr_row) + 1, column=ci + 1, value=new_row[ci])
            ci += 1
        changed = True
        _log(u"copy_sheets_col_rename: лист «%s» — заголовки обновлены" % nm)
    if changed:
        wb.save(path)
    try:
        wb.close()
    except Exception:
        pass


def _direct_ods_copy_sheets_column_rename(path, settings, trim_map=None):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell

    block_by_name = (settings or {}).get("copy_sheets_block_by_name") or {}
    file_index_by_name = (settings or {}).get("copy_sheets_file_index_by_name") or {}
    per_file = (settings or {}).get("per_file_col_rename_items") or []
    sheet_names = (settings or {}).get("direct_result_sheet_names") or list(block_by_name.keys())
    trim_map = trim_map or {}
    if not sheet_names or not per_file:
        return
    doc = load(path)
    changed = False
    for table in doc.spreadsheet.getElementsByType(Table):
        nm = unicode(table.getAttribute("name") or u"").strip()
        if nm not in sheet_names:
            continue
        fi = file_index_by_name.get(nm)
        rename_items = []
        if fi is not None and fi < len(per_file):
            rename_items = per_file[fi] or []
        if not rename_items:
            continue
        blk = block_by_name.get(nm) or {}
        hdr_row = _direct_hdr_row_after_copy_trim(blk, trim_map.get(nm))
        if hdr_row < 0:
            continue
        rows = table.getElementsByType(TableRow)
        if hdr_row < 0 or hdr_row >= len(rows):
            continue
        cells = rows[hdr_row].getElementsByType(TableCell)
        row_vals = []
        ci = 0
        while ci < len(cells):
            row_vals.append(unicode(_ods_cell_text(cells[ci])).strip())
            ci += 1
        new_row, n = _direct_apply_col_rename_to_header_row(
            row_vals,
            rename_items,
            start_col=int(blk.get("start_col", 0) or 0),
            col_count=blk.get("col_count"),
        )
        if n <= 0:
            continue
        from odf.text import P
        ci = 0
        while ci < len(new_row):
            if ci < len(cells):
                cell = cells[ci]
                txt = unicode(new_row[ci] or u"")
                try:
                    for p in list(cell.getElementsByType(P)):
                        cell.removeChild(p)
                except Exception:
                    pass
                try:
                    cell.setAttribute("valuetype", "string")
                    cell.setAttribute("stringvalue", txt)
                except Exception:
                    pass
                try:
                    cell.addElement(P(text=txt))
                except Exception:
                    pass
            ci += 1
        changed = True
        _log(u"copy_sheets_col_rename: лист «%s» — заголовки обновлены" % nm)
    if changed:
        doc.save(path)


def direct_apply_copy_sheets_column_rename(path, settings, trim_map=None):
    """
    Переименование заголовков на скопированных листах в файле результата.

    Для xml_excel_ods: после обрезки строк, до постобработки.
    """
    if not path or not settings:
        return
    per_file = (settings or {}).get("per_file_col_rename_items") or []
    if not per_file:
        return
    if not any(per_file):
        return
    if _is_ods(path):
        _direct_ods_copy_sheets_column_rename(path, settings, trim_map=trim_map)
    elif _is_xlsx_target(path):
        _direct_xlsx_copy_sheets_column_rename(path, settings, trim_map=trim_map)


def _direct_strip_leading_apostrophe(text):
    """Снять ведущий апостроф Calc/Excel ('02.02.1976 → 02.02.1976)."""
    s = unicode(text or u"")
    s = s.strip()
    while s and s[0] in (u"'", u"\u2019", u"`"):
        s = s[1:].lstrip()
    return s


def _direct_is_plain_string_for_datetime(v):
    if v is None:
        return False
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return False
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return False
    return isinstance(v, (str, unicode))


def _direct_build_datetime_column_set(get_row_map, hdr_row0, end_col, end_row):
    """
    Столбцы, где ≥ _DIRECT_COPY_DT_COL_MATCH_MIN строк данных распознаны как дата/время.

    get_row_map(row0) → {col_index: raw_value}.
    """
    if end_col < 0 or end_row <= hdr_row0:
        return set()
    counts = {}
    confirmed = set()
    data_r0 = int(hdr_row0) + 1
    rr = data_r0
    while rr <= end_row:
        if (rr - data_r0) % _DIRECT_COPY_DT_ROW_BATCH == 0:
            _direct_ui_yield()
        row_map = get_row_map(rr) or {}
        col = 0
        while col <= end_col:
            if col not in confirmed:
                val = row_map.get(col)
                if _direct_is_plain_string_for_datetime(val):
                    if _direct_try_parse_datetime_text(val) is not None:
                        counts[col] = counts.get(col, 0) + 1
                        if counts[col] >= _DIRECT_COPY_DT_COL_MATCH_MIN:
                            confirmed.add(col)
            col += 1
        rr += 1
    return confirmed


def _direct_build_number_column_set( get_row_map, hdr_row0, end_col, end_row, exclude_cols=None ):
    """
    Столбцы, где ≥ _DIRECT_COPY_DT_COL_MATCH_MIN строк данных распознаны как число.

    Строки, успешно разбираемые как дата/время, не учитываются.
    exclude_cols — уже обработанные столбцы (например, даты).
    """
    if end_col < 0 or end_row <= hdr_row0:
        return set()
    skip = set(exclude_cols or ())
    counts = {}
    confirmed = set()
    data_r0 = int(hdr_row0) + 1
    rr = data_r0
    while rr <= end_row:
        if (rr - data_r0) % _DIRECT_COPY_DT_ROW_BATCH == 0:
            _direct_ui_yield()
        row_map = get_row_map(rr) or {}
        col = 0
        while col <= end_col:
            if col not in skip and col not in confirmed:
                val = row_map.get(col)
                if _direct_is_plain_string_for_datetime(val):
                    if _direct_try_parse_number_text(val) is not None:
                        counts[col] = counts.get(col, 0) + 1
                        if counts[col] >= _DIRECT_COPY_DT_COL_MATCH_MIN:
                            confirmed.add(col)
            col += 1
        rr += 1
    return confirmed


def _direct_build_convert_column_sets_probe( get_row_map, hdr_row0, end_col, end_row, probe_rows=None ):
    """Столбцы дат/чисел по первым probe_rows строкам данных (≥3 совпадений)."""
    if end_col < 0 or end_row <= hdr_row0:
        return set(), set()
    if _direct_convert_auto_probe():
        probe = int(probe_rows or _DIRECT_COPY_COL_PROBE_ROWS)
        scan_end = min(int(end_row), int(hdr_row0) + max(probe, _DIRECT_COPY_DT_COL_MATCH_MIN))
        date_cols = _direct_build_datetime_column_set(
            get_row_map, hdr_row0, end_col, scan_end
        )
        num_cols = _direct_build_number_column_set(
            get_row_map, hdr_row0, end_col, scan_end, exclude_cols=date_cols
        )
    else:
        date_cols = set()
        num_cols = set()
    try:
        from libre_macros_param_codec import columns_map_apply_convert_policies

        items = (_direct_options or {}).get("column_map_items") or []
        if items:
            hdr_map = get_row_map(int(hdr_row0)) or {}
            headers_by_col = {}
            col = 0
            while col <= end_col:
                headers_by_col[col] = unicode(hdr_map.get(col) or u"").strip()
                col += 1
            date_cols, num_cols, _ef = columns_map_apply_convert_policies(
                date_cols, num_cols, headers_by_col, items, output_headers=False
            )
    except Exception:
        pass
    return date_cols, num_cols


def _direct_copy_sheet_target_keys(sheet_names):
    target_keys = set()
    for nm in sheet_names or []:
        try:
            target_keys.add(_merge_identity_key(nm))
        except Exception:
            target_keys.add(unicode(nm or u"").strip().casefold())
    return target_keys


def _direct_copy_sheet_name_matches(nm, sheet_names, target_keys):
    if nm == u"":
        return False
    try:
        if _merge_identity_key(nm) in target_keys:
            return True
    except Exception:
        pass
    return nm in (sheet_names or [])


def _direct_ods_trim_table_rows(table, spec):
    """Обрезка строк table по spec; возвращает (removed_top, removed_after)."""
    from odf.table import TableRow

    try:
        top = max(0, int(spec.get("top") or 0))
    except Exception:
        top = 0
    try:
        after = max(0, int(spec.get("after_header") or 0))
    except Exception:
        after = 0
    if top <= 0 and after <= 0:
        return 0, 0
    removed = 0
    while removed < top:
        row_nodes = table.getElementsByType(TableRow)
        if not row_nodes:
            break
        try:
            table.removeChild(row_nodes[0])
        except Exception:
            break
        removed += 1
    after_removed = 0
    while after_removed < after:
        row_nodes = table.getElementsByType(TableRow)
        if len(row_nodes) < 2:
            break
        try:
            table.removeChild(row_nodes[1])
        except Exception:
            break
        after_removed += 1
    return removed, after_removed


def _direct_ods_convert_table_rows(table_name, rows, hdr_row0, style_cache):
    """Конвертация текстовых дат/чисел на листе ODS (даты, затем числа)."""
    end_col, end_row = _direct_scan_bounds_ods(rows, hdr_row0, 0)
    if end_col < 0 or end_row <= hdr_row0:
        return False

    def _row_map(rr):
        if rr < 0 or rr >= len(rows):
            return {}
        return dict(_ods_row_col_values(rows[rr]))

    date_cols, num_cols = _direct_build_convert_column_sets_probe(
        _row_map, hdr_row0, end_col, end_row
    )
    if not date_cols and not num_cols:
        return False
    nm = unicode(table_name or u"")
    if date_cols:
        _log(
            "copy_sheets_convert_dt: лист «%s» — столбцов %d (%s)"
            % (nm, len(date_cols), sorted(date_cols))
        )
    if num_cols:
        _log(
            "copy_sheets_convert_num: лист «%s» — столбцов %d (%s)"
            % (nm, len(num_cols), sorted(num_cols))
        )
    converted_dt = 0
    converted_num = 0
    if date_cols:
        rr = hdr_row0 + 1
        while rr <= end_row:
            if (rr - hdr_row0) % _DIRECT_COPY_DT_ROW_BATCH == 0:
                _direct_ui_yield()
            if rr < len(rows):
                converted_dt += _direct_ods_convert_row_datetimes(
                    rows[rr], date_cols, end_col, style_cache
                )
            rr += 1
        _log(
            "copy_sheets_convert_dt: лист «%s» — ячеек %d" % (nm, int(converted_dt))
        )
    if num_cols:
        rr = hdr_row0 + 1
        while rr <= end_row:
            if (rr - hdr_row0) % _DIRECT_COPY_DT_ROW_BATCH == 0:
                _direct_ui_yield()
            if rr < len(rows):
                converted_num += _direct_ods_convert_row_numbers(
                    rows[rr], num_cols, end_col, style_cache
                )
            rr += 1
        _log(
            "copy_sheets_convert_num: лист «%s» — ячеек %d"
            % (nm, int(converted_num))
        )
    return True


def _direct_ods_prebuilt_cell_style(cache, parsed):
    """Имя stylename для ячейки после convert (кэш на проход по листу)."""
    if cache is None or parsed is None:
        return None
    if isinstance(parsed, (datetime.date, datetime.datetime)):
        fmt = _direct_date_format_for_parsed(parsed, for_ods=True)
        key = (u"dt", fmt)
    elif isinstance(parsed, (int, float)):
        num_fmt = _direct_default_number_format()
        if num_fmt == u"":
            return None
        key = (u"num", num_fmt)
    else:
        return None
    pre = cache.setdefault(u"converted", {})
    if key in pre:
        return pre[key]
    sn = _direct_ods_cell_style_name(
        cache, {u"data_style": u"custom:" + key[1]}
    )
    pre[key] = sn
    return sn


def _direct_ods_replace_cell_at_col(row_elem, col_index, new_cell):
    """
    Заменить логический столбец col_index без полного expand строки.
    При numbercolumnsrepeated — split только затронутого блока.
    """
    from odf.table import TableCell
    from odf.text import P

    if row_elem is None or new_cell is None or int(col_index) < 0:
        return False
    col_index = int(col_index)
    col = 0
    for cell in list(row_elem.getElementsByType(TableCell)):
        try:
            rep = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        except Exception:
            rep = 1
        if rep < 1:
            rep = 1
        if col_index < col + rep:
            offset = col_index - col
            if rep == 1:
                try:
                    row_elem.insertBefore(new_cell, cell)
                    row_elem.removeChild(cell)
                except Exception:
                    return False
                return True
            try:
                cell.removeAttribute("numbercolumnsrepeated")
            except Exception:
                pass
            parts = []
            if offset > 0:
                left = _direct_ods_clone_empty_like(cell)
                if offset > 1:
                    left.setAttribute("numbercolumnsrepeated", str(offset))
                parts.append(left)
            parts.append(new_cell)
            right_count = rep - offset - 1
            if right_count > 0:
                right = _direct_ods_clone_empty_like(cell)
                if right_count > 1:
                    right.setAttribute("numbercolumnsrepeated", str(right_count))
                parts.append(right)
            try:
                for part in reversed(parts):
                    row_elem.insertBefore(part, cell)
                row_elem.removeChild(cell)
            except Exception:
                return False
            return True
        col += rep
    while col < col_index:
        pad = TableCell()
        pad.addElement(P(text=u""))
        try:
            row_elem.addElement(pad)
        except Exception:
            return False
        col += 1
    try:
        row_elem.addElement(new_cell)
    except Exception:
        return False
    return True


def _direct_ods_convert_row_numbers(row_elem, num_cols, end_col0, style_cache=None):
    """Заменить текстовые числа в num_cols (без полного expand строки)."""
    _unused = end_col0
    if not num_cols or row_elem is None:
        return 0
    row_map = dict(_ods_row_col_values(row_elem))
    n = 0
    for col in sorted(num_cols):
        val = row_map.get(col)
        if not _direct_is_plain_string_for_datetime(val):
            continue
        parsed = _direct_try_parse_number_text(val)
        if parsed is None:
            continue
        new_cell = _ods_make_cell(parsed)
        sn = _direct_ods_prebuilt_cell_style(style_cache, parsed)
        if sn:
            try:
                new_cell.setAttribute(u"stylename", sn)
            except Exception:
                pass
        elif style_cache is not None:
            _direct_ods_apply_converted_cell_style(style_cache, new_cell, parsed)
        if _direct_ods_replace_cell_at_col(row_elem, col, new_cell):
            n += 1
    return n


def _direct_ods_convert_row_datetimes(row_elem, date_cols, end_col0, style_cache=None):
    """Заменить текстовые даты в date_cols (без полного expand строки)."""
    _unused = end_col0
    if not date_cols or row_elem is None:
        return 0
    row_map = dict(_ods_row_col_values(row_elem))
    n = 0
    for col in sorted(date_cols):
        val = row_map.get(col)
        if not _direct_is_plain_string_for_datetime(val):
            continue
        parsed = _direct_try_parse_datetime_text(val)
        if parsed is None:
            continue
        new_cell = _ods_make_cell(parsed)
        sn = _direct_ods_prebuilt_cell_style(style_cache, parsed)
        if sn:
            try:
                new_cell.setAttribute(u"stylename", sn)
            except Exception:
                pass
        elif style_cache is not None:
            _direct_ods_apply_converted_cell_style(style_cache, new_cell, parsed)
        if _direct_ods_replace_cell_at_col(row_elem, col, new_cell):
            n += 1
    return n


def _direct_ods_copy_sheets_convert_datetimes(path, sheet_names, hdr_row0=0):
    _direct_ods_copy_sheets_trim_and_convert(
        path, {}, sheet_names, hdr_row0=hdr_row0, do_convert=True
    )


def _direct_xlsx_copy_sheets_convert_datetimes(path, sheet_names, hdr_row0=0):
    _direct_xlsx_copy_sheets_trim_and_convert(
        path, {}, sheet_names, hdr_row0=hdr_row0, do_convert=True
    )


def direct_apply_copy_sheets_convert_datetimes(path, sheet_names, header_row=0):
    """
    xml_excel_ods + копирование листов: текстовые даты и числа → типы в файле.

    Только при MERGE_XML_CONVERT_TO_NUMBERS=Да. Столбец — дата/время или число,
    если ≥3 строк данных распознаны (даты — MERGE_XML_DATETIME_FORMATS / базовый набор;
    числа — пробел/апостроф между разрядами, дробная , или .).
    Ведущий апостроф снимается перед разбором.
    """
    if not _direct_wants_convert_to_numbers():
        return
    direct_apply_copy_sheets_row_trim_and_convert(
        path, {}, sheet_names, header_row=header_row, convert=True
    )


def _direct_copy_sheet_source_name(sheet_name):
    nm = unicode(sheet_name or u"").strip()
    m = re.match(r"^\d+_(.+)$", nm)
    if m:
        return m.group(1)
    return nm


def _direct_copy_sheet_labels(settings, sheet_name):
    """Метки #Путь и #ДатаВремяФайла для скопированного листа."""
    path_mode = str((settings or {}).get("path_display") or "none").strip().lower()
    datetime_mode = str(
        (settings or {}).get("datetime_display") or "none"
    ).strip().lower()
    if path_mode == "none" and datetime_mode == "none":
        return u"", u""
    file_index_by_name = (settings or {}).get("copy_sheets_file_index_by_name") or {}
    fi = file_index_by_name.get(sheet_name)
    if fi is None:
        try:
            fi = int(re.match(r"^(\d+)_", unicode(sheet_name or u"")).group(1)) - 1
        except Exception:
            fi = 0
    files = (settings or {}).get("files") or []
    flags = (settings or {}).get("is_current_book_flags") or []
    is_curr = bool(flags[fi]) if fi is not None and fi < len(flags) else False
    source_sheet = _direct_copy_sheet_source_name(sheet_name)
    path_label = u""
    mtime_str = u""
    if path_mode != "none":
        if is_curr:
            path_label = _direct_service_path_value(
                {
                    "is_current_book": True,
                    "src_path": u"",
                    "src_sheet_name": source_sheet,
                },
                path_mode,
            )
        elif fi is not None and fi < len(files):
            path_label = _service_path_value(files[fi], path_mode, source_sheet)
    if datetime_mode != "none" and not is_curr and fi is not None and fi < len(files):
        mtime_str = _service_datetime_value(files[fi], datetime_mode)
    return path_label, mtime_str


def _direct_row_has_service_headers(values, path_mode, datetime_mode):
    idx = 0
    if path_mode != "none":
        if idx >= len(values):
            return False
        if _merge_identity_key(unicode(values[idx] or u"")) != _merge_identity_key(
            u"#Путь"
        ):
            return False
        idx += 1
    if datetime_mode != "none":
        if idx >= len(values):
            return False
        if _merge_identity_key(unicode(values[idx] or u"")) != _merge_identity_key(
            u"#ДатаВремяФайла"
        ):
            return False
    return True


def _direct_xlsx_copy_sheets_service_columns(path, sheet_names, settings, header_row=0):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    path_mode = str((settings or {}).get("path_display") or "none").strip().lower()
    datetime_mode = str(
        (settings or {}).get("datetime_display") or "none"
    ).strip().lower()
    n_insert = (1 if path_mode != "none" else 0) + (
        1 if datetime_mode != "none" else 0
    )
    if n_insert <= 0:
        return
    target_keys = set()
    for nm in sheet_names:
        try:
            target_keys.add(_merge_identity_key(nm))
        except Exception:
            target_keys.add(unicode(nm or u"").strip().casefold())
    wb = load_workbook(path)
    changed = False
    try:
        for ws in wb.worksheets:
            nm = unicode(ws.title or u"").strip()
            try:
                if _merge_identity_key(nm) not in target_keys:
                    continue
            except Exception:
                if nm not in sheet_names:
                    continue
            hdr_1 = int(header_row or 0) + 1
            hdr_vals = []
            c = 1
            while c <= min(n_insert + 3, ws.max_column or 1):
                hdr_vals.append(ws.cell(row=hdr_1, column=c).value)
                c += 1
            if _direct_row_has_service_headers(hdr_vals, path_mode, datetime_mode):
                continue
            path_label, mtime_str = _direct_copy_sheet_labels(settings, nm)
            ws.insert_cols(1, n_insert)
            col = 1
            if path_mode != "none":
                ws.cell(row=hdr_1, column=col).value = u"#Путь"
                col += 1
            if datetime_mode != "none":
                ws.cell(row=hdr_1, column=col).value = u"#ДатаВремяФайла"
            r = hdr_1 + 1
            max_row = ws.max_row or r
            while r <= max_row:
                col = 1
                if path_mode != "none":
                    ws.cell(row=r, column=col).value = path_label
                    col += 1
                if datetime_mode != "none":
                    dt_val = mtime_str
                    ws.cell(row=r, column=col).value = dt_val
                    if isinstance(dt_val, DirectServiceDatetime):
                        fmt = _direct_normalize_excel_date_format(
                            _direct_service_datetime_format(dt_val)
                        )
                        ws.cell(row=r, column=col).number_format = fmt
                r += 1
            changed = True
            _log(
                "copy_sheets_service_cols: лист «%s» — добавлено %d"
                % (nm, int(n_insert))
            )
        if changed:
            wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_ods_replace_row_values(row_elem, values):
    from odf.table import TableCell

    for cell in list(row_elem.getElementsByType(TableCell)):
        try:
            row_elem.removeChild(cell)
        except Exception:
            pass
    for val in values:
        row_elem.addElement(_ods_make_cell(val))


def _direct_ods_copy_sheets_service_columns(path, sheet_names, settings, header_row=0):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    path_mode = str((settings or {}).get("path_display") or "none").strip().lower()
    datetime_mode = str(
        (settings or {}).get("datetime_display") or "none"
    ).strip().lower()
    n_insert = (1 if path_mode != "none" else 0) + (
        1 if datetime_mode != "none" else 0
    )
    if n_insert <= 0:
        return
    target_keys = set()
    for nm in sheet_names:
        try:
            target_keys.add(_merge_identity_key(nm))
        except Exception:
            target_keys.add(unicode(nm or u"").strip().casefold())
    doc = load(path)
    changed = False
    hdr_row0 = int(header_row or 0)
    for table in doc.spreadsheet.getElementsByType(Table):
        nm = unicode(table.getAttribute("name") or u"").strip()
        try:
            if _merge_identity_key(nm) not in target_keys:
                continue
        except Exception:
            if nm not in sheet_names:
                continue
        rows = table.getElementsByType(TableRow)
        if hdr_row0 < 0 or hdr_row0 >= len(rows):
            continue
        hdr_map = dict(_ods_row_col_values(rows[hdr_row0]))
        hdr_vals = [hdr_map.get(i) for i in range(max(3, n_insert))]
        if _direct_row_has_service_headers(hdr_vals, path_mode, datetime_mode):
            continue
        path_label, mtime_str = _direct_copy_sheet_labels(settings, nm)
        end_col, end_row = _direct_scan_bounds_ods(rows, hdr_row0, 0)
        rr = 0
        while rr <= end_row and rr < len(rows):
            row_map = dict(_ods_row_col_values(rows[rr]))
            old_vals = []
            c = 0
            while c <= end_col:
                old_vals.append(row_map.get(c))
                c += 1
            new_vals = []
            if rr == hdr_row0:
                if path_mode != "none":
                    new_vals.append(u"#Путь")
                if datetime_mode != "none":
                    new_vals.append(u"#ДатаВремяФайла")
            else:
                if path_mode != "none":
                    new_vals.append(path_label)
                if datetime_mode != "none":
                    new_vals.append(mtime_str)
            new_vals.extend(old_vals)
            _direct_ods_replace_row_values(rows[rr], new_vals)
            rr += 1
        changed = True
        _log(
            "copy_sheets_service_cols: лист «%s» — добавлено %d"
            % (nm, int(n_insert))
        )
    if changed:
        doc.save(path)


def direct_apply_copy_sheets_service_columns( path, sheet_names, settings, header_row=0 ):
    """
    xml_excel_ods + копирование листов: вставить #Путь / #ДатаВремяФайла слева.

    Не используется: в режиме «Копирование листов» служебные колонки отключены.
    Оставлено для совместимости вызовов; no-op.
    """
    return


def direct_apply_delete_top_rows(path, delete_map, skipped_top_by_sheet=None):
    """Удалить верхние строки на листах файла результата (после прямого переноса)."""
    if not delete_map:
        return
    if _is_ods(path):
        _direct_ods_delete_top_rows(path, delete_map, skipped_top_by_sheet)
    elif _is_xlsx_target(path):
        _direct_xlsx_delete_top_rows(path, delete_map, skipped_top_by_sheet)


# --- постобработка_xml (фаза 2): манипуляции с файлом до reopen UNO ---


def _direct_pp_fn_key(name):
    try:
        from libre_macros_param_codec import normalize_fn_key

        return normalize_fn_key(name)
    except Exception:
        return unicode(name or u"").strip().casefold()


def _direct_pp_is_apply_formula(name):
    return _direct_pp_fn_key(name) == u"применить_формулу"


def _direct_pp_is_freeze_header(name):
    return _direct_pp_fn_key(name) == u"закрепить_заголовок"


def _direct_pp_is_skip_empty_rows(name):
    # normalize_fn_key: пропуск_пустых_строк → удаление_строк
    return _direct_pp_fn_key(name) == u"удаление_строк"


def _direct_pp_is_header_plus_height(name):
    return _direct_pp_fn_key(name) == u"заголовок_плюс_высота"


def _direct_pp_is_file_visual(name):
    """Визуальные шаги, которые применяются в файле (без UNO)."""
    return _direct_pp_fn_key(name) in _DIRECT_XML_FILE_VISUAL_KEYS


def _direct_pp_is_file_structural(name):
    """Структурные шаги, которые применяются в файле (без UNO)."""
    return _direct_pp_fn_key(name) in _DIRECT_XML_FILE_STRUCTURAL_KEYS


def _direct_pp_file_structural_applies(fn_key, extra_raw, sheet_name):
    """Нужно ли применять file-structural на данном листе (как file-visual)."""
    key = _direct_pp_fn_key(fn_key)
    if key == u"сводная_таблица":
        return _direct_pp_pivot_source_matches(extra_raw, sheet_name)
    return _direct_pp_file_visual_applies(fn_key, extra_raw, sheet_name)


def _direct_pp_is_uno_pilot(name):
    """Шаги → очередь UNO после reopen (autofit / заголовок_плюс_высота / …)."""
    return _direct_pp_fn_key(name) in _DIRECT_XML_UNO_PILOT_KEYS


def _direct_pp_sheet_list_applies(fn_key, extra_raw, sheet_name):
    """Фильтр sheets для sheet-list fn (тонкая_сетка, авто_*, …)."""
    raw = unicode(extra_raw or u"").strip()
    if raw == u"":
        return True
    try:
        from libre_macros_param_codec import param_decode
    except Exception:
        return True
    try:
        blocks = param_decode(fn_key, raw)
    except Exception:
        blocks = []
    if not blocks:
        return True
    block = blocks[0] if isinstance(blocks[0], dict) else {}
    names = block.get("sheets")
    if not names:
        return True
    try:
        names_set = set(
            unicode(x or u"").strip() for x in names if unicode(x or u"").strip() != u""
        )
    except Exception:
        names_set = set()
    if not names_set:
        return True
    if _direct_pp_sheet_name_matches(sheet_name, names_set):
        return True
    m = re.match(r"^(\d+)_(.+)$", unicode(sheet_name or u"").strip())
    if m is not None and _direct_pp_sheet_name_matches(m.group(2), names_set):
        return True
    return False


def _direct_pp_file_visual_applies(fn_key, extra_raw, sheet_name):
    """Нужно ли применять file-visual / UNO-пилот на данном листе."""
    key = _direct_pp_fn_key(fn_key)
    raw = unicode(extra_raw or u"").strip()
    if key in _DIRECT_XML_SHEET_LIST_KEYS:
        return _direct_pp_sheet_list_applies(key, raw, sheet_name)
    # sheet-block: пустой C — все листы; иначе нужен блок для листа (или fallback без sheet)
    if raw == u"":
        return True
    block = _direct_pp_pick_sheet_block(key, raw, sheet_name)
    return block is not None


def _direct_pp_uno_pilot_applies(fn_key, extra_raw, sheet_name):
    """Нужно ли ставить шаг в очередь UNO для данного листа."""
    return _direct_pp_file_visual_applies(fn_key, extra_raw, sheet_name)


def _direct_pp_pick_sheet_block(fn_key, extra_raw, sheet_name):
    """Первый JSON-блок fn_key для листа (поле sheet) или fallback без sheet."""
    blocks = _direct_pp_blocks_for_sheet(fn_key, extra_raw, sheet_name)
    if blocks is None:
        return None
    if len(blocks) == 0:
        raw = unicode(extra_raw or u"").strip()
        if raw == u"":
            return {}
        return None
    return blocks[0]


def _direct_pp_blocks_for_sheet(fn_key, extra_raw, sheet_name):
    """
    JSON-блоки fn_key только для текущего листа.
    Конкретные sheet/sheets приоритетнее блоков без листа.
    Пустой C → [{}]; JSON без блока для листа → [].
    None — ошибка decode/import (как прежний pick → None на пустом списке блоков).
    """
    raw = unicode(extra_raw or u"").strip()
    if raw == u"":
        return [{}]
    try:
        from libre_macros_param_codec import param_decode
        from libre_macros_lib import _lm_pp_parse_sheet_names_from_spec
    except Exception:
        return None
    try:
        blocks = param_decode(fn_key, raw)
    except Exception:
        blocks = []
    if not blocks:
        return None
    key = _direct_pp_fn_key(fn_key)
    sheets_is_source_list = key in _DIRECT_PP_SHEETS_IS_SOURCE_LIST_KEYS
    specific = []
    global_blocks = []
    bi = 0
    while bi < len(blocks):
        block = blocks[bi]
        bi += 1
        if not isinstance(block, dict):
            continue
        sheet_ref = unicode(block.get("sheet") or u"").strip()
        sheets_ref = block.get("sheets")
        sheet_names = []
        if sheets_ref is not None and not sheets_is_source_list:
            if isinstance(sheets_ref, (list, tuple)):
                sheet_names = [
                    unicode(x or u"").strip()
                    for x in sheets_ref
                    if unicode(x or u"").strip() != u""
                ]
            else:
                one = unicode(sheets_ref or u"").strip()
                sheet_names = [one] if one else []
        if sheet_ref == u"" and not sheet_names:
            global_blocks.append(block)
            continue
        matched = False
        if sheet_names:
            if _direct_pp_sheet_name_matches(sheet_name, sheet_names):
                matched = True
            else:
                m = re.match(r"^(\d+)_(.+)$", unicode(sheet_name or u"").strip())
                if m is not None and _direct_pp_sheet_name_matches(
                    m.group(2), sheet_names
                ):
                    matched = True
        else:
            try:
                names = _lm_pp_parse_sheet_names_from_spec(sheet_ref)
            except Exception:
                names = [sheet_ref]
            if _direct_pp_sheet_name_matches(sheet_name, names):
                matched = True
            else:
                m = re.match(r"^(\d+)_(.+)$", unicode(sheet_name or u"").strip())
                if m is not None and _direct_pp_sheet_name_matches(m.group(2), names):
                    matched = True
        if matched:
            specific.append(block)
    if len(specific) > 0:
        return specific
    return global_blocks


def _direct_pp_resolve_col_tokens(tokens, header_h2i, end_col0):
    """Токены столбцов (1-based / имя / буквы / 'точное') → список 0-based индексов.
    Заголовок важнее букв: иначе GUID/ID читаются как индекс столбца.
    """
    out = []
    seen = set()
    unwrap = None
    clean_fn = None
    try:
        from libre_macros_param_codec import unwrap_column_name_token as unwrap
        from libre_macros_param_codec import _clean_column_match_text as clean_fn
    except Exception:
        unwrap = None
        clean_fn = None
    for tok in tokens or []:
        t_raw = unicode(tok or u"").strip()
        if t_raw == u"":
            continue
        exact = False
        t = t_raw
        if unwrap is not None:
            try:
                t, exact = unwrap(t_raw)
            except Exception:
                t, exact = t_raw, False
        if clean_fn is not None:
            try:
                t = clean_fn(t)
            except Exception:
                pass
        t = unicode(t or u"").strip()
        if t == u"":
            continue
        idx = None
        if (not exact) and t.isdigit():
            idx = int(t) - 1
        else:
            idx = header_h2i.get(_merge_identity_key(t))
            if idx is None and not exact:
                needle = _merge_identity_key(t)
                for hk, hi in header_h2i.items():
                    if needle in hk or hk in needle:
                        idx = hi
                        break
            if idx is None and (not exact) and _is_col_letters(t):
                idx = _col_letters_to_index(t)
        if idx is None:
            continue
        idx = int(idx)
        if idx < 0 or (end_col0 >= 0 and idx > end_col0):
            continue
        if idx not in seen:
            seen.add(idx)
            out.append(idx)
    return out


def _direct_pp_resolve_col_matches(token, header_h2i, end_col0, titles=None):
    """Все 0-based индексы столбцов по токену (слева направо).
    Заголовок важнее букв столбца. Токен в '…' — только полное совпадение.
    """
    t_raw = unicode(token or u"").strip()
    if t_raw == u"":
        return []
    exact = False
    t = t_raw
    match_fn = None
    try:
        from libre_macros_param_codec import unwrap_column_name_token
        from libre_macros_param_codec import _clean_column_match_text
        from libre_macros_param_codec import column_header_matches_token as match_fn
        t, exact = unwrap_column_name_token(t_raw)
        t = _clean_column_match_text(t)
    except Exception:
        t, exact = t_raw, False
        try:
            from libre_macros_param_codec import _clean_column_match_text
            t = _clean_column_match_text(t)
        except Exception:
            pass
    t = unicode(t or u"").strip()
    if t == u"":
        return []
    if (not exact) and t.isdigit():
        idx = int(t) - 1
        if 0 <= idx <= end_col0:
            return [idx]
        return []
    if (not exact) and u"*" in t and titles is not None:
        out = []
        ci = 0
        while ci < len(titles):
            if _direct_pp_header_matches_marker(titles[ci], t):
                if ci not in out:
                    out.append(ci)
            ci += 1
        return out
    if exact and titles is not None and match_fn is not None:
        out = []
        ci = 0
        while ci < len(titles) and ci <= end_col0:
            try:
                ok = match_fn(titles[ci], t_raw)
            except Exception:
                ok = False
            if ok and ci not in out:
                out.append(ci)
            ci += 1
        if out:
            return out
    idx = header_h2i.get(_merge_identity_key(t))
    if idx is not None and 0 <= int(idx) <= end_col0:
        return [int(idx)]
    if exact:
        return []
    needle = _merge_identity_key(t)
    out = []
    for hk, hi in header_h2i.items():
        if needle in hk or hk in needle:
            hi = int(hi)
            if hi not in out and hi <= end_col0:
                out.append(hi)
    if out:
        out.sort()
        return out
    if _is_col_letters(t):
        idx = _col_letters_to_index(t)
        if 0 <= idx <= end_col0:
            return [idx]
        return []
    return []


def _direct_pp_resolve_col_unique(token, header_h2i, end_col0, titles=None):
    t = unicode(token or u"").strip()
    if t == u"":
        return None, u"пустой токен"
    if u"*" not in t:
        one = _direct_pp_resolve_col_tokens([t], header_h2i, end_col0)
        if one:
            return one[0], None
    matches = _direct_pp_resolve_col_matches(token, header_h2i, end_col0, titles)
    label = t
    try:
        from libre_macros_param_codec import unwrap_column_name_token
        from libre_macros_param_codec import _clean_column_match_text
        bare, _q = unwrap_column_name_token(t)
        bare = _clean_column_match_text(bare)
        if bare != u"":
            label = bare
    except Exception:
        pass
    if not matches:
        return None, u"столбец «%s» не найден" % label
    if len(matches) > 1:
        return None, u"столбец «%s» не уникален (%d совпадений)" % (
            label,
            len(matches),
        )
    return matches[0], None


def _direct_pp_resolve_move_cols(tokens, header_h2i, end_col0, titles=None):
    out = []
    seen = set()
    for tok in tokens or []:
        if isinstance(tok, int):
            idx = int(tok) - 1 if int(tok) > 0 else int(tok)
            if 0 <= idx <= end_col0 and idx not in seen:
                seen.add(idx)
                out.append(idx)
            continue
        t = unicode(tok or u"").strip()
        if t == u"":
            continue
        for idx in _direct_pp_resolve_col_matches(t, header_h2i, end_col0, titles):
            if idx not in seen:
                seen.add(idx)
                out.append(idx)
    return out


def _direct_pp_normalize_reorder_position(raw):
    try:
        from libre_macros_param_codec import _normalize_reorder_position

        return _normalize_reorder_position(raw)
    except Exception:
        key = unicode(raw or u"").strip().casefold()
        if key in (u"start", u"end", u"before", u"after"):
            return key
        return u""


def _direct_pp_build_column_reorder(all_cols, moving_cols, position, relative_col):
    moving_set = set(moving_cols)
    rest = [c for c in all_cols if c not in moving_set]
    if position == u"start":
        insert_at = 0
    elif position == u"end":
        insert_at = len(rest)
    elif position in (u"before", u"after"):
        if relative_col is None:
            return None, u"не задан опорный столбец"
        if relative_col not in rest:
            return None, u"опорный столбец не в оставшихся"
        insert_at = rest.index(relative_col)
        if position == u"after":
            insert_at += 1
    else:
        return None, u"неизвестная position: %s" % position
    return rest[:insert_at] + list(moving_cols) + rest[insert_at:], None


def _direct_pp_build_column_copy(all_cols, moving_cols, position, relative_col):
    rest = list(all_cols)
    if position == u"start":
        insert_at = 0
    elif position == u"end":
        insert_at = len(rest)
    elif position in (u"before", u"after"):
        if relative_col is None:
            return None, None, u"не задан опорный столбец"
        if relative_col not in rest:
            return None, None, u"опорный столбец не в диапазоне"
        insert_at = rest.index(relative_col)
        if position == u"after":
            insert_at += 1
    else:
        return None, None, u"неизвестная position: %s" % position
    return rest[:insert_at] + list(moving_cols) + rest[insert_at:], insert_at, None


def _direct_pp_block_is_create_new_columns(block):
    """True — создать пустые столбцы по new_names (флаг «новые»)."""
    if block is None:
        return False
    if u"create_new" in block:
        try:
            return bool(block.get(u"create_new"))
        except Exception:
            return False
    for key in (u"новые", u"new_empty", u"blank_columns", u"create_empty"):
        if key in block:
            try:
                return bool(block.get(key))
            except Exception:
                return False
    return False


def _direct_pp_block_is_copy_columns(block):
    if block is None:
        return False
    if _direct_pp_block_is_create_new_columns(block):
        return False
    if u"copy" in block:
        try:
            return bool(block.get(u"copy"))
        except Exception:
            pass
    mode = unicode(block.get(u"mode") or u"").strip().casefold()
    if mode in (
        u"copy",
        u"копировать",
        u"_копировать_",
        u"дублировать",
        u"duplicate",
    ):
        return True
    return False


def _direct_pp_resolve_insert_at(all_cols, position, relative_col):
    """Индекс вставки (0-based) для copy/create_new."""
    rest = list(all_cols)
    if position == u"start":
        return 0, None
    if position == u"end":
        return len(rest), None
    if position in (u"before", u"after"):
        if relative_col is None:
            return None, u"не задан опорный столбец"
        if relative_col not in rest:
            return None, u"опорный столбец не в диапазоне"
        insert_at = rest.index(relative_col)
        if position == u"after":
            insert_at += 1
        return insert_at, None
    return None, u"неизвестная position: %s" % position


def _direct_pp_xlsx_insert_empty_columns(ws, insert_at, names, hdr_row_1):
    """Вставить пустые столбцы openpyxl и задать заголовки."""
    if ws is None or not names or insert_at is None:
        return False
    n = len(names)
    if n <= 0:
        return False
    try:
        ws.insert_cols(int(insert_at) + 1, amount=n)
    except Exception:
        return False
    i = 0
    while i < n:
        title = unicode(names[i] or u"").strip()
        if title == u"":
            title = u"Столбец_%d" % (i + 1)
        try:
            ws.cell(row=int(hdr_row_1), column=int(insert_at) + i + 1).value = title
        except Exception:
            pass
        i += 1
    return True


def _direct_pp_parse_copy_new_names(block):
    raw = None
    if block is not None:
        raw = block.get(u"new_names")
        if raw is None:
            raw = block.get(u"copy_names")
        if raw is None:
            raw = block.get(u"new_columns")
    if raw is None or raw == u"":
        return []
    if isinstance(raw, (list, tuple)):
        return [unicode(x).strip() for x in raw]
    text = unicode(raw).strip()
    if text == u"":
        return []
    return [p.strip() for p in re.split(r"[,;]", text)]


def _direct_pp_unique_header_name(base, existing_cf):
    base = unicode(base or u"").strip()
    if base == u"":
        base = u"_Столбец"
    candidate = base
    n = 2
    while _merge_identity_key(candidate) in existing_cf:
        candidate = u"%s_%d" % (base, n)
        n += 1
        if n > 200:
            break
    return candidate


def _direct_pp_apply_copied_headers_xlsx( ws, hdr_row_1, insert_at, moving, new_names, src_titles, end_col_new ):
    if insert_at is None or not moving:
        return
    existing = set()
    c = 0
    while c <= int(end_col_new):
        try:
            val = ws.cell(row=int(hdr_row_1), column=c + 1).value
        except Exception:
            val = None
        t = unicode(val or u"").strip()
        if t != u"":
            existing.add(_merge_identity_key(t))
        c += 1
    names = list(new_names or [])
    titles = list(src_titles or [])
    i = 0
    while i < len(moving):
        desired = names[i] if i < len(names) else u""
        desired = unicode(desired or u"").strip()
        if desired == u"":
            base = titles[i] if i < len(titles) else u""
            base = unicode(base or u"").strip()
            if base == u"":
                base = u"Столбец_%d" % (moving[i] + 1)
            desired = base
        # временно убрать текущий заголовок копии из множества, чтобы авто-имя
        # не конфликтовало само с собой
        col_1 = insert_at + i + 1
        try:
            cur = unicode(
                ws.cell(row=int(hdr_row_1), column=col_1).value or u""
            ).strip()
        except Exception:
            cur = u""
        if cur != u"":
            existing.discard(_merge_identity_key(cur))
        name = _direct_pp_unique_header_name(desired, existing)
        ws.cell(row=int(hdr_row_1), column=col_1).value = name
        existing.add(_merge_identity_key(name))
        i += 1


def _direct_pp_apply_copied_headers_ods( rows, hdr_row0, insert_at, moving, new_names, src_titles, end_col_new ):
    if insert_at is None or not moving:
        return
    existing = set()
    titles_now = _direct_pp_ods_header_titles_list(rows, hdr_row0, end_col_new)
    for t in titles_now:
        tt = unicode(t or u"").strip()
        if tt != u"":
            existing.add(_merge_identity_key(tt))
    names = list(new_names or [])
    titles = list(src_titles or [])
    while hdr_row0 >= len(rows):
        from odf.table import TableRow

        rows.append(TableRow())
    i = 0
    while i < len(moving):
        desired = names[i] if i < len(names) else u""
        desired = unicode(desired or u"").strip()
        if desired == u"":
            base = titles[i] if i < len(titles) else u""
            base = unicode(base or u"").strip()
            if base == u"":
                base = u"Столбец_%d" % (moving[i] + 1)
            desired = base
        col = insert_at + i
        if col < len(titles_now):
            cur = unicode(titles_now[col] or u"").strip()
            if cur != u"":
                existing.discard(_merge_identity_key(cur))
        name = _direct_pp_unique_header_name(desired, existing)
        _direct_ods_set_cell_at_col(rows[hdr_row0], col, _ods_make_cell(name))
        existing.add(_merge_identity_key(name))
        i += 1


def _direct_pp_plan_column_reorder(block, header_h2i, end_col0, titles=None):
    """
    План перестановки/копии/новых пустых столбцов.

    Возвращает (new_order, insert_at, moving, err, create_names).
    create_names — список заголовков для create_new; иначе None.
    insert_at — для copy/create_new; иначе None.
    """
    sheet_filter = unicode(block.get("sheet") or u"").strip()
    if sheet_filter == u"" or _merge_identity_key(sheet_filter) == u"все":
        return None, None, None, u"нужен конкретный лист в sheet", None
    tokens = block.get("columns") or []
    position = _direct_pp_normalize_reorder_position(block.get("position"))
    if position == u"":
        return None, None, None, u"не задана position", None
    do_create = _direct_pp_block_is_create_new_columns(block)
    do_copy = _direct_pp_block_is_copy_columns(block)
    relative_col = None
    rel_token = unicode(block.get("relative_column") or u"").strip()
    if position in (u"before", u"after"):
        if rel_token == u"":
            return None, None, None, u"для before/after нужен relative_column", None
        relative_col, rel_err = _direct_pp_resolve_col_unique(
            rel_token, header_h2i, end_col0, titles
        )
        if rel_err:
            return None, None, None, rel_err, None
    all_cols = list(range(end_col0 + 1))
    if do_create:
        names = _direct_pp_parse_copy_new_names(block)
        if not names:
            return None, None, None, u"для «новые» нужны имена в new_names", None
        insert_at, ierr = _direct_pp_resolve_insert_at(
            all_cols, position, relative_col
        )
        if ierr:
            return None, None, None, ierr, None
        # new_order: -1 = пустой новый столбец (для ODS rebuild)
        new_order = list(all_cols)
        i = 0
        while i < len(names):
            new_order.insert(int(insert_at) + i, -1)
            i += 1
        return new_order, insert_at, [], None, names
    if not tokens:
        return None, None, None, u"не заданы columns", None
    moving = _direct_pp_resolve_move_cols(tokens, header_h2i, end_col0, titles)
    if not moving:
        return None, None, None, u"столбцы для переноса не найдены", None
    if position in (u"before", u"after"):
        if (not do_copy) and relative_col in moving:
            moving = [c for c in moving if c != relative_col]
        if not moving:
            return None, None, None, u"после исключения опорного не осталось столбцов", None
    if do_copy:
        new_order, insert_at, err = _direct_pp_build_column_copy(
            all_cols, moving, position, relative_col
        )
        return new_order, insert_at, moving, err, None
    new_order, err = _direct_pp_build_column_reorder(
        all_cols, moving, position, relative_col
    )
    return new_order, None, moving, err, None


def _direct_pp_xlsx_copy_cell_style(src_cell, dst_cell):
    """Скопировать стиль ячейки openpyxl (best-effort)."""
    if src_cell is None or dst_cell is None:
        return
    try:
        from copy import copy as _copy
    except Exception:
        _copy = None
    try:
        dst_cell.number_format = src_cell.number_format
    except Exception:
        pass
    if _copy is None:
        return
    for attr in ("font", "border", "fill", "protection", "alignment"):
        try:
            setattr(dst_cell, attr, _copy(getattr(src_cell, attr)))
        except Exception:
            pass


def _direct_pp_xlsx_copy_columns_insert( ws, end_row_1, end_col0, moving, insert_at, hdr_row_1 ):
    """
    Копия столбцов: insert_cols + значение/стиль/ширина (без переписывания всего листа).
    moving — 0-based индексы ДО вставки; insert_at — 0-based позиция вставки.
    """
    if ws is None or not moving or insert_at is None:
        return False
    n = len(moving)
    if n <= 0:
        return False
    try:
        from openpyxl.utils import get_column_letter
    except Exception:
        get_column_letter = None
    # openpyxl: 1-based
    try:
        ws.insert_cols(int(insert_at) + 1, amount=n)
    except Exception:
        return False

    def _map_src(old):
        if int(old) >= int(insert_at):
            return int(old) + n
        return int(old)

    max_row = max(int(end_row_1), int(ws.max_row or end_row_1))
    i = 0
    while i < n:
        src0 = _map_src(moving[i])
        dst0 = int(insert_at) + i
        src_1 = src0 + 1
        dst_1 = dst0 + 1
        r = 1
        while r <= max_row:
            try:
                src_cell = ws.cell(row=r, column=src_1)
                dst_cell = ws.cell(row=r, column=dst_1)
                dst_cell.value = src_cell.value
                _direct_pp_xlsx_copy_cell_style(src_cell, dst_cell)
            except Exception:
                pass
            r += 1
        if get_column_letter is not None:
            try:
                src_letter = get_column_letter(src_1)
                dst_letter = get_column_letter(dst_1)
                src_w = ws.column_dimensions[src_letter].width
                if src_w is not None:
                    ws.column_dimensions[dst_letter].width = src_w
            except Exception:
                pass
        i += 1
    return True


def _direct_pp_xlsx_apply_column_order(ws, end_row_1, end_col0, new_order):
    if not new_order:
        return False
    ncol = end_col0 + 1
    ncol_new = len(new_order)
    if ncol_new != ncol:
        return False
    if new_order == list(range(ncol)):
        return True
    max_row = max(end_row_1, ws.max_row or end_row_1)
    r = 1
    while r <= max_row:
        row_vals = []
        c = 0
        while c < ncol:
            row_vals.append(ws.cell(row=r, column=c + 1).value)
            c += 1
        nc = 0
        while nc < len(new_order):
            old = new_order[nc]
            ws.cell(row=r, column=nc + 1).value = (
                row_vals[old] if old < len(row_vals) else None
            )
            nc += 1
        r += 1
    return True


def _direct_pp_ods_apply_column_order(rows, sheet, end_row0, end_col0, new_order):
    if not new_order:
        return False
    ncol = end_col0 + 1
    ncol_new = len(new_order)
    if ncol_new < ncol:
        return False
    if ncol_new == ncol and new_order == list(range(ncol)):
        return True
    r = 0
    while r <= end_row0:
        while r >= len(rows):
            from odf.table import TableRow

            rows.append(TableRow())
            sheet.addElement(rows[-1])
        by_col = {}
        for cidx, val in _ods_row_col_values(rows[r]):
            by_col[int(cidx)] = val
        _direct_ods_expand_row_to_cols(rows[r], max(end_col0, ncol_new - 1))
        nc = 0
        while nc < len(new_order):
            old = new_order[nc]
            val = by_col.get(old)
            _direct_ods_set_cell_at_col(
                rows[r], nc, _ods_make_cell(u"" if val is None else unicode(val))
            )
            nc += 1
        r += 1
    return True


def _direct_pp_ods_reorder_columns_rebuild( doc, sheet, sheet_name, end_row0, end_col0, new_order ):
    """
    Быстрый путь: новый Table с ячейками в new_order, замена старого листа.
    Без expand_row/set_cell_at_col на каждую ячейку (медленно на тысячах строк).
    new_order может быть длиннее end_col0+1 (режим копирования).
    """
    from odf.table import Table, TableRow

    if not new_order:
        return False
    ncol = int(end_col0) + 1
    ncol_new = len(new_order)
    if ncol_new < ncol:
        return False
    if ncol_new == ncol and new_order == list(range(ncol)):
        return True
    rows = sheet.getElementsByType(TableRow)
    new_table = Table(name=unicode(sheet_name))
    t0 = _direct_now()
    r = 0
    while r <= int(end_row0):
        if r % 200 == 0:
            _direct_ui_yield()
        if r % 1000 == 0 and r > 0:
            _direct_trace(u"переставить_столбцы/ods", u"rebuild rows=%d" % r)
        by_col = {}
        if r < len(rows):
            for cidx, val in _ods_row_col_values(rows[r]):
                by_col[int(cidx)] = val
        nr = TableRow()
        nc = 0
        while nc < len(new_order):
            src_i = new_order[nc]
            try:
                src_i = int(src_i)
            except (TypeError, ValueError):
                src_i = -1
            if src_i < 0:
                val = u""
            else:
                val = by_col.get(src_i)
            nr.addElement(_ods_make_cell(val if val is not None else u""))
            nc += 1
        new_table.addElement(nr)
        r += 1
    try:
        doc.spreadsheet.removeChild(sheet)
    except Exception:
        pass
    doc.spreadsheet.addElement(new_table)
    _direct_trace(
        u"переставить_столбцы/ods",
        u"rebuild %d строк за %.1fs" % (r, _direct_now() - t0),
    )
    return True


def _direct_pp_color_to_rgb_hex(token):
    """Цвет из параметров → 'RRGGBB' или None."""
    raw = unicode(token or u"").strip()
    if raw == u"":
        return None
    try:
        from libre_macros_lib import _lm_pp_zebra_resolve_color_token

        rgb_int = _lm_pp_zebra_resolve_color_token(raw)
        if rgb_int is None:
            return None
        return u"%06X" % (int(rgb_int) & 0xFFFFFF)
    except Exception:
        pass
    if raw.startswith(u"#") and len(raw) == 7:
        return raw[1:].upper()
    return None


def _direct_pp_skip_empty_rows_xlsx(path, sheet_name, extra_raw):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    block = _direct_pp_pick_sheet_block(u"удаление_строк", extra_raw, sheet_name)
    if block is None:
        return False, u"лист не в фильтре sheet"
    mode = unicode(block.get("mode") or u"columns").strip().casefold()
    if mode == u"formula":
        return (
            False,
            u"режим formula в xml_excel_ods не поддерживается (только columns)",
        )
    if mode in (u"none", u""):
        return False, u"mode=none"
    cols = block.get("columns") or []
    if not cols:
        return False, u"не заданы columns"
    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    wb = load_workbook(path)
    try:
        if sheet_name not in wb.sheetnames:
            return False, u"лист не найден"
        ws = wb[sheet_name]
        hdr_row_1 = int(bounds["hdr_row"]) + 1
        header_h2i = _direct_xlsx_header_h2i(ws, hdr_row_1=hdr_row_1)
        end_col0 = int(bounds["end_col"])
        col_indices = _direct_pp_resolve_col_tokens(cols, header_h2i, end_col0)
        if not col_indices:
            return False, u"столбцы не найдены: %s" % u",".join(unicode(c) for c in cols)
        sr = int(bounds["first_data"]) + 1  # 1-based
        er = int(bounds["last_data"]) + 1
        deleted = 0
        rr = er
        while rr >= sr:
            all_empty = True
            for ci in col_indices:
                v = ws.cell(row=rr, column=ci + 1).value
                if not _direct_value_empty(v):
                    all_empty = False
                    break
            if all_empty:
                ws.delete_rows(rr, 1)
                deleted += 1
            rr -= 1
        if deleted > 0:
            wb.save(path)
        return True, u"удалено строк=%d (columns)" % deleted
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_pp_skip_empty_rows_ods(path, sheet_name, extra_raw):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    block = _direct_pp_pick_sheet_block(u"удаление_строк", extra_raw, sheet_name)
    if block is None:
        return False, u"лист не в фильтре sheet"
    mode = unicode(block.get("mode") or u"columns").strip().casefold()
    if mode == u"formula":
        return (
            False,
            u"режим formula в xml_excel_ods не поддерживается (только columns)",
        )
    if mode in (u"none", u""):
        return False, u"mode=none"
    cols = block.get("columns") or []
    if not cols:
        return False, u"не заданы columns"
    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    doc = load(path)
    sheet = None
    for t in doc.spreadsheet.getElementsByType(Table):
        if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
            sheet = t
            break
    if sheet is None:
        return False, u"лист не найден"
    rows = list(sheet.getElementsByType(TableRow))
    header_h2i = _direct_ods_header_h2i(rows, hdr_row=int(bounds["hdr_row"]))
    end_col0 = int(bounds["end_col"])
    col_indices = _direct_pp_resolve_col_tokens(cols, header_h2i, end_col0)
    if not col_indices:
        return False, u"столбцы не найдены: %s" % u",".join(unicode(c) for c in cols)
    sr = int(bounds["first_data"])
    er = int(bounds["last_data"])
    deleted = 0
    rr = er
    while rr >= sr:
        if rr >= len(rows):
            rr -= 1
            continue
        by_col = {}
        for col_idx, val in _ods_row_col_values(rows[rr]):
            by_col[int(col_idx)] = val
        all_empty = True
        for ci in col_indices:
            if not _direct_value_empty(by_col.get(ci)):
                all_empty = False
                break
        if all_empty:
            try:
                sheet.removeChild(rows[rr])
            except Exception:
                try:
                    sheet.childNodes.remove(rows[rr])
                except Exception:
                    pass
            del rows[rr]
            deleted += 1
        rr -= 1
    if deleted > 0:
        doc.save(path)
    return True, u"удалено строк=%d (columns)" % deleted


def _direct_pp_header_height_cfg(block):
    """Блок кодека → cfg для оформления заголовка в direct."""
    if not isinstance(block, dict):
        block = {}
    try:
        height_mm = float(block.get("height_mm"))
    except Exception:
        height_mm = 12.7
    return {
        u"height_mm": height_mm,
        u"h_align": unicode(block.get("h_align") or u"center").strip().lower() or u"center",
        u"v_align": unicode(block.get("v_align") or u"center").strip().lower() or u"center",
        u"bold": bool(block.get("bold")),
        u"font_name": unicode(block.get("font_name") or u"").strip(),
        u"font_size": block.get("font_size"),
        u"fill_color": unicode(block.get("fill_color") or u"").strip(),
        u"font_color": unicode(block.get("font_color") or u"").strip(),
        u"font_color_auto": bool(block.get("font_color_auto", True)),
    }


def _direct_pp_header_plus_height_xlsx(path, sheet_name, extra_raw):
    """
    xlsx: в файл высоту не пишем — UNO после reopen делает pad один раз.
    (Раньше openpyxl ставил абсолютную высоту, затем UNO добавлял pad → ×2.)
    Только проверка фильтра sheet.
    """
    _unused = path
    block = _direct_pp_pick_sheet_block(u"заголовок_плюс_высота", extra_raw, sheet_name)
    if block is None and unicode(extra_raw or u"").strip() != u"":
        return False, u"лист не в фильтре sheet"
    cfg = _direct_pp_header_height_cfg(block if block is not None else {})
    return True, u"в очереди на UNO (≈%.1f мм)" % float(cfg[u"height_mm"])


def _direct_pp_header_plus_height_ods(path, sheet_name, extra_raw):
    """
    ODS: высота/стили заголовка через odfpy ограничены.
    Ставим в очередь на UNO после reopen (как freeze).
    """
    _unused = path
    block = _direct_pp_pick_sheet_block(u"заголовок_плюс_высота", extra_raw, sheet_name)
    if block is None and unicode(extra_raw or u"").strip() != u"":
        return False, u"лист не в фильтре sheet"
    cfg = _direct_pp_header_height_cfg(block if block is not None else {})
    return True, u"в очереди на UNO (≈%.1f мм)" % float(cfg[u"height_mm"])


def _direct_pp_sheet_name_matches(sheet_name, spec_names):
    try:
        from libre_macros_lib import _lm_pp_sheet_name_matches_target

        return _lm_pp_sheet_name_matches_target(sheet_name, spec_names)
    except Exception:
        return True


def _direct_pp_pick_formula_block(extra_raw, sheet_name):
    raw = unicode(extra_raw or u"").strip()
    if raw == u"":
        return None
    try:
        from libre_macros_param_codec import param_decode
        from libre_macros_lib import _lm_pp_parse_sheet_names_from_spec

        blocks = param_decode(u"применить_формулу", raw)
    except Exception:
        blocks = []
    if not blocks:
        return None
    fallback = None
    bi = 0
    while bi < len(blocks):
        block = blocks[bi]
        bi += 1
        if not isinstance(block, dict):
            continue
        sheet_ref = unicode(block.get("sheet") or u"").strip()
        if sheet_ref == u"":
            fallback = block
            continue
        names = _lm_pp_parse_sheet_names_from_spec(sheet_ref)
        if _direct_pp_sheet_name_matches(sheet_name, names):
            return block
        m = re.match(r"^(\d+)_(.+)$", unicode(sheet_name or u"").strip())
        if m is not None and _direct_pp_sheet_name_matches(m.group(2), names):
            return block
    return fallback


def _direct_pp_parse_as_values(block):
    if not isinstance(block, dict) or "as_values" not in block:
        return False
    try:
        from libre_macros_lib import lm_parse_bool_param

        parsed = lm_parse_bool_param(block.get("as_values"), default=False)
        return bool(parsed)
    except Exception:
        return bool(block.get("as_values"))


def _direct_pp_formula_spec(extra_raw, sheet_name):
    block = _direct_pp_pick_formula_block(extra_raw, sheet_name)
    if block is None:
        return None
    col_name = unicode(block.get("column") or u"").strip()
    formula = unicode(block.get("formula") or u"").strip()
    if col_name == u"" or formula == u"":
        return None
    # as_values в Постобработка_xml запрещён: материализация через прямой доступ кривая.
    # Даже при as_values=true в JSON всегда формулы в файле (без очереди materialize).
    return col_name, formula, False


def direct_pp_formula_column_for_sheet(extra_raw, sheet_name):
    """Имя целевого столбца «применить_формулу» для листа (или None)."""
    spec = _direct_pp_formula_spec(extra_raw, sheet_name)
    if spec is None:
        return None
    return spec[0]


def _direct_pp_freeze_header_applies(extra_raw, sheet_name):
    """
    Фильтр sheets для «закрепить_заголовок» (как в UNO sheet-list fn).
    Пустой extra/без sheets — применять на всех листах.
    """
    raw = unicode(extra_raw or u"").strip()
    if raw == u"":
        return True
    try:
        from libre_macros_param_codec import param_decode
    except Exception:
        return True
    try:
        blocks = param_decode(u"закрепить_заголовок", raw)
    except Exception:
        blocks = []
    if not blocks:
        return True
    block = blocks[0] if isinstance(blocks[0], dict) else {}
    names = block.get("sheets")
    if not names:
        return True
    try:
        names_set = set(unicode(x or u"").strip() for x in names if unicode(x or u"").strip() != u"")
    except Exception:
        names_set = set()
    if not names_set:
        return True
    if _direct_pp_sheet_name_matches(sheet_name, names_set):
        return True
    m = re.match(r"^(\d+)_(.+)$", unicode(sheet_name or u"").strip())
    if m is not None and _direct_pp_sheet_name_matches(m.group(2), names_set):
        return True
    return False


def _direct_pp_expand_formula(template, row_0, col_0, sheet_name=u""):
    """Развернуть шаблон формулы; подставить <<Переменные.…>> из runtime-карты сбора."""
    try:
        import libre_macros_collect_cfg as _cw_cfg
        import libre_macros_lib as lm
        from libre_macros_lib import _lm_pp_expand_apply_formula_template
    except Exception as err:
        _log("formula expand import error: %s" % err)
        return u""
    prev_ctx = lm.lm_pp_active_context()
    ctx = dict(prev_ctx or {})
    # Всегда брать актуальный снимок карты сбора (не кэш из prev_ctx).
    ctx["source_variables_map"] = getattr(_cw_cfg, "_MERGE_SOURCE_VARIABLES_MAP", {}) or {}
    ctx["source_variables_order"] = getattr(_cw_cfg, "_MERGE_SOURCE_VARIABLES_ORDER", []) or []
    lm.lm_pp_set_active_context(ctx)
    try:
        text = unicode(template or u"")
        hint = unicode(sheet_name or u"").strip()
        if u"<<" in text:
            if not ctx["source_variables_map"]:
                _log(
                    u"применить_формулу: карта переменных пуста при подстановке <<Переменные…>> "
                    u"(sheet=%s)" % hint
                )
            text = lm._lm_pp_expand_source_variables_template(
                text,
                doc=None,
                sheet=None,
                row_0=row_0,
                header_row=0,
                log_fn_key=u"применить_формулу",
                for_formula=True,
                sheet_name_hint=hint,
            )
        return _lm_pp_expand_apply_formula_template(text, row_0, col_0, doc=None)
    except Exception as err:
        _log("formula expand error: %s" % err)
        return u""
    finally:
        lm.lm_pp_set_active_context(prev_ctx)


def _direct_pp_prepare_formula_for_file(formula_text, for_xlsx=False):
    """
    Имена EN; для XLSX разделители аргументов ; → ,
    (иначе LO/AO при чтении OOXML превращает ; в | — разделитель строк массива).
    ODS OpenFormula оставляем с ;.
    """
    ftxt = unicode(formula_text or u"").strip()
    if ftxt == u"":
        return ftxt
    try:
        from libre_macros_lib import _lm_pp_formula_funcs_to_api

        ftxt = unicode(_lm_pp_formula_funcs_to_api(ftxt) or ftxt)
    except Exception:
        pass
    if for_xlsx:
        try:
            from libre_macros_lib import _lm_pp_formula_separators_to_api

            ftxt = unicode(_lm_pp_formula_separators_to_api(ftxt) or ftxt)
        except Exception:
            pass
    return ftxt


def _direct_pp_expand_source_variables_text(text, sheet_name=u""):
    """Подстановка <<Переменные.…>> в тексте (переименовать_лист и др.) для xml/direct."""
    raw = unicode(text or u"")
    if raw.strip() == u"":
        return raw
    try:
        import libre_macros_lib as lm
        import libre_macros_collect_cfg as _cw_cfg
    except ImportError:
        return raw
    prev_ctx = lm.lm_pp_active_context()
    ctx = dict(prev_ctx or {})
    ctx["source_variables_map"] = getattr(_cw_cfg, "_MERGE_SOURCE_VARIABLES_MAP", {}) or {}
    ctx["source_variables_order"] = getattr(_cw_cfg, "_MERGE_SOURCE_VARIABLES_ORDER", []) or []
    lm.lm_pp_set_active_context(ctx)
    try:
        expanded = lm._lm_pp_expand_source_variables_template(
            raw,
            doc=None,
            sheet=None,
            log_fn_key=u"переименовать_лист",
            for_formula=False,
            sheet_name_hint=unicode(sheet_name or u""),
        )
    finally:
        lm.lm_pp_set_active_context(prev_ctx)
    return unicode(expanded or u"")


def _direct_pp_expand_rename_name_template( template, sheet_name=u"", old_name=u"", header_row=0, used_sr=0, used_er=-1, cell_text_fn=None, log_fn_key=u"переименовать_лист"):
    """Полный разворот шаблона имени (карта + [Старое_имя]/[Заголовок]/[Строка])."""
    text = unicode(template or u"")
    if text.strip() == u"":
        return u""
    try:
        import libre_macros_lib as lm
        import libre_macros_collect_cfg as _cw_cfg
    except ImportError:
        return text
    prev_ctx = lm.lm_pp_active_context()
    ctx = dict(prev_ctx or {})
    ctx["source_variables_map"] = getattr(_cw_cfg, "_MERGE_SOURCE_VARIABLES_MAP", {}) or {}
    ctx["source_variables_order"] = getattr(_cw_cfg, "_MERGE_SOURCE_VARIABLES_ORDER", []) or []
    lm.lm_pp_set_active_context(ctx)
    try:
        expanded = lm._lm_pp_expand_rename_name_template(
            text,
            doc=None,
            sheet=None,
            old_name=unicode(old_name or u""),
            header_row=int(header_row or 0),
            used_sr=int(used_sr or 0),
            used_er=int(used_er if used_er is not None else -1),
            cell_text_fn=cell_text_fn,
            log_fn_key=log_fn_key,
            sheet_name_hint=unicode(sheet_name or u""),
        )
    finally:
        lm.lm_pp_set_active_context(prev_ctx)
    return unicode(expanded or u"")


def _direct_pp_xlsx_cell_text(ws, row_0, col_0):
    try:
        val = ws.cell(row=int(row_0) + 1, column=int(col_0) + 1).value
    except Exception:
        return u""
    if val is None:
        return u""
    return unicode(val).strip()


def _direct_pp_xlsx_used_rows(ws):
    try:
        mr = int(ws.max_row or 0)
    except Exception:
        mr = 0
    if mr < 1:
        return 0, -1
    return 0, mr - 1


def _direct_pp_ods_cell_text(rows, row_0, col_0):
    try:
        r = int(row_0)
        c = int(col_0)
    except (TypeError, ValueError):
        return u""
    if r < 0 or c < 0 or r >= len(rows):
        return u""
    try:
        for cidx, val in _ods_row_col_values(rows[r]):
            if int(cidx) == c:
                if val is None:
                    return u""
                return unicode(val).strip()
    except Exception:
        return u""
    return u""


def _direct_pp_ods_used_rows(rows):
    if not rows:
        return 0, -1
    return 0, len(rows) - 1


def _direct_pp_finalize_sheet_name(existing_names, base_name, exclude_name=None):
    try:
        import libre_macros_lib as lm
    except ImportError:
        return _direct_pp_unique_sheet_name(existing_names, base_name, exclude_name=exclude_name)
    exclude = unicode(exclude_name or u"").strip()
    taken = set()
    for n in existing_names or []:
        nn = unicode(n or u"").strip()
        if nn == u"":
            continue
        if exclude != u"" and _merge_identity_key(nn) == _merge_identity_key(exclude):
            continue
        taken.add(_merge_identity_key(nn))

    def _is_taken(cand):
        return _merge_identity_key(cand) in taken

    return lm._lm_pp_finalize_sheet_name(base_name, is_taken=_is_taken)


def _direct_pp_expand_rename_sheet_name( new_name, sheet_name=u"", old_name=u"", header_row=0, used_sr=0, used_er=-1, cell_text_fn=None):
    """Развернуть шаблон имени листа и вернуть сырое имя (без unique)."""
    text = unicode(new_name or u"").strip()
    if text == u"":
        return u""
    return _direct_pp_expand_rename_name_template(
        text,
        sheet_name=sheet_name,
        old_name=old_name if old_name != u"" else sheet_name,
        header_row=header_row,
        used_sr=used_sr,
        used_er=used_er,
        cell_text_fn=cell_text_fn,
        log_fn_key=u"переименовать_лист",
    ).strip()
_DIRECT_FORMULA_BRACKET_REF = re.compile(r"\[([^\]]+)\]\{([^}]+)\}")
_DIRECT_FORMULA_QUOTED_BRACKET_REF = re.compile(r"\[\s*\"([^\"]+)\"\s*\]\{([^}]+)\}")
_DIRECT_FORMULA_CURLY_HEADER_REF = re.compile(
    r"\{\s*([^{}]+?)\s*\}\s*\{\s*([^{}]*(?:row_id|row)[^{}]*)\s*\}",
    re.IGNORECASE,
)


def _direct_expand_header_refs(formula_text, header_h2i):
    """
    Подставить буквы столбцов вместо ссылок на заголовок в шаблоне формулы:
      [Имя]{row} / ["Имя"]{row} / {Имя}{row}
    """
    if not header_h2i:
        return unicode(formula_text or u"")
    text = unicode(formula_text or u"")

    def _h2letters(name):
        key = _merge_identity_key(unicode(name or u""))
        idx = header_h2i.get(key)
        if idx is None:
            return None
        return _col_index_to_letters(idx)

    def _bracket_replace(match):
        letters = _h2letters(match.group(1).strip())
        if letters is None:
            return match.group(0)
        return letters + u"{" + unicode(match.group(2)) + u"}"

    def _quoted_bracket_replace(match):
        letters = _h2letters(match.group(1).strip())
        if letters is None:
            return match.group(0)
        return letters + u"{" + unicode(match.group(2)) + u"}"

    def _curly_pair_replace(match):
        letters = _h2letters(match.group(1).strip())
        if letters is None:
            return match.group(0)
        return letters + u"{" + unicode(match.group(2)) + u"}"

    text = _DIRECT_FORMULA_CURLY_HEADER_REF.sub(_curly_pair_replace, text)
    text = _DIRECT_FORMULA_BRACKET_REF.sub(_bracket_replace, text)
    text = _DIRECT_FORMULA_QUOTED_BRACKET_REF.sub(_quoted_bracket_replace, text)
    return text


def _direct_xlsx_header_h2i(ws, hdr_row_1=1, max_cols=500):
    """Карта заголовок(casefold) → 0-based индекс столбца на листе XLSX."""
    out = {}
    if ws is None:
        return out
    clean_fn = None
    try:
        from libre_macros_param_codec import _clean_column_match_text as clean_fn
    except Exception:
        clean_fn = None
    try:
        max_c = int(ws.max_column or 1)
    except Exception:
        max_c = 1
    max_c = min(max_c, int(max_cols))
    ci = 1
    while ci <= max_c:
        try:
            v = ws.cell(row=int(hdr_row_1), column=int(ci)).value
        except Exception:
            v = None
        name = unicode(v if v is not None else u"").strip()
        if clean_fn is not None and name != u"":
            try:
                name = unicode(clean_fn(name) or u"").strip()
            except Exception:
                pass
        if name != u"":
            key = _merge_identity_key(name)
            if key not in out:
                out[key] = int(ci) - 1
        ci += 1
    return out


def _direct_ods_header_h2i(rows, hdr_row=0, max_cols=500):
    """Карта заголовок(casefold) → 0-based индекс столбца на листе ODS."""
    out = {}
    if rows is None or int(hdr_row) < 0 or int(hdr_row) >= len(rows):
        return out
    clean_fn = None
    try:
        from libre_macros_param_codec import _clean_column_match_text as clean_fn
    except Exception:
        clean_fn = None
    for col_idx, val in _ods_row_col_values(rows[int(hdr_row)]):
        if int(col_idx) >= int(max_cols):
            break
        name = unicode(val if val is not None else u"").strip()
        if clean_fn is not None and name != u"":
            try:
                name = unicode(clean_fn(name) or u"").strip()
            except Exception:
                pass
        if name == u"":
            continue
        key = _merge_identity_key(name)
        if key not in out:
            out[key] = int(col_idx)
    return out


def _direct_result_sheet_names_from_file(path, one_sheet):
    """
    Актуальные имена листов результата из файла (без кэша settings).

    Всегда читает файл: после сводная_таблица / копировать_лист появляются
    новые вкладки — они должны попасть в итерацию пайплайна.
    При one_sheet лист Merge (если есть) ставится первым.
    """
    if _is_ods(path):
        names = _ods_sheet_names(path)
    else:
        names = _xlsx_sheet_names(path)
    out = []
    for n in names:
        if _direct_is_service_sheet_name(n):
            continue
        out.append(n)
    if one_sheet:
        merge_key = _merge_identity_key(_DIRECT_MERGE_ONE_SHEET)
        merge_name = None
        rest = []
        for n in out:
            if merge_name is None and _merge_identity_key(n) == merge_key:
                merge_name = n
            else:
                rest.append(n)
        if merge_name is not None:
            return [merge_name] + rest
        if out:
            return out
        return [_DIRECT_MERGE_ONE_SHEET]
    return out


def _direct_result_sheet_names(path, one_sheet, settings=None):
    """
    Листы результата для Постобработка_xml.

    Если в settings есть direct_result_sheet_names — сопоставить с файлом.
    После переименовать_лист / сводная_таблица: устаревшие имена и новые
    вкладки подтягиваются из файла (кэш обновляется).
    """
    explicit = []
    try:
        explicit = list((settings or {}).get("direct_result_sheet_names") or [])
    except Exception:
        explicit = []
    names_file = _direct_result_sheet_names_from_file(path, one_sheet)
    if not explicit:
        return names_file
    keys_all = {}
    for nm in names_file:
        keys_all[_merge_identity_key(nm)] = nm
    out_explicit = []
    missing = False
    seen = set()
    for nm in explicit:
        key = _merge_identity_key(unicode(nm or u""))
        found = keys_all.get(key)
        if found is not None:
            out_explicit.append(found)
            seen.add(key)
        else:
            missing = True
    # Новые листы (сводная, копия), которых не было в кэше — добавить в конец.
    added = False
    for nm in names_file:
        key = _merge_identity_key(nm)
        if key not in seen:
            out_explicit.append(nm)
            seen.add(key)
            added = True
    if missing or not out_explicit or added:
        if settings is not None:
            try:
                settings["direct_result_sheet_names"] = list(
                    names_file if (missing or not out_explicit) else out_explicit
                )
            except Exception:
                pass
        if missing or not out_explicit:
            return names_file
    return out_explicit


def _direct_pp_sync_result_sheet_names(path, one_sheet, settings, sheet_names_buf):
    """После rename/copy/pivot: перечитать имена из файла и обновить буфер + settings."""
    names = _direct_result_sheet_names_from_file(path, one_sheet)
    try:
        sheet_names_buf[:] = names
    except Exception:
        pass
    if settings is not None:
        try:
            settings["direct_result_sheet_names"] = list(names)
        except Exception:
            pass
    return names


def _direct_result_sheet_bounds(path, sheet_name):
    if _is_ods(path):
        import odf_bundled  # noqa: F401
        from odf.opendocument import load
        from odf.table import Table, TableRow

        doc = load(path)
        tables = doc.spreadsheet.getElementsByType(Table)
        sheet = None
        for t in tables:
            if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
                sheet = t
                break
        if sheet is None:
            return None
        rows = sheet.getElementsByType(TableRow)
        if len(rows) == 0:
            return None
        ec, er = _direct_scan_bounds_ods(rows, 0, 0)
        if er < 1:
            return None
        return {"hdr_row": 0, "first_data": 1, "last_data": int(er), "end_col": int(ec)}
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return None
        idx = wb.sheetnames.index(sheet_name)
        ec, er = _direct_resolve_sheet_used_bounds(path, "xlsx", idx, 0, 0)
        if er < 1:
            return None
        return {"hdr_row": 0, "first_data": 1, "last_data": int(er), "end_col": int(ec)}
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_xlsx_header_col_index(ws, col_name, hdr_row_1=1):
    hkey = _merge_identity_key(col_name)
    max_c = int(ws.max_column or 1)
    end_col = 0
    ci = 1
    while ci <= max_c:
        v = ws.cell(row=hdr_row_1, column=ci).value
        if _merge_identity_key(unicode(v if v is not None else u"")) == hkey:
            return ci - 1
        if not _direct_value_empty(v):
            end_col = ci - 1
        ci += 1
    return end_col + 1


def _direct_ods_header_col_index(rows, col_name, hdr_row=0):
    hkey = _merge_identity_key(col_name)
    end_col = 0
    found = -1
    if hdr_row < len(rows):
        for col_idx, val in _ods_row_col_values(rows[hdr_row]):
            if _merge_identity_key(unicode(val if val is not None else u"")) == hkey:
                found = int(col_idx)
                break
            if not _direct_value_empty(val):
                end_col = int(col_idx)
    if found >= 0:
        return found
    return end_col + 1


def _ods_make_formula_cell(formula_text):
    from odf.table import TableCell
    from odf.text import P

    fo = unicode(formula_text or u"").strip()
    if fo == u"":
        return _ods_make_cell(u"")
    if not fo.startswith(u"="):
        fo = u"=" + fo
    oformula = u"of:" + fo
    cell = TableCell(formula=oformula, valuetype="float", value="0")
    cell.addElement(P(text=u""))
    return cell


def _direct_ods_row_cell_count(row_elem):
    from odf.table import TableCell

    col = 0
    for cell in row_elem.getElementsByType(TableCell):
        try:
            rep = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        except Exception:
            rep = 1
        col += max(1, rep)
    return col


def _direct_ods_set_cell_at_col(row_elem, col_index, new_cell):
    from odf.table import TableCell
    from odf.text import P

    cells = list(row_elem.getElementsByType(TableCell))
    while len(cells) <= col_index:
        pad = TableCell()
        pad.addElement(P(text=u""))
        row_elem.addElement(pad)
        cells.append(pad)
    old = cells[col_index]
    try:
        row_elem.insertBefore(new_cell, old)
        row_elem.removeChild(old)
    except Exception:
        try:
            row_elem.addElement(new_cell)
        except Exception:
            pass


def _direct_pp_apply_formula_ods(path, sheet_name, col_name, formula_template, as_values):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    doc = load(path)
    sheet = None
    for t in doc.spreadsheet.getElementsByType(Table):
        if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
            sheet = t
            break
    if sheet is None:
        return False, u"лист не найден"
    rows = sheet.getElementsByType(TableRow)
    header_h2i = _direct_ods_header_h2i(rows, hdr_row=0)
    target_col = _direct_ods_header_col_index(rows, col_name, hdr_row=0)
    sr = bounds["first_data"]
    er = bounds["last_data"]
    rr = 0
    while rr <= er:
        if rr >= len(rows):
            rows.append(TableRow())
            sheet.addElement(rows[-1])
        if rr == 0:
            _direct_ods_set_cell_at_col(rows[rr], target_col, _ods_make_cell(col_name))
        else:
            expanded = _direct_expand_header_refs(formula_template, header_h2i)
            ftxt = _direct_pp_expand_formula(expanded, rr, target_col, sheet_name=sheet_name)
            ftxt = _direct_pp_prepare_formula_for_file(ftxt, for_xlsx=False)
            _direct_ods_set_cell_at_col(rows[rr], target_col, _ods_make_formula_cell(ftxt))
        rr += 1
    doc.save(path)
    if as_values:
        return True, u"формулы в файле; as_values — после открытия Calc"
    return True, u"столбец «%s», строки %d..%d" % (col_name, sr + 1, er + 1)


def _direct_pp_apply_formula_xlsx(path, sheet_name, col_name, formula_template, as_values):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    wb = load_workbook(path)
    try:
        if sheet_name not in wb.sheetnames:
            return False, u"лист не найден"
        ws = wb[sheet_name]
        header_h2i = _direct_xlsx_header_h2i(ws, hdr_row_1=1)
        target_col = _direct_xlsx_header_col_index(ws, col_name, hdr_row_1=1)
        if target_col < 0 or target_col > 300:
            return False, u"не удалось добавить столбец"
        ws.cell(row=1, column=target_col + 1, value=col_name)
        sr = bounds["first_data"]
        er = bounds["last_data"]
        rr = sr
        expanded = _direct_expand_header_refs(formula_template, header_h2i)
        while rr <= er:
            ftxt = _direct_pp_expand_formula(expanded, rr, target_col, sheet_name=sheet_name)
            ftxt = _direct_pp_prepare_formula_for_file(ftxt, for_xlsx=True)
            ws.cell(row=rr + 1, column=target_col + 1, value=ftxt)
            rr += 1
        wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if as_values:
        return True, u"формулы в файле; as_values — после открытия Calc"
    return True, u"столбец «%s», строки %d..%d" % (col_name, sr + 1, er + 1)


# =============================================================================
# File-visual Постобработка_xml (xlsx openpyxl + ods odfpy, одинаковый набор шагов)
# =============================================================================

_DIRECT_PP_MM_TO_PT = 2.83465
_DIRECT_PP_DEFAULT_FONT_NAME = u"PT Sans"
_DIRECT_PP_DEFAULT_FONT_SIZE = 11.0
_DIRECT_PP_MONEY_FMT = u"#,##0.00"
_DIRECT_PP_DATE_FMT = u"DD.MM.YYYY"
_DIRECT_PP_MONEY_MARKERS = (u"сумм", u"руб", u"amount", u"price", u"стоим")
_DIRECT_PP_DATE_MARKERS = (u"дата", u"date", u"время", u"time")
_DIRECT_PP_FORMAT_PROBE_ROWS = 3
_DIRECT_PP_GRID_WIDTH_ALIASES = {
    u"тонкая": u"thin",
    u"thin": u"thin",
    u"средняя": u"medium",
    u"medium": u"medium",
    u"толстая": u"thick",
    u"thick": u"thick",
}
# Excel thin/medium/thick → pt для fo:border в ODS
_DIRECT_PP_ODS_BORDER_PT = {
    u"thin": u"0.5pt",
    u"medium": u"1pt",
    u"thick": u"1.5pt",
}
_DIRECT_PP_ODS_INDENT_CM_PER_STEP = 0.353  # ≈ UNO ParaIndent 353 (1/100 мм)
# мм → единицы ширины столбца Excel (как у сводной: max_w = mm / 1.8)
_DIRECT_PP_MM_TO_EXCEL_COL_WIDTH = 1.8
_DIRECT_PP_DEFAULT_COL_WIDTH_MM = 30.0
_DIRECT_PP_COLORIZE_DEFAULT_BG = 0xD9EAF7  # голубой как UNO
_DIRECT_PP_COLORIZE_OUTLINE_DEFAULT = u"1F4E79"


def _direct_pp_track_file_visual_sheet(sheet_name, fn_key=None):
    key = unicode(sheet_name or u"").strip()
    if key == u"":
        return
    fk = _direct_pp_fn_key(fn_key) if fn_key else u""
    if fk in _DIRECT_XML_PRESERVE_UNO_DECORATION_KEYS:
        return
    if key not in _direct_xml_pp_file_visual_sheets:
        _direct_xml_pp_file_visual_sheets.append(key)


def _direct_pp_queue_format_uno_reapply(fn_key, sheet_name, extra):
    """
    Повтор format_* в UNO после встроенного оформления (см. _merge_apply_xml_uno_pending).

    Стили чисел, записанные в ODS/xlsx на этапе file-visual, сбрасываются шагами
    apply_result_sheet_formatting; format_* намеренно не отключает defaults целиком.
    """
    global _direct_xml_pp_uno_pending
    fk = _direct_pp_fn_key(fn_key)
    if fk not in _DIRECT_XML_FORMAT_UNO_REAPPLY_KEYS:
        return False
    sn = unicode(sheet_name or u"").strip()
    if sn == u"":
        return False
    ex = unicode(extra or u"")
    for item in _direct_xml_pp_uno_pending or []:
        if not isinstance(item, dict):
            continue
        if (
            _direct_pp_fn_key(item.get(u"name")) == fk
            and unicode(item.get(u"sheet") or u"").strip() == sn
            and unicode(item.get(u"extra") or u"") == ex
        ):
            return False
    _direct_xml_pp_uno_pending.append(
        {u"name": fk, u"sheet": sn, u"extra": ex}
    )
    return True


def _direct_pp_track_format_header_cols(sheet_name, col_indices):
    """Запомнить столбцы, где format_столбцы оформил строку заголовка в файле."""
    key = unicode(sheet_name or u"").strip()
    if key == u"":
        return
    cols = _direct_xml_pp_format_header_cols.setdefault(key, set())
    for ci in col_indices or []:
        try:
            cols.add(int(ci))
        except (TypeError, ValueError):
            pass


def _direct_pp_grid_side_style(width_token):
    key = unicode(width_token or u"").strip().casefold()
    return _DIRECT_PP_GRID_WIDTH_ALIASES.get(key, u"thin")


def _direct_pp_excel_number_format(fmt_raw):
    raw = unicode(fmt_raw or u"").strip()
    if raw == u"":
        return u"General"
    low = raw.casefold()
    if low in (u"общий", u"general"):
        return u"General"
    try:
        from libre_macros_lib import _lm_pp_normalize_calc_number_format

        raw = unicode(_lm_pp_normalize_calc_number_format(raw, lang=u"ru") or raw)
    except Exception:
        raw = _direct_translate_cyrillic_date_format(raw)
    return raw


def _direct_pp_format_is_general(fmt_raw):
    """True — «не менять числовой формат» (пусто / General). «0» — валидный код Calc (целое)."""
    raw = unicode(fmt_raw or u"").strip()
    if raw == u"":
        return True
    return raw.casefold() in (u"общий", u"general")


def _direct_pp_format_columns_ods_props_from_rule(rule):
    """Свойства ODS-ячейки: только явно заданные в правиле (без «Общий»)."""
    if not isinstance(rule, dict):
        return {}
    props = {}
    if not _direct_pp_format_is_general(rule.get("format")):
        fmt = _direct_pp_excel_number_format(rule.get("format"))
        props[u"data_style"] = u"custom:" + fmt
    fill_hex = _direct_pp_fill_hex_from_rule(rule)
    font_hex = _direct_pp_font_hex_from_rule(rule)
    if not font_hex and rule.get("font_auto") and fill_hex:
        font_hex = _direct_pp_colorize_auto_font_hex(fill_hex)
    h_align, v_align = _direct_pp_aligns_from_rule(rule)
    if fill_hex:
        props[u"fill"] = fill_hex
    if font_hex:
        props[u"font_color"] = font_hex
    if h_align:
        props[u"h_align"] = h_align
    if v_align:
        props[u"v_align"] = v_align
    bw = unicode(rule.get("border_width") or u"").strip()
    bc = unicode(rule.get("border_color") or u"").strip()
    if bw or bc:
        side = _direct_pp_grid_side_style(bw or u"тонкая")
        props[u"border_pt"] = _DIRECT_PP_ODS_BORDER_PT.get(side, u"0.5pt")
        props[u"border_color"] = _direct_pp_color_to_rgb_hex(bc) or u"666666"
    return props


def _direct_pp_format_columns_rows_1based(rule, hdr_row_1, first_data_1, last_data_1):
    """Диапазон строк 1-based для правила формат_столбцы."""
    if bool(rule.get("header_only")):
        return hdr_row_1, hdr_row_1
    if bool(rule.get("include_header")):
        return hdr_row_1, last_data_1
    return first_data_1, last_data_1


def _direct_pp_format_columns_rows_0based(rule, hdr_row0, first_data0, last_data0):
    if bool(rule.get("header_only")):
        return hdr_row0, hdr_row0
    if bool(rule.get("include_header")):
        return hdr_row0, last_data0
    return first_data0, last_data0


def _direct_pp_resolve_format_columns_indices(col_spec, header_h2i, end_col0, titles=None):
    """column из правила формат_столбцы → отсортированные 0-based индексы столбцов."""
    if col_spec is None:
        return []
    if col_spec == u"all":
        return list(range(0, end_col0 + 1))
    if isinstance(col_spec, (str, unicode)):
        low = unicode(col_spec).strip().casefold()
        if low in (u"all", u"все", u"*"):
            return list(range(0, end_col0 + 1))
    tokens = []
    if isinstance(col_spec, (list, tuple)):
        for item in col_spec:
            if isinstance(item, int):
                tokens.append(unicode(item))
            else:
                t = unicode(item or u"").strip()
                if t:
                    tokens.append(t)
    else:
        text = unicode(col_spec).strip()
        try:
            from libre_macros_param_codec import split_column_name_tokens

            parts = split_column_name_tokens(text)
            if parts:
                tokens = [unicode(p).strip() for p in parts if unicode(p).strip()]
            elif u"," in text:
                tokens = [p.strip() for p in text.split(u",") if p.strip()]
            else:
                tokens = [text]
        except Exception:
            if u"," in text:
                tokens = [p.strip() for p in text.split(u",") if p.strip()]
            else:
                tokens = [text]
    out = []
    seen = set()
    for tok in tokens:
        low = unicode(tok).strip().casefold()
        if low in (u"all", u"все", u"*"):
            ci = 0
            while ci <= end_col0:
                if ci not in seen:
                    seen.add(ci)
                    out.append(ci)
                ci += 1
            continue
        # '…' = exact заголовок; иначе как раньше (токены / *).
        try:
            from libre_macros_param_codec import unwrap_column_name_token

            bare, exact = unwrap_column_name_token(tok)
        except Exception:
            bare, exact = unicode(tok).strip(), False
        if exact and titles is not None:
            bare_l = unicode(bare or u"").strip().casefold()
            ti = 0
            while ti < len(titles):
                if unicode(titles[ti] or u"").strip().casefold() == bare_l:
                    if ti not in seen and 0 <= ti <= end_col0:
                        seen.add(ti)
                        out.append(ti)
                    break
                ti += 1
            continue
        resolve_tok = bare if exact else tok
        if u"*" in unicode(resolve_tok) and titles is not None:
            matches = _direct_pp_resolve_col_matches(resolve_tok, header_h2i, end_col0, titles)
        else:
            matches = _direct_pp_resolve_col_tokens([resolve_tok], header_h2i, end_col0)
        for ci in matches:
            if ci not in seen:
                seen.add(ci)
                out.append(ci)
    out.sort()
    return out


def _direct_pp_border_style_from_rule(rule):
    """(style_name, color_hex) для xlsx Border/Side или (None, None)."""
    if not isinstance(rule, dict):
        return None, None
    width = unicode(rule.get("border_width") or u"").strip()
    color = unicode(rule.get("border_color") or u"").strip()
    if width == u"" and color == u"":
        return None, None
    side = _direct_pp_grid_side_style(width or u"тонкая")
    style_name = side  # thin/medium/thick
    color_hex = _direct_pp_color_to_rgb_hex(color) if color else u"666666"
    return style_name, color_hex


def _direct_pp_fill_hex_from_rule(rule):
    if not isinstance(rule, dict):
        return None
    fill = rule.get("fill")
    if fill is None or unicode(fill).strip() == u"":
        return None
    return _direct_pp_color_to_rgb_hex(fill)


def _direct_pp_font_hex_from_rule(rule):
    if not isinstance(rule, dict):
        return None
    if rule.get("font_auto") is True:
        return None
    font = rule.get("font")
    if font is None or unicode(font).strip() == u"":
        return None
    return _direct_pp_color_to_rgb_hex(font)


def _direct_pp_aligns_from_rule(rule):
    if not isinstance(rule, dict):
        return None, None
    h = unicode(rule.get("h_align") or u"").strip().lower()
    if h not in (u"left", u"center", u"right"):
        h = None
    v = unicode(rule.get("v_align") or u"").strip().lower()
    if v not in (u"top", u"center", u"bottom"):
        v = None
    return h, v


def _direct_pp_format_wants_text_probe(block, kind):
    if isinstance(block, dict) and block.get("convert_existing"):
        return True
    if kind == u"number" and _direct_wants_convert_to_numbers():
        return True
    return False


def _direct_pp_resolve_format_number_fmt(block):
    fmt = unicode((block or {}).get("format") or u"").strip().replace(u"\xa0", u" ")
    return fmt if fmt else _direct_default_number_format()


def _direct_pp_resolve_format_date_fmt(block):
    fmt = unicode((block or {}).get("format") or u"").strip()
    return fmt if fmt else _direct_default_date_format()


def _direct_pp_xlsx_number_format_for_pp(block):
    raw = _direct_pp_resolve_format_number_fmt(block)
    if not raw:
        return u""
    raw = raw.replace(u"\xa0", u" ")
    if u";" in raw or u"[" in raw or u"$" in raw or u"₽" in raw:
        return raw
    out = raw
    m = re.search(r",(\d+)$", out)
    if m:
        out = out[: m.start()] + u"." + m.group(1)
    fmt = _direct_pp_excel_number_format(out)
    if fmt == u"General":
        return u""
    return fmt


def _direct_pp_ods_number_style_for_pp(block):
    raw = _direct_pp_resolve_format_number_fmt(block)
    if not raw:
        return None
    return u"custom:" + raw


def _direct_pp_xlsx_date_format_for_pp(block):
    raw = _direct_pp_resolve_format_date_fmt(block)
    if not raw:
        return u""
    return _direct_normalize_excel_date_format(raw)


def _direct_pp_ods_date_style_for_pp(block):
    raw = _direct_pp_resolve_format_date_fmt(block)
    if not raw:
        return None
    return u"custom:" + _direct_normalize_calc_date_format(raw)


def _direct_pp_columns_from_markers(titles_cf, markers_cf):
    cols = set()
    ci = 0
    while ci < len(titles_cf):
        title = titles_cf[ci]
        hit = False
        for mk in markers_cf:
            if mk and mk in title:
                hit = True
                break
        if hit:
            cols.add(ci)
        ci += 1
    return cols


def _direct_pp_probe_number_columns( get_cell, first_data0, end_col0, last_data0, want_text ):
    probe_end = min(
        int(last_data0),
        int(first_data0) + _DIRECT_PP_FORMAT_PROBE_ROWS - 1,
    )
    if probe_end < int(first_data0):
        return set()
    need = min(_DIRECT_COPY_DT_COL_MATCH_MIN, probe_end - int(first_data0) + 1)
    cols = set()
    ci = 0
    while ci <= end_col0:
        native = False
        text_hits = 0
        rr = int(first_data0)
        while rr <= probe_end:
            val = get_cell(rr, ci)
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                native = True
            elif want_text and _direct_is_plain_string_for_datetime(val):
                if _direct_try_parse_number_text(val) is not None:
                    text_hits += 1
            rr += 1
        if native or (want_text and text_hits >= need):
            cols.add(ci)
        ci += 1
    return cols


def _direct_pp_probe_date_columns( get_cell, first_data0, end_col0, last_data0, want_text, formats_raw=None ):
    probe_end = min(
        int(last_data0),
        int(first_data0) + _DIRECT_PP_FORMAT_PROBE_ROWS - 1,
    )
    if probe_end < int(first_data0):
        return set()
    need = min(_DIRECT_COPY_DT_COL_MATCH_MIN, probe_end - int(first_data0) + 1)
    allowed = None
    raw = unicode(formats_raw or u"").strip()
    if raw:
        allowed = _direct_datetime_formats_from_raw(raw)
    cols = set()
    ci = 0
    while ci <= end_col0:
        native = False
        text_hits = 0
        rr = int(first_data0)
        while rr <= probe_end:
            val = get_cell(rr, ci)
            if isinstance(val, (datetime.date, datetime.datetime)):
                native = True
            elif want_text and _direct_is_plain_string_for_datetime(val):
                if allowed is not None:
                    parsed = _direct_try_parse_datetime_text(
                        val, allowed_dt=allowed, formats_raw=raw
                    )
                else:
                    parsed = _direct_try_parse_datetime_text(val)
                if parsed is not None:
                    text_hits += 1
            rr += 1
        if native or (want_text and text_hits >= need):
            cols.add(ci)
        ci += 1
    return cols


def _direct_pp_resolve_money_columns( block, header_h2i, end_col0, titles_cf, get_cell, first_data0, last_data0 ):
    tokens = (block or {}).get("columns")
    if tokens:
        idxs = _direct_pp_resolve_col_tokens(tokens, header_h2i, end_col0)
        return set(idxs) if idxs else set()
    markers = (block or {}).get("markers")
    if markers:
        markers_cf = [
            unicode(m or u"").strip().casefold()
            for m in markers
            if unicode(m or u"").strip()
        ]
        if markers_cf:
            return _direct_pp_columns_from_markers(titles_cf, markers_cf)
    want_text = _direct_pp_format_wants_text_probe(block, u"number")
    return _direct_pp_probe_number_columns(
        get_cell, first_data0, end_col0, last_data0, want_text
    )


def _direct_pp_resolve_date_columns( block, header_h2i, end_col0, titles_cf, get_cell, first_data0, last_data0 ):
    tokens = (block or {}).get("columns")
    if tokens:
        idxs = _direct_pp_resolve_col_tokens(tokens, header_h2i, end_col0)
        return set(idxs) if idxs else set()
    markers = (block or {}).get("markers")
    if markers:
        markers_cf = [
            unicode(m or u"").strip().casefold()
            for m in markers
            if unicode(m or u"").strip()
        ]
        if markers_cf:
            return _direct_pp_columns_from_markers(titles_cf, markers_cf)
    want_text = _direct_pp_format_wants_text_probe(block, u"date")
    parse_raw = unicode((block or {}).get("parse_formats") or u"").strip()
    return _direct_pp_probe_date_columns(
        get_cell,
        first_data0,
        end_col0,
        last_data0,
        want_text,
        formats_raw=parse_raw if parse_raw else None,
    )


def _direct_pp_load_format_block(fn_key, extra, sheet_name, normalize_fn):
    """Нормализованный блок формат_* для листа или None (лист вне фильтра)."""
    block = _direct_pp_pick_sheet_block(fn_key, extra, sheet_name)
    if block is None:
        return None
    try:
        block = normalize_fn(block)
    except Exception:
        pass
    if not isinstance(block, dict):
        return {}
    return block


def _direct_pp_load_format_blocks(fn_key, extra, sheet_name, normalize_fn):
    """Все нормализованные блоки формат_* для текущего листа (пусто = [])."""
    blocks = _direct_pp_blocks_for_sheet(fn_key, extra, sheet_name)
    if blocks is None:
        return []
    out = []
    for block in blocks:
        try:
            nb = normalize_fn(block)
        except Exception:
            nb = block
        if isinstance(nb, dict):
            out.append(nb)
    return out


def _direct_pp_zebra_styles_from_extra(extra_raw, sheet_name):
    """roles header/even/odd → {role: {fill_hex, font_hex}} с дефолтом odd=F2F2F2."""
    block = _direct_pp_pick_sheet_block(u"зебра_диапазон", extra_raw, sheet_name)
    if block is None:
        block = {}
    try:
        from libre_macros_param_codec import _normalize_zebra_block

        block = _normalize_zebra_block(block)
    except Exception:
        pass
    roles_in = block.get("roles") or {}
    out = {u"header": None, u"even": None, u"odd": None}
    for role in (u"header", u"even", u"odd"):
        spec = roles_in.get(role)
        if not isinstance(spec, dict):
            continue
        fill_hex = _direct_pp_color_to_rgb_hex(spec.get("fill"))
        font_hex = _direct_pp_color_to_rgb_hex(spec.get("font"))
        if fill_hex or font_hex:
            out[role] = {u"fill": fill_hex, u"font": font_hex}
    has_data = out[u"even"] is not None or out[u"odd"] is not None
    if not has_data and out[u"header"] is None:
        out[u"odd"] = {u"fill": u"F2F2F2", u"font": None}
    elif out[u"even"] is None and has_data:
        out[u"even"] = {u"fill": u"F2F2F2", u"font": None}
    return out


def _direct_pp_zebra_rel_is_even_role(rel):
    """Как UNO: rel=0 odd, rel=1 even."""
    if rel < 0:
        return False
    return rel % 2 == 1


def _direct_pp_rgb_int_to_hex(rgb_int):
    try:
        return u"%06X" % (int(rgb_int) & 0xFFFFFF)
    except Exception:
        return None


def _direct_pp_interpolate_rgb_hex(hex_min, hex_max, norm):
    """Линейная RGB-интерполяция (как UNO _lm_pp_interpolate_color)."""
    t = max(0.0, min(1.0, float(norm)))

    def _parts(h):
        s = unicode(h or u"000000").lstrip(u"#").upper()
        if len(s) != 6:
            s = u"000000"
        return int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)

    r0, g0, b0 = _parts(hex_min)
    r1, g1, b1 = _parts(hex_max)
    return u"%02X%02X%02X" % (
        int(round(r0 + (r1 - r0) * t)),
        int(round(g0 + (g1 - g0) * t)),
        int(round(b0 + (b1 - b0) * t)),
    )


def _direct_pp_cell_as_float(val):
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        try:
            return float(val)
        except Exception:
            return None
    s = unicode(val).strip().replace(u",", u".")
    if s == u"":
        return None
    try:
        return float(s)
    except Exception:
        return None


def _direct_pp_colorize_trim_key(val):
    if val is None:
        return u""
    if isinstance(val, bool):
        return unicode(val)
    if isinstance(val, (int, float)):
        try:
            if isinstance(val, float) and val == int(val):
                return unicode(int(val))
        except Exception:
            pass
        return unicode(val)
    return unicode(val).strip()


def _direct_pp_colorize_composite_key(row_vals, key_cols0):
    parts = []
    for ci in key_cols0:
        try:
            v = row_vals[ci] if ci < len(row_vals) else None
        except Exception:
            v = None
        parts.append(_direct_pp_colorize_trim_key(v).casefold())
    return u"\x1f".join(parts)


def _direct_pp_colorize_auto_font_hex(bg_hex):
    s = unicode(bg_hex or u"FFFFFF").lstrip(u"#").upper()
    if len(s) != 6:
        return u"000000"
    r, g, b = int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255.0
    return u"000000" if luminance > 0.55 else u"FFFFFF"


def _direct_pp_colorize_blend_hex(hex_a, hex_b, t):
    return _direct_pp_interpolate_rgb_hex(hex_a, hex_b, t)


def _direct_pp_colorize_bg_hex_for_group(color_hexes, group_index, total_groups):
    colors = list(color_hexes or [])
    if not colors:
        colors = [_direct_pp_rgb_int_to_hex(_DIRECT_PP_COLORIZE_DEFAULT_BG)]
    if len(colors) == 1:
        if total_groups <= 1:
            t = 0.42
        else:
            t = 0.14 + 0.72 * (float(group_index) / float(total_groups - 1))
        return _direct_pp_colorize_blend_hex(u"FFFFFF", colors[0], t)
    return colors[group_index % len(colors)]


def _direct_pp_gradient_spec_from_block(block):
    """Параметры градиента из JSON-блока (цвета → hex)."""
    if block is None:
        block = {}
    try:
        from libre_macros_param_codec import _normalize_gradient_block

        block = _normalize_gradient_block(block)
    except Exception:
        pass
    marker = unicode(block.get("marker") or u"сумм").strip() or u"сумм"
    hex_min = _direct_pp_color_to_rgb_hex(block.get("color_min") or u"зеленый")
    hex_max = _direct_pp_color_to_rgb_hex(block.get("color_max") or u"красный")
    if hex_min is None:
        hex_min = u"00FF00"
    if hex_max is None:
        hex_max = u"FF0000"
    whole_row = bool(block.get("whole_row"))
    sort_enabled = bool(block.get("sort"))
    sort_asc = True
    if sort_enabled:
        sort_dir = unicode(block.get("sort_dir") or u"по_возр").strip().lower()
        sort_asc = sort_dir not in (
            u"по_убыв",
            u"убыв",
            u"desc",
            u"-",
            u"убывание",
        )
    return {
        u"marker": marker,
        u"hex_min": hex_min,
        u"hex_max": hex_max,
        u"whole_row": whole_row,
        u"sort": sort_enabled,
        u"sort_asc": sort_asc,
    }


def _direct_pp_colorize_spec_from_block(block):
    if block is None:
        block = {}
    try:
        from libre_macros_param_codec import _normalize_colorize_block

        block = _normalize_colorize_block(block)
    except Exception:
        pass
    key_cols = list(block.get("key_columns") or [])
    color_hexes = []
    for tok in block.get("colors") or []:
        hx = _direct_pp_color_to_rgb_hex(tok)
        if hx:
            color_hexes.append(hx)
    sort_val = block.get("sort")
    sort_asc = None
    if sort_val == u"asc":
        sort_asc = True
    elif sort_val == u"desc":
        sort_asc = False
    outline = bool(block.get("outline_blocks"))
    outline_hex = _direct_pp_color_to_rgb_hex(block.get("outline_color"))
    if outline_hex is None:
        outline_hex = _DIRECT_PP_COLORIZE_OUTLINE_DEFAULT
    return {
        u"key_columns": key_cols,
        u"colors": color_hexes,
        u"sort_asc": sort_asc,
        u"outline": outline,
        u"outline_hex": outline_hex,
    }


def _direct_pp_column_width_spec_from_block(block):
    if block is None:
        block = {}
    try:
        from libre_macros_param_codec import _normalize_column_width_block

        block = _normalize_column_width_block(block)
    except Exception:
        pass
    try:
        width_mm = float(block.get("width_mm", _DIRECT_PP_DEFAULT_COL_WIDTH_MM))
    except (TypeError, ValueError):
        width_mm = _DIRECT_PP_DEFAULT_COL_WIDTH_MM
    cols = block.get("columns", u"all")
    return {u"columns": cols, u"width_mm": width_mm}


def _direct_ods_column_style_name(cache, width_mm, hidden=False):
    from odf.style import Style, TableColumnProperties

    if u"col" not in cache:
        cache[u"col"] = {}
    key = (float(width_mm), bool(hidden))
    if key in cache[u"col"]:
        return cache[u"col"][key]
    name = _direct_ods_next_style_name(cache, u"K")
    st = Style(name=name, family=u"table-column")
    if hidden:
        st.addElement(
            TableColumnProperties(columnwidth=u"0mm", useoptimalcolumnwidth=u"false")
        )
    else:
        st.addElement(
            TableColumnProperties(
                columnwidth=u"%.2fmm" % float(width_mm),
                useoptimalcolumnwidth=u"false",
            )
        )
    cache[u"doc"].automaticstyles.addElement(st)
    cache[u"col"][key] = name
    return name


def _direct_ods_ensure_table_columns(sheet, end_col0):
    """Гарантировать TableColumn элементы 0..end_col0."""
    from odf.table import TableColumn

    cols = list(sheet.getElementsByType(TableColumn))
    while len(cols) <= int(end_col0):
        sheet.addElement(TableColumn())
        cols = list(sheet.getElementsByType(TableColumn))
    return cols


def _direct_pp_file_visual(path, sheet_name, fn_key, extra):
    """Диспетчер file-visual: одинаковый набор шагов для xlsx и ods."""
    key = _direct_pp_fn_key(fn_key)
    if _is_ods(path):
        return _direct_pp_file_visual_ods(path, sheet_name, key, extra)
    return _direct_pp_file_visual_xlsx(path, sheet_name, key, extra)


def _direct_pp_file_visual_xlsx(path, sheet_name, fn_key, extra):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    hdr_row0 = int(bounds["hdr_row"])
    first_data0 = int(bounds["first_data"])
    last_data0 = int(bounds["last_data"])
    end_col0 = int(bounds["end_col"])
    hdr_row_1 = hdr_row0 + 1
    first_data_1 = first_data0 + 1
    last_data_1 = last_data0 + 1
    end_col_1 = end_col0 + 1

    wb = load_workbook(path)
    try:
        if sheet_name not in wb.sheetnames:
            return False, u"лист не найден"
        ws = wb[sheet_name]
        header_h2i = _direct_xlsx_header_h2i(ws, hdr_row_1=hdr_row_1)

        def _apply_border_range(r0_1, r1_1, style_name, color_hex):
            side = Side(style=style_name, color=color_hex)
            rr = r0_1
            while rr <= r1_1:
                cc = 1
                while cc <= end_col_1:
                    cell = ws.cell(row=rr, column=cc)
                    cur = cell.border
                    cell.border = Border(
                        left=side,
                        right=side,
                        top=side,
                        bottom=side,
                    )
                    cc += 1
                rr += 1

        def _set_align(cell, **kwargs):
            try:
                cur = cell.alignment
                if cur is None:
                    cell.alignment = Alignment(**kwargs)
                else:
                    merged = {}
                    for attr in (
                        "horizontal",
                        "vertical",
                        "textRotation",
                        "wrap_text",
                        "shrink_to_fit",
                        "indent",
                    ):
                        merged[attr] = getattr(cur, attr, None)
                    merged.update(kwargs)
                    cell.alignment = Alignment(**merged)
            except Exception:
                cell.alignment = Alignment(**kwargs)

        def _set_font(cell, name=None, size=None, color_hex=None, bold=None):
            try:
                cur = cell.font
                kwargs = {}
                if name is not None:
                    kwargs["name"] = name
                elif cur is not None and cur.name:
                    kwargs["name"] = cur.name
                if size is not None:
                    kwargs["size"] = size
                elif cur is not None and cur.size:
                    kwargs["size"] = cur.size
                if bold is not None:
                    kwargs["bold"] = bold
                elif cur is not None:
                    kwargs["bold"] = cur.bold
                if color_hex:
                    kwargs["color"] = color_hex
                elif cur is not None and getattr(cur, "color", None) is not None:
                    kwargs["color"] = cur.color
                if cur is not None:
                    for attr in (
                        "italic",
                        "underline",
                        "strike",
                        "vertAlign",
                        "charset",
                        "scheme",
                        "family",
                    ):
                        val = getattr(cur, attr, None)
                        if val is not None:
                            kwargs[attr] = val
                cell.font = Font(**kwargs)
            except Exception:
                kwargs = {}
                if name is not None:
                    kwargs["name"] = name
                if size is not None:
                    kwargs["size"] = size
                if bold is not None:
                    kwargs["bold"] = bold
                if color_hex:
                    kwargs["color"] = color_hex
                cell.font = Font(**kwargs)

        note = u""
        if fn_key in (u"тонкая_сетка", u"толстая_сетка", u"сетка"):
            if fn_key == u"тонкая_сетка":
                style_name, color_hex, include_header = u"thin", u"666666", True
            elif fn_key == u"толстая_сетка":
                style_name, color_hex, include_header = u"thick", u"333333", True
            else:
                block = _direct_pp_pick_sheet_block(u"сетка", extra, sheet_name)
                if block is None:
                    block = {}
                try:
                    from libre_macros_param_codec import _normalize_grid_block

                    block = _normalize_grid_block(block)
                except Exception:
                    pass
                style_name = _direct_pp_grid_side_style(block.get("width") or u"тонкая")
                color_hex = _direct_pp_color_to_rgb_hex(block.get("color")) or u"666666"
                include_header = bool(block.get("include_header"))
            r0 = hdr_row_1 if include_header else first_data_1
            if last_data_1 < r0:
                return False, u"пустой диапазон"
            _apply_border_range(r0, last_data_1, style_name, color_hex)
            note = u"сетка %s/%s, строк %d..%d" % (style_name, color_hex, r0, last_data_1)

        elif fn_key == u"перенос":
            block = _direct_pp_pick_sheet_block(u"перенос", extra, sheet_name)
            if block is None:
                block = {}
            try:
                from libre_macros_param_codec import _normalize_wrap_block

                block = _normalize_wrap_block(u"перенос", block)
            except Exception:
                pass
            wrap = True if block.get("wrap") is not False else False
            include_header = bool(block.get("include_header"))
            r0 = hdr_row_1 if include_header else first_data_1
            rr = r0
            while rr <= last_data_1:
                cc = 1
                while cc <= end_col_1:
                    _set_align(ws.cell(row=rr, column=cc), wrap_text=wrap)
                    cc += 1
                rr += 1
            note = u"wrap=%s, строк %d..%d" % (wrap, r0, last_data_1)

        elif fn_key == u"высота_строки":
            block = _direct_pp_pick_sheet_block(u"высота_строки", extra, sheet_name)
            if block is None:
                block = {}
            try:
                from libre_macros_param_codec import _normalize_row_height_block

                block = _normalize_row_height_block(block)
            except Exception:
                pass
            try:
                height_mm = float(block.get("height_mm", 5.0))
            except (TypeError, ValueError):
                height_mm = 5.0
            rows_val = block.get("rows", u"all")
            target_rows_1 = []
            if rows_val in (None, u"", u"all"):
                rr = first_data_1
                while rr <= last_data_1:
                    target_rows_1.append(rr)
                    rr += 1
            elif isinstance(rows_val, (list, tuple)):
                for r in rows_val:
                    try:
                        target_rows_1.append(int(r))
                    except (TypeError, ValueError):
                        pass
            if height_mm == 0.0:
                for rr in target_rows_1:
                    ws.row_dimensions[rr].hidden = True
                note = u"скрыто строк=%d" % len(target_rows_1)
            else:
                pt = float(height_mm) * _DIRECT_PP_MM_TO_PT
                for rr in target_rows_1:
                    dim = ws.row_dimensions[rr]
                    dim.hidden = False
                    dim.height = pt
                note = u"высота %.2f мм (%d строк)" % (height_mm, len(target_rows_1))

        elif fn_key == u"левое_выравнивание":
            rr = first_data_1
            while rr <= last_data_1:
                cc = 1
                while cc <= end_col_1:
                    _set_align(ws.cell(row=rr, column=cc), horizontal=u"left")
                    cc += 1
                rr += 1
            note = u"left, строк %d..%d" % (first_data_1, last_data_1)

        elif fn_key == u"отступ":
            block = _direct_pp_pick_sheet_block(u"отступ", extra, sheet_name)
            if block is None:
                block = {}
            try:
                from libre_macros_param_codec import _normalize_indent_block

                block = _normalize_indent_block(block)
            except Exception:
                pass
            h_align = unicode(block.get("h_align") or u"left").strip().lower()
            if h_align not in (u"left", u"center", u"right"):
                h_align = u"left"
            try:
                steps = max(0, int(block.get("steps", 1)))
            except (TypeError, ValueError):
                steps = 1
            cols = block.get("columns")
            if cols:
                col_indices = _direct_pp_resolve_col_tokens(cols, header_h2i, end_col0)
            else:
                col_indices = list(range(0, end_col0 + 1))
            if not col_indices:
                return False, u"столбцы не найдены"
            for ci in col_indices:
                rr = first_data_1
                while rr <= last_data_1:
                    _set_align(
                        ws.cell(row=rr, column=ci + 1),
                        horizontal=h_align,
                        indent=steps,
                    )
                    rr += 1
            note = u"%s/%d, столбцов=%d" % (h_align, steps, len(col_indices))

        elif fn_key == u"шрифт":
            block = _direct_pp_pick_sheet_block(u"шрифт", extra, sheet_name)
            if block is None:
                block = {}
            try:
                from libre_macros_param_codec import _normalize_font_block

                block = _normalize_font_block(block)
            except Exception:
                pass
            raw_extra = unicode(extra or u"").strip()
            # Пустой C → PT Sans / 11 на данные и заголовок
            if raw_extra == u"":
                font_name = _DIRECT_PP_DEFAULT_FONT_NAME
                size_data = _DIRECT_PP_DEFAULT_FONT_SIZE
                size_header = _DIRECT_PP_DEFAULT_FONT_SIZE
            else:
                font_name = block.get("name")
                size_data = block.get("size_data")
                size_header = block.get("size_header")
                if font_name in (None, u""):
                    font_name = None
                else:
                    font_name = unicode(font_name).strip()
            rr = first_data_1
            while rr <= last_data_1:
                cc = 1
                while cc <= end_col_1:
                    _set_font(
                        ws.cell(row=rr, column=cc),
                        name=font_name,
                        size=size_data,
                    )
                    cc += 1
                rr += 1
            cc = 1
            while cc <= end_col_1:
                _set_font(
                    ws.cell(row=hdr_row_1, column=cc),
                    name=font_name,
                    size=size_header,
                )
                cc += 1
            note = u"name=%s data=%s header=%s" % (
                font_name or u"—",
                size_data if size_data is not None else u"—",
                size_header if size_header is not None else u"—",
            )

        elif fn_key == u"зебра_диапазон":
            styles = _direct_pp_zebra_styles_from_extra(extra, sheet_name)
            header_st = styles.get(u"header")
            if header_st is not None:
                fill_hex = header_st.get(u"fill")
                font_hex = header_st.get(u"font")
                fill = None
                if fill_hex:
                    fill = PatternFill(
                        patternType=u"solid", fgColor=fill_hex, bgColor=fill_hex
                    )
                cc = 1
                while cc <= end_col_1:
                    cell = ws.cell(row=hdr_row_1, column=cc)
                    if fill is not None:
                        cell.fill = fill
                    _set_font(
                        cell,
                        name=_DIRECT_PP_DEFAULT_FONT_NAME,
                        size=_DIRECT_PP_DEFAULT_FONT_SIZE,
                        color_hex=font_hex,
                        bold=True,
                    )
                    cc += 1
            rel = 0
            rr = first_data_1
            while rr <= last_data_1:
                role = u"even" if _direct_pp_zebra_rel_is_even_role(rel) else u"odd"
                st = styles.get(role)
                if st is not None:
                    fill_hex = st.get(u"fill")
                    font_hex = st.get(u"font")
                    fill = None
                    if fill_hex:
                        fill = PatternFill(
                            patternType=u"solid",
                            fgColor=fill_hex,
                            bgColor=fill_hex,
                        )
                    cc = 1
                    while cc <= end_col_1:
                        cell = ws.cell(row=rr, column=cc)
                        if fill is not None:
                            cell.fill = fill
                        if font_hex:
                            _set_font(cell, color_hex=font_hex)
                        cc += 1
                rel += 1
                rr += 1
            note = u"зебра строк %d..%d" % (first_data_1, last_data_1)

        elif fn_key == u"формат_деньги":
            from libre_macros_param_codec import _normalize_format_money_block

            fmt_blocks = _direct_pp_load_format_blocks(
                u"формат_деньги", extra, sheet_name, _normalize_format_money_block
            )
            if len(fmt_blocks) == 0:
                note = u"лист не в фильтре sheet"
            else:
                titles = []
                cc = 1
                while cc <= end_col_1:
                    v = ws.cell(row=hdr_row_1, column=cc).value
                    titles.append(unicode(v if v is not None else u"").casefold())
                    cc += 1
                first_data0 = first_data_1 - 1
                last_data0 = last_data_1 - 1

                def _xlsx_cell(rr0, ci0):
                    return ws.cell(row=rr0 + 1, column=ci0 + 1).value

                matched = 0
                for block in fmt_blocks:
                    cols = _direct_pp_resolve_money_columns(
                        block,
                        header_h2i,
                        end_col0,
                        titles,
                        _xlsx_cell,
                        first_data0,
                        last_data0,
                    )
                    num_fmt = _direct_pp_xlsx_number_format_for_pp(block)
                    for ci in sorted(cols):
                        if num_fmt:
                            rr = first_data_1
                            while rr <= last_data_1:
                                ws.cell(row=rr, column=ci + 1).number_format = num_fmt
                                rr += 1
                        matched += 1
                note = u"денежных столбцов=%d" % matched

        elif fn_key == u"формат_даты":
            from libre_macros_param_codec import _normalize_format_date_block

            fmt_blocks = _direct_pp_load_format_blocks(
                u"формат_даты", extra, sheet_name, _normalize_format_date_block
            )
            if len(fmt_blocks) == 0:
                note = u"лист не в фильтре sheet"
            else:
                titles = []
                cc = 1
                while cc <= end_col_1:
                    v = ws.cell(row=hdr_row_1, column=cc).value
                    titles.append(unicode(v if v is not None else u"").casefold())
                    cc += 1
                first_data0 = first_data_1 - 1
                last_data0 = last_data_1 - 1

                def _xlsx_date_cell(rr0, ci0):
                    return ws.cell(row=rr0 + 1, column=ci0 + 1).value

                matched = 0
                for block in fmt_blocks:
                    cols = _direct_pp_resolve_date_columns(
                        block,
                        header_h2i,
                        end_col0,
                        titles,
                        _xlsx_date_cell,
                        first_data0,
                        last_data0,
                    )
                    date_fmt = _direct_pp_xlsx_date_format_for_pp(block)
                    for ci in sorted(cols):
                        if date_fmt:
                            rr = first_data_1
                            while rr <= last_data_1:
                                ws.cell(row=rr, column=ci + 1).number_format = date_fmt
                                rr += 1
                        matched += 1
                note = u"датных столбцов=%d" % matched

        elif fn_key == u"формат_столбцы":
            fmt_blocks = _direct_pp_blocks_for_sheet(
                u"формат_столбцы", extra, sheet_name
            )
            if not fmt_blocks:
                note = u"лист не в фильтре sheet"
            else:
                try:
                    from libre_macros_param_codec import _normalize_format_columns_block
                except Exception:
                    _normalize_format_columns_block = None
                applied = 0
                titles = []
                cc = 1
                while cc <= end_col_1:
                    v = ws.cell(row=hdr_row_1, column=cc).value
                    titles.append(unicode(v if v is not None else u"").casefold())
                    cc += 1
                for block in fmt_blocks:
                    if _normalize_format_columns_block is not None:
                        try:
                            block = _normalize_format_columns_block(block)
                        except Exception:
                            pass
                    if not isinstance(block, dict):
                        continue
                    rules = block.get("rules") or []
                    for rule in rules:
                        if not isinstance(rule, dict):
                            continue
                        col_tok = rule.get("column")
                        if col_tok is None:
                            continue
                        idxs = _direct_pp_resolve_format_columns_indices(
                            col_tok, header_h2i, end_col0, titles
                        )
                        if not idxs:
                            continue
                        apply_fmt = not _direct_pp_format_is_general(rule.get("format"))
                        fmt = (
                            _direct_pp_excel_number_format(rule.get("format"))
                            if apply_fmt
                            else None
                        )
                        fill_hex = _direct_pp_fill_hex_from_rule(rule)
                        font_hex = _direct_pp_font_hex_from_rule(rule)
                        if not font_hex and rule.get("font_auto") and fill_hex:
                            font_hex = _direct_pp_colorize_auto_font_hex(fill_hex)
                        h_align, v_align = _direct_pp_aligns_from_rule(rule)
                        bstyle, bhex = _direct_pp_border_style_from_rule(rule)
                        fill = None
                        if fill_hex:
                            fill = PatternFill(
                                patternType=u"solid",
                                fgColor=fill_hex,
                                bgColor=fill_hex,
                            )
                        side = Side(style=bstyle, color=bhex) if bstyle else None
                        if (
                            not apply_fmt
                            and fill is None
                            and not font_hex
                            and h_align is None
                            and v_align is None
                            and side is None
                        ):
                            continue
                        r0, r1 = _direct_pp_format_columns_rows_1based(
                            rule, hdr_row_1, first_data_1, last_data_1
                        )
                        if (
                            bool(rule.get("header_only"))
                            or bool(rule.get("include_header"))
                        ) and r0 <= hdr_row_1 <= r1:
                            _direct_pp_track_format_header_cols(sheet_name, idxs)
                        for ci in idxs:
                            rr = r0
                            while rr <= r1:
                                cell = ws.cell(row=rr, column=ci + 1)
                                if apply_fmt and fmt:
                                    cell.number_format = fmt
                                if fill is not None:
                                    cell.fill = fill
                                if font_hex:
                                    _set_font(cell, color_hex=font_hex)
                                if h_align is not None or v_align is not None:
                                    kwargs = {}
                                    if h_align is not None:
                                        kwargs["horizontal"] = h_align
                                    if v_align is not None:
                                        kwargs["vertical"] = v_align
                                    _set_align(cell, **kwargs)
                                if side is not None:
                                    cur = cell.border
                                    cell.border = Border(
                                        left=side
                                        if getattr(cur, "left", None) is None
                                        else side,
                                        right=side
                                        if getattr(cur, "right", None) is None
                                        else side,
                                        top=side
                                        if getattr(cur, "top", None) is None
                                        else side,
                                        bottom=side
                                        if getattr(cur, "bottom", None) is None
                                        else side,
                                    )
                                rr += 1
                        applied += 1
                note = u"правил применено=%d" % applied

        elif fn_key == u"ширина_столбцов":
            from openpyxl.utils import get_column_letter

            block = _direct_pp_pick_sheet_block(u"ширина_столбцов", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            spec = _direct_pp_column_width_spec_from_block(block)
            width_mm = float(spec[u"width_mm"])
            cols_val = spec[u"columns"]
            if cols_val in (None, u"", u"all"):
                col_indices = list(range(0, end_col0 + 1))
            else:
                col_indices = _direct_pp_resolve_col_tokens(
                    cols_val if isinstance(cols_val, (list, tuple)) else [cols_val],
                    header_h2i,
                    end_col0,
                )
            if not col_indices:
                return False, u"столбцы не найдены"
            excel_w = float(width_mm) / _DIRECT_PP_MM_TO_EXCEL_COL_WIDTH
            for ci in col_indices:
                letter = get_column_letter(ci + 1)
                dim = ws.column_dimensions[letter]
                if width_mm == 0.0:
                    dim.hidden = True
                else:
                    dim.hidden = False
                    dim.width = excel_w
            note = (
                u"скрыто столбцов=%d" % len(col_indices)
                if width_mm == 0.0
                else u"ширина %.2f мм (%d столбцов)" % (width_mm, len(col_indices))
            )

        elif fn_key == u"градиент":
            block = _direct_pp_pick_sheet_block(u"градиент", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            spec = _direct_pp_gradient_spec_from_block(block)
            marker = spec[u"marker"]
            titles = []
            cc = 1
            while cc <= end_col_1:
                v = ws.cell(row=hdr_row_1, column=cc).value
                titles.append(unicode(v if v is not None else u"").casefold())
                cc += 1
            needle = marker.casefold()
            target_col0 = None
            ci = 0
            while ci < len(titles):
                if needle in titles[ci]:
                    target_col0 = ci
                    break
                ci += 1
            if target_col0 is None:
                return False, u"столбец не найден по маркеру «%s»" % marker
            if last_data_1 < first_data_1:
                return False, u"нет строк данных"
            if spec[u"sort"]:
                rows_data = []
                rr = first_data_1
                while rr <= last_data_1:
                    row_vals = []
                    cci = 0
                    while cci <= end_col0:
                        row_vals.append(ws.cell(row=rr, column=cci + 1).value)
                        cci += 1
                    rows_data.append(row_vals)
                    rr += 1
                sorted_rows = _direct_pp_sort_rows_stable(
                    rows_data, [(target_col0, not spec[u"sort_asc"])]
                )
                rr = first_data_1
                for row_vals in sorted_rows:
                    cci = 0
                    while cci <= end_col0:
                        ws.cell(row=rr, column=cci + 1).value = (
                            row_vals[cci] if cci < len(row_vals) else None
                        )
                        cci += 1
                    rr += 1
            values = []
            rr = first_data_1
            while rr <= last_data_1:
                fv = _direct_pp_cell_as_float(
                    ws.cell(row=rr, column=target_col0 + 1).value
                )
                if fv is not None:
                    values.append((rr, fv))
                rr += 1
            if not values:
                return False, u"нет числовых значений в столбце «%s»" % marker
            min_val = min(v[1] for v in values)
            max_val = max(v[1] for v in values)
            val_range = max_val - min_val if max_val != min_val else 1.0
            paint_cols = (
                list(range(0, end_col0 + 1))
                if spec[u"whole_row"]
                else [target_col0]
            )
            for row_1, val in values:
                norm = (val - min_val) / val_range if val_range != 0 else 0.5
                bg = _direct_pp_interpolate_rgb_hex(
                    spec[u"hex_min"], spec[u"hex_max"], norm
                )
                fill = PatternFill(
                    patternType=u"solid", fgColor=bg, bgColor=bg
                )
                for ci in paint_cols:
                    ws.cell(row=row_1, column=ci + 1).fill = fill
            extras = []
            if spec[u"whole_row"]:
                extras.append(u"вся строка")
            if spec[u"sort"]:
                extras.append(
                    u"сорт. %s" % (u"возр" if spec[u"sort_asc"] else u"убыв")
                )
            note = u"маркер «%s», ячеек %d%s" % (
                marker,
                len(values),
                (u", " + u", ".join(extras)) if extras else u"",
            )

        elif fn_key == u"раскрасить_блоки":
            from openpyxl.styles import Border, Side

            block = _direct_pp_pick_sheet_block(u"раскрасить_блоки", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            spec = _direct_pp_colorize_spec_from_block(block)
            if not spec[u"key_columns"]:
                return False, u"не заданы key_columns"
            key_cols0 = _direct_pp_resolve_col_tokens(
                spec[u"key_columns"], header_h2i, end_col0
            )
            if not key_cols0:
                return False, u"столбцы ключа не найдены"
            if last_data_1 < first_data_1:
                return False, u"нет строк данных"
            rows_data = []
            rr = first_data_1
            while rr <= last_data_1:
                row_vals = []
                cci = 0
                while cci <= end_col0:
                    row_vals.append(ws.cell(row=rr, column=cci + 1).value)
                    cci += 1
                rows_data.append(row_vals)
                rr += 1
            keyed = []
            for row_vals in rows_data:
                keyed.append(
                    (_direct_pp_colorize_composite_key(row_vals, key_cols0), row_vals)
                )
            if spec[u"sort_asc"] is not None:
                keyed.sort(key=lambda item: item[0], reverse=not spec[u"sort_asc"])
                rr = first_data_1
                for _k, row_vals in keyed:
                    cci = 0
                    while cci <= end_col0:
                        ws.cell(row=rr, column=cci + 1).value = (
                            row_vals[cci] if cci < len(row_vals) else None
                        )
                        cci += 1
                    rr += 1
            distinct = []
            prev = None
            for k, _rv in keyed:
                if k != prev:
                    distinct.append(k)
                    prev = k
            total_groups = len(distinct) if distinct else 1
            key_to_group = {}
            gi = 0
            while gi < len(distinct):
                key_to_group[distinct[gi]] = gi
                gi += 1
            rr = first_data_1
            pi = 0
            while pi < len(keyed):
                k = keyed[pi][0]
                gidx = key_to_group.get(k, 0)
                bg = _direct_pp_colorize_bg_hex_for_group(
                    spec[u"colors"], gidx, total_groups
                )
                fg = _direct_pp_colorize_auto_font_hex(bg)
                fill = PatternFill(
                    patternType=u"solid", fgColor=bg, bgColor=bg
                )
                cci = 0
                while cci <= end_col0:
                    cell = ws.cell(row=rr, column=cci + 1)
                    cell.fill = fill
                    _set_font(cell, color_hex=fg)
                    cci += 1
                rr += 1
                pi += 1
            if spec[u"outline"]:
                side = Side(style=u"medium", color=spec[u"outline_hex"])
                gi = 0
                while gi < len(keyed):
                    k = keyed[gi][0]
                    block_start = first_data_1 + gi
                    gj = gi + 1
                    while gj < len(keyed) and keyed[gj][0] == k:
                        gj += 1
                    block_end = first_data_1 + gj - 1
                    r_out = block_start
                    while r_out <= block_end:
                        c_out = 1
                        while c_out <= end_col_1:
                            cell = ws.cell(row=r_out, column=c_out)
                            left = side if c_out == 1 else None
                            right = side if c_out == end_col_1 else None
                            top = side if r_out == block_start else None
                            bottom = side if r_out == block_end else None
                            if left or right or top or bottom:
                                cur = cell.border
                                cell.border = Border(
                                    left=left or (cur.left if cur else None),
                                    right=right or (cur.right if cur else None),
                                    top=top or (cur.top if cur else None),
                                    bottom=bottom or (cur.bottom if cur else None),
                                )
                            c_out += 1
                        r_out += 1
                    gi = gj
            note = u"строк %d, групп %d, ключи %s" % (
                len(keyed),
                total_groups,
                u",".join(unicode(c + 1) for c in key_cols0),
            )

        else:
            return False, u"неизвестный file-visual: %s" % fn_key

        wb.save(path)
        return True, note
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_ods_clone_empty_like(cell):
    """Пустая ячейка с теми же valuetype/stylename (для разворота repeated)."""
    from odf.table import TableCell
    from odf.text import P

    clone = TableCell()
    for attr in ("valuetype", "stylename"):
        try:
            av = cell.getAttribute(attr)
        except Exception:
            av = None
        if av is not None and unicode(av) != u"":
            try:
                clone.setAttribute(attr, av)
            except Exception:
                pass
    try:
        fv = getattr(cell, u"_lm_fv_props", None)
        if isinstance(fv, dict):
            clone._lm_fv_props = dict(fv)
    except Exception:
        pass
    clone.addElement(P(text=u""))
    return clone


def _direct_ods_expand_row_to_cols(row_elem, end_col0):
    """
    Разворачивает numbercolumnsrepeated до отдельных ячеек 0..end_col0.
    Возвращает список TableCell длины end_col0+1.

    Идempotent: повторный вызов на уже развёрнутой строке не пересобирает DOM.
    """
    from odf.table import TableCell
    from odf.text import P

    need = int(end_col0) + 1
    if need < 1:
        return []

    logical = []
    has_repeat = False
    for cell in list(row_elem.getElementsByType(TableCell)):
        try:
            rep = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        except Exception:
            rep = 1
        if rep != 1:
            has_repeat = True
        if rep < 1:
            rep = 1
        logical.append(cell)
        ri = 1
        while ri < rep:
            logical.append(_direct_ods_clone_empty_like(cell))
            ri += 1

    if not has_repeat and len(logical) >= need:
        return logical[:need]
    if not has_repeat:
        while len(logical) < need:
            pad = TableCell()
            pad.addElement(P(text=u""))
            logical.append(pad)
            try:
                row_elem.addElement(pad)
            except Exception:
                pass
        return logical[:need]

    expanded = []
    for cell in list(row_elem.getElementsByType(TableCell)):
        try:
            rep = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        except Exception:
            rep = 1
        if rep < 1:
            rep = 1
        try:
            if cell.getAttribute("numbercolumnsrepeated") is not None:
                cell.removeAttribute("numbercolumnsrepeated")
        except Exception:
            pass
        expanded.append(cell)
        ri = 1
        while ri < rep:
            expanded.append(_direct_ods_clone_empty_like(cell))
            ri += 1
    while len(expanded) < need:
        pad = TableCell()
        pad.addElement(P(text=u""))
        expanded.append(pad)
    expanded = expanded[:need]

    for cell in list(row_elem.getElementsByType(TableCell)):
        try:
            row_elem.removeChild(cell)
        except Exception:
            pass
    for cell in expanded:
        parent = getattr(cell, "parentNode", None)
        if parent is not None:
            try:
                parent.removeChild(cell)
            except Exception:
                try:
                    cell.parentNode = None
                except Exception:
                    pass
        try:
            row_elem.addElement(cell)
        except Exception:
            clone = _direct_ods_clone_empty_like(cell)
            row_elem.addElement(clone)
            cell = clone
    return expanded


def _direct_ods_style_cache_new(doc):
    """Кэш стилей на один load/save; seq стартует выше уже существующих LM* в файле."""
    seq = 0
    try:
        from odf.style import Style
        from odf.number import NumberStyle, DateStyle

        for container in (doc.automaticstyles, doc.styles):
            for st in container.getElementsByType(Style):
                nm = unicode(st.getAttribute("name") or u"")
                if nm.startswith(u"LM"):
                    tail = nm[2:]
                    # LMC12 / LMR3 / LMN5
                    digits = u""
                    i = len(tail) - 1
                    while i >= 0 and tail[i].isdigit():
                        digits = tail[i] + digits
                        i -= 1
                    if digits:
                        try:
                            seq = max(seq, int(digits))
                        except Exception:
                            pass
            for st in container.getElementsByType(NumberStyle):
                nm = unicode(st.getAttribute("name") or u"")
                if nm.startswith(u"LMN"):
                    try:
                        seq = max(seq, int(nm[3:]))
                    except Exception:
                        pass
            for st in container.getElementsByType(DateStyle):
                nm = unicode(st.getAttribute("name") or u"")
                if nm.startswith(u"LMN"):
                    try:
                        seq = max(seq, int(nm[3:]))
                    except Exception:
                        pass
    except Exception:
        seq = 0
    return {
        u"doc": doc,
        u"cell": {},  # key -> Style name
        u"row": {},
        u"num": {},
        u"seq": seq,
    }


def _direct_ods_next_style_name(cache, prefix):
    cache[u"seq"] = int(cache.get(u"seq") or 0) + 1
    return u"LM%s%d" % (prefix, cache[u"seq"])


def _direct_ods_add_date_style_from_format(doc, name, fmt_raw):
    """DateStyle ODF из строки формата Calc (DD.MM.YYYY NNN, …)."""
    from odf.number import (
        DateStyle,
        Day,
        Month,
        Year,
        DayOfWeek,
        Hours,
        Minutes,
        Seconds,
        Text,
    )

    fmt = _direct_normalize_calc_date_format(fmt_raw)
    ds = DateStyle(name=name)
    tokens = (
        (u"YYYY", lambda: Year(style=u"long")),
        (u"YY", lambda: Year(style=u"short")),
        (u"NNNN", lambda: DayOfWeek(style=u"long")),
        (u"NNN", lambda: DayOfWeek(style=u"short")),
        (u"DD", lambda: Day(style=u"long")),
        (u"MM", lambda: Month(style=u"long")),
        (u"HH:MM:SS", None),
        (u"HH:MM", None),
        (u"HH", lambda: Hours(style=u"long")),
        (u"SS", lambda: Seconds(style=u"long")),
    )
    pos = 0
    while pos < len(fmt):
        matched = False
        for token, factory in tokens:
            chunk = fmt[pos : pos + len(token)]
            if chunk.upper() != token:
                continue
            if token == u"HH:MM:SS":
                ds.addElement(Hours(style=u"long"))
                ds.addElement(Text(text=u":"))
                ds.addElement(Minutes(style=u"long"))
                ds.addElement(Text(text=u":"))
                ds.addElement(Seconds(style=u"long"))
            elif token == u"HH:MM":
                ds.addElement(Hours(style=u"long"))
                ds.addElement(Text(text=u":"))
                ds.addElement(Minutes(style=u"long"))
            elif factory is not None:
                ds.addElement(factory())
            pos += len(token)
            matched = True
            break
        if matched:
            continue
        ds.addElement(Text(text=fmt[pos]))
        pos += 1
    doc.styles.addElement(ds)


def _direct_ods_parse_calc_format_spec(fmt_raw):
    """Разбор строки формата Calc для ODS number/currency-style."""
    raw = unicode(fmt_raw or u"").replace(u"\xa0", u" ").strip()
    if raw == u"":
        return None
    parts = [p.strip() for p in raw.split(u";")]
    pos = parts[0] if parts else raw
    neg = parts[1] if len(parts) > 1 else u""
    neg_red = u"[RED]" in neg.upper()
    currency_sym = u""
    lang = u"ru"
    country = u"RU"
    cm = re.search(r"\[\$([^-\]]*)-(\d+)\]", pos + u";" + neg)
    if cm:
        currency_sym = unicode(cm.group(1) or u"").strip()
    if not currency_sym:
        for sym in (u"₽", u"р.", u"руб"):
            if sym in pos or sym in neg:
                currency_sym = sym
                break
    dec = 0
    for chunk in (pos, neg):
        chunk_clean = re.sub(r"\[[^\]]*\]", u"", chunk)
        m = re.search(r"0[,\.](0+)", chunk_clean)
        if m:
            dec = max(dec, len(m.group(1)))
    grouping = bool(
        re.search(r"#\s*[#0]", pos)
        or re.search(r"#\s*[#0]", neg)
        or u"#,##" in pos.replace(u" ", u"")
    )
    low = raw.casefold()
    is_date = (
        u"dd" in low
        or u"yyyy" in low
        or u"yy" in low
        or u"nnn" in low
        or u"дд" in low
        or u"гг" in low
    )
    return {
        u"decimal_places": dec,
        u"grouping": grouping,
        u"currency_symbol": currency_sym,
        u"currency_lang": lang,
        u"currency_country": country,
        u"negative_red": neg_red,
        u"is_currency": bool(currency_sym),
        u"has_negative_part": neg != u"",
        u"is_date": is_date,
    }


def _direct_ods_make_number_element(spec):
    from odf.number import Number, EmbeddedText

    dec = int(spec.get(u"decimal_places") or 0)
    attrs = {
        u"decimalplaces": unicode(dec),
        u"minintegerdigits": u"1",
    }
    if spec.get(u"grouping"):
        attrs[u"grouping"] = u"true"
    num = Number(**attrs)
    if spec.get(u"grouping"):
        try:
            num.addElement(EmbeddedText(position=u"3", text=u" "))
        except Exception:
            pass
    return num


def _direct_ods_append_currency_body(style_el, spec, with_minus=False):
    from odf.number import Text, CurrencySymbol

    if with_minus:
        style_el.addElement(Text(text=u"-"))
    style_el.addElement(_direct_ods_make_number_element(spec))
    sym = unicode(spec.get(u"currency_symbol") or u"").strip()
    if sym != u"":
        style_el.addElement(Text(text=u" "))
        style_el.addElement(
            CurrencySymbol(
                language=spec.get(u"currency_lang") or u"ru",
                country=spec.get(u"currency_country") or u"RU",
                text=sym,
            )
        )


def _direct_ods_add_currency_style_pair(doc, base_name, spec):
    from odf.number import CurrencyStyle
    from odf.style import Map, TextProperties

    pos_name = base_name + u"P0"
    p0 = CurrencyStyle(name=pos_name)
    try:
        p0.setAttribute(u"volatile", u"true")
    except Exception:
        pass
    _direct_ods_append_currency_body(p0, spec, with_minus=False)
    doc.styles.addElement(p0)

    main = CurrencyStyle(name=base_name)
    if spec.get(u"negative_red"):
        main.addElement(TextProperties(color=u"#ff0000"))
    _direct_ods_append_currency_body(
        main, spec, with_minus=bool(spec.get(u"has_negative_part"))
    )
    main.addElement(Map(condition=u"value()>=0", applystylename=pos_name))
    doc.styles.addElement(main)


def _direct_ods_add_number_style_from_calc_format(doc, style_name, fmt_raw):
    spec = _direct_ods_parse_calc_format_spec(fmt_raw)
    if spec is None:
        return
    if spec.get(u"is_date"):
        _direct_ods_add_date_style_from_format(doc, style_name, fmt_raw)
        return
    if spec.get(u"is_currency"):
        _direct_ods_add_currency_style_pair(doc, style_name, spec)
        return
    from odf.number import NumberStyle

    ns = NumberStyle(name=style_name)
    ns.addElement(_direct_ods_make_number_element(spec))
    doc.styles.addElement(ns)


def _direct_ods_ensure_number_style(cache, kind):
    """kind: money | date | general | custom:<fmt> → data style name."""
    key = unicode(kind or u"general")
    if key in cache[u"num"]:
        return cache[u"num"][key]
    doc = cache[u"doc"]
    from odf.number import (
        NumberStyle,
        Number,
        Text,
        DateStyle,
        Year,
        Month,
        Day,
    )

    name = _direct_ods_next_style_name(cache, u"N")

    def _add_default_date_style():
        _direct_ods_add_date_style_from_format(doc, name, u"DD.MM.YYYY")

    if key == u"money":
        ns = NumberStyle(name=name)
        ns.addElement(Number(decimalplaces="2", minintegerdigits="1", grouping="true"))
        doc.styles.addElement(ns)
    elif key.startswith(u"reuse:"):
        existing = unicode(key[6:] or u"").strip()
        if existing != u"":
            cache[u"num"][key] = existing
            return existing
    elif key == u"date":
        _add_default_date_style()
    elif key.startswith(u"custom:"):
        fmt = key[7:]
        spec = _direct_ods_parse_calc_format_spec(fmt)
        if spec is not None and (
            spec.get(u"is_currency")
            or spec.get(u"is_date")
            or spec.get(u"decimal_places", 0) > 0
            or spec.get(u"grouping")
        ):
            _direct_ods_add_number_style_from_calc_format(doc, name, fmt)
        else:
            low = fmt.casefold()
            if u"0.00" in low or u"#,##0" in low or u"# ##0" in low:
                ns = NumberStyle(name=name)
                ns.addElement(
                    Number(decimalplaces="2", minintegerdigits="1", grouping="true")
                )
                doc.styles.addElement(ns)
            elif (
                u"dd" in low
                or u"yyyy" in low
                or u"yy" in low
                or u"nnn" in low
                or u"hh" in low
                or u"mm" in low
                or u"дд" in low
                or u"гг" in low
            ):
                _direct_ods_add_date_style_from_format(doc, name, fmt)
            else:
                ns = NumberStyle(name=name)
                ns.addElement(Number(decimalplaces="0", minintegerdigits="1"))
                doc.styles.addElement(ns)
    else:
        ns = NumberStyle(name=name)
        ns.addElement(Number(decimalplaces="0", minintegerdigits="1"))
        doc.styles.addElement(ns)
    cache[u"num"][key] = name
    return name


def _direct_ods_cell_style_name(cache, props, parent_style=None):
    """
    props dict:
      border_pt, border_color (#RRGGBB), fill (#RRGGBB), wrap (bool),
      h_align (left|center|right), v_align (top|center|bottom), indent_steps (int),
      font_name, font_size, font_color (#RRGGBB), bold (bool),
      data_style (money|date|general|custom:…)
    """
    from odf.style import (
        Style,
        TableCellProperties,
        TextProperties,
        ParagraphProperties,
    )

    border_pt = props.get(u"border_pt")
    border_color = props.get(u"border_color")
    fill = props.get(u"fill")
    wrap = props.get(u"wrap")
    h_align = props.get(u"h_align")
    v_align = props.get(u"v_align")
    indent_steps = props.get(u"indent_steps")
    font_name = props.get(u"font_name")
    font_size = props.get(u"font_size")
    font_color = props.get(u"font_color")
    bold = props.get(u"bold")
    data_style = props.get(u"data_style")
    parent_u = unicode(parent_style or u"").strip()

    key = (
        border_pt,
        border_color,
        fill,
        wrap,
        h_align,
        v_align,
        indent_steps,
        font_name,
        font_size,
        font_color,
        bold,
        data_style,
        parent_u,
    )
    if key in cache[u"cell"]:
        return cache[u"cell"][key]

    name = _direct_ods_next_style_name(cache, u"C")
    kwargs = {u"name": name, u"family": u"table-cell"}
    if parent_u:
        kwargs[u"parentstylename"] = parent_u
    if data_style:
        kwargs[u"datastylename"] = _direct_ods_ensure_number_style(cache, data_style)
    st = Style(**kwargs)

    tcp_kwargs = {}
    if border_pt and border_color:
        tcp_kwargs[u"border"] = u"%s solid #%s" % (border_pt, border_color)
    if fill:
        tcp_kwargs[u"backgroundcolor"] = u"#%s" % fill
    if wrap is True:
        tcp_kwargs[u"wrapoption"] = u"wrap"
    elif wrap is False:
        tcp_kwargs[u"wrapoption"] = u"no-wrap"
    if tcp_kwargs:
        st.addElement(TableCellProperties(**tcp_kwargs))

    pp_kwargs = {}
    if h_align in (u"left", u"center", u"right"):
        align_map = {u"left": u"start", u"center": u"center", u"right": u"end"}
        pp_kwargs[u"textalign"] = align_map[h_align]
    if v_align in (u"top", u"center", u"bottom"):
        valign_map = {u"top": u"top", u"center": u"middle", u"bottom": u"bottom"}
        pp_kwargs[u"verticalalign"] = valign_map[v_align]
    if indent_steps is not None and int(indent_steps) > 0:
        pp_kwargs[u"marginleft"] = u"%.3fcm" % (
            float(indent_steps) * _DIRECT_PP_ODS_INDENT_CM_PER_STEP
        )
    if pp_kwargs:
        st.addElement(ParagraphProperties(**pp_kwargs))

    tp_kwargs = {}
    if font_name:
        tp_kwargs[u"fontfamily"] = font_name
    if font_size is not None:
        try:
            tp_kwargs[u"fontsize"] = u"%spt" % (
                unicode(font_size).rstrip(u"0").rstrip(u".")
                if u"." in unicode(float(font_size))
                else unicode(int(float(font_size)))
            )
            # проще всегда с одним десятичным
            tp_kwargs[u"fontsize"] = u"%.1fpt" % float(font_size)
        except Exception:
            pass
    if font_color:
        tp_kwargs[u"color"] = u"#%s" % font_color
    if bold:
        tp_kwargs[u"fontweight"] = u"bold"
    if tp_kwargs:
        st.addElement(TextProperties(**tp_kwargs))

    cache[u"doc"].automaticstyles.addElement(st)
    cache[u"cell"][key] = name
    return name


def _direct_ods_row_style_name(cache, height_mm, hidden=False):
    from odf.style import Style, TableRowProperties

    key = (height_mm, bool(hidden))
    if key in cache[u"row"]:
        return cache[u"row"][key]
    name = _direct_ods_next_style_name(cache, u"R")
    st = Style(name=name, family=u"table-row")
    if hidden:
        st.addElement(
            TableRowProperties(rowheight=u"0mm", useoptimalrowheight=u"false")
        )
    else:
        st.addElement(
            TableRowProperties(
                rowheight=u"%.2fmm" % float(height_mm),
                useoptimalrowheight=u"false",
            )
        )
    cache[u"doc"].automaticstyles.addElement(st)
    cache[u"row"][key] = name
    return name


def _direct_ods_merge_style_props(base, overlay):
    out = dict(base or {})
    for k, v in (overlay or {}).items():
        if v is not None:
            out[k] = v
    return out


def _direct_ods_find_style_by_name(doc, style_name):
    if not style_name:
        return None
    name = unicode(style_name)
    try:
        from odf.style import Style

        for st in doc.automaticstyles.getElementsByType(Style):
            if unicode(st.getAttribute("name") or u"") == name:
                return st
        for st in doc.styles.getElementsByType(Style):
            if unicode(st.getAttribute("name") or u"") == name:
                return st
    except Exception:
        pass
    return None


def _direct_ods_style_name_chain(doc, style_name):
    """Имена стилей от родителя к потомку (для наследования ODF)."""
    names = []
    seen = set()
    cur = unicode(style_name or u"").strip()
    while cur and cur not in seen:
        seen.add(cur)
        names.append(cur)
        st = _direct_ods_find_style_by_name(doc, cur)
        if st is None:
            break
        parent = None
        for attr in (u"parentstylename", u"parent-style-name"):
            try:
                parent = st.getAttribute(attr)
            except Exception:
                parent = None
            if parent:
                break
        cur = unicode(parent or u"").strip()
    names.reverse()
    return names


def _direct_ods_merge_style_props_from_element(st, out, cache=None):
    """Дописать в out свойства одного odf Style (без затирания уже заданных)."""
    if st is None:
        return
    try:
        ds = st.getAttribute(u"datastylename")
        if ds:
            ds_u = unicode(ds).strip()
            if ds_u != u"":
                resolved = None
                if cache is not None:
                    for k, v in (cache.get(u"num") or {}).items():
                        if unicode(v) == ds_u:
                            resolved = k
                            break
                if resolved is not None:
                    out[u"data_style"] = resolved
                else:
                    out[u"data_style"] = u"reuse:" + ds_u
    except Exception:
        pass
    try:
        from odf.style import TableCellProperties, TextProperties, ParagraphProperties

        for tcp in st.getElementsByType(TableCellProperties):
            border = tcp.getAttribute(u"border")
            if border:
                parts = unicode(border).split()
                if len(parts) >= 3:
                    out[u"border_pt"] = parts[0]
                    col = parts[-1].lstrip(u"#").upper()
                    if len(col) == 6:
                        out[u"border_color"] = col
            bg = tcp.getAttribute(u"backgroundcolor")
            if bg:
                out[u"fill"] = unicode(bg).lstrip(u"#").upper()
            wrap = tcp.getAttribute(u"wrapoption")
            if wrap == u"wrap":
                out[u"wrap"] = True
            elif wrap == u"no-wrap":
                out[u"wrap"] = False
        for pp in st.getElementsByType(ParagraphProperties):
            ta = pp.getAttribute(u"textalign")
            if ta in (u"start", u"left"):
                out[u"h_align"] = u"left"
            elif ta == u"center":
                out[u"h_align"] = u"center"
            elif ta in (u"end", u"right"):
                out[u"h_align"] = u"right"
            va = pp.getAttribute(u"verticalalign")
            if va == u"top":
                out[u"v_align"] = u"top"
            elif va in (u"middle", u"center"):
                out[u"v_align"] = u"center"
            elif va == u"bottom":
                out[u"v_align"] = u"bottom"
            ml = pp.getAttribute(u"marginleft")
            if ml:
                try:
                    s = unicode(ml).strip().lower().replace(u"cm", u"")
                    cm = float(s)
                    out[u"indent_steps"] = max(
                        0, int(round(cm / _DIRECT_PP_ODS_INDENT_CM_PER_STEP))
                    )
                except Exception:
                    pass
        for tp in st.getElementsByType(TextProperties):
            ff = tp.getAttribute(u"fontfamily")
            if ff:
                out[u"font_name"] = unicode(ff)
            fs = tp.getAttribute(u"fontsize")
            if fs:
                try:
                    out[u"font_size"] = float(
                        unicode(fs).lower().replace(u"pt", u"").strip()
                    )
                except Exception:
                    pass
            fc = tp.getAttribute(u"color")
            if fc:
                out[u"font_color"] = unicode(fc).lstrip(u"#").upper()
            fw = tp.getAttribute(u"fontweight")
            if fw and unicode(fw).casefold() == u"bold":
                out[u"bold"] = True
    except Exception:
        pass


def _direct_ods_read_cell_style_props(doc, cell, cache=None):
    """Прочитать накопленные визуальные props из stylename ячейки (+ цепочка родителей)."""
    try:
        sn = cell.getAttribute(u"stylename")
    except Exception:
        sn = None
    if not sn:
        return {}
    out = {}
    for nm in _direct_ods_style_name_chain(doc, unicode(sn)):
        st = _direct_ods_find_style_by_name(doc, nm)
        _direct_ods_merge_style_props_from_element(st, out, cache=cache)
    return out


def _direct_ods_apply_cell_props(cache, cell, props):
    """Смержить props с уже назначенным стилем ячейки и записать новый stylename."""
    parent_style = None
    try:
        parent_style = cell.getAttribute(u"stylename")
    except Exception:
        parent_style = None
    prev = getattr(cell, u"_lm_fv_props", None)
    if prev is None:
        prev = _direct_ods_read_cell_style_props(cache[u"doc"], cell, cache=cache)
    merged = _direct_ods_merge_style_props(prev, props)
    cell._lm_fv_props = merged
    sn = _direct_ods_cell_style_name(cache, merged, parent_style=parent_style)
    try:
        cell.setAttribute(u"stylename", sn)
    except Exception:
        pass


def _direct_ods_header_titles(rows, hdr_row0, end_col0):
    titles = []
    if hdr_row0 < 0 or hdr_row0 >= len(rows):
        return [u""] * (end_col0 + 1)
    by_col = {}
    for col_idx, val in _ods_row_col_values(rows[hdr_row0]):
        by_col[int(col_idx)] = unicode(val if val is not None else u"")
    ci = 0
    while ci <= end_col0:
        titles.append(by_col.get(ci, u"").casefold())
        ci += 1
    return titles


def _direct_pp_file_visual_ods(path, sheet_name, fn_key, extra):
    """
    ODS file-visual — тот же набор шагов, что и xlsx (odfpy styles).
    """
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    hdr_row0 = int(bounds["hdr_row"])
    first_data0 = int(bounds["first_data"])
    last_data0 = int(bounds["last_data"])
    end_col0 = int(bounds["end_col"])

    doc = load(path)
    sheet = None
    for t in doc.spreadsheet.getElementsByType(Table):
        if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
            sheet = t
            break
    if sheet is None:
        return False, u"лист не найден"
    rows = list(sheet.getElementsByType(TableRow))
    # добить строки до last_data
    while len(rows) <= last_data0:
        nr = TableRow()
        sheet.addElement(nr)
        rows.append(nr)

    cache = _direct_ods_style_cache_new(doc)
    header_h2i = _direct_ods_header_h2i(rows, hdr_row=hdr_row0)

    def _cells_for_row(r0):
        return _direct_ods_expand_row_to_cols(rows[r0], end_col0)

    note = u""
    if fn_key in (u"тонкая_сетка", u"толстая_сетка", u"сетка"):
        if fn_key == u"тонкая_сетка":
            side, color_hex, include_header = u"thin", u"666666", True
        elif fn_key == u"толстая_сетка":
            side, color_hex, include_header = u"thick", u"333333", True
        else:
            block = _direct_pp_pick_sheet_block(u"сетка", extra, sheet_name)
            if block is None:
                block = {}
            try:
                from libre_macros_param_codec import _normalize_grid_block

                block = _normalize_grid_block(block)
            except Exception:
                pass
            side = _direct_pp_grid_side_style(block.get("width") or u"тонкая")
            color_hex = _direct_pp_color_to_rgb_hex(block.get("color")) or u"666666"
            include_header = bool(block.get("include_header"))
        border_pt = _DIRECT_PP_ODS_BORDER_PT.get(side, u"0.5pt")
        r0 = hdr_row0 if include_header else first_data0
        if last_data0 < r0:
            return False, u"пустой диапазон"
        rr = r0
        while rr <= last_data0:
            cells = _cells_for_row(rr)
            ci = 0
            while ci <= end_col0:
                _direct_ods_apply_cell_props(
                    cache,
                    cells[ci],
                    {u"border_pt": border_pt, u"border_color": color_hex},
                )
                ci += 1
            rr += 1
        note = u"сетка %s/%s, строк %d..%d" % (side, color_hex, r0 + 1, last_data0 + 1)

    elif fn_key == u"перенос":
        block = _direct_pp_pick_sheet_block(u"перенос", extra, sheet_name)
        if block is None:
            block = {}
        try:
            from libre_macros_param_codec import _normalize_wrap_block

            block = _normalize_wrap_block(u"перенос", block)
        except Exception:
            pass
        wrap = True if block.get("wrap") is not False else False
        include_header = bool(block.get("include_header"))
        r0 = hdr_row0 if include_header else first_data0
        rr = r0
        while rr <= last_data0:
            cells = _cells_for_row(rr)
            ci = 0
            while ci <= end_col0:
                _direct_ods_apply_cell_props(cache, cells[ci], {u"wrap": wrap})
                ci += 1
            rr += 1
        note = u"wrap=%s, строк %d..%d" % (wrap, r0 + 1, last_data0 + 1)

    elif fn_key == u"высота_строки":
        block = _direct_pp_pick_sheet_block(u"высота_строки", extra, sheet_name)
        if block is None:
            block = {}
        try:
            from libre_macros_param_codec import _normalize_row_height_block

            block = _normalize_row_height_block(block)
        except Exception:
            pass
        try:
            height_mm = float(block.get("height_mm", 5.0))
        except (TypeError, ValueError):
            height_mm = 5.0
        rows_val = block.get("rows", u"all")
        target_rows_0 = []
        if rows_val in (None, u"", u"all"):
            rr = first_data0
            while rr <= last_data0:
                target_rows_0.append(rr)
                rr += 1
        elif isinstance(rows_val, (list, tuple)):
            for r in rows_val:
                try:
                    target_rows_0.append(int(r) - 1)
                except (TypeError, ValueError):
                    pass
        hidden = height_mm == 0.0
        sn = _direct_ods_row_style_name(
            cache, 5.0 if hidden else height_mm, hidden=hidden
        )
        for rr in target_rows_0:
            if rr < 0 or rr >= len(rows):
                continue
            try:
                rows[rr].setAttribute(u"stylename", sn)
            except Exception:
                pass
        note = (
            u"скрыто строк=%d" % len(target_rows_0)
            if hidden
            else u"высота %.2f мм (%d строк)" % (height_mm, len(target_rows_0))
        )

    elif fn_key == u"левое_выравнивание":
        rr = first_data0
        while rr <= last_data0:
            cells = _cells_for_row(rr)
            ci = 0
            while ci <= end_col0:
                _direct_ods_apply_cell_props(
                    cache, cells[ci], {u"h_align": u"left"}
                )
                ci += 1
            rr += 1
        note = u"left, строк %d..%d" % (first_data0 + 1, last_data0 + 1)

    elif fn_key == u"отступ":
        block = _direct_pp_pick_sheet_block(u"отступ", extra, sheet_name)
        if block is None:
            block = {}
        try:
            from libre_macros_param_codec import _normalize_indent_block

            block = _normalize_indent_block(block)
        except Exception:
            pass
        h_align = unicode(block.get("h_align") or u"left").strip().lower()
        if h_align not in (u"left", u"center", u"right"):
            h_align = u"left"
        try:
            steps = max(0, int(block.get("steps", 1)))
        except (TypeError, ValueError):
            steps = 1
        cols = block.get("columns")
        if cols:
            col_indices = _direct_pp_resolve_col_tokens(cols, header_h2i, end_col0)
        else:
            col_indices = list(range(0, end_col0 + 1))
        if not col_indices:
            return False, u"столбцы не найдены"
        for ci in col_indices:
            rr = first_data0
            while rr <= last_data0:
                cells = _cells_for_row(rr)
                _direct_ods_apply_cell_props(
                    cache,
                    cells[ci],
                    {u"h_align": h_align, u"indent_steps": steps},
                )
                rr += 1
        note = u"%s/%d, столбцов=%d" % (h_align, steps, len(col_indices))

    elif fn_key == u"шрифт":
        block = _direct_pp_pick_sheet_block(u"шрифт", extra, sheet_name)
        if block is None:
            block = {}
        try:
            from libre_macros_param_codec import _normalize_font_block

            block = _normalize_font_block(block)
        except Exception:
            pass
        raw_extra = unicode(extra or u"").strip()
        if raw_extra == u"":
            font_name = _DIRECT_PP_DEFAULT_FONT_NAME
            size_data = _DIRECT_PP_DEFAULT_FONT_SIZE
            size_header = _DIRECT_PP_DEFAULT_FONT_SIZE
        else:
            font_name = block.get("name")
            size_data = block.get("size_data")
            size_header = block.get("size_header")
            if font_name in (None, u""):
                font_name = None
            else:
                font_name = unicode(font_name).strip()
        rr = first_data0
        while rr <= last_data0:
            cells = _cells_for_row(rr)
            ci = 0
            while ci <= end_col0:
                _direct_ods_apply_cell_props(
                    cache,
                    cells[ci],
                    {u"font_name": font_name, u"font_size": size_data},
                )
                ci += 1
            rr += 1
        cells = _cells_for_row(hdr_row0)
        ci = 0
        while ci <= end_col0:
            _direct_ods_apply_cell_props(
                cache,
                cells[ci],
                {u"font_name": font_name, u"font_size": size_header},
            )
            ci += 1
        note = u"name=%s data=%s header=%s" % (
            font_name or u"—",
            size_data if size_data is not None else u"—",
            size_header if size_header is not None else u"—",
        )

    elif fn_key == u"зебра_диапазон":
        styles = _direct_pp_zebra_styles_from_extra(extra, sheet_name)
        header_st = styles.get(u"header")
        if header_st is not None:
            cells = _cells_for_row(hdr_row0)
            ci = 0
            while ci <= end_col0:
                _direct_ods_apply_cell_props(
                    cache,
                    cells[ci],
                    {
                        u"fill": header_st.get(u"fill"),
                        u"font_color": header_st.get(u"font"),
                        u"font_name": _DIRECT_PP_DEFAULT_FONT_NAME,
                        u"font_size": _DIRECT_PP_DEFAULT_FONT_SIZE,
                        u"bold": True,
                    },
                )
                ci += 1
        rel = 0
        rr = first_data0
        while rr <= last_data0:
            role = u"even" if _direct_pp_zebra_rel_is_even_role(rel) else u"odd"
            st = styles.get(role)
            if st is not None:
                cells = _cells_for_row(rr)
                ci = 0
                while ci <= end_col0:
                    props = {}
                    if st.get(u"fill"):
                        props[u"fill"] = st.get(u"fill")
                    if st.get(u"font"):
                        props[u"font_color"] = st.get(u"font")
                    if props:
                        _direct_ods_apply_cell_props(cache, cells[ci], props)
                    ci += 1
            rel += 1
            rr += 1
        note = u"зебра строк %d..%d" % (first_data0 + 1, last_data0 + 1)

    elif fn_key == u"формат_деньги":
        from libre_macros_param_codec import _normalize_format_money_block

        fmt_blocks = _direct_pp_load_format_blocks(
            u"формат_деньги", extra, sheet_name, _normalize_format_money_block
        )
        if len(fmt_blocks) == 0:
            return False, u"лист не в фильтре sheet"
        titles = _direct_ods_header_titles(rows, hdr_row0, end_col0)

        def _ods_cell(rr0, ci0):
            row_cells = _cells_for_row(rr0)
            if ci0 < 0 or ci0 >= len(row_cells):
                return None
            return _ods_cell_value(row_cells[ci0])

        matched = 0
        for block in fmt_blocks:
            cols = _direct_pp_resolve_money_columns(
                block,
                header_h2i,
                end_col0,
                titles,
                _ods_cell,
                first_data0,
                last_data0,
            )
            dstyle = _direct_pp_ods_number_style_for_pp(block)
            for ci in sorted(cols):
                if dstyle:
                    rr = first_data0
                    while rr <= last_data0:
                        cells = _cells_for_row(rr)
                        _direct_ods_apply_cell_props(
                            cache, cells[ci], {u"data_style": dstyle}
                        )
                        rr += 1
                matched += 1
        note = u"денежных столбцов=%d" % matched

    elif fn_key == u"формат_даты":
        from libre_macros_param_codec import _normalize_format_date_block

        fmt_blocks = _direct_pp_load_format_blocks(
            u"формат_даты", extra, sheet_name, _normalize_format_date_block
        )
        if len(fmt_blocks) == 0:
            return False, u"лист не в фильтре sheet"
        titles = _direct_ods_header_titles(rows, hdr_row0, end_col0)

        def _ods_date_cell(rr0, ci0):
            row_cells = _cells_for_row(rr0)
            if ci0 < 0 or ci0 >= len(row_cells):
                return None
            return _ods_cell_value(row_cells[ci0])

        matched = 0
        for block in fmt_blocks:
            cols = _direct_pp_resolve_date_columns(
                block,
                header_h2i,
                end_col0,
                titles,
                _ods_date_cell,
                first_data0,
                last_data0,
            )
            dstyle = _direct_pp_ods_date_style_for_pp(block)
            for ci in sorted(cols):
                if dstyle:
                    rr = first_data0
                    while rr <= last_data0:
                        cells = _cells_for_row(rr)
                        _direct_ods_apply_cell_props(
                            cache, cells[ci], {u"data_style": dstyle}
                        )
                        rr += 1
                matched += 1
        note = u"датных столбцов=%d" % matched

    elif fn_key == u"формат_столбцы":
        fmt_blocks = _direct_pp_blocks_for_sheet(u"формат_столбцы", extra, sheet_name)
        if not fmt_blocks:
            return False, u"лист не в фильтре sheet"
        try:
            from libre_macros_param_codec import _normalize_format_columns_block
        except Exception:
            _normalize_format_columns_block = None
        applied = 0
        titles = _direct_ods_header_titles(rows, hdr_row0, end_col0)
        for block in fmt_blocks:
            if _normalize_format_columns_block is not None:
                try:
                    block = _normalize_format_columns_block(block)
                except Exception:
                    pass
            if not isinstance(block, dict):
                continue
            rules = block.get("rules") or []
            for rule in rules:
                if not isinstance(rule, dict):
                    continue
                col_tok = rule.get("column")
                if col_tok is None:
                    continue
                idxs = _direct_pp_resolve_format_columns_indices(
                    col_tok, header_h2i, end_col0, titles
                )
                if not idxs:
                    continue
                props_base = _direct_pp_format_columns_ods_props_from_rule(rule)
                if not props_base:
                    continue
                r0, r1 = _direct_pp_format_columns_rows_0based(
                    rule, hdr_row0, first_data0, last_data0
                )
                if (
                    bool(rule.get("header_only")) or bool(rule.get("include_header"))
                ) and r0 <= hdr_row0 <= r1:
                    _direct_pp_track_format_header_cols(sheet_name, idxs)
                for ci in idxs:
                    rr = r0
                    while rr <= r1:
                        cells = _cells_for_row(rr)
                        _direct_ods_apply_cell_props(cache, cells[ci], props_base)
                        rr += 1
                applied += 1
        note = u"правил применено=%d" % applied

    elif fn_key == u"ширина_столбцов":
        block = _direct_pp_pick_sheet_block(u"ширина_столбцов", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        spec = _direct_pp_column_width_spec_from_block(block)
        width_mm = float(spec[u"width_mm"])
        cols_val = spec[u"columns"]
        if cols_val in (None, u"", u"all"):
            col_indices = list(range(0, end_col0 + 1))
        else:
            col_indices = _direct_pp_resolve_col_tokens(
                cols_val if isinstance(cols_val, (list, tuple)) else [cols_val],
                header_h2i,
                end_col0,
            )
        if not col_indices:
            return False, u"столбцы не найдены"
        cols = _direct_ods_ensure_table_columns(sheet, end_col0)
        sn = _direct_ods_column_style_name(
            cache, 0.0 if width_mm == 0.0 else width_mm, hidden=(width_mm == 0.0)
        )
        for ci in col_indices:
            if ci < 0 or ci >= len(cols):
                continue
            try:
                cols[ci].setAttribute(u"stylename", sn)
            except Exception:
                pass
        note = (
            u"скрыто столбцов=%d" % len(col_indices)
            if width_mm == 0.0
            else u"ширина %.2f мм (%d столбцов)" % (width_mm, len(col_indices))
        )

    elif fn_key == u"градиент":
        block = _direct_pp_pick_sheet_block(u"градиент", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        spec = _direct_pp_gradient_spec_from_block(block)
        marker = spec[u"marker"]
        titles = _direct_ods_header_titles(rows, hdr_row0, end_col0)
        needle = marker.casefold()
        target_col0 = None
        ci = 0
        while ci < len(titles):
            if needle in titles[ci]:
                target_col0 = ci
                break
            ci += 1
        if target_col0 is None:
            return False, u"столбец не найден по маркеру «%s»" % marker
        if last_data0 < first_data0:
            return False, u"нет строк данных"
        if spec[u"sort"]:
            rows_data = []
            rr = first_data0
            while rr <= last_data0:
                cells = _cells_for_row(rr)
                row_vals = []
                cci = 0
                while cci <= end_col0:
                    row_vals.append(_ods_cell_value(cells[cci]))
                    cci += 1
                rows_data.append(row_vals)
                rr += 1
            sorted_rows = _direct_pp_sort_rows_stable(
                rows_data, [(target_col0, not spec[u"sort_asc"])]
            )
            rr = first_data0
            for row_vals in sorted_rows:
                cci = 0
                while cci <= end_col0:
                    _direct_ods_set_cell_at_col(
                        rows[rr],
                        cci,
                        _ods_make_cell(
                            row_vals[cci] if cci < len(row_vals) else None
                        ),
                    )
                    cci += 1
                rr += 1
        values = []
        rr = first_data0
        while rr <= last_data0:
            cells = _cells_for_row(rr)
            fv = _direct_pp_cell_as_float(_ods_cell_value(cells[target_col0]))
            if fv is not None:
                values.append((rr, fv))
            rr += 1
        if not values:
            return False, u"нет числовых значений в столбце «%s»" % marker
        min_val = min(v[1] for v in values)
        max_val = max(v[1] for v in values)
        val_range = max_val - min_val if max_val != min_val else 1.0
        paint_cols = (
            list(range(0, end_col0 + 1)) if spec[u"whole_row"] else [target_col0]
        )
        for row_0, val in values:
            norm = (val - min_val) / val_range if val_range != 0 else 0.5
            bg = _direct_pp_interpolate_rgb_hex(
                spec[u"hex_min"], spec[u"hex_max"], norm
            )
            cells = _cells_for_row(row_0)
            for ci in paint_cols:
                _direct_ods_apply_cell_props(cache, cells[ci], {u"fill": bg})
        extras = []
        if spec[u"whole_row"]:
            extras.append(u"вся строка")
        if spec[u"sort"]:
            extras.append(u"сорт. %s" % (u"возр" if spec[u"sort_asc"] else u"убыв"))
        note = u"маркер «%s», ячеек %d%s" % (
            marker,
            len(values),
            (u", " + u", ".join(extras)) if extras else u"",
        )

    elif fn_key == u"раскрасить_блоки":
        block = _direct_pp_pick_sheet_block(u"раскрасить_блоки", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        spec = _direct_pp_colorize_spec_from_block(block)
        if not spec[u"key_columns"]:
            return False, u"не заданы key_columns"
        key_cols0 = _direct_pp_resolve_col_tokens(
            spec[u"key_columns"], header_h2i, end_col0
        )
        if not key_cols0:
            return False, u"столбцы ключа не найдены"
        if last_data0 < first_data0:
            return False, u"нет строк данных"
        rows_data = []
        rr = first_data0
        while rr <= last_data0:
            cells = _cells_for_row(rr)
            row_vals = []
            cci = 0
            while cci <= end_col0:
                row_vals.append(_ods_cell_value(cells[cci]))
                cci += 1
            rows_data.append(row_vals)
            rr += 1
        keyed = []
        for row_vals in rows_data:
            keyed.append(
                (_direct_pp_colorize_composite_key(row_vals, key_cols0), row_vals)
            )
        if spec[u"sort_asc"] is not None:
            keyed.sort(key=lambda item: item[0], reverse=not spec[u"sort_asc"])
            rr = first_data0
            for _k, row_vals in keyed:
                cci = 0
                while cci <= end_col0:
                    _direct_ods_set_cell_at_col(
                        rows[rr],
                        cci,
                        _ods_make_cell(
                            row_vals[cci] if cci < len(row_vals) else None
                        ),
                    )
                    cci += 1
                rr += 1
        distinct = []
        prev = None
        for k, _rv in keyed:
            if k != prev:
                distinct.append(k)
                prev = k
        total_groups = len(distinct) if distinct else 1
        key_to_group = {}
        gi = 0
        while gi < len(distinct):
            key_to_group[distinct[gi]] = gi
            gi += 1
        border_pt = _DIRECT_PP_ODS_BORDER_PT.get(u"medium", u"1pt")
        rr = first_data0
        pi = 0
        while pi < len(keyed):
            k = keyed[pi][0]
            gidx = key_to_group.get(k, 0)
            bg = _direct_pp_colorize_bg_hex_for_group(
                spec[u"colors"], gidx, total_groups
            )
            fg = _direct_pp_colorize_auto_font_hex(bg)
            cells = _cells_for_row(rr)
            cci = 0
            while cci <= end_col0:
                _direct_ods_apply_cell_props(
                    cache, cells[cci], {u"fill": bg, u"font_color": fg}
                )
                cci += 1
            rr += 1
            pi += 1
        if spec[u"outline"]:
            gi = 0
            while gi < len(keyed):
                k = keyed[gi][0]
                block_start = first_data0 + gi
                gj = gi + 1
                while gj < len(keyed) and keyed[gj][0] == k:
                    gj += 1
                block_end = first_data0 + gj - 1
                r_out = block_start
                while r_out <= block_end:
                    cells = _cells_for_row(r_out)
                    c_out = 0
                    while c_out <= end_col0:
                        # внешний контур: border на крайних ячейках блока
                        if (
                            r_out in (block_start, block_end)
                            or c_out in (0, end_col0)
                        ):
                            _direct_ods_apply_cell_props(
                                cache,
                                cells[c_out],
                                {
                                    u"border_pt": border_pt,
                                    u"border_color": spec[u"outline_hex"],
                                },
                            )
                        c_out += 1
                    r_out += 1
                gi = gj
        note = u"строк %d, групп %d, ключи %s" % (
            len(keyed),
            total_groups,
            u",".join(unicode(c + 1) for c in key_cols0),
        )

    else:
        return False, u"неизвестный file-visual: %s" % fn_key

    doc.save(path)
    return True, note


# --- file-structural (Постобработка_xml): удаление/конкат/замена/fill/sort/rename/copy ---


def _direct_pp_header_matches_marker(title, marker):
    """Как UNO `_lm_pp_header_matches_marker`: подстрока или glob с *."""
    pattern_raw = unicode(marker or u"").strip()
    text = unicode(title or u"").strip()
    if pattern_raw == u"" or text == u"":
        return False

    # Exact-режим: 'Имя' → равенство без подстроки/шаблонов.
    exact = False
    bare = pattern_raw
    try:
        from libre_macros_param_codec import unwrap_column_name_token

        bare, exact = unwrap_column_name_token(pattern_raw)
    except Exception:
        bare = pattern_raw
        exact = False

    bare_u = unicode(bare or u"").strip()
    if bare_u == u"":
        return False

    if exact:
        return bare_u.casefold() == text.casefold()

    if u"*" in bare_u:
        return _text_match(text, bare_u)
    return bare_u.casefold() in text.casefold()


def _direct_pp_resolve_delete_marker_columns(marker, end_col0, titles):
    """
    Один токен удалить_столбцы → 0-based индексы.
    Число (1-based) / буква A/AA — без отсечения по «пустым» хвостам used-area
    (верхняя граница — max(end_col0, len(titles)-1, 1023));
    иначе все заголовки по header_matches_marker.
    """
    t = unicode(marker or u"").strip()
    if t == u"":
        return []
    max_col = int(end_col0) if end_col0 is not None else -1
    try:
        if titles is not None and len(titles) > 0:
            max_col = max(max_col, len(titles) - 1)
    except Exception:
        pass
    if max_col < 0:
        max_col = 1023
    else:
        max_col = max(max_col, 1023)
    try:
        idx = int(t) - 1
        if 0 <= idx <= max_col:
            return [idx]
        return []
    except (TypeError, ValueError):
        pass
    out = []
    ci = 0
    while ci <= int(end_col0) and ci < len(titles or []):
        if _direct_pp_header_matches_marker(titles[ci], t):
            out.append(ci)
        ci += 1
    if out:
        return out
    if _is_col_letters(t):
        idx = _col_letters_to_index(t)
        if 0 <= idx <= max_col:
            return [idx]
        return []
    return []


def _direct_pp_delete_markers_col_set(markers, end_col0, titles):
    """Множество столбцов к удалению по списку смешанных маркеров."""
    del_set = set()
    toks = _direct_pp_normalize_markers_tokens(markers)
    for mk in toks:
        for ci in _direct_pp_resolve_delete_marker_columns(mk, end_col0, titles):
            del_set.add(ci)
    return del_set


def _direct_pp_normalize_markers_tokens(markers):
    """markers из блока → список строк (строку «11,12» разбивает)."""
    if markers is None:
        return []
    if isinstance(markers, (str, unicode)):
        return [
            p.strip()
            for p in unicode(markers).replace(u";", u",").split(u",")
            if p.strip()
        ]
    if isinstance(markers, (list, tuple)):
        out = []
        for m in markers:
            if isinstance(m, (str, unicode)) and (u"," in unicode(m) or u";" in unicode(m)):
                out.extend(_direct_pp_normalize_markers_tokens(m))
            else:
                s = unicode(m or u"").strip()
                if s:
                    out.append(s)
        return out
    s = unicode(markers).strip()
    return [s] if s else []


def _direct_pp_replace_normalize_text(s, squeeze_spaces=False):
    if s is None:
        return u""
    t = unicode(s)
    if squeeze_spaces:
        t = re.sub(r"\s+", u" ", t.strip())
    # squeeze_spaces=false — краевые пробелы значимы
    return t


def _direct_pp_cell_is_plain_string(v):
    """True, если значение похоже на текст (не число/дата/bool)."""
    if v is None:
        return False
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return False
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return False
    return isinstance(v, (str, unicode))


def _direct_pp_replace_values_apply_block(block, matrix, headers, col_indices):
    """
    Замена_значений по матрице в памяти (фильтр строк + match_pick).
    Возвращает (ok, stats_or_error).
    """
    try:
        from libre_macros_lib import (
            LM_PP_REPLACE_EMPTY,
            _lm_pp_replace_values_align_find_repl,
            _lm_pp_replace_values_eval_replace_expr,
            _lm_pp_replace_values_expand_list,
            _lm_pp_replace_values_is_empty_token,
            _lm_pp_replace_values_normalize_text,
            _lm_pp_replace_values_parse_replace_mode,
            _lm_pp_replace_values_replace_expr_text,
            _lm_pp_replace_values_split_list,
            lm_pp_replace_values_apply_matrix,
            lm_pp_replace_values_prepare_filters,
        )
        from libre_macros_lambda_column_lib import make_rows_accessor, make_upd_rows_store
    except Exception as err:
        return False, u"нет libre_macros_lib: %s" % err

    find_list = _lm_pp_replace_values_split_list(
        block.get("find") or block.get("search") or block.get("patterns")
    )
    find_list = _lm_pp_replace_values_expand_list(find_list)
    if not find_list:
        return False, u"не заданы find"
    case_insensitive = bool(
        block.get("case_insensitive")
        or block.get("без_учета_регистра")
        or block.get("ignore_case")
    )
    squeeze_spaces = bool(
        block.get("squeeze_spaces")
        or block.get("сжать_пробелы")
        or block.get("trim_spaces")
    )
    prepared = lm_pp_replace_values_prepare_filters(block)
    if prepared.get("error"):
        return False, prepared.get("error")

    replace_mode = _lm_pp_replace_values_parse_replace_mode(block)
    if replace_mode == u"lambda":
        rows_accessor = make_rows_accessor(matrix, headers, start_col=0)
        repl_list, rerr = _lm_pp_replace_values_eval_replace_expr(
            _lm_pp_replace_values_replace_expr_text(block),
            rows_accessor=rows_accessor,
            upd_store=make_upd_rows_store(),
            row_id=0,
        )
        if rerr:
            return False, rerr
    else:
        repl_list = _lm_pp_replace_values_split_list(
            block.get("replace")
            if "replace" in block
            else (block.get("replacements") or block.get("to"))
        )
        repl_list = _lm_pp_replace_values_expand_list(repl_list)
    find_list, repl_list, align_warn = _lm_pp_replace_values_align_find_repl(
        find_list, repl_list
    )
    if not find_list:
        return False, u"после согласования find/replace нет пар"
    find_norm = [
        _lm_pp_replace_values_normalize_text(x, squeeze_spaces=squeeze_spaces)
        for x in find_list
    ]
    repl_norm = [
        _lm_pp_replace_values_normalize_text(x, squeeze_spaces=squeeze_spaces)
        for x in repl_list
    ]
    ri = 0
    while ri < len(repl_norm):
        if _lm_pp_replace_values_is_empty_token(repl_norm[ri]):
            repl_norm[ri] = LM_PP_REPLACE_EMPTY
        ri += 1
    stats = lm_pp_replace_values_apply_matrix(
        matrix,
        headers,
        col_indices,
        find_norm,
        repl_norm,
        case_insensitive=case_insensitive,
        squeeze_spaces=squeeze_spaces,
        start_col=0,
        row_fn=prepared.get("row_fn"),
        row_glo=prepared.get("row_glo"),
        match_pick=prepared.get("match_pick") or u"all",
        match_fn=prepared.get("match_fn"),
        match_glo=prepared.get("match_glo"),
    )
    if align_warn:
        try:
            stats = dict(stats or {})
            stats[u"align_warn"] = align_warn
        except Exception:
            pass
    return True, stats


def _direct_pp_normalize_resolve_key(block, wb=None, sheet_name=u""):
    """Ключ шифра для XML-фазы (литерал / key_file / ячейка той же книги)."""
    from libre_macros_normalize_lib import (
        NormalizeOpError,
        needs_crypto_key,
        parse_key_cell_ref,
        resolve_crypto_key,
    )

    ops = block.get("ops") if isinstance(block, dict) else None
    if not needs_crypto_key(ops):
        return None
    # сначала литерал / файл без UNO
    try:
        key = resolve_crypto_key(block, doc=None, default_sheet=None, read_file=True)
        if key not in (None, u""):
            return key
    except NormalizeOpError:
        pass
    # key_cell из openpyxl книги
    ref = block.get("key_cell")
    if ref in (None, u"", {}):
        from libre_macros_normalize_lib import coerce_ops_list

        for op in coerce_ops_list(ops):
            if op.get("op") in (u"encrypt", u"decrypt") and op.get("key_cell"):
                ref = op.get("key_cell")
                break
    if ref in (None, u"", {}) or wb is None:
        raise NormalizeOpError(
            u"для encrypt/decrypt нужен key / key_cell / key_file"
        )
    sn, a1 = parse_key_cell_ref(ref)
    sn = sn or sheet_name
    try:
        ws = wb[sn] if sn in wb.sheetnames else wb.active
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

        col_letter, row = coordinate_from_string(a1)
        cell = ws.cell(row=row, column=column_index_from_string(col_letter))
        v = cell.value
        if v is None:
            return u""
        return unicode(v).strip()
    except Exception as err:
        raise NormalizeOpError(u"key_cell «%s»: %s" % (unicode(ref), err))


def _direct_pp_normalize_apply_cell( v, ops, crypto_key, non_text, on_error, empty_cells, row=None, row_filters=None, row_id=None, rows_accessor=None, upd_store=None, compute_fns=None):
    from libre_macros_normalize_lib import (
        NormalizeOpError,
        apply_ops,
        cell_text_for_ops,
    )

    try:
        text0 = cell_text_for_ops(
            v, non_text=non_text, ops=ops, empty_cells=empty_cells
        )
    except NormalizeOpError:
        if on_error == u"empty":
            return u"", True, True
        return v, False, True
    if text0 is None:
        return v, False, False
    new_text, err, _n = apply_ops(
        text0,
        ops,
        key=crypto_key,
        on_error=on_error,
        cell_value=v,
        row=row,
        row_filters=row_filters,
        row_id=row_id,
        rows_accessor=rows_accessor,
        upd_store=upd_store,
        compute_fns=compute_fns,
    )
    changed = new_text != text0
    return new_text if changed else v, changed, err is not None


def _direct_pp_unique_sheet_name(existing_names, base_name, exclude_name=None):
    """Уникальное имя листа (макс. 31), суффиксы _2,_3…"""
    base = unicode(base_name or u"").strip()
    if base == u"":
        return u""
    max_len = 31
    exclude = unicode(exclude_name or u"").strip()
    taken = set()
    for n in existing_names or []:
        nn = unicode(n or u"").strip()
        if nn == u"":
            continue
        if exclude != u"" and _merge_identity_key(nn) == _merge_identity_key(exclude):
            continue
        taken.add(_merge_identity_key(nn))

    def _ok(name):
        return _merge_identity_key(name) not in taken

    cand = base[:max_len]
    if _ok(cand):
        return cand
    suffix = 2
    while suffix <= 999:
        tail = u"_%d" % suffix
        room = max_len - len(tail)
        if room < 1:
            return u""
        cand = base[:room] + tail
        if _ok(cand):
            return cand
        suffix += 1
    return u""


def _direct_pp_find_sheet_name(names, needle):
    """Найти имя листа по совпадению (в т.ч. с префиксом N_)."""
    needle = unicode(needle or u"").strip()
    if needle == u"":
        return None
    for n in names or []:
        if _direct_pp_sheet_name_matches(n, [needle]):
            return n
        m = re.match(r"^(\d+)_(.+)$", unicode(n or u"").strip())
        if m is not None and _direct_pp_sheet_name_matches(m.group(2), [needle]):
            return n
    return None


def _direct_pp_xlsx_header_titles(ws, hdr_row_1, end_col0):
    titles = []
    ci = 0
    while ci <= int(end_col0):
        try:
            v = ws.cell(row=int(hdr_row_1), column=ci + 1).value
        except Exception:
            v = None
        titles.append(unicode(v if v is not None else u"").strip())
        ci += 1
    return titles


def _direct_pp_ods_header_titles_list(rows, hdr_row0, end_col0):
    titles = [u""] * (int(end_col0) + 1)
    if rows is None or int(hdr_row0) < 0 or int(hdr_row0) >= len(rows):
        return titles
    for col_idx, val in _ods_row_col_values(rows[int(hdr_row0)]):
        if col_idx > int(end_col0):
            break
        titles[int(col_idx)] = unicode(val if val is not None else u"").strip()
    return titles


def _direct_pp_sort_key_tuple(row_vals, key_specs):
    """Ключ сортировки: список (is_empty, typed_or_str) по key_specs."""
    parts = []
    for ci, desc in key_specs:
        try:
            v = row_vals[ci] if ci < len(row_vals) else None
        except Exception:
            v = None
        empty = _direct_value_empty(v)
        if empty:
            typed = u""
        elif isinstance(v, bool):
            typed = (0, 1 if v else 0)
        elif isinstance(v, (int, float)):
            typed = (1, float(v))
        elif isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
            typed = (2, unicode(v))
        else:
            typed = (3, unicode(v if v is not None else u"").casefold())
        # desc: инвертируем через обёртку; пустые всегда в конце при asc
        parts.append((empty, typed, desc))
    return parts


def _direct_pp_sort_rows_stable(rows_data, key_specs):
    """Стабильная multi-key сортировка списков значений (data rows)."""
    if not rows_data or not key_specs:
        return list(rows_data)

    indexed = list(enumerate(rows_data))

    def _cmp_pair(a, b):
        # a/b: (orig_i, row_vals)
        ka = _direct_pp_sort_key_tuple(a[1], key_specs)
        kb = _direct_pp_sort_key_tuple(b[1], key_specs)
        i = 0
        while i < len(ka):
            ea, ta, desc = ka[i]
            eb, tb, _d = kb[i]
            if ea != eb:
                # пустые в конец
                return 1 if ea else -1
            if ta < tb:
                return 1 if desc else -1
            if ta > tb:
                return -1 if desc else 1
            i += 1
        return a[0] - b[0]

    try:
        from functools import cmp_to_key

        indexed.sort(key=cmp_to_key(_cmp_pair))
    except Exception:
        # py2-style fallback: bubble (редко)
        n = len(indexed)
        i = 0
        while i < n:
            j = 0
            while j < n - 1:
                if _cmp_pair(indexed[j], indexed[j + 1]) > 0:
                    indexed[j], indexed[j + 1] = indexed[j + 1], indexed[j]
                j += 1
            i += 1
    return [row for _i, row in indexed]


def _direct_pp_fill_down_calc_normalize_template(template):
    """R{-1}C → R[-1]C (как UNO заполнение_вниз_вычислить)."""
    s = unicode(template or u"").strip()
    if s == u"":
        return s
    s = re.sub(r"R\{([+-]?\d+)\}C", r"R[\1]C", s, flags=re.IGNORECASE)
    s = re.sub(
        r"R(\[([+-]?\d+)\]|(\d+))?C\{([+-]?\d+)\}",
        lambda m: u"R%sC[%s]" % (m.group(1) or u"", m.group(4)),
        s,
        flags=re.IGNORECASE,
    )
    return s


def _direct_pp_fill_down_calc_pairs(block, header_h2i, end_col0):
    """
    JSON-блок → [(col_token, formula, col_idx, expand_to_right), …].
    Как UNO _lm_pp_fill_down_calculate_resolve_block_pairs, без sheet.
    """
    if not isinstance(block, dict):
        return []
    try:
        from libre_macros_param_codec import _normalize_fill_down_calculate_block

        block = _normalize_fill_down_calculate_block(block)
    except Exception:
        pass
    columns_in = block.get("columns")
    has_columns = isinstance(columns_in, (list, tuple)) and len(columns_in) > 0
    rules = block.get("rules") or []
    default_formula = unicode(block.get("formula") or u"").strip()

    formula_by_col = {}
    expand_by_col = {}
    for r in rules:
        if not isinstance(r, dict):
            continue
        col_t = r.get("column")
        formula = unicode(r.get("formula") or u"").strip()
        if col_t is None or formula == u"":
            continue
        idxs = _direct_pp_resolve_col_tokens([col_t], header_h2i, end_col0)
        if not idxs:
            continue
        formula_by_col[int(idxs[0])] = formula
        if r.get("expand_to_right"):
            expand_by_col[int(idxs[0])] = True

    pairs = []
    if has_columns:
        for col_t in columns_in:
            idxs = _direct_pp_resolve_col_tokens([col_t], header_h2i, end_col0)
            if not idxs:
                pairs.append((col_t, None, None, False))
                continue
            idx = int(idxs[0])
            formula = formula_by_col.get(idx)
            if formula is None and default_formula:
                formula = default_formula
            if formula is None and len(rules) == 1 and isinstance(rules[0], dict):
                formula = unicode(rules[0].get("formula") or u"").strip()
            expand = bool(expand_by_col.get(idx, False))
            if formula is None and len(rules) == 1 and isinstance(rules[0], dict):
                expand = expand or bool(rules[0].get("expand_to_right"))
            pairs.append((col_t, formula or None, idx, expand))
        return pairs

    for r in rules:
        if not isinstance(r, dict):
            continue
        col_t = r.get("column")
        formula = unicode(r.get("formula") or u"").strip()
        if col_t is None or formula == u"":
            continue
        idxs = _direct_pp_resolve_col_tokens([col_t], header_h2i, end_col0)
        idx = idxs[0] if idxs else None
        pairs.append((col_t, formula, idx, bool(r.get("expand_to_right"))))
    return pairs


def _direct_pp_fill_down_calc_write_value(template, row_0, col_0, header_h2i, for_xlsx=False, sheet_name=u""):
    """
    Шаблон → значение ячейки для файла.
    Формула (=…) — развернуть {row}/{col}/[Заголовок]{row}/R1C1 + <<Переменные…>>.
    Иначе — константа (строка).
    """
    raw = _direct_pp_fill_down_calc_normalize_template(template)
    if raw == u"":
        return u"", False
    if not raw.startswith(u"="):
        return raw, False
    expanded = _direct_expand_header_refs(raw, header_h2i)
    ftxt = _direct_pp_expand_formula(expanded, row_0, col_0, sheet_name=sheet_name)
    if ftxt == u"":
        ftxt = expanded
    if ftxt and not unicode(ftxt).startswith(u"="):
        ftxt = u"=" + unicode(ftxt)
    ftxt = _direct_pp_prepare_formula_for_file(ftxt, for_xlsx=bool(for_xlsx))
    return unicode(ftxt), True


# --- сводная_таблица (file): материализованный отчёт без DataPilot ---


def _direct_pp_pivot_parse_config(extra_raw):
    """JSON сводной → dict (как lm_pp_pivot_config_from_json)."""
    raw = unicode(extra_raw or u"").strip()
    if raw == u"":
        return None
    try:
        from libre_macros_pivot_lib import lm_pp_pivot_config_from_json

        return lm_pp_pivot_config_from_json(raw)
    except Exception:
        pass
    try:
        from libre_macros_param_codec import param_decode

        blocks = param_decode(u"сводная_таблица", raw)
        if blocks and isinstance(blocks[0], dict):
            return blocks[0]
    except Exception:
        pass
    return None


def _direct_pp_pivot_source_matches(extra_raw, sheet_name):
    """Фильтр source_sheet (как UNO _lm_pp_pivot_source_sheet_matches)."""
    cfg = _direct_pp_pivot_parse_config(extra_raw)
    if cfg is None:
        return False
    spec = unicode(cfg.get("source_sheet") or u"").strip()
    if spec == u"":
        return True
    actual = unicode(sheet_name or u"").strip()
    for part in re.split(r"[;,|]", spec):
        p = unicode(part or u"").strip()
        if p != u"" and _text_match(actual, p):
            return True
        m = re.match(r"^(\d+)_(.+)$", actual)
        if m is not None and p != u"" and _text_match(m.group(2), p):
            return True
    return False


def _direct_pp_pivot_parse_corner(spec):
    """A1 / K100 / 2:1 / цифра-строка → (col0|None, row0|None) или None."""
    s = unicode(spec or u"").strip()
    if s == u"":
        return None
    m = re.match(r"^([A-Za-z]+)(\d+)$", s, re.IGNORECASE)
    if m:
        return _col_letters_to_index(m.group(1)), int(m.group(2)) - 1
    m = re.match(r"^(\d+):(\d+)$", s)
    if m:
        return int(m.group(2)) - 1, int(m.group(1)) - 1
    if s.isdigit():
        return None, int(s) - 1
    return None


def _direct_pp_pivot_header_row_index(header_row):
    hr = unicode(header_row or u"").strip()
    if hr.isdigit():
        return max(int(hr) - 1, 0)
    return 0


def _direct_pp_pivot_resolve_field_index(titles, field_name):
    """Имя поля → индекс в titles (casefold / text_match)."""
    needle = unicode(field_name or u"").strip()
    if needle == u"":
        return None
    nkey = _merge_identity_key(needle)
    i = 0
    while i < len(titles):
        if _merge_identity_key(titles[i]) == nkey:
            return i
        i += 1
    i = 0
    while i < len(titles):
        if _text_match(titles[i], needle):
            return i
        i += 1
    return None


def _direct_pp_pivot_norm_key_part(v):
    """Ключ группировки: пусто/None → '' (для шапки → [пусто])."""
    if v is None:
        return u""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            if float(v) == int(v):
                return unicode(int(v))
        except Exception:
            pass
        return unicode(v)
    s = unicode(v).strip()
    return s


def _direct_pp_pivot_header_format_value(text):
    t = unicode(text or u"").strip()
    if t == u"" or t == u"0":
        return _DIRECT_PP_PIVOT_HEADER_EMPTY
    return u" ".join(t.split())


def _direct_pp_pivot_build_col_header(col_fields, col_key):
    """Метка: значение\\Метка2: значение2 — как UNO collapse."""
    parts = []
    i = 0
    while i < len(col_fields):
        fname = unicode(col_fields[i] or u"").strip()
        val = col_key[i] if i < len(col_key) else u""
        parts.append(
            u"%s: %s" % (fname, _direct_pp_pivot_header_format_value(val))
        )
        i += 1
    return _DIRECT_PP_PIVOT_HEADER_DELIM.join(parts)


def _direct_pp_pivot_to_number(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = unicode(v).strip().replace(u" ", u"").replace(u",", u".")
    if s == u"":
        return None
    try:
        return float(s)
    except Exception:
        return None


def _direct_pp_pivot_aggregate(func_name, values):
    """Агрегация списка чисел (None уже отфильтрованы где нужно)."""
    fn = unicode(func_name or u"SUM").strip().upper()
    if fn in (u"STD_DEV",):
        fn = u"STDDEV"
    nums = []
    for v in values:
        n = _direct_pp_pivot_to_number(v)
        if n is not None:
            nums.append(n)
    if fn == u"COUNT":
        # COUNT как LO: число непустых исходных значений
        return float(len([x for x in values if not _direct_value_empty(x)]))
    if not nums:
        return None
    if fn == u"SUM":
        return float(sum(nums))
    if fn == u"AVERAGE":
        return float(sum(nums)) / float(len(nums))
    if fn == u"MAX":
        return float(max(nums))
    if fn == u"MIN":
        return float(min(nums))
    if fn == u"PRODUCT":
        p = 1.0
        for n in nums:
            p *= n
        return p
    if fn == u"STDDEV":
        if len(nums) < 2:
            return None
        mean = sum(nums) / float(len(nums))
        var = sum((x - mean) ** 2 for x in nums) / float(len(nums) - 1)
        return var ** 0.5
    if fn == u"VAR":
        if len(nums) < 2:
            return None
        mean = sum(nums) / float(len(nums))
        return sum((x - mean) ** 2 for x in nums) / float(len(nums) - 1)
    return float(sum(nums))


def _direct_pp_pivot_sort_key(parts):
    """
    Сортировка ключей: числа как числа, строки как строки, пусто в конец.
    Тег типа обязателен: иначе Python 3 падает на сравнении float < str
    (смешанные значения в одном поле, напр. «Поставщик»).
    """
    out = []
    for p in parts:
        s = unicode(p)
        if s == u"":
            # (пусто?, kind, payload) — пустые после непустых
            out.append((1, 0, u""))
            continue
        n = _direct_pp_pivot_to_number(s)
        if n is not None:
            out.append((0, 0, n))  # kind 0 = число
        else:
            out.append((0, 1, s.casefold()))  # kind 1 = строка
    return tuple(out)


def _direct_pp_pivot_compute(headers, rows, config):
    """
    Материализованная сводная: headers[list], rows[list of list] →
    (header_row[list], data_rows[list of list], n_row_fields, note).
    Макет как после UNO as_values + collapse: одна строка шапки, tabular, grand total.
    """
    cfg = config or {}
    row_fields = [unicode(x).strip() for x in (cfg.get("row_fields") or []) if unicode(x).strip()]
    col_fields = [
        unicode(x).strip() for x in (cfg.get("column_fields") or []) if unicode(x).strip()
    ]
    filter_fields = [
        unicode(x).strip() for x in (cfg.get("filter_fields") or []) if unicode(x).strip()
    ]
    data_fields_in = cfg.get("data_fields") or []
    measures = []
    for item in data_fields_in:
        if not isinstance(item, dict):
            continue
        field = unicode(item.get("field") or u"").strip()
        if field == u"":
            continue
        func = unicode(item.get("function") or item.get("fn") or u"SUM").strip()
        measures.append((field, func or u"SUM"))
    if not measures:
        return None, None, 0, u"нет data_fields"

    h2i = {}
    hi = 0
    while hi < len(headers):
        h2i[_merge_identity_key(headers[hi])] = hi
        hi += 1

    def _idx(name):
        return _direct_pp_pivot_resolve_field_index(headers, name)

    row_idxs = []
    for nm in row_fields:
        ix = _idx(nm)
        if ix is None:
            return None, None, 0, u"поле строк «%s» не найдено" % nm
        row_idxs.append(ix)
    col_idxs = []
    for nm in col_fields:
        ix = _idx(nm)
        if ix is None:
            return None, None, 0, u"поле столбцов «%s» не найдено" % nm
        col_idxs.append(ix)
    filter_idxs = []
    for nm in filter_fields:
        ix = _idx(nm)
        if ix is None:
            return None, None, 0, u"поле фильтра «%s» не найдено" % nm
        filter_idxs.append(ix)
    measure_idxs = []
    for field, func in measures:
        ix = _idx(field)
        if ix is None:
            return None, None, 0, u"поле данных «%s» не найдено" % field
        measure_idxs.append((ix, field, func))

    # filter_fields: в файле без Page — строки с пустым значением фильтра пропускаем
    # (как «все значения» Page; пустые ключи фильтра не отбрасываем — только отсутствие столбца уже проверено)

    # buckets[(row_key, col_key)][mi] = list of raw values
    buckets = {}
    row_keys_set = set()
    col_keys_set = set()
    for row in rows:
        if filter_idxs:
            # не фильтруем по значению — Page в LO показывает все; оставляем все строки
            pass
        rkey = tuple(
            _direct_pp_pivot_norm_key_part(row[i] if i < len(row) else None)
            for i in row_idxs
        )
        ckey = tuple(
            _direct_pp_pivot_norm_key_part(row[i] if i < len(row) else None)
            for i in col_idxs
        )
        row_keys_set.add(rkey)
        col_keys_set.add(ckey)
        key = (rkey, ckey)
        if key not in buckets:
            buckets[key] = [[] for _ in measure_idxs]
        mi = 0
        while mi < len(measure_idxs):
            ix, _f, _fn = measure_idxs[mi]
            buckets[key][mi].append(row[ix] if ix < len(row) else None)
            mi += 1

    row_keys = sorted(row_keys_set, key=_direct_pp_pivot_sort_key)
    col_keys = sorted(col_keys_set, key=_direct_pp_pivot_sort_key)
    if not col_fields:
        col_keys = [()]

    n_row = len(row_fields)
    n_meas = len(measures)
    # шапка: row field names + column headers (× measures)
    header = list(row_fields)
    for ck in col_keys:
        if col_fields:
            base = _direct_pp_pivot_build_col_header(col_fields, ck)
        else:
            base = u""
        mi = 0
        while mi < n_meas:
            field, func = measures[mi]
            if n_meas > 1:
                label = u"%s (%s)" % (field, unicode(func).upper())
                if base:
                    hdr = base + _DIRECT_PP_PIVOT_HEADER_DELIM + label
                else:
                    hdr = label
            else:
                hdr = base if base else field
            header.append(hdr)
            mi += 1

    def _cell_value(rkey, ckey, mi):
        bucket = buckets.get((rkey, ckey))
        if bucket is None:
            return None
        return _direct_pp_pivot_aggregate(measure_idxs[mi][2], bucket[mi])

    data_out = []
    for rkey in row_keys:
        line = list(rkey)
        # tabular: пустые ключи оставляем как есть (''), в ячейке — '')
        for ck in col_keys:
            mi = 0
            while mi < n_meas:
                line.append(_cell_value(rkey, ck, mi))
                mi += 1
        data_out.append(line)

    # Grand total — только при полях строк. Без row_fields единственная data-строка
    # уже является полным итогом; метка «Итого» давала фантомный сдвиг столбцов.
    if data_out and n_row > 0:
        total = [_DIRECT_PP_PIVOT_TOTAL_LABEL]
        ti = 1
        while ti < n_row:
            total.append(u"")
            ti += 1
        for ck in col_keys:
            mi = 0
            while mi < n_meas:
                all_vals = []
                for rkey in row_keys:
                    bucket = buckets.get((rkey, ck))
                    if bucket is not None:
                        all_vals.extend(bucket[mi])
                total.append(
                    _direct_pp_pivot_aggregate(measure_idxs[mi][2], all_vals)
                )
                mi += 1
        data_out.append(total)

    note = u"строк=%d, столбцов=%d, мер=%d" % (
        len(row_keys),
        max(1, len(col_keys)) * n_meas,
        n_meas,
    )
    if filter_fields:
        note += u"; filter_fields учтены как Page(все)"
    return header, data_out, n_row, note


def _direct_pp_pivot_read_xlsx_matrix(ws, config, bounds):
    """Прочитать headers+rows с листа xlsx по config/bounds."""
    hdr_row0 = int(bounds["hdr_row"])
    first_data0 = int(bounds["first_data"])
    last_data0 = int(bounds["last_data"])
    end_col0 = int(bounds["end_col"])
    sc, sr, ec, er = 0, hdr_row0, end_col0, last_data0
    cfg_start = unicode(config.get("data_start") or u"").strip()
    cfg_end = unicode(config.get("data_end") or u"").strip()
    cfg_hdr = unicode(config.get("header_row") or u"").strip()
    if cfg_start or cfg_end or cfg_hdr:
        start = _direct_pp_pivot_parse_corner(cfg_start)
        end = _direct_pp_pivot_parse_corner(cfg_end)
        hr = _direct_pp_pivot_header_row_index(cfg_hdr)
        if start is not None:
            if start[0] is not None:
                sc = int(start[0])
            if start[1] is not None:
                sr = int(start[1])
        if end is not None:
            if end[0] is not None:
                ec = int(end[0])
            if end[1] is not None:
                er = int(end[1])
        if cfg_hdr:
            sr = min(sr, hr)
            hdr_row0 = hr
        else:
            hdr_row0 = sr
    else:
        hdr_row0 = sr
    if ec < sc or er < hdr_row0:
        return None, None, u"нет данных"
    headers = []
    ci = sc
    while ci <= ec:
        v = ws.cell(row=hdr_row0 + 1, column=ci + 1).value
        headers.append(unicode(v if v is not None else u"").strip())
        ci += 1
    rows = []
    rr = hdr_row0 + 1
    while rr <= er:
        line = []
        ci = sc
        empty = True
        while ci <= ec:
            v = ws.cell(row=rr + 1, column=ci + 1).value
            if not _direct_value_empty(v):
                empty = False
            line.append(v)
            ci += 1
        if not empty:
            rows.append(line)
        rr += 1
    if not rows:
        return headers, [], u"нет строк данных"
    return headers, rows, u""


def _direct_pp_pivot_read_ods_matrix(rows_elems, config, bounds):
    """Прочитать headers+rows с ODS TableRow list."""
    hdr_row0 = int(bounds["hdr_row"])
    last_data0 = int(bounds["last_data"])
    end_col0 = int(bounds["end_col"])
    sc, sr, ec, er = 0, hdr_row0, end_col0, last_data0
    cfg_start = unicode(config.get("data_start") or u"").strip()
    cfg_end = unicode(config.get("data_end") or u"").strip()
    cfg_hdr = unicode(config.get("header_row") or u"").strip()
    if cfg_start or cfg_end or cfg_hdr:
        start = _direct_pp_pivot_parse_corner(cfg_start)
        end = _direct_pp_pivot_parse_corner(cfg_end)
        hr = _direct_pp_pivot_header_row_index(cfg_hdr)
        if start is not None:
            if start[0] is not None:
                sc = int(start[0])
            if start[1] is not None:
                sr = int(start[1])
        if end is not None:
            if end[0] is not None:
                ec = int(end[0])
            if end[1] is not None:
                er = int(end[1])
        if cfg_hdr:
            sr = min(sr, hr)
            hdr_row0 = hr
        else:
            hdr_row0 = sr
    else:
        hdr_row0 = sr
    if ec < sc or er < hdr_row0 or hdr_row0 >= len(rows_elems):
        return None, None, u"нет данных"
    by_hdr = dict(_ods_row_col_values(rows_elems[hdr_row0]))
    headers = []
    ci = sc
    while ci <= ec:
        v = by_hdr.get(ci)
        headers.append(unicode(v if v is not None else u"").strip())
        ci += 1
    rows = []
    rr = hdr_row0 + 1
    while rr <= er and rr < len(rows_elems):
        by_col = dict(_ods_row_col_values(rows_elems[rr]))
        line = []
        empty = True
        ci = sc
        while ci <= ec:
            v = by_col.get(ci)
            if not _direct_value_empty(v):
                empty = False
            line.append(v)
            ci += 1
        if not empty:
            rows.append(line)
        rr += 1
    if not rows:
        return headers, [], u"нет строк данных"
    return headers, rows, u""


def _direct_pp_pivot_apply_format_xlsx(ws, n_cols, n_data_rows, n_row_fields):
    """Оформление как UNO materialized: сетка, wrap, bold header/итог, числа, ширина."""
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    side = Side(style=u"thin", color=u"666666")
    border = Border(left=side, right=side, top=side, bottom=side)
    header_font = Font(bold=True)
    header_align = Alignment(horizontal=u"center", vertical=u"center", wrap_text=True)
    data_align = Alignment(vertical=u"center", wrap_text=True)
    last_row = 1 + int(n_data_rows)
    last_col = int(n_cols)
    rr = 1
    while rr <= last_row:
        cc = 1
        while cc <= last_col:
            cell = ws.cell(row=rr, column=cc)
            cell.border = border
            if rr == 1:
                cell.font = header_font
                cell.alignment = header_align
            else:
                cell.alignment = data_align
                if cc > int(n_row_fields):
                    cell.number_format = _DIRECT_PP_PIVOT_NUM_FMT
            cc += 1
        rr += 1
    if last_row > 1:
        # жирный итог
        first = unicode(ws.cell(row=last_row, column=1).value or u"").strip().casefold()
        if first in (
            u"итого",
            u"итог",
            u"total",
            u"grand total",
            u"total result",
            u"общий итог",
            u"всего",
            u"результат",
        ):
            cc = 1
            while cc <= last_col:
                ws.cell(row=last_row, column=cc).font = Font(bold=True)
                cc += 1
    # ширина ≤ 50 мм (~28 excel units)
    max_w = _DIRECT_PP_PIVOT_MAX_COL_WIDTH / 1.8
    cc = 1
    while cc <= last_col:
        letter = get_column_letter(cc)
        try:
            cur = ws.column_dimensions[letter].width
            if cur is None or cur > max_w:
                ws.column_dimensions[letter].width = max_w
        except Exception:
            ws.column_dimensions[letter].width = max_w
        cc += 1
    try:
        ws.auto_filter.ref = u"A1:%s%d" % (get_column_letter(last_col), last_row)
    except Exception:
        pass
    # freeze — в очередь UNO (openpyxl freeze часто игнорируется Calc)
    try:
        ws.freeze_panes = u"A2"
    except Exception:
        pass


def _direct_pp_pivot_apply_format_ods(doc, sheet, rows, n_cols, n_data_rows, n_row_fields):
    """ODS-оформление материализованной сводной."""
    cache = _direct_ods_style_cache_new(doc)
    last_row0 = int(n_data_rows)  # 0=header
    end_col0 = int(n_cols) - 1
    border_pt = _DIRECT_PP_ODS_BORDER_PT.get(u"thin", u"0.5pt")
    rr = 0
    while rr <= last_row0 and rr < len(rows):
        cells = _direct_ods_expand_row_to_cols(rows[rr], end_col0)
        ci = 0
        while ci <= end_col0:
            props = {
                u"border_pt": border_pt,
                u"border_color": u"666666",
                u"wrap": True,
            }
            if rr == 0:
                props[u"bold"] = True
                props[u"h_align"] = u"center"
            elif ci >= int(n_row_fields):
                props[u"data_style"] = u"money"
            _direct_ods_apply_cell_props(cache, cells[ci], props)
            ci += 1
        rr += 1
    if last_row0 > 0 and last_row0 < len(rows):
        by0 = dict(_ods_row_col_values(rows[last_row0]))
        first = unicode(by0.get(0) or u"").strip().casefold()
        if first in (
            u"итого",
            u"итог",
            u"total",
            u"grand total",
            u"total result",
            u"общий итог",
            u"всего",
            u"результат",
        ):
            cells = _direct_ods_expand_row_to_cols(rows[last_row0], end_col0)
            ci = 0
            while ci <= end_col0:
                _direct_ods_apply_cell_props(cache, cells[ci], {u"bold": True})
                ci += 1


def _direct_pp_xlsx_cell_is_empty(cell):
    """Пустая для заполнение_вниз_вычислить: нет значения и нет формулы."""
    if cell is None:
        return True
    try:
        v = cell.value
    except Exception:
        return True
    return _direct_value_empty(v)


def _direct_pp_ods_cell_is_empty(row_elem, col_index):
    """Пустая ODS-ячейка: нет текста/значения и нет formula=."""
    from odf.table import TableCell

    cells = list(row_elem.getElementsByType(TableCell))
    if col_index < 0 or col_index >= len(cells):
        return True
    cell = cells[col_index]
    try:
        fo = unicode(cell.getAttribute("formula") or u"").strip()
    except Exception:
        fo = u""
    if fo != u"":
        return False
    return _direct_value_empty(_ods_cell_value(cell))


def _direct_pp_file_structural(path, sheet_name, fn_key, extra):
    """Диспетчер file-structural: xlsx / ods."""
    key = _direct_pp_fn_key(fn_key)
    _direct_trace(u"file_structural", u"%s ctx=%s file=%s" % (key, sheet_name, path))
    if _is_ods(path):
        return _direct_pp_file_structural_ods(path, sheet_name, key, extra)
    return _direct_pp_file_structural_xlsx(path, sheet_name, key, extra)


def _direct_pp_file_structural_xlsx(path, sheet_name, fn_key, extra):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    key = _direct_pp_fn_key(fn_key)

    # rename / copy не требуют bounds данных
    if key == u"объединить_листы_в_один":
        block = _direct_pp_pick_sheet_block(u"объединить_листы_в_один", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        return _direct_pp_merge_sheets_into_one(path, block)

    if key == u"переименовать_лист":
        block = _direct_pp_pick_sheet_block(u"переименовать_лист", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        raw_name = unicode(block.get("new_name") or u"").strip()
        if raw_name == u"":
            return False, u"пустое new_name"
        wb = load_workbook(path)
        try:
            if sheet_name not in wb.sheetnames:
                return False, u"лист не найден"
            ws = wb[sheet_name]
            used_sr, used_er = _direct_pp_xlsx_used_rows(ws)
            hdr_row0 = 0
            first_data0 = max(0, hdr_row0 + 1)
            last_data0 = max(first_data0, used_er)
            try:
                bounds = _direct_result_sheet_bounds(path, sheet_name)
                if bounds is not None:
                    hdr_row0 = int(bounds.get("hdr_row") or 0)
                    first_data0 = int(bounds.get("first_data") or (hdr_row0 + 1))
                    last_data0 = int(bounds.get("last_data") or first_data0)
            except Exception:
                hdr_row0 = 0
                first_data0 = max(0, hdr_row0 + 1)
                last_data0 = max(first_data0, used_er)
            new_name = _direct_pp_expand_rename_sheet_name(
                raw_name,
                sheet_name=sheet_name,
                old_name=sheet_name,
                header_row=hdr_row0,
                used_sr=first_data0,
                used_er=last_data0,
                cell_text_fn=lambda r, c, _ws=ws: _direct_pp_xlsx_cell_text(_ws, r, c),
            )
            if new_name == u"":
                return False, u"пустое имя после подстановки"
            target = _direct_pp_finalize_sheet_name(
                wb.sheetnames, new_name, exclude_name=sheet_name
            )
            if target == u"":
                return False, u"пустое имя"
            if target == sheet_name:
                return True, u"имя без изменений"
            ws.title = target
            wb.save(path)
            return True, u"«%s» → «%s»" % (sheet_name, target)
        finally:
            try:
                wb.close()
            except Exception:
                pass

    if key == u"копировать_лист" or key == u"копировать_переместить_лист":
        block = _direct_pp_pick_sheet_block(u"копировать_лист", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        source = unicode(
            block.get("source_sheet") or block.get("source") or u""
        ).strip()
        dest = unicode(
            block.get("dest_sheet")
            or block.get("dest")
            or block.get("target_sheet")
            or u""
        ).strip()
        if source == u"" or dest == u"":
            return False, u"нужны source_sheet и dest_sheet"
        # columns/skip_empty/header_row: для XLSX копируем целиком и удаляем лишнее.
        _direct_trace(
            u"копировать_лист/xlsx",
            u"ctx=%s source=%s dest=%s file=%s" % (sheet_name, source, dest, path),
        )
        t0 = _direct_now()
        _direct_trace(u"копировать_лист/xlsx", u"load_workbook…")
        wb = load_workbook(path)
        _direct_trace(
            u"копировать_лист/xlsx",
            u"загружено за %.1fs, листов=%d" % (_direct_now() - t0, len(wb.sheetnames)),
        )
        try:
            src_actual = _direct_pp_find_sheet_name(wb.sheetnames, source)
            if src_actual is None:
                return False, u"лист-источник «%s» не найден" % source
            target = _direct_pp_unique_sheet_name(wb.sheetnames, dest, exclude_name=None)
            if target == u"":
                return False, u"пустое имя копии"
            if _merge_identity_key(src_actual) == _merge_identity_key(target):
                return True, u"«%s» уже существует" % target
            _direct_trace(
                u"копировать_лист/xlsx",
                u"copy_worksheet «%s» → «%s»…" % (src_actual, target),
            )
            t1 = _direct_now()
            new_ws = wb.copy_worksheet(wb[src_actual])
            _direct_trace(
                u"копировать_лист/xlsx",
                u"copy_worksheet за %.1fs" % (_direct_now() - t1),
            )
            new_ws.title = target
            # Если задан columns — оставляем только эти столбцы (быстро через delete_cols).
            columns_spec = block.get("columns") or []
            skip_empty = bool(block.get("skip_empty", False))
            hdr1 = block.get("header_row")
            hdr_row_1 = 1
            if hdr1 not in (None, u""):
                try:
                    hdr_row_1 = max(1, int(hdr1))
                except Exception:
                    hdr_row_1 = 1
            if columns_spec:
                try:
                    end_now = max(0, int(new_ws.max_column or 1) - 1)
                except Exception:
                    end_now = 0
                header_h2i = _direct_xlsx_header_h2i(new_ws, hdr_row_1=hdr_row_1, max_cols=500)
                keep0 = set(_direct_pp_resolve_col_tokens(columns_spec, header_h2i, end_now))
                if keep0:
                    _direct_trace(
                        u"копировать_лист/xlsx",
                        u"columns keep=%d (max_col=%d)" % (len(keep0), end_now),
                    )
                    # удалить лишние столбцы справа налево
                    col0 = end_now
                    while col0 >= 0:
                        if col0 not in keep0:
                            try:
                                new_ws.delete_cols(col0 + 1, 1)
                            except Exception:
                                pass
                        col0 -= 1
                    # skip_empty: после delete_cols индексы keep0 устарели — пересчитать по новому заголовку
                    if skip_empty:
                        try:
                            end_after = max(0, int(new_ws.max_column or 1) - 1)
                        except Exception:
                            end_after = 0
                        hdr_h2i = _direct_xlsx_header_h2i(
                            new_ws, hdr_row_1=hdr_row_1, max_cols=500
                        )
                        skip_tokens = block.get("skip_empty_columns") or []
                        if skip_tokens:
                            skip_cols_0 = _direct_pp_resolve_col_tokens(
                                skip_tokens, hdr_h2i, end_after
                            )
                        else:
                            ordered = _direct_pp_resolve_col_tokens(
                                columns_spec, hdr_h2i, end_after
                            )
                            skip_cols_0 = [ordered[0]] if ordered else [0]
                        if skip_cols_0:
                            try:
                                max_r = int(new_ws.max_row or hdr_row_1)
                            except Exception:
                                max_r = hdr_row_1
                            rr = max_r
                            while rr > hdr_row_1:
                                any_val = False
                                for c0 in skip_cols_0:
                                    try:
                                        v = new_ws.cell(
                                            row=rr, column=c0 + 1
                                        ).value
                                    except Exception:
                                        v = None
                                    if not _direct_value_empty(v) and unicode(
                                        v if v is not None else u""
                                    ).strip() != u"":
                                        any_val = True
                                        break
                                if not any_val:
                                    try:
                                        new_ws.delete_rows(rr, 1)
                                    except Exception:
                                        pass
                                rr -= 1
                else:
                    _direct_trace(u"копировать_лист/xlsx", u"columns: не нашли ни одного столбца")
            _direct_pp_track_created_sheet(target, header_row=0, source=src_actual)
            _direct_trace(u"копировать_лист/xlsx", u"save…")
            t2 = _direct_now()
            wb.save(path)
            _direct_trace(
                u"копировать_лист/xlsx",
                u"save за %.1fs, всего %.1fs" % (_direct_now() - t2, _direct_now() - t0),
            )
            return True, u"«%s» → «%s»" % (src_actual, target)
        finally:
            try:
                wb.close()
            except Exception:
                pass

    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    hdr_row0 = int(bounds["hdr_row"])
    first_data0 = int(bounds["first_data"])
    last_data0 = int(bounds["last_data"])
    end_col0 = int(bounds["end_col"])
    hdr_row_1 = hdr_row0 + 1
    first_data_1 = first_data0 + 1
    last_data_1 = last_data0 + 1

    wb = load_workbook(path)
    try:
        if sheet_name not in wb.sheetnames:
            return False, u"лист не найден"
        ws = wb[sheet_name]
        header_h2i = _direct_xlsx_header_h2i(ws, hdr_row_1=hdr_row_1)

        if key == u"переставить_столбцы":
            block = _direct_pp_pick_sheet_block(
                u"переставить_столбцы", extra, sheet_name
            )
            if block is None:
                return False, u"лист не в фильтре sheet"
            try:
                max_c = int(ws.max_column or 1)
            except Exception:
                max_c = end_col0 + 1
            end_now = max(end_col0, max(0, max_c - 1))
            titles = _direct_pp_xlsx_header_titles(ws, hdr_row_1, end_now)
            new_order, insert_at, moving, plan_err, create_names = _direct_pp_plan_column_reorder(
                block, header_h2i, end_now, titles
            )
            if plan_err:
                return False, plan_err
            if _direct_debug_enabled():
                print(
                    u"  [xml переставить_столбцы] %s order=%s copy=%s create=%s"
                    % (
                        sheet_name,
                        new_order,
                        _direct_pp_block_is_copy_columns(block),
                        bool(create_names),
                    )
                )
            src_titles = []
            if insert_at is not None and moving:
                for mc in moving:
                    if titles is not None and mc < len(titles):
                        src_titles.append(unicode(titles[mc] or u""))
                    else:
                        src_titles.append(u"")
            last_row_1 = max(last_data_1, ws.max_row or last_data_1)
            if create_names:
                ok = _direct_pp_xlsx_insert_empty_columns(
                    ws, insert_at, create_names, hdr_row_1
                )
                action = u"новые пустые"
                n_done = len(create_names)
            elif insert_at is not None and moving:
                ok = _direct_pp_xlsx_copy_columns_insert(
                    ws, last_row_1, end_now, moving, insert_at, hdr_row_1
                )
                if ok:
                    _direct_pp_apply_copied_headers_xlsx(
                        ws,
                        hdr_row_1,
                        insert_at,
                        moving,
                        _direct_pp_parse_copy_new_names(block),
                        src_titles,
                        end_now + len(moving),
                    )
                action = u"скопировано"
                n_done = len(moving)
            else:
                ok = _direct_pp_xlsx_apply_column_order(
                    ws, last_row_1, end_now, new_order
                )
                action = u"переставлено"
                n_done = len(block.get("columns") or [])
            wb.save(path)
            return ok, u"%s столбцов: %d" % (action, n_done)

        if key == u"переименовать_столбцы":
            block = _direct_pp_pick_sheet_block(
                u"переименовать_столбцы", extra, sheet_name
            )
            if block is None:
                return False, u"лист не в фильтре sheet"
            mappings = block.get("mappings") or []
            if not mappings:
                return False, u"пустые mappings"
            try:
                from libre_macros_param_codec import expand_rename_columns_mapping_pairs
                mappings = expand_rename_columns_mapping_pairs(mappings)
            except Exception:
                pass
            if not mappings:
                return False, u"пустые mappings"
            try:
                import libre_macros_lib as lm
            except ImportError:
                lm = None
            used_sr, used_er = _direct_pp_xlsx_used_rows(ws)
            first_data0 = max(0, hdr_row_1)
            last_data0 = max(first_data0, used_er)
            existing = {}
            ci = 0
            while ci <= end_col0:
                title = _direct_pp_xlsx_cell_text(ws, hdr_row_1 - 1, ci)
                key_h = unicode(title or u"").strip().casefold()
                if key_h != u"":
                    existing[key_h] = existing.get(key_h, 0) + 1
                ci += 1
            renamed = 0
            for item in mappings:
                if not isinstance(item, dict):
                    continue
                old_name = unicode(item.get("old") or u"").strip()
                new_tmpl = unicode(item.get("new") or u"").strip()
                if old_name == u"" or new_tmpl == u"":
                    continue
                col_idx, col_err = _direct_pp_resolve_col_unique(
                    old_name, header_h2i, end_col0
                )
                if col_err:
                    return False, col_err
                actual_old = _direct_pp_xlsx_cell_text(ws, hdr_row_1 - 1, col_idx)
                if actual_old == u"":
                    actual_old = old_name
                expanded = _direct_pp_expand_rename_name_template(
                    new_tmpl,
                    sheet_name=sheet_name,
                    old_name=actual_old,
                    header_row=hdr_row_1 - 1,
                    used_sr=first_data0,
                    used_er=last_data0,
                    cell_text_fn=lambda r, c, _ws=ws: _direct_pp_xlsx_cell_text(_ws, r, c),
                    log_fn_key=u"переименовать_столбцы",
                ).strip()
                old_key = unicode(actual_old or u"").strip().casefold()

                def _col_taken(cand, _old_key=old_key, _existing=existing):
                    k = unicode(cand or u"").strip().casefold()
                    if k == u"":
                        return True
                    if k == _old_key:
                        return False
                    return _existing.get(k, 0) > 0

                if lm is not None:
                    new_name = lm._lm_pp_finalize_column_name(
                        expanded, is_taken=_col_taken
                    )
                else:
                    new_name = expanded
                if new_name == u"":
                    continue
                ws.cell(row=hdr_row_1, column=col_idx + 1).value = new_name
                if old_key != u"" and existing.get(old_key, 0) > 0:
                    existing[old_key] = existing.get(old_key, 1) - 1
                    if existing[old_key] <= 0:
                        existing.pop(old_key, None)
                nk = unicode(new_name).strip().casefold()
                existing[nk] = existing.get(nk, 0) + 1
                renamed += 1
            wb.save(path)
            return True, u"переименовано: %d" % renamed

        if key == u"удалить_столбцы":
            block = _direct_pp_pick_sheet_block(u"удалить_столбцы", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            markers = _direct_pp_normalize_markers_tokens(block.get("markers"))
            if not markers:
                markers = [u"маркер", u"marker"]
            try:
                max_c = int(ws.max_column or 1)
            except Exception:
                max_c = end_col0 + 1
            end_now = max(0, max_c - 1)
            titles = _direct_pp_xlsx_header_titles(ws, hdr_row_1, end_now)
            del_cols = sorted(
                _direct_pp_delete_markers_col_set(markers, end_now, titles),
                reverse=True,
            )
            total_deleted = 0
            for col_to_delete in del_cols:
                ws.delete_cols(col_to_delete + 1, 1)
                total_deleted += 1
            wb.save(path)
            return True, u"удалено столбцов: %d, маркеры: %s" % (
                total_deleted,
                u", ".join(unicode(m) for m in markers),
            )

        if key == u"конкатенация_столбцов":
            block = _direct_pp_pick_sheet_block(
                u"конкатенация_столбцов", extra, sheet_name
            )
            if block is None:
                return False, u"лист не в фильтре sheet"
            col_name = unicode(block.get("new_column") or u"").strip()
            # separator: не strip — пробелы могут быть значимы
            delim = unicode(block.get("separator") or u"")
            indices = block.get("columns") or []
            if col_name == u"" or not indices:
                return False, u"нужны new_column и columns"
            src_cols0 = _direct_pp_resolve_col_tokens(indices, header_h2i, end_col0)
            if not src_cols0:
                return False, u"столбцы не найдены"
            target_col0 = _direct_xlsx_header_col_index(
                ws, col_name, hdr_row_1=hdr_row_1
            )
            ws.cell(row=hdr_row_1, column=target_col0 + 1).value = col_name
            rr = first_data_1
            n = 0
            while rr <= last_data_1:
                parts = []
                for c0 in src_cols0:
                    v = ws.cell(row=rr, column=c0 + 1).value
                    parts.append(u"" if v is None else unicode(v))
                ws.cell(row=rr, column=target_col0 + 1).value = delim.join(parts)
                n += 1
                rr += 1
            wb.save(path)
            return True, u"столбец %d «%s», строк %d" % (target_col0 + 1, col_name, n)

        if key == u"замена_значений":
            block = _direct_pp_pick_sheet_block(u"замена_значений", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            cols_in = block.get("columns") or []
            col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
            if not col_indices:
                return False, u"столбцы не найдены"
            need_end = max(int(end_col0), max(col_indices) if col_indices else 0)
            titles = _direct_pp_xlsx_header_titles(ws, hdr_row_1, need_end)
            width = need_end + 1
            headers = list(titles)
            while len(headers) < width:
                headers.append(u"")
            matrix = []
            rr = first_data_1
            while rr <= last_data_1:
                row = []
                c = 0
                while c < width:
                    try:
                        row.append(ws.cell(row=rr, column=c + 1).value)
                    except Exception:
                        row.append(None)
                    c += 1
                matrix.append(row)
                rr += 1
            orig = []
            for ci in col_indices:
                snap = []
                ri = 0
                while ri < len(matrix):
                    try:
                        snap.append(matrix[ri][ci] if ci < len(matrix[ri]) else None)
                    except Exception:
                        snap.append(None)
                    ri += 1
                orig.append((ci, snap))
            ok, payload = _direct_pp_replace_values_apply_block(
                block, matrix, headers, col_indices
            )
            if not ok:
                return False, payload
            changed_cells = 0
            for ci, snap in orig:
                ri = 0
                while ri < len(matrix):
                    try:
                        new_v = matrix[ri][ci] if ci < len(matrix[ri]) else None
                    except Exception:
                        new_v = None
                    old_v = snap[ri] if ri < len(snap) else None
                    if new_v != old_v:
                        try:
                            ws.cell(
                                row=first_data_1 + ri, column=ci + 1
                            ).value = new_v
                        except Exception:
                            pass
                        changed_cells += 1
                    ri += 1
            wb.save(path)
            note = u"изменено ячеек: %d, столбцов: %d" % (
                changed_cells,
                len(col_indices),
            )
            try:
                aw = (payload or {}).get(u"align_warn")
                if aw:
                    note = u"%s; предупреждение: %s" % (note, aw)
            except Exception:
                pass
            return True, note

        if key == u"текстовые_операции":
            from libre_macros_normalize_lib import (
                NormalizeOpError,
                build_row_dict,
                coerce_ops_list,
                compile_ops_compute_fns,
                compile_ops_row_filters,
            )

            block = _direct_pp_pick_sheet_block(u"текстовые_операции", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            cols_in = block.get("columns") or []
            col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
            if not col_indices:
                return False, u"столбцы не найдены"
            ops = coerce_ops_list(block.get("ops"))
            if not ops:
                return False, u"пустой стек ops"
            non_text = unicode(block.get("non_text") or u"skip").strip() or u"skip"
            on_error = unicode(block.get("on_error") or u"keep").strip() or u"keep"
            empty_cells = unicode(block.get("empty_cells") or u"skip").strip() or u"skip"
            try:
                crypto_key = _direct_pp_normalize_resolve_key(
                    block, wb=wb, sheet_name=sheet_name
                )
            except NormalizeOpError as err:
                return False, unicode(err)
            try:
                row_filters = compile_ops_row_filters(ops)
            except Exception as err:
                return False, u"фильтр строк: %s" % err
            try:
                compute_fns = compile_ops_compute_fns(ops)
            except Exception as err:
                return False, u"вычислить: %s" % err
            headers = _direct_pp_xlsx_header_titles(ws, hdr_row_1, end_col0)
            from libre_macros_lambda_column_lib import make_rows_accessor, make_upd_rows_store

            matrix = []
            rr0 = first_data_1
            while rr0 <= last_data_1:
                values = []
                c = 0
                while c <= end_col0:
                    values.append(ws.cell(row=rr0, column=c + 1).value)
                    c += 1
                matrix.append(values)
                rr0 += 1
            rows_accessor = make_rows_accessor(matrix, headers, start_col=0)
            changed_cells = 0
            error_cells = 0
            for ci in col_indices:
                upd_store = make_upd_rows_store()
                rr = first_data_1
                while rr <= last_data_1:
                    cell = ws.cell(row=rr, column=ci + 1)
                    v = cell.value
                    if non_text == u"skip" and not _direct_pp_cell_is_plain_string(v):
                        if v is not None and unicode(v).strip() != u"":
                            rr += 1
                            continue
                    row_i = rr - first_data_1
                    values = matrix[row_i] if 0 <= row_i < len(matrix) else []
                    row_dict = build_row_dict(headers, values, row_i)
                    try:
                        new_v, changed, had_err = _direct_pp_normalize_apply_cell(
                            v,
                            ops,
                            crypto_key,
                            non_text,
                            on_error,
                            empty_cells,
                            row=row_dict,
                            row_filters=row_filters,
                            row_id=row_i,
                            rows_accessor=rows_accessor,
                            upd_store=upd_store,
                            compute_fns=compute_fns,
                        )
                    except NormalizeOpError as err:
                        return False, unicode(err)
                    if had_err:
                        error_cells += 1
                        if on_error == u"stop":
                            wb.save(path)
                            return False, u"ошибка op (on_error=stop)"
                    if changed:
                        cell.value = new_v
                        try:
                            if 0 <= row_i < len(matrix) and ci < len(matrix[row_i]):
                                matrix[row_i][ci] = new_v
                        except Exception:
                            pass
                        changed_cells += 1
                    upd_store.commit(row_i)
                    rr += 1
            wb.save(path)
            return True, u"изменено ячеек: %d, ошибок: %d, столбцов: %d" % (
                changed_cells,
                error_cells,
                len(col_indices),
            )

        if key == u"заполнение_вниз":
            block = _direct_pp_pick_sheet_block(u"заполнение_вниз", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            cols_in = block.get("columns") or []
            col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
            if not col_indices:
                return False, u"столбцы не найдены"
            if last_data_1 <= first_data_1:
                return False, u"нет строк данных"
            filled = 0
            for ci in col_indices:
                last_val = None
                has_last = False
                rr = first_data_1
                while rr <= last_data_1:
                    cell = ws.cell(row=rr, column=ci + 1)
                    v = cell.value
                    if _direct_value_empty(v):
                        if has_last:
                            cell.value = last_val
                            filled += 1
                    else:
                        last_val = v
                        has_last = True
                    rr += 1
            wb.save(path)
            return True, u"столбцов: %d, заполнено ячеек: %d" % (
                len(col_indices),
                filled,
            )

        if key == u"заполнить_вверх":
            block = _direct_pp_pick_sheet_block(u"заполнить_вверх", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            cols_in = block.get("columns") or []
            col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
            if not col_indices:
                return False, u"столбцы не найдены"
            if last_data_1 <= first_data_1:
                return False, u"нет строк данных"
            filled = 0
            for ci in col_indices:
                last_val = None
                has_last = False
                rr = last_data_1
                while rr >= first_data_1:
                    cell = ws.cell(row=rr, column=ci + 1)
                    v = cell.value
                    if _direct_value_empty(v):
                        if has_last:
                            cell.value = last_val
                            filled += 1
                    else:
                        last_val = v
                        has_last = True
                    rr -= 1
            wb.save(path)
            return True, u"столбцов: %d, заполнено ячеек: %d" % (
                len(col_indices),
                filled,
            )

        if key == u"заполнение_вниз_вычислить":
            block = _direct_pp_pick_sheet_block(
                u"заполнение_вниз_вычислить", extra, sheet_name
            )
            if block is None:
                return False, u"лист не в фильтре sheet"
            if last_data_1 < first_data_1:
                return False, u"нет строк данных"
            pairs = _direct_pp_fill_down_calc_pairs(block, header_h2i, end_col0)
            if not pairs:
                return False, u"нет заданий (columns/rules)"
            filled = 0
            skipped = 0
            for col_t, formula, idx, expand in pairs:
                if idx is None or formula is None:
                    skipped += 1
                    continue
                rr = first_data_1
                while rr <= last_data_1:
                    cell = ws.cell(row=rr, column=idx + 1)
                    if not _direct_pp_xlsx_cell_is_empty(cell):
                        rr += 1
                        continue
                    val, _is_f = _direct_pp_fill_down_calc_write_value(
                        formula, rr - 1, idx, header_h2i, for_xlsx=True, sheet_name=sheet_name
                    )
                    if val == u"" and formula.strip().startswith(u"="):
                        rr += 1
                        continue
                    cell.value = val
                    filled += 1
                    if expand:
                        c0 = idx + 1
                        while c0 <= end_col0:
                            ccell = ws.cell(row=rr, column=c0 + 1)
                            if not _direct_pp_xlsx_cell_is_empty(ccell):
                                c0 += 1
                                continue
                            cval, _is_cf = _direct_pp_fill_down_calc_write_value(
                                formula, rr - 1, c0, header_h2i, for_xlsx=True, sheet_name=sheet_name
                            )
                            if cval == u"" and formula.strip().startswith(u"="):
                                c0 += 1
                                continue
                            ccell.value = cval
                            filled += 1
                            c0 += 1
                    rr += 1
            wb.save(path)
            note = u"столбцов: %d, заполнено ячеек: %d" % (
                len(pairs) - skipped,
                filled,
            )
            if skipped:
                note += u", пропущено столбцов: %d" % skipped
            return True, note

        if key == u"сводная_таблица":
            config = _direct_pp_pivot_parse_config(extra)
            if config is None:
                return False, u"пустой/неверный JSON"
            headers, src_rows, err = _direct_pp_pivot_read_xlsx_matrix(
                ws, config, bounds
            )
            if err and not src_rows:
                return False, err
            hdr, data_rows, n_row_f, note = _direct_pp_pivot_compute(
                headers, src_rows, config
            )
            if hdr is None:
                return False, note
            dest_base = unicode(config.get("sheet_name") or u"Сводная").strip()
            if dest_base == u"":
                dest_base = u"Сводная"
            # заменить одноимённый лист, если есть
            existing = _direct_pp_find_sheet_name(wb.sheetnames, dest_base)
            if existing is not None and _merge_identity_key(existing) != _merge_identity_key(
                sheet_name
            ):
                try:
                    del wb[existing]
                except Exception:
                    pass
            dest = _direct_pp_unique_sheet_name(
                wb.sheetnames, dest_base, exclude_name=None
            )
            if dest == u"":
                return False, u"пустое имя листа сводной"
            if dest in wb.sheetnames:
                try:
                    del wb[dest]
                except Exception:
                    pass
            out_ws = wb.create_sheet(title=dest)
            cc = 0
            while cc < len(hdr):
                out_ws.cell(row=1, column=cc + 1).value = hdr[cc]
                cc += 1
            ri = 0
            while ri < len(data_rows):
                line = data_rows[ri]
                cc = 0
                while cc < len(line):
                    out_ws.cell(row=ri + 2, column=cc + 1).value = line[cc]
                    cc += 1
                # добить пустые до ширины шапки
                while cc < len(hdr):
                    out_ws.cell(row=ri + 2, column=cc + 1).value = None
                    cc += 1
                ri += 1
            _direct_pp_pivot_apply_format_xlsx(
                out_ws, len(hdr), len(data_rows), n_row_f
            )
            # freeze надёжнее через UNO после reopen
            try:
                if dest not in _direct_xml_pp_freeze_pending:
                    _direct_xml_pp_freeze_pending.append(dest)
            except Exception:
                pass
            _direct_pp_track_created_sheet(dest, header_row=0, source=sheet_name)
            wb.save(path)
            return True, u"лист «%s»; %s" % (dest, note)

        if key == u"сортировка":
            block = _direct_pp_pick_sheet_block(u"сортировка", extra, sheet_name)
            if block is None:
                return False, u"лист не в фильтре sheet"
            keys_in = block.get("keys") or []
            key_specs = []
            for item in keys_in:
                if not isinstance(item, dict):
                    continue
                col_tok = unicode(item.get("column") or u"").strip()
                if col_tok == u"":
                    continue
                idxs = _direct_pp_resolve_col_tokens([col_tok], header_h2i, end_col0)
                if not idxs:
                    continue
                key_specs.append((idxs[0], bool(item.get("desc"))))
            if not key_specs:
                return False, u"нет ключей сортировки"
            if last_data_1 < first_data_1:
                return False, u"нет строк данных"
            # собрать строки значений 0..end_col0
            rows_data = []
            rr = first_data_1
            while rr <= last_data_1:
                row_vals = []
                ci = 0
                while ci <= end_col0:
                    row_vals.append(ws.cell(row=rr, column=ci + 1).value)
                    ci += 1
                rows_data.append(row_vals)
                rr += 1
            sorted_rows = _direct_pp_sort_rows_stable(rows_data, key_specs)
            rr = first_data_1
            for row_vals in sorted_rows:
                ci = 0
                while ci <= end_col0:
                    ws.cell(row=rr, column=ci + 1).value = (
                        row_vals[ci] if ci < len(row_vals) else None
                    )
                    ci += 1
                rr += 1
            wb.save(path)
            return True, u"отсортировано строк: %d, ключей: %d" % (
                len(sorted_rows),
                len(key_specs),
            )

        return False, u"неизвестный file-structural: %s" % key
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_pp_ods_remove_col_from_row(row_elem, col_index):
    """Удалить логический столбец col_index из строки ODS (с expand)."""
    from odf.table import TableCell

    # узнать ширину
    max_col = 0
    for cidx, _v in _ods_row_col_values(row_elem):
        if cidx > max_col:
            max_col = cidx
    if col_index > max_col:
        return
    cells = _direct_ods_expand_row_to_cols(row_elem, max_col)
    if col_index < 0 or col_index >= len(cells):
        return
    # пересобрать без col_index
    for cell in list(row_elem.getElementsByType(TableCell)):
        try:
            row_elem.removeChild(cell)
        except Exception:
            pass
    ci = 0
    while ci < len(cells):
        if ci != col_index:
            row_elem.addElement(cells[ci])
        ci += 1


def _direct_pp_ods_clone_table(src_table, new_name):
    """Клон Table (deepcopy или ручное копирование строк/ячеек)."""
    import copy
    from odf.table import Table, TableRow, TableCell, TableColumn

    _direct_trace(u"копировать_лист/ods", u"deepcopy «%s»" % new_name)
    try:
        new_table = copy.deepcopy(src_table)
        new_table.setAttribute(u"name", new_name)
        _direct_trace(u"копировать_лист/ods", u"deepcopy ok")
        return new_table
    except Exception:
        _direct_trace(u"копировать_лист/ods", u"deepcopy failed — лёгкое копирование")
    new_table = Table(name=new_name)
    # колонки
    try:
        for col in src_table.getElementsByType(TableColumn):
            try:
                new_table.addElement(copy.deepcopy(col))
            except Exception:
                new_table.addElement(TableColumn())
    except Exception:
        pass

    def _copy_elem_attrs(src_elem, dst_elem):
        try:
            attrs = getattr(src_elem, "attributes", None)
            if not attrs:
                return
            for k, v in attrs.items():
                try:
                    dst_elem.setAttribute(k, v)
                except Exception:
                    pass
        except Exception:
            pass

    def _clone_cell_light(src_cell):
        """
        Копирование ячейки без deepcopy (быстро и безопасно).
        Сохраняем основные атрибуты (style/value/type/formula), текст — через _ods_cell_text.
        """
        from odf.table import TableCell
        from odf.text import P

        dst = TableCell()
        _copy_elem_attrs(src_cell, dst)
        try:
            # текст (один параграф) — достаточен для наших целей
            dst.addElement(P(text=unicode(_ods_cell_text(src_cell))))
        except Exception:
            pass
        return dst

    row_n = 0
    for row in src_table.getElementsByType(TableRow):
        if row_n == 0:
            _direct_trace(u"копировать_лист/ods", u"копирование строк…")
        if row_n % 50 == 0:
            _direct_ui_yield()
            if row_n > 0:
                _direct_trace(u"копировать_лист/ods", u"строк скопировано: %d" % row_n)
        new_row = TableRow()
        _copy_elem_attrs(row, new_row)
        for cell in row.getElementsByType(TableCell):
            try:
                new_row.addElement(_clone_cell_light(cell))
            except Exception:
                val = _ods_cell_value(cell)
                new_row.addElement(_ods_make_cell(val))
        new_table.addElement(new_row)
        row_n += 1
    _direct_trace(u"копировать_лист/ods", u"построчно: %d строк" % row_n)
    return new_table


def _direct_pp_ods_resolve_copy_columns(src_rows, hdr_row0, columns_spec):
    """
    columns_spec: list токенов (индекс 1-based/буквы/заголовок/шаблон).
    Возвращает list[(col_idx, title)] в порядке исходного листа.
    """
    if not columns_spec:
        return None
    if hdr_row0 < 0 or hdr_row0 >= len(src_rows):
        return None
    header_pairs = []
    for cidx, val in _ods_row_col_values(src_rows[hdr_row0]):
        header_pairs.append((int(cidx), unicode(val or u"").strip()))
    if not header_pairs:
        return None
    out = []
    seen = set()
    for tok in columns_spec:
        t = unicode(tok or u"").strip()
        if t == u"":
            continue
        # индекс 1-based
        if t.isdigit():
            try:
                ci0 = int(t) - 1
            except Exception:
                continue
            if ci0 < 0:
                continue
            if ci0 not in seen:
                title = u""
                for hc, h in header_pairs:
                    if hc == ci0:
                        title = h
                        break
                if title == u"":
                    title = u"COL_%d" % (ci0 + 1)
                out.append((ci0, title))
                seen.add(ci0)
            continue
        if _is_col_letters(t):
            ci0 = _col_letters_to_index(t)
            if ci0 >= 0 and ci0 not in seen:
                title = u""
                for hc, h in header_pairs:
                    if hc == ci0:
                        title = h
                        break
                if title == u"":
                    title = u"COL_%d" % (ci0 + 1)
                out.append((ci0, title))
                seen.add(ci0)
            continue
        # по заголовку / шаблону
        for hc, h in header_pairs:
            if hc in seen:
                continue
            if h != u"" and (_merge_identity_key(h) == _merge_identity_key(t) or _text_match(h, t)):
                out.append((hc, h))
                seen.add(hc)
                break
    return out if out else None


def _direct_pp_ods_copy_sheet_filtered( path, src_actual, target, header_row0, columns_spec, skip_empty, skip_empty_columns=None):
    """
    ODS: копия листа с выбором столбцов и пропуском пустых строк.
    """
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow

    skip_empty_columns = skip_empty_columns or []
    _direct_trace(
        u"копировать_лист/ods",
        u"filtered columns=%d skip_empty=%s skip_cols=%d"
        % (len(columns_spec or []), bool(skip_empty), len(skip_empty_columns or [])),
    )
    doc = load(path)
    tables = doc.spreadsheet.getElementsByType(Table)
    src_table = None
    for t in tables:
        if unicode(t.getAttribute("name") or u"") == unicode(src_actual):
            src_table = t
            break
    if src_table is None:
        return False, u"лист-источник «%s» не найден" % src_actual
    src_rows = src_table.getElementsByType(TableRow)
    hdr0 = int(header_row0)
    if hdr0 < 0:
        hdr0 = 0
    cols = _direct_pp_ods_resolve_copy_columns(src_rows, hdr0, columns_spec)
    if cols is None:
        return False, u"копировать_лист: columns не найдены по заголовкам"

    # skip_empty_columns: токены относительно результата (индексы/буквы/заголовки в *итоговом* наборе)
    def _resolve_skip_cols(tokens):
        if not tokens:
            return []
        h2i = {}
        i = 0
        while i < len(cols):
            h2i[_merge_identity_key(cols[i][1])] = i
            i += 1
        out_idx = []
        for tok in tokens:
            t = unicode(tok or u"").strip()
            if t == u"":
                continue
            if t.isdigit():
                try:
                    i0 = int(t) - 1
                except Exception:
                    continue
                if 0 <= i0 < len(cols) and i0 not in out_idx:
                    out_idx.append(i0)
                continue
            if _is_col_letters(t):
                i0 = _col_letters_to_index(t)
                if 0 <= i0 < len(cols) and i0 not in out_idx:
                    out_idx.append(i0)
                continue
            key = _merge_identity_key(t)
            if key in h2i:
                i0 = h2i[key]
                if i0 not in out_idx:
                    out_idx.append(i0)
        return out_idx

    skip_res_idx = _resolve_skip_cols(skip_empty_columns)
    if skip_empty and not skip_res_idx:
        skip_res_idx = [0]  # default: первая колонка результата

    def _copy_elem_attrs(src_elem, dst_elem):
        try:
            attrs = getattr(src_elem, "attributes", None)
            if not attrs:
                return
            for k, v in attrs.items():
                try:
                    dst_elem.setAttribute(k, v)
                except Exception:
                    pass
        except Exception:
            pass

    def _clone_cell_light(src_cell):
        """Клонировать ячейку с ODF-атрибутами (стили/тип/значение) — сохраняет форматы чисел/дат."""
        from odf.table import TableCell
        from odf.text import P

        dst = TableCell()
        _copy_elem_attrs(src_cell, dst)
        try:
            dst.addElement(P(text=unicode(_ods_cell_text(src_cell))))
        except Exception:
            pass
        return dst

    # создать новую таблицу
    new_table = Table(name=target)
    # header row
    hdr_row = TableRow()
    for _ci0, title in cols:
        hdr_row.addElement(_ods_make_cell(title))
    new_table.addElement(hdr_row)

    total_rows = 0
    r = hdr0 + 1
    while r < len(src_rows):
        if total_rows % 200 == 0:
            _direct_ui_yield()
        if total_rows % 500 == 0 and total_rows > 0:
            _direct_trace(u"копировать_лист/ods", u"filtered rows=%d" % total_rows)
        # expand row once to max kept col
        max_ci0 = max(ci0 for ci0, _t in cols) if cols else 0
        cells = _direct_ods_expand_row_to_cols(src_rows[r], max_ci0)
        # пустота строки: по выбранным колонкам результата (по умолчанию первая)
        any_val = True
        if skip_empty:
            any_val = False
            for res_i in skip_res_idx:
                src_ci0 = cols[res_i][0]
                if 0 <= src_ci0 < len(cells):
                    v = _ods_cell_value(cells[src_ci0])
                    if not _direct_value_empty(v) and unicode(v).strip() != u"":
                        any_val = True
                        break
        if (not skip_empty) or any_val:
            nr = TableRow()
            # копируем ячейки со стилями/типами
            for ci0, _title in cols:
                if 0 <= ci0 < len(cells):
                    nr.addElement(_clone_cell_light(cells[ci0]))
                else:
                    nr.addElement(_ods_make_cell(u""))
            new_table.addElement(nr)
            total_rows += 1
        r += 1

    doc.spreadsheet.addElement(new_table)
    doc.save(path)
    return True, u"filtered: cols=%d, rows=%d" % (len(cols), total_rows)


def _direct_pp_resolve_merge_sheet_names(path, sheets_spec):
    if isinstance(sheets_spec, (list, tuple)):
        parts = [unicode(x or u"").strip() for x in sheets_spec]
    else:
        parts = []
        for chunk in unicode(sheets_spec or u"").replace(u";", u",").split(u","):
            st = chunk.strip()
            if st != u"":
                parts.append(st)
    parts = [p for p in parts if p != u""]
    if not parts:
        return []
    try:
        all_names = (
            _ods_sheet_names(path) if _is_ods(path) else _xlsx_sheet_names(path)
        )
    except Exception:
        all_names = []
    out = []
    seen = set()
    pi = 0
    while pi < len(parts):
        part = parts[pi]
        pi += 1
        if u"*" in part or u"?" in part:
            ni = 0
            while ni < len(all_names):
                nm = all_names[ni]
                if _text_match(nm, part):
                    key = _merge_identity_key(nm)
                    if key not in seen:
                        seen.add(key)
                        out.append(nm)
                ni += 1
            continue
        actual = _direct_pp_find_sheet_name(all_names, part)
        if actual is not None:
            key = _merge_identity_key(actual)
            if key not in seen:
                seen.add(key)
                out.append(actual)
    return out


def _direct_pp_xlsx_empty_first_col(val):
    if val is None:
        return True
    if isinstance(val, (bytes, str, unicode)):
        return unicode(val).strip() == u""
    return False


def _ods_merge_header_cols(hdr_row_elem, end_col0, start_col0=0):
    """
    Столбцы заголовка ODS без разворота numbercolumnsrepeated.

    Иначе одна «пустая» ячейка с repeat=16380 даёт 16384 логических колонки.
    """
    cols = []
    if hdr_row_elem is None:
        return cols
    try:
        end_col0 = int(end_col0)
        start_col0 = int(start_col0)
    except (TypeError, ValueError):
        return cols
    if end_col0 < start_col0:
        return cols
    cell_map = _ods_row_cell_by_col(hdr_row_elem)
    ci = start_col0
    while ci <= end_col0:
        if ci in cell_map:
            val = _ods_cell_value(cell_map[ci])
        else:
            val = u""
        h = unicode(val or u"").strip()
        if h == u"":
            h = u"COL_%d" % (ci + 1)
        cols.append((ci, h))
        ci += 1
    return cols


def _direct_pp_xlsx_copy_merge_cell(src_cell, dst_cell, with_formatting=False):
    """
    Перенос ячейки при объединить_листы_в_один (xlsx / xml_excel_ods).

    Всегда: value + number_format (даты, деньги, числа).
    with_formatting: font/fill/border/alignment/protection.
    """
    if src_cell is None or dst_cell is None:
        return
    try:
        dst_cell.value = src_cell.value
    except Exception:
        pass
    try:
        nf = src_cell.number_format
        if nf is not None and unicode(nf).strip() != u"":
            dst_cell.number_format = nf
    except Exception:
        pass
    if not with_formatting:
        return
    try:
        from copy import copy as _copy

        if getattr(src_cell, "has_style", False):
            try:
                dst_cell._style = _copy(src_cell._style)
            except Exception:
                pass
        for attr in ("font", "fill", "border", "alignment", "protection"):
            try:
                setattr(dst_cell, attr, _copy(getattr(src_cell, attr)))
            except Exception:
                pass
    except Exception:
        pass


def _direct_pp_merge_sheets_into_one(path, block):
    sheets_spec = (
        (block or {}).get("sheets")
        or (block or {}).get("source_sheets")
        or (block or {}).get("sheet_list")
        or []
    )
    dest = unicode(
        (block or {}).get("dest_sheet")
        or (block or {}).get("result_sheet")
        or (block or {}).get("target_sheet")
        or u""
    ).strip()
    with_formatting = False
    try:
        from libre_macros_lib import lm_parse_bool_param

        with_formatting = bool(
            lm_parse_bool_param((block or {}).get("with_formatting"), default=False)
        )
    except Exception:
        try:
            with_formatting = bool((block or {}).get("with_formatting"))
        except Exception:
            with_formatting = False
    _direct_trace(
        u"объединить_листы_в_один",
        u"file=%s dest=%r sheets_spec=%r header_row=%r with_formatting=%s"
        % (
            path,
            dest,
            sheets_spec,
            (block or {}).get("header_row"),
            with_formatting,
        ),
    )
    if dest == u"":
        return False, u"нужен dest_sheet"
    source_names = _direct_pp_resolve_merge_sheet_names(path, sheets_spec)
    if not source_names:
        _direct_trace(u"объединить_листы_в_один", u"нет листов по sheets")
        return False, u"нет листов по sheets"
    _direct_trace(
        u"объединить_листы_в_один",
        u"resolved sources: %s" % u",".join(source_names),
    )
    try:
        hdr0 = int((block or {}).get("header_row") or 1) - 1
    except (TypeError, ValueError):
        hdr0 = 0
    if hdr0 < 0:
        hdr0 = 0
    hdr_1 = hdr0 + 1
    _direct_trace(
        u"объединить_листы_в_один",
        u"header_row 1-based=%d (0-based=%d)" % (hdr_1, hdr0),
    )

    if _is_ods(path):
        import odf_bundled  # noqa: F401
        from odf.opendocument import load
        from odf.table import Table, TableRow, TableCell
        from odf.text import P

        _direct_trace(
            u"объединить_листы_в_один/ods",
            u"sources=%s dest=%s with_formatting=%s"
            % (u",".join(source_names), dest, with_formatting),
        )
        doc = load(path)
        tables = doc.spreadsheet.getElementsByType(Table)
        names = [unicode(t.getAttribute("name") or u"") for t in tables]
        dest_actual = _direct_pp_unique_sheet_name(names, dest, exclude_name=None)
        dest_table = None
        for t in tables:
            if unicode(t.getAttribute("name") or u"") == dest_actual:
                dest_table = t
                break
        if dest_table is None:
            dest_table = Table(name=dest_actual)
            doc.spreadsheet.addElement(dest_table)
            _direct_trace(
                u"объединить_листы_в_один/ods",
                u"создан лист «%s»" % dest_actual,
            )
        else:
            rows = dest_table.getElementsByType(TableRow)
            _direct_trace(
                u"объединить_листы_в_один/ods",
                u"очистка «%s», было строк=%d" % (dest_actual, len(rows)),
            )
            ri = len(rows) - 1
            while ri >= 0:
                try:
                    dest_table.removeChild(rows[ri])
                except Exception:
                    pass
                ri -= 1

        def _ods_empty_cell():
            cell = TableCell()
            cell.addElement(P(text=u""))
            return cell

        def _ods_dest_header_row():
            dest_rows = dest_table.getElementsByType(TableRow)
            while len(dest_rows) <= hdr0:
                dest_table.addElement(TableRow())
                dest_rows = dest_table.getElementsByType(TableRow)
            return dest_rows[hdr0]

        h2i = {}
        headers_order = []
        dest_hdr_row = None
        total_rows = 0
        si = 0
        while si < len(source_names):
            src_name = source_names[si]
            src_table = None
            for t in tables:
                if unicode(t.getAttribute("name") or u"") == src_name:
                    src_table = t
                    break
            if src_table is None:
                return False, u"лист «%s» не найден" % src_name
            src_rows = src_table.getElementsByType(TableRow)
            end_col, end_row = _direct_scan_bounds_ods(src_rows, hdr0, 0)
            cols = []
            if end_col >= 0 and hdr0 < len(src_rows):
                cols = _ods_merge_header_cols(src_rows[hdr0], end_col, 0)
            if not cols:
                _direct_trace(
                    u"объединить_листы_в_один/ods",
                    u"«%s» пропуск — нет столбцов" % src_name,
                )
                si += 1
                continue
            first_col = cols[0][0]
            col_hdr = u",".join(c[1] for c in cols[:10])
            if len(cols) > 10:
                col_hdr += u",…"
            _direct_trace(
                u"объединить_листы_в_один/ods",
                u"«%s» cols=%d end_col=%d end_row=%d headers=[%s]"
                % (src_name, len(cols), end_col, end_row, col_hdr),
            )
            sheet_rows = 0
            if si == 0:
                ci = 0
                while ci < len(cols):
                    _c1, hname = cols[ci]
                    k = _merge_identity_key(hname)
                    if k not in h2i:
                        h2i[k] = len(headers_order)
                        headers_order.append(hname)
                    ci += 1
                dest_hdr_row = _ods_dest_header_row()
                for cell in list(dest_hdr_row.getElementsByType(TableCell)):
                    try:
                        dest_hdr_row.removeChild(cell)
                    except Exception:
                        pass
                ci = 0
                while ci < len(headers_order):
                    dest_hdr_row.addElement(_ods_make_cell(headers_order[ci]))
                    ci += 1
            elif dest_hdr_row is None:
                dest_hdr_row = _ods_dest_header_row()

            r = hdr0 + 1
            yield_n = 0
            while r <= end_row:
                if r >= len(src_rows):
                    break
                src_cell_map = _ods_row_cell_by_col(src_rows[r])
                fc = src_cell_map.get(first_col)
                if fc is None:
                    fc_empty = True
                else:
                    fc_empty = _direct_pp_xlsx_empty_first_col(_ods_cell_value(fc))
                if fc_empty:
                    r += 1
                    continue
                row_slots = [None] * len(headers_order)
                ci = 0
                while ci < len(cols):
                    c1, hname = cols[ci]
                    k = _merge_identity_key(hname)
                    if k not in h2i:
                        h2i[k] = len(headers_order)
                        headers_order.append(hname)
                        row_slots.append(None)
                        if dest_hdr_row is not None:
                            dest_hdr_row.addElement(_ods_make_cell(hname))
                    dest_c = h2i[k]
                    while len(row_slots) <= dest_c:
                        row_slots.append(None)
                    if c1 in src_cell_map:
                        row_slots[dest_c] = _direct_pp_ods_clone_cell_light(
                            src_cell_map[c1]
                        )
                    ci += 1
                new_row = TableRow()
                ci = 0
                while ci < len(row_slots):
                    cell = row_slots[ci]
                    if cell is None:
                        new_row.addElement(_ods_empty_cell())
                    else:
                        new_row.addElement(cell)
                    ci += 1
                dest_table.addElement(new_row)
                total_rows += 1
                sheet_rows += 1
                r += 1
                yield_n += 1
                if yield_n >= 100:
                    _direct_ui_yield()
                    yield_n = 0
            _direct_trace(
                u"объединить_листы_в_один/ods",
                u"«%s» скопировано строк=%d" % (src_name, sheet_rows),
            )
            si += 1
        _direct_pp_track_created_sheet(dest_actual, header_row=hdr0, source=u",".join(source_names))
        doc.save(path)
        note = u"объединено %d лист(ов), %d строк → «%s»%s" % (
            len(source_names),
            total_rows,
            dest_actual,
            u" +формат" if with_formatting else u"",
        )
        _direct_trace(u"объединить_листы_в_один/ods", note)
        return True, note

    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    _direct_trace(
        u"объединить_листы_в_один/xlsx",
        u"sources=%s dest=%s with_formatting=%s"
        % (u",".join(source_names), dest, with_formatting),
    )
    wb = load_workbook(path)
    try:
        dest_actual = _direct_pp_unique_sheet_name(wb.sheetnames, dest, exclude_name=None)
        if dest_actual in wb.sheetnames:
            ws_dest = wb[dest_actual]
            if ws_dest.max_row:
                _direct_trace(
                    u"объединить_листы_в_один/xlsx",
                    u"очистка «%s», было строк=%d"
                    % (dest_actual, int(ws_dest.max_row or 0)),
                )
                ws_dest.delete_rows(1, ws_dest.max_row)
        else:
            ws_dest = wb.create_sheet(title=dest_actual)
            _direct_trace(
                u"объединить_листы_в_один/xlsx",
                u"создан лист «%s»" % dest_actual,
            )
        h2i = {}
        headers_order = []
        out_data_1 = hdr_1 + 1
        total_rows = 0
        si = 0
        while si < len(source_names):
            src_name = source_names[si]
            if src_name not in wb.sheetnames:
                return False, u"лист «%s» не найден" % src_name
            ws = wb[src_name]
            max_c = int(ws.max_column or 1)
            max_r = int(ws.max_row or hdr_1)
            cols = []
            ci = 1
            while ci <= max_c:
                v = ws.cell(row=hdr_1, column=ci).value
                h = _direct_clean_header_text(v) if v is not None else u""
                if h == u"":
                    h = u"COL_%d" % ci
                cols.append((ci, h))
                ci += 1
            if not cols:
                _direct_trace(
                    u"объединить_листы_в_один/xlsx",
                    u"«%s» пропуск — нет столбцов" % src_name,
                )
                si += 1
                continue
            first_col = cols[0][0]
            col_hdr = u",".join(c[1] for c in cols[:10])
            if len(cols) > 10:
                col_hdr += u",…"
            _direct_trace(
                u"объединить_листы_в_один/xlsx",
                u"«%s» cols=%d max_r=%d headers=[%s]"
                % (src_name, len(cols), max_r, col_hdr),
            )
            sheet_rows = 0
            if si == 0:
                ci = 0
                while ci < len(cols):
                    _c1, hname = cols[ci]
                    k = _merge_identity_key(hname)
                    if k not in h2i:
                        h2i[k] = len(headers_order)
                        headers_order.append(hname)
                        ws_dest.cell(row=hdr_1, column=len(headers_order), value=hname)
                    ci += 1
            r = hdr_1 + 1
            while r <= max_r:
                if _direct_pp_xlsx_empty_first_col(
                    ws.cell(row=r, column=first_col).value
                ):
                    r += 1
                    continue
                ci = 0
                while ci < len(cols):
                    c1, hname = cols[ci]
                    k = _merge_identity_key(hname)
                    if k not in h2i:
                        h2i[k] = len(headers_order)
                        headers_order.append(hname)
                        ws_dest.cell(
                            row=hdr_1, column=len(headers_order), value=hname
                        )
                    dest_c = h2i[k] + 1
                    _direct_pp_xlsx_copy_merge_cell(
                        ws.cell(row=r, column=c1),
                        ws_dest.cell(row=out_data_1, column=dest_c),
                        with_formatting=with_formatting,
                    )
                    ci += 1
                out_data_1 += 1
                total_rows += 1
                sheet_rows += 1
                r += 1
            _direct_trace(
                u"объединить_листы_в_один/xlsx",
                u"«%s» скопировано строк=%d" % (src_name, sheet_rows),
            )
            si += 1
        _direct_pp_track_created_sheet(dest_actual, header_row=hdr0, source=u",".join(source_names))
        wb.save(path)
        note = u"объединено %d лист(ов), %d строк → «%s»%s" % (
            len(source_names),
            total_rows,
            dest_actual,
            u" +формат" if with_formatting else u"",
        )
        _direct_trace(u"объединить_листы_в_один/xlsx", note)
        return True, note
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_pp_file_structural_ods(path, sheet_name, fn_key, extra):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableColumn

    key = _direct_pp_fn_key(fn_key)

    def _load_sheet():
        doc = load(path)
        sheet = None
        tables = doc.spreadsheet.getElementsByType(Table)
        for t in tables:
            if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
                sheet = t
                break
        return doc, sheet, tables

    if key == u"объединить_листы_в_один":
        block = _direct_pp_pick_sheet_block(u"объединить_листы_в_один", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        return _direct_pp_merge_sheets_into_one(path, block)

    if key == u"переименовать_лист":
        block = _direct_pp_pick_sheet_block(u"переименовать_лист", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        raw_name = unicode(block.get("new_name") or u"").strip()
        if raw_name == u"":
            return False, u"пустое new_name"
        doc, sheet, tables = _load_sheet()
        if sheet is None:
            return False, u"лист не найден"
        rows = sheet.getElementsByType(TableRow)
        used_sr, used_er = _direct_pp_ods_used_rows(rows)
        hdr_row0 = 0
        first_data0 = max(0, hdr_row0 + 1)
        last_data0 = max(first_data0, used_er)
        try:
            bounds = _direct_result_sheet_bounds(path, sheet_name)
            if bounds is not None:
                hdr_row0 = int(bounds.get("hdr_row") or 0)
                first_data0 = int(bounds.get("first_data") or (hdr_row0 + 1))
                last_data0 = int(bounds.get("last_data") or first_data0)
        except Exception:
            hdr_row0 = 0
            first_data0 = max(0, hdr_row0 + 1)
            last_data0 = max(first_data0, used_er)
        new_name = _direct_pp_expand_rename_sheet_name(
            raw_name,
            sheet_name=sheet_name,
            old_name=sheet_name,
            header_row=hdr_row0,
            used_sr=first_data0,
            used_er=last_data0,
            cell_text_fn=lambda r, c, _rows=rows: _direct_pp_ods_cell_text(_rows, r, c),
        )
        if new_name == u"":
            return False, u"пустое имя после подстановки"
        names = [unicode(t.getAttribute("name") or u"") for t in tables]
        target = _direct_pp_finalize_sheet_name(names, new_name, exclude_name=sheet_name)
        if target == u"":
            return False, u"пустое имя"
        if target == sheet_name:
            return True, u"имя без изменений"
        sheet.setAttribute(u"name", target)
        doc.save(path)
        return True, u"«%s» → «%s»" % (sheet_name, target)

    if key == u"копировать_лист" or key == u"копировать_переместить_лист":
        block = _direct_pp_pick_sheet_block(u"копировать_лист", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        source = unicode(
            block.get("source_sheet") or block.get("source") or u""
        ).strip()
        dest = unicode(
            block.get("dest_sheet")
            or block.get("dest")
            or block.get("target_sheet")
            or u""
        ).strip()
        if source == u"" or dest == u"":
            return False, u"нужны source_sheet и dest_sheet"
        _direct_trace(
            u"копировать_лист/ods",
            u"ctx=%s source=%s dest=%s file=%s" % (sheet_name, source, dest, path),
        )
        t0 = _direct_now()
        _direct_trace(u"копировать_лист/ods", u"load…")
        doc = load(path)
        _direct_trace(
            u"копировать_лист/ods",
            u"загружено за %.1fs" % (_direct_now() - t0),
        )
        tables = doc.spreadsheet.getElementsByType(Table)
        names = [unicode(t.getAttribute("name") or u"") for t in tables]
        src_actual = _direct_pp_find_sheet_name(names, source)
        if src_actual is None:
            return False, u"лист-источник «%s» не найден" % source
        src_table = None
        for t in tables:
            if unicode(t.getAttribute("name") or u"") == src_actual:
                src_table = t
                break
        if src_table is None:
            return False, u"лист-источник «%s» не найден" % source
        target = _direct_pp_unique_sheet_name(names, dest, exclude_name=None)
        if target == u"":
            return False, u"пустое имя копии"
        if _merge_identity_key(src_actual) == _merge_identity_key(target):
            return True, u"«%s» уже существует" % target
        columns_spec = block.get("columns") or []
        skip_empty = bool(block.get("skip_empty", False))
        hdr1 = block.get("header_row")
        hdr0 = 0
        if hdr1 not in (None, u""):
            try:
                hdr0 = max(0, int(hdr1) - 1)
            except Exception:
                hdr0 = 0
        if columns_spec:
            ok, note = _direct_pp_ods_copy_sheet_filtered(
                path,
                src_actual,
                target,
                hdr0,
                columns_spec,
                skip_empty,
                block.get("skip_empty_columns") or [],
            )
            if ok:
                _direct_pp_track_created_sheet(target, header_row=hdr0, source=src_actual)
                return True, u"«%s» → «%s» (%s)" % (src_actual, target, note)
            return False, note
        t1 = _direct_now()
        new_table = _direct_pp_ods_clone_table(src_table, target)
        _direct_trace(
            u"копировать_лист/ods",
            u"клон за %.1fs" % (_direct_now() - t1),
        )
        doc.spreadsheet.addElement(new_table)
        _direct_pp_track_created_sheet(target, header_row=0, source=src_actual)
        _direct_trace(u"копировать_лист/ods", u"save…")
        t2 = _direct_now()
        doc.save(path)
        _direct_trace(
            u"копировать_лист/ods",
            u"save за %.1fs, всего %.1fs" % (_direct_now() - t2, _direct_now() - t0),
        )
        return True, u"«%s» → «%s»" % (src_actual, target)

    # ВАЖНО: для некоторых шагов не нужен bounds (и bounds читает файл отдельно, что дорого).
    # Обрабатываем их ДО _direct_result_sheet_bounds.
    if key == u"удалить_столбцы":
        block = _direct_pp_pick_sheet_block(u"удалить_столбцы", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        markers = _direct_pp_normalize_markers_tokens(block.get("markers"))
        if not markers:
            markers = [u"маркер", u"marker"]
        doc, sheet, _tables = _load_sheet()
        if sheet is None:
            return False, u"лист не найден"
        rows = sheet.getElementsByType(TableRow)
        if not rows:
            doc.save(path)
            return True, u"лист пуст"
        # header_row: для ODS в direct это 0 по умолчанию
        hdr_row0 = 0
        ec = 0
        if hdr_row0 < len(rows):
            for cidx, val in _ods_row_col_values(rows[hdr_row0]):
                if not _direct_value_empty(val):
                    ec = max(ec, int(cidx))
        titles = _direct_pp_ods_header_titles_list(rows, hdr_row0, ec)
        del_set = _direct_pp_delete_markers_col_set(markers, ec, titles)
        keep_cols = []
        keep_set = set()
        ci = 0
        while ci <= ec and ci < len(titles):
            if ci not in del_set:
                keep_cols.append((ci, titles[ci]))
                keep_set.add(ci)
            ci += 1
        _direct_trace(
            u"удалить_столбцы/ods",
            u"sheet=%s markers=%s (rebuild keep=%d del=%d)"
            % (
                sheet_name,
                u", ".join(unicode(m) for m in markers),
                len(keep_cols),
                len(del_set),
            ),
        )
        if not keep_cols:
            doc.save(path)
            return True, u"удалено столбцов: %d (всё), маркеры: %s" % (
                len(del_set),
                u", ".join(unicode(m) for m in markers),
            )
        # собрать новый лист с тем же именем (внутри одного doc)
        new_table = Table(name=unicode(sheet_name))
        hdr_row = TableRow()
        for _ci0, title in keep_cols:
            hdr_row.addElement(_ods_make_cell(title))
        new_table.addElement(hdr_row)
        total_rows = 0
        r = hdr_row0 + 1
        # читаем строку как sparse values; заполняем только keep_cols
        while r < len(rows):
            if total_rows % 200 == 0:
                _direct_ui_yield()
            if total_rows % 1000 == 0 and total_rows > 0:
                _direct_trace(u"удалить_столбцы/ods", u"rows=%d" % total_rows)
            row_vals = {}
            for cidx, val in _ods_row_col_values(rows[r]):
                ci0 = int(cidx)
                if ci0 in keep_set:
                    row_vals[ci0] = val
            nr = TableRow()
            for ci0, _title in keep_cols:
                nr.addElement(_ods_make_cell(row_vals.get(ci0, u"")))
            new_table.addElement(nr)
            total_rows += 1
            r += 1
        try:
            doc.spreadsheet.removeChild(sheet)
        except Exception:
            pass
        doc.spreadsheet.addElement(new_table)
        doc.save(path)
        return True, u"удалено столбцов: %d, осталось: %d, строк: %d" % (
            len(del_set),
            len(keep_cols),
            total_rows,
        )

    bounds = _direct_result_sheet_bounds(path, sheet_name)
    if bounds is None:
        return False, u"нет данных на листе"
    hdr_row0 = int(bounds["hdr_row"])
    first_data0 = int(bounds["first_data"])
    last_data0 = int(bounds["last_data"])
    end_col0 = int(bounds["end_col"])

    doc, sheet, tables = _load_sheet()
    if sheet is None:
        return False, u"лист не найден"
    rows = sheet.getElementsByType(TableRow)
    header_h2i = _direct_ods_header_h2i(rows, hdr_row=hdr_row0)

    if key == u"переставить_столбцы":
        block = _direct_pp_pick_sheet_block(
            u"переставить_столбцы", extra, sheet_name
        )
        if block is None:
            return False, u"лист не в фильтре sheet"
        ec = end_col0
        titles = _direct_pp_ods_header_titles_list(rows, hdr_row0, ec)
        new_order, insert_at, moving, plan_err, create_names = _direct_pp_plan_column_reorder(
            block, header_h2i, ec, titles
        )
        if plan_err:
            return False, plan_err
        if _direct_debug_enabled():
            print(
                u"  [xml переставить_столбцы] %s order=%s copy=%s create=%s"
                % (
                    sheet_name,
                    new_order,
                    _direct_pp_block_is_copy_columns(block),
                    bool(create_names),
                )
            )
        src_titles = []
        if insert_at is not None and moving:
            for mc in moving:
                if titles is not None and mc < len(titles):
                    src_titles.append(unicode(titles[mc] or u""))
                else:
                    src_titles.append(u"")
        _direct_trace(
            u"переставить_столбцы/ods",
            u"sheet=%s rebuild order=%s rows=0..%d cols=0..%d"
            % (sheet_name, new_order, int(last_data0), int(ec)),
        )
        t0 = _direct_now()
        ok = _direct_pp_ods_reorder_columns_rebuild(
            doc, sheet, sheet_name, last_data0, ec, new_order
        )
        if ok and create_names and insert_at is not None:
            sheet_new = None
            for t in doc.spreadsheet.getElementsByType(Table):
                if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
                    sheet_new = t
                    break
            if sheet_new is not None:
                rows_new = sheet_new.getElementsByType(TableRow)
                # apply headers for empty slots (moving-like indices = insert_at..+)
                fake_moving = list(range(len(create_names)))
                _direct_pp_apply_copied_headers_ods(
                    rows_new,
                    hdr_row0,
                    insert_at,
                    fake_moving,
                    create_names,
                    [u""] * len(create_names),
                    len(new_order) - 1,
                )
        elif ok and insert_at is not None and moving:
            # rebuild заменяет Table в doc (ещё не на диске)
            sheet_new = None
            for t in doc.spreadsheet.getElementsByType(Table):
                if unicode(t.getAttribute("name") or u"") == unicode(sheet_name):
                    sheet_new = t
                    break
            if sheet_new is not None:
                rows_new = sheet_new.getElementsByType(TableRow)
                _direct_pp_apply_copied_headers_ods(
                    rows_new,
                    hdr_row0,
                    insert_at,
                    moving,
                    _direct_pp_parse_copy_new_names(block),
                    src_titles,
                    len(new_order) - 1,
                )
        _direct_trace(
            u"переставить_столбцы/ods",
            u"готово за %.1fs" % (_direct_now() - t0),
        )
        doc.save(path)
        if create_names:
            action = u"новые пустые"
            n_done = len(create_names)
        elif _direct_pp_block_is_copy_columns(block):
            action = u"скопировано"
            n_done = len(moving or [])
        else:
            action = u"переставлено"
            n_done = len(block.get("columns") or [])
        return ok, u"%s столбцов: %d" % (action, n_done)

    if key == u"переименовать_столбцы":
        block = _direct_pp_pick_sheet_block(
            u"переименовать_столбцы", extra, sheet_name
        )
        if block is None:
            return False, u"лист не в фильтре sheet"
        mappings = block.get("mappings") or []
        if not mappings:
            return False, u"пустые mappings"
        try:
            from libre_macros_param_codec import expand_rename_columns_mapping_pairs
            mappings = expand_rename_columns_mapping_pairs(mappings)
        except Exception:
            pass
        if not mappings:
            return False, u"пустые mappings"
        try:
            import libre_macros_lib as lm
        except ImportError:
            lm = None
        used_sr, used_er = _direct_pp_ods_used_rows(rows)
        first_data0 = int(first_data0)
        last_data0 = int(last_data0)
        existing = {}
        ci = 0
        while ci <= end_col0:
            title = _direct_pp_ods_cell_text(rows, hdr_row0, ci)
            key_h = unicode(title or u"").strip().casefold()
            if key_h != u"":
                existing[key_h] = existing.get(key_h, 0) + 1
            ci += 1
        renamed = 0
        for item in mappings:
            if not isinstance(item, dict):
                continue
            old_name = unicode(item.get("old") or u"").strip()
            new_tmpl = unicode(item.get("new") or u"").strip()
            if old_name == u"" or new_tmpl == u"":
                continue
            col_idx, col_err = _direct_pp_resolve_col_unique(
                old_name, header_h2i, end_col0, titles=None
            )
            if col_err:
                return False, col_err
            actual_old = _direct_pp_ods_cell_text(rows, hdr_row0, col_idx)
            if actual_old == u"":
                actual_old = old_name
            expanded = _direct_pp_expand_rename_name_template(
                new_tmpl,
                sheet_name=sheet_name,
                old_name=actual_old,
                header_row=hdr_row0,
                used_sr=first_data0,
                used_er=last_data0,
                cell_text_fn=lambda r, c, _rows=rows: _direct_pp_ods_cell_text(_rows, r, c),
                log_fn_key=u"переименовать_столбцы",
            ).strip()
            old_key = unicode(actual_old or u"").strip().casefold()

            def _col_taken(cand, _old_key=old_key, _existing=existing):
                k = unicode(cand or u"").strip().casefold()
                if k == u"":
                    return True
                if k == _old_key:
                    return False
                return _existing.get(k, 0) > 0

            if lm is not None:
                new_name = lm._lm_pp_finalize_column_name(expanded, is_taken=_col_taken)
            else:
                new_name = expanded
            if new_name == u"":
                continue
            if hdr_row0 < len(rows):
                # end_col0, не col_idx: иначе строка заголовка обрезается до переименуемого столбца
                _direct_ods_expand_row_to_cols(rows[hdr_row0], end_col0)
                _direct_ods_set_cell_at_col(
                    rows[hdr_row0], col_idx, _ods_make_cell(new_name)
                )
            if old_key != u"" and existing.get(old_key, 0) > 0:
                existing[old_key] = existing.get(old_key, 1) - 1
                if existing[old_key] <= 0:
                    existing.pop(old_key, None)
            nk = unicode(new_name).strip().casefold()
            existing[nk] = existing.get(nk, 0) + 1
            renamed += 1
        doc.save(path)
        return True, u"переименовано: %d" % renamed

    if key == u"удалить_столбцы":
        block = _direct_pp_pick_sheet_block(u"удалить_столбцы", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        markers = _direct_pp_normalize_markers_tokens(block.get("markers"))
        if not markers:
            markers = [u"маркер", u"marker"]
        # Быстрый путь: пересобрать лист с нужными столбцами (без построчного удаления).
        _direct_trace(
            u"удалить_столбцы/ods",
            u"sheet=%s markers=%s (rebuild)" % (sheet_name, u", ".join(unicode(m) for m in markers)),
        )
        rows = sheet.getElementsByType(TableRow)
        if not rows:
            return True, u"лист пуст"
        ec = 0
        if hdr_row0 < len(rows):
            for cidx, val in _ods_row_col_values(rows[hdr_row0]):
                if not _direct_value_empty(val):
                    ec = max(ec, int(cidx))
        titles = _direct_pp_ods_header_titles_list(rows, hdr_row0, ec)
        del_set = _direct_pp_delete_markers_col_set(markers, ec, titles)
        keep_cols = []
        ci = 0
        while ci <= ec and ci < len(titles):
            if ci not in del_set:
                keep_cols.append((ci, titles[ci]))
            ci += 1
        if not keep_cols:
            doc.save(path)
            return True, u"удалено столбцов: %d (всё), маркеры: %s" % (
                len(del_set),
                u", ".join(unicode(m) for m in markers),
            )
        # собрать новый лист с тем же именем
        from odf.table import Table, TableRow

        new_table = Table(name=unicode(sheet_name))
        # заголовок
        hdr_row = TableRow()
        for _ci0, title in keep_cols:
            hdr_row.addElement(_ods_make_cell(title))
        new_table.addElement(hdr_row)
        total_rows = 0
        r = hdr_row0 + 1
        while r <= max(last_data0, len(rows) - 1) and r < len(rows):
            if total_rows % 200 == 0:
                _direct_ui_yield()
            row_map = {}
            for cidx, val in _ods_row_col_values(rows[r]):
                row_map[int(cidx)] = val
            nr = TableRow()
            for ci0, _title in keep_cols:
                nr.addElement(_ods_make_cell(row_map.get(ci0, u"")))
            new_table.addElement(nr)
            total_rows += 1
            r += 1
        # заменить старый table
        try:
            doc.spreadsheet.removeChild(sheet)
        except Exception:
            pass
        doc.spreadsheet.addElement(new_table)
        doc.save(path)
        return True, u"удалено столбцов: %d, осталось: %d, строк: %d" % (
            len(del_set),
            len(keep_cols),
            total_rows,
        )

    if key == u"конкатенация_столбцов":
        block = _direct_pp_pick_sheet_block(u"конкатенация_столбцов", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        col_name = unicode(block.get("new_column") or u"").strip()
        # separator: не strip — пробелы могут быть значимы
        delim = unicode(block.get("separator") or u"")
        indices = block.get("columns") or []
        if col_name == u"" or not indices:
            return False, u"нужны new_column и columns"
        src_cols0 = _direct_pp_resolve_col_tokens(indices, header_h2i, end_col0)
        if not src_cols0:
            return False, u"столбцы не найдены"
        target_col0 = _direct_ods_header_col_index(rows, col_name, hdr_row=hdr_row0)
        need_end = max(end_col0, target_col0, max(src_cols0) if src_cols0 else 0)
        # заголовок
        if hdr_row0 < len(rows):
            _direct_ods_expand_row_to_cols(rows[hdr_row0], need_end)
            _direct_ods_set_cell_at_col(
                rows[hdr_row0], target_col0, _ods_make_cell(col_name)
            )
        n = 0
        rr = first_data0
        while rr <= last_data0:
            while rr >= len(rows):
                rows.append(TableRow())
                sheet.addElement(rows[-1])
            _direct_ods_expand_row_to_cols(rows[rr], need_end)
            by_col = {}
            for cidx, val in _ods_row_col_values(rows[rr]):
                by_col[int(cidx)] = val
            parts = []
            for c0 in src_cols0:
                v = by_col.get(c0)
                parts.append(u"" if v is None else unicode(v))
            _direct_ods_set_cell_at_col(
                rows[rr], target_col0, _ods_make_cell(delim.join(parts))
            )
            n += 1
            rr += 1
        doc.save(path)
        return True, u"столбец %d «%s», строк %d" % (target_col0 + 1, col_name, n)

    if key == u"замена_значений":
        block = _direct_pp_pick_sheet_block(u"замена_значений", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        cols_in = block.get("columns") or []
        col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
        if not col_indices:
            return False, u"столбцы не найдены"
        need_end = max(int(end_col0), max(col_indices) if col_indices else 0)
        headers = _direct_pp_ods_header_titles_list(rows, hdr_row0, need_end)
        width = need_end + 1
        while len(headers) < width:
            headers.append(u"")
        matrix = []
        rr = first_data0
        while rr <= last_data0:
            if rr >= len(rows):
                break
            by_col = dict(_ods_row_col_values(rows[rr]))
            row = []
            c = 0
            while c < width:
                row.append(by_col.get(c))
                c += 1
            matrix.append(row)
            rr += 1
        orig = []
        for ci in col_indices:
            snap = []
            ri = 0
            while ri < len(matrix):
                try:
                    snap.append(matrix[ri][ci] if ci < len(matrix[ri]) else None)
                except Exception:
                    snap.append(None)
                ri += 1
            orig.append((ci, snap))
        ok, payload = _direct_pp_replace_values_apply_block(
            block, matrix, headers, col_indices
        )
        if not ok:
            return False, payload
        changed_cells = 0
        for ci, snap in orig:
            ri = 0
            while ri < len(matrix):
                sheet_rr = first_data0 + ri
                if sheet_rr >= len(rows):
                    break
                try:
                    new_v = matrix[ri][ci] if ci < len(matrix[ri]) else None
                except Exception:
                    new_v = None
                old_v = snap[ri] if ri < len(snap) else None
                if new_v != old_v:
                    _direct_ods_expand_row_to_cols(rows[sheet_rr], need_end)
                    _direct_ods_set_cell_at_col(
                        rows[sheet_rr],
                        ci,
                        _ods_make_cell(u"" if new_v is None else unicode(new_v)),
                    )
                    changed_cells += 1
                ri += 1
        doc.save(path)
        note = u"изменено ячеек: %d, столбцов: %d" % (
            changed_cells,
            len(col_indices),
        )
        try:
            aw = (payload or {}).get(u"align_warn")
            if aw:
                note = u"%s; предупреждение: %s" % (note, aw)
        except Exception:
            pass
        return True, note

    if key == u"текстовые_операции":
        from libre_macros_normalize_lib import (
            NormalizeOpError,
            build_row_dict,
            coerce_ops_list,
            compile_ops_compute_fns,
            compile_ops_row_filters,
            needs_crypto_key,
            resolve_crypto_key,
        )

        block = _direct_pp_pick_sheet_block(u"текстовые_операции", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        cols_in = block.get("columns") or []
        col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
        if not col_indices:
            return False, u"столбцы не найдены"
        ops = coerce_ops_list(block.get("ops"))
        if not ops:
            return False, u"пустой стек ops"
        non_text = unicode(block.get("non_text") or u"skip").strip() or u"skip"
        on_error = unicode(block.get("on_error") or u"keep").strip() or u"keep"
        empty_cells = unicode(block.get("empty_cells") or u"skip").strip() or u"skip"
        crypto_key = None
        if needs_crypto_key(ops):
            try:
                crypto_key = resolve_crypto_key(
                    block, doc=None, default_sheet=None, read_file=True
                )
            except NormalizeOpError:
                crypto_key = None
            if crypto_key in (None, u""):
                return False, (
                    u"для encrypt/decrypt в ODS xml укажите key или key_file "
                    u"(key_cell в XML-ODS пока не поддержан)"
                )
        try:
            row_filters = compile_ops_row_filters(ops)
        except Exception as err:
            return False, u"фильтр строк: %s" % err
        try:
            compute_fns = compile_ops_compute_fns(ops)
        except Exception as err:
            return False, u"вычислить: %s" % err
        need_end = max(end_col0, max(col_indices) if col_indices else 0)
        headers = []
        if first_data0 > 0 and first_data0 - 1 < len(rows):
            _direct_ods_expand_row_to_cols(rows[first_data0 - 1], need_end)
            by_h = dict(_ods_row_col_values(rows[first_data0 - 1]))
            c = 0
            while c <= need_end:
                headers.append(by_h.get(c))
                c += 1
        from libre_macros_lambda_column_lib import make_rows_accessor, make_upd_rows_store

        matrix = []
        rr0 = first_data0
        while rr0 <= last_data0:
            if rr0 >= len(rows):
                break
            _direct_ods_expand_row_to_cols(rows[rr0], need_end)
            by_col = dict(_ods_row_col_values(rows[rr0]))
            values = []
            c = 0
            while c <= need_end:
                values.append(by_col.get(c))
                c += 1
            matrix.append(values)
            rr0 += 1
        rows_accessor = make_rows_accessor(matrix, headers, start_col=0)
        changed_cells = 0
        error_cells = 0
        for ci in col_indices:
            upd_store = make_upd_rows_store()
            rr = first_data0
            while rr <= last_data0:
                if rr >= len(rows):
                    break
                _direct_ods_expand_row_to_cols(rows[rr], need_end)
                by_col = dict(_ods_row_col_values(rows[rr]))
                v = by_col.get(ci)
                if non_text == u"skip" and not _direct_pp_cell_is_plain_string(v):
                    if not _direct_value_empty(v):
                        rr += 1
                        continue
                row_i = rr - first_data0
                values = matrix[row_i] if 0 <= row_i < len(matrix) else []
                row_dict = build_row_dict(headers, values, row_i)
                try:
                    new_v, changed, had_err = _direct_pp_normalize_apply_cell(
                        v,
                        ops,
                        crypto_key,
                        non_text,
                        on_error,
                        empty_cells,
                        row=row_dict,
                        row_filters=row_filters,
                        row_id=row_i,
                        rows_accessor=rows_accessor,
                        upd_store=upd_store,
                        compute_fns=compute_fns,
                    )
                except NormalizeOpError as err:
                    return False, unicode(err)
                if had_err:
                    error_cells += 1
                    if on_error == u"stop":
                        doc.save(path)
                        return False, u"ошибка op (on_error=stop)"
                if changed:
                    _direct_ods_set_cell_at_col(
                        rows[rr], ci, _ods_make_cell(new_v)
                    )
                    try:
                        if 0 <= row_i < len(matrix) and ci < len(matrix[row_i]):
                            matrix[row_i][ci] = new_v
                    except Exception:
                        pass
                    changed_cells += 1
                upd_store.commit(row_i)
                rr += 1
        doc.save(path)
        return True, u"изменено ячеек: %d, ошибок: %d, столбцов: %d" % (
            changed_cells,
            error_cells,
            len(col_indices),
        )

    if key == u"заполнение_вниз":
        block = _direct_pp_pick_sheet_block(u"заполнение_вниз", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        cols_in = block.get("columns") or []
        col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
        if not col_indices:
            return False, u"столбцы не найдены"
        if last_data0 < first_data0:
            return False, u"нет строк данных"
        need_end = max(end_col0, max(col_indices) if col_indices else 0)
        filled = 0
        for ci in col_indices:
            last_val = None
            has_last = False
            rr = first_data0
            while rr <= last_data0:
                if rr >= len(rows):
                    break
                _direct_ods_expand_row_to_cols(rows[rr], need_end)
                by_col = dict(_ods_row_col_values(rows[rr]))
                v = by_col.get(ci)
                if _direct_value_empty(v):
                    if has_last:
                        _direct_ods_set_cell_at_col(
                            rows[rr], ci, _ods_make_cell(last_val)
                        )
                        filled += 1
                else:
                    last_val = v
                    has_last = True
                rr += 1
        doc.save(path)
        return True, u"столбцов: %d, заполнено ячеек: %d" % (len(col_indices), filled)

    if key == u"заполнить_вверх":
        block = _direct_pp_pick_sheet_block(u"заполнить_вверх", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        cols_in = block.get("columns") or []
        col_indices = _direct_pp_resolve_col_tokens(cols_in, header_h2i, end_col0)
        if not col_indices:
            return False, u"столбцы не найдены"
        if last_data0 < first_data0:
            return False, u"нет строк данных"
        need_end = max(end_col0, max(col_indices) if col_indices else 0)
        filled = 0
        for ci in col_indices:
            last_val = None
            has_last = False
            rr = last_data0
            while rr >= first_data0:
                if rr >= len(rows):
                    rr -= 1
                    continue
                _direct_ods_expand_row_to_cols(rows[rr], need_end)
                by_col = dict(_ods_row_col_values(rows[rr]))
                v = by_col.get(ci)
                if _direct_value_empty(v):
                    if has_last:
                        _direct_ods_set_cell_at_col(
                            rows[rr], ci, _ods_make_cell(last_val)
                        )
                        filled += 1
                else:
                    last_val = v
                    has_last = True
                rr -= 1
        doc.save(path)
        return True, u"столбцов: %d, заполнено ячеек: %d" % (len(col_indices), filled)

    if key == u"заполнение_вниз_вычислить":
        block = _direct_pp_pick_sheet_block(
            u"заполнение_вниз_вычислить", extra, sheet_name
        )
        if block is None:
            return False, u"лист не в фильтре sheet"
        if last_data0 < first_data0:
            return False, u"нет строк данных"
        pairs = _direct_pp_fill_down_calc_pairs(block, header_h2i, end_col0)
        if not pairs:
            return False, u"нет заданий (columns/rules)"
        col_idxs = [p[2] for p in pairs if p[2] is not None]
        need_end = max(end_col0, max(col_idxs) if col_idxs else 0)
        filled = 0
        skipped = 0
        for col_t, formula, idx, expand in pairs:
            if idx is None or formula is None:
                skipped += 1
                continue
            rr = first_data0
            while rr <= last_data0:
                while rr >= len(rows):
                    rows.append(TableRow())
                    sheet.addElement(rows[-1])
                _direct_ods_expand_row_to_cols(rows[rr], need_end)
                if not _direct_pp_ods_cell_is_empty(rows[rr], idx):
                    rr += 1
                    continue
                val, is_f = _direct_pp_fill_down_calc_write_value(
                    formula, rr, idx, header_h2i, for_xlsx=False, sheet_name=sheet_name
                )
                if val == u"" and unicode(formula).strip().startswith(u"="):
                    rr += 1
                    continue
                if is_f:
                    _direct_ods_set_cell_at_col(
                        rows[rr], idx, _ods_make_formula_cell(val)
                    )
                else:
                    _direct_ods_set_cell_at_col(rows[rr], idx, _ods_make_cell(val))
                filled += 1
                if expand:
                    c0 = idx + 1
                    while c0 <= end_col0:
                        _direct_ods_expand_row_to_cols(rows[rr], max(need_end, c0))
                        if not _direct_pp_ods_cell_is_empty(rows[rr], c0):
                            c0 += 1
                            continue
                        cval, cis_f = _direct_pp_fill_down_calc_write_value(
                            formula, rr, c0, header_h2i, for_xlsx=False, sheet_name=sheet_name
                        )
                        if cval == u"" and unicode(formula).strip().startswith(u"="):
                            c0 += 1
                            continue
                        if cis_f:
                            _direct_ods_set_cell_at_col(
                                rows[rr], c0, _ods_make_formula_cell(cval)
                            )
                        else:
                            _direct_ods_set_cell_at_col(rows[rr], c0, _ods_make_cell(cval))
                        filled += 1
                        c0 += 1
                rr += 1
        doc.save(path)
        note = u"столбцов: %d, заполнено ячеек: %d" % (len(pairs) - skipped, filled)
        if skipped:
            note += u", пропущено столбцов: %d" % skipped
        return True, note

    if key == u"сводная_таблица":
        config = _direct_pp_pivot_parse_config(extra)
        if config is None:
            return False, u"пустой/неверный JSON"
        headers, src_rows, err = _direct_pp_pivot_read_ods_matrix(rows, config, bounds)
        if err and not src_rows:
            return False, err
        hdr, data_rows, n_row_f, note = _direct_pp_pivot_compute(
            headers, src_rows, config
        )
        if hdr is None:
            return False, note
        dest_base = unicode(config.get("sheet_name") or u"Сводная").strip()
        if dest_base == u"":
            dest_base = u"Сводная"
        names = [unicode(t.getAttribute("name") or u"") for t in tables]
        # удалить одноимённый лист (не источник)
        existing = _direct_pp_find_sheet_name(names, dest_base)
        if existing is not None and _merge_identity_key(existing) != _merge_identity_key(
            sheet_name
        ):
            for t in list(tables):
                if unicode(t.getAttribute("name") or u"") == existing:
                    try:
                        doc.spreadsheet.removeChild(t)
                    except Exception:
                        pass
                    break
            tables = doc.spreadsheet.getElementsByType(Table)
            names = [unicode(t.getAttribute("name") or u"") for t in tables]
        dest = _direct_pp_unique_sheet_name(names, dest_base, exclude_name=None)
        if dest == u"":
            return False, u"пустое имя листа сводной"
        if dest in names:
            for t in list(tables):
                if unicode(t.getAttribute("name") or u"") == dest:
                    try:
                        doc.spreadsheet.removeChild(t)
                    except Exception:
                        pass
                    break
        out_table = Table(name=dest)
        _ods_add_row(out_table, hdr)
        for line in data_rows:
            vals = list(line)
            while len(vals) < len(hdr):
                vals.append(None)
            _ods_add_row(out_table, vals[: len(hdr)])
        doc.spreadsheet.addElement(out_table)
        out_rows = out_table.getElementsByType(TableRow)
        _direct_pp_pivot_apply_format_ods(
            doc, out_table, out_rows, len(hdr), len(data_rows), n_row_f
        )
        try:
            if dest not in _direct_xml_pp_freeze_pending:
                _direct_xml_pp_freeze_pending.append(dest)
        except Exception:
            pass
        _direct_pp_track_created_sheet(dest, header_row=0, source=sheet_name)
        doc.save(path)
        return True, u"лист «%s»; %s" % (dest, note)

    if key == u"сортировка":
        block = _direct_pp_pick_sheet_block(u"сортировка", extra, sheet_name)
        if block is None:
            return False, u"лист не в фильтре sheet"
        keys_in = block.get("keys") or []
        key_specs = []
        for item in keys_in:
            if not isinstance(item, dict):
                continue
            col_tok = unicode(item.get("column") or u"").strip()
            if col_tok == u"":
                continue
            idxs = _direct_pp_resolve_col_tokens([col_tok], header_h2i, end_col0)
            if not idxs:
                continue
            key_specs.append((idxs[0], bool(item.get("desc"))))
        if not key_specs:
            return False, u"нет ключей сортировки"
        if last_data0 < first_data0:
            return False, u"нет строк данных"
        # собрать data TableRow + значения
        data_elems = []
        data_vals = []
        rr = first_data0
        while rr <= last_data0 and rr < len(rows):
            row_elem = rows[rr]
            by_col = dict(_ods_row_col_values(row_elem))
            row_vals = []
            ci = 0
            while ci <= end_col0:
                row_vals.append(by_col.get(ci))
                ci += 1
            data_elems.append(row_elem)
            data_vals.append(row_vals)
            rr += 1
        if not data_elems:
            return False, u"нет строк данных"
        indexed = list(enumerate(data_vals))

        def _cmp_pair(a, b):
            ka = _direct_pp_sort_key_tuple(a[1], key_specs)
            kb = _direct_pp_sort_key_tuple(b[1], key_specs)
            i = 0
            while i < len(ka):
                ea, ta, desc = ka[i]
                eb, tb, _d = kb[i]
                if ea != eb:
                    return 1 if ea else -1
                if ta < tb:
                    return 1 if desc else -1
                if ta > tb:
                    return -1 if desc else 1
                i += 1
            return a[0] - b[0]

        try:
            from functools import cmp_to_key

            indexed.sort(key=cmp_to_key(_cmp_pair))
        except Exception:
            n = len(indexed)
            i = 0
            while i < n:
                j = 0
                while j < n - 1:
                    if _cmp_pair(indexed[j], indexed[j + 1]) > 0:
                        indexed[j], indexed[j + 1] = indexed[j + 1], indexed[j]
                    j += 1
                i += 1
        # переставить TableRow в DOM: удалить data rows, вставить в новом порядке
        # после заголовка
        parent = sheet
        # удалить data rows из parent (сохраняя ссылки)
        for row_elem in data_elems:
            try:
                parent.removeChild(row_elem)
            except Exception:
                pass
        # найти точку вставки: после header row
        insert_after = None
        if hdr_row0 < len(rows):
            insert_after = rows[hdr_row0]
        # строки после data (если были) — оставить как есть
        for _orig_i, _vals in indexed:
            row_elem = data_elems[_orig_i]
            if insert_after is not None:
                try:
                    # insertAfter: next sibling of insert_after
                    nxt = insert_after.nextSibling
                    if nxt is not None:
                        parent.insertBefore(row_elem, nxt)
                    else:
                        parent.addElement(row_elem)
                except Exception:
                    parent.addElement(row_elem)
                insert_after = row_elem
            else:
                parent.addElement(row_elem)
                insert_after = row_elem
        doc.save(path)
        return True, u"отсортировано строк: %d, ключей: %d" % (
            len(indexed),
            len(key_specs),
        )

    return False, u"неизвестный file-structural: %s" % key


def direct_run_xml_postprocess(target_path, settings, one_sheet):
    """
    direct-пайплайн xml_excel_ods в файле до reopen UNO:
      Постобработка_xml + ВПР_XML в порядке строк листа параметров.
    """
    global _direct_xml_pp_formula_done, _direct_xml_pp_as_values_pending
    global _direct_xml_pp_freeze_pending, _direct_xml_pp_uno_pending
    global _direct_xml_pp_file_visual_sheets, _direct_xml_pp_created_sheets
    global _direct_xml_pp_format_header_cols
    _direct_xml_pp_formula_done = set()
    _direct_xml_pp_as_values_pending = []
    _direct_xml_pp_freeze_pending = []
    _direct_xml_pp_uno_pending = []
    _direct_xml_pp_file_visual_sheets = []
    _direct_xml_pp_format_header_cols = {}
    _direct_xml_pp_created_sheets = []
    pipeline = (settings or {}).get("xml_postprocess_pipeline") or []
    if not pipeline:
        _direct_journal(u"", u"", u"", u"", u"инфо", u"Постобработка_xml: шаги не заданы")
        return
    sheet_names = _direct_result_sheet_names(target_path, one_sheet, settings=settings)
    try:
        fsize = os.path.getsize(target_path)
    except Exception:
        fsize = -1
    _direct_trace(
        u"Постобработка_xml",
        u"start file=%s size=%s steps=%d sheets=%d"
        % (target_path, fsize, len(pipeline), len(sheet_names)),
    )
    _direct_journal(
        u"",
        u"",
        u"",
        u"",
        u"Постобработка_xml",
        u"этап: шаги в файле; шагов=%d, листов=%d (%s)"
        % (
            len(pipeline),
            len(sheet_names),
            u", ".join(sheet_names[:8])
            + (u"…" if len(sheet_names) > 8 else u""),
        ),
    )
    vlookup_run_n = 0
    # Предупреждение: если копировать_лист уже с columns, то удалить_столбцы на тот же лист обычно лишний.
    copy_sheet_dest_with_columns = set()
    step_i = 0
    for step in pipeline:
        step_i += 1
        step_kind = (step or {}).get("kind") or u"pp"
        step_name = (step or {}).get("name") or u""
        _direct_trace(
            u"Постобработка_xml",
            u"шаг %d/%d kind=%s name=%s"
            % (step_i, len(pipeline), step_kind, step_name),
        )
        if (step or {}).get("kind") == "vlookup":
            vlookup_run_n += 1
            try:
                _direct_run_vlookup_step(
                    target_path, (step or {}).get("spec") or {}, vlookup_run_n
                )
            except Exception as err:
                _direct_journal(
                    u"",
                    u"",
                    u"",
                    u"",
                    u"ошибка",
                    u"ВПР_XML: %s (продолжаем сбор)" % unicode(err),
                )
            continue
        if (step or {}).get("kind") == "merge_sheets":
            spec = (step or {}).get("spec") or {}
            try:
                ok, note = _direct_pp_merge_sheets_into_one(target_path, spec)
            except Exception as err:
                ok, note = False, unicode(err)
            _direct_journal(
                u"",
                unicode(spec.get("dest_sheet") or u""),
                u"",
                u"",
                u"ok" if ok else u"ошибка",
                u"объединить_листы_в_один: %s" % note,
            )
            if ok:
                _direct_pp_sync_result_sheet_names(
                    target_path, one_sheet, settings, sheet_names
                )
            continue
        if (step or {}).get("kind") == "split_sheets":
            _direct_journal(
                u"",
                u"",
                u"",
                u"",
                u"ошибка",
                u"разделить_листы: при xml_excel_ods только run_phase=after_final "
                u"(шаг в файле не выполняется)",
            )
            continue
        name = step.get("name") or u""
        extra = step.get("formula_extra")
        if extra is None:
            extra = (
                step.get("extra")
                or step.get("grid_extra")
                or step.get("header_height_extra")
                or step.get("sort_extra")
                or step.get("colorize_extra")
                or step.get("color_scale_extra")
                or step.get("pivot_extra")
                or u""
            )
        if extra is None:
            extra = u""
        fn_key = _direct_pp_fn_key(name)

        # накопить dest_sheet из копировать_лист/копировать_переместить_лист, где задан columns
        if fn_key == u"копировать_лист" or fn_key == u"копировать_переместить_лист":
            try:
                from libre_macros_param_codec import param_decode

                # Алиас «копировать_лист» нормализуется в «копировать_переместить_лист».
                blocks = param_decode(u"копировать_лист", extra)
                for b in blocks:
                    if not isinstance(b, dict):
                        continue
                    cols = b.get("columns") or []
                    if not cols:
                        continue
                    dest = unicode(b.get("dest_sheet") or u"").strip()
                    if dest != u"":
                        copy_sheet_dest_with_columns.add(_merge_identity_key(dest))
            except Exception:
                pass

        if _direct_pp_is_skip_empty_rows(name):
            for sheet_name in sheet_names:
                _direct_ui_notify_source(target_path, sheet_name)
                try:
                    if _is_ods(target_path):
                        ok, note = _direct_pp_skip_empty_rows_ods(
                            target_path, sheet_name, extra
                        )
                    else:
                        ok, note = _direct_pp_skip_empty_rows_xlsx(
                            target_path, sheet_name, extra
                        )
                except Exception as err:
                    ok, note = False, unicode(err)
                _direct_journal(
                    u"",
                    sheet_name,
                    u"",
                    u"",
                    u"ok" if ok else u"пропуск",
                    u"пропуск_пустых_строк: %s" % note,
                )
            continue

        if _direct_pp_is_file_structural(name):
            applied_n = 0
            tried_n = 0
            _direct_trace(fn_key, u"file_structural begin")
            for sheet_name in list(sheet_names):
                if not _direct_pp_file_structural_applies(fn_key, extra, sheet_name):
                    continue
                if (
                    fn_key == u"удалить_столбцы"
                    and copy_sheet_dest_with_columns
                    and _merge_identity_key(sheet_name) in copy_sheet_dest_with_columns
                ):
                    _direct_journal(
                        u"",
                        sheet_name,
                        u"",
                        u"",
                        u"инфо",
                        u"удалить_столбцы: возможно лишний шаг — лист получен через копировать_лист(columns=...)",
                    )
                tried_n += 1
                _direct_ui_notify_source(target_path, sheet_name)
                _direct_trace(fn_key, u"лист=%s apply…" % sheet_name)
                try:
                    ok, note = _direct_pp_file_structural(
                        target_path, sheet_name, fn_key, extra
                    )
                except Exception as err:
                    ok, note = False, unicode(err)
                _direct_trace(
                    fn_key,
                    u"лист=%s %s: %s" % (sheet_name, u"ok" if ok else u"err", note),
                )
                if ok:
                    applied_n += 1
                    if fn_key in (
                        u"переименовать_лист",
                        u"копировать_лист",
                        u"сводная_таблица",
                    ):
                        _direct_pp_sync_result_sheet_names(
                            target_path, one_sheet, settings, sheet_names
                        )
                        _direct_journal(
                            u"",
                            u"",
                            u"",
                            u"",
                            u"Постобработка_xml",
                            u"листы после %s: %s"
                            % (fn_key, u", ".join(sheet_names)),
                        )
                _direct_journal(
                    u"",
                    sheet_name,
                    u"",
                    u"",
                    u"ok" if ok else u"ошибка",
                    u"%s: %s" % (fn_key, note),
                )
            if applied_n == 0 and tried_n == 0:
                _direct_journal(
                    u"",
                    u"",
                    u"",
                    u"",
                    u"пропуск",
                    u"%s: нет листов по фильтру sheet/sheets" % fn_key,
                )
            continue

        if _direct_pp_is_file_visual(name):
            applied_n = 0
            for sheet_name in sheet_names:
                if not _direct_pp_file_visual_applies(fn_key, extra, sheet_name):
                    continue
                _direct_ui_notify_source(target_path, sheet_name)
                try:
                    ok, note = _direct_pp_file_visual(
                        target_path, sheet_name, fn_key, extra
                    )
                except Exception as err:
                    ok, note = False, unicode(err)
                if ok:
                    _direct_pp_track_file_visual_sheet(sheet_name, fn_key=fn_key)
                    if _direct_pp_queue_format_uno_reapply(fn_key, sheet_name, extra):
                        note = u"%s; UNO повтор после оформления" % note
                    applied_n += 1
                _direct_journal(
                    u"",
                    sheet_name,
                    u"",
                    u"",
                    u"ok" if ok else u"ошибка",
                    u"%s: %s" % (fn_key, note),
                )
            if applied_n == 0:
                _direct_journal(
                    u"",
                    u"",
                    u"",
                    u"",
                    u"пропуск",
                    u"%s: нет листов по фильтру sheet/sheets" % fn_key,
                )
            continue

        if _direct_pp_is_uno_pilot(name):
            # Autofit / заголовок_плюс_высота / автофильтр → UNO после reopen.
            applied_n = 0
            for sheet_name in sheet_names:
                if not _direct_pp_uno_pilot_applies(fn_key, extra, sheet_name):
                    continue
                _direct_ui_notify_source(target_path, sheet_name)
                _direct_xml_pp_uno_pending.append(
                    {
                        u"name": fn_key,
                        u"sheet": unicode(sheet_name or u"").strip(),
                        u"extra": unicode(extra or u""),
                    }
                )
                applied_n += 1
                _direct_journal(
                    u"",
                    sheet_name,
                    u"",
                    u"",
                    u"ok",
                    u"%s: в очереди на UNO после открытия" % fn_key,
                )
            if applied_n == 0:
                _direct_journal(
                    u"",
                    u"",
                    u"",
                    u"",
                    u"пропуск",
                    u"%s: нет листов по фильтру sheet/sheets" % fn_key,
                )
            continue

        if not _direct_pp_is_apply_formula(name):
            if _direct_pp_is_freeze_header(name):
                applied_n = 0
                for sheet_name in sheet_names:
                    if not _direct_pp_freeze_header_applies(extra, sheet_name):
                        continue
                    _direct_ui_notify_source(target_path, sheet_name)
                    # Не пишем freeze в файл: Calc часто игнорирует openpyxl freeze_panes.
                    # Ставим в очередь — UNO закрепит после показа окна.
                    key = unicode(sheet_name or u"").strip()
                    if key and key not in _direct_xml_pp_freeze_pending:
                        _direct_xml_pp_freeze_pending.append(key)
                    applied_n += 1
                    _direct_journal(
                        u"",
                        sheet_name,
                        u"",
                        u"",
                        u"ok",
                        u"закрепить_заголовок: в очереди на UNO после показа окна",
                    )
                if applied_n == 0:
                    _direct_journal(
                        u"",
                        u"",
                        u"",
                        u"",
                        u"пропуск",
                        u"закрепить_заголовок: нет листов по фильтру sheets",
                    )
                continue
            _direct_journal(
                u"",
                u"",
                u"",
                u"",
                u"инфо",
                u"Постобработка_xml: «%s» — в файле пока не реализована (пропуск)"
                % name,
            )
            continue
        for sheet_name in sheet_names:
            block = _direct_pp_pick_formula_block(extra, sheet_name)
            spec = _direct_pp_formula_spec(extra, sheet_name)
            if spec is None:
                continue
            col_name, formula, as_values = spec
            if block is not None and _direct_pp_parse_as_values(block):
                _direct_journal(
                    u"",
                    sheet_name,
                    col_name,
                    u"",
                    u"инфо",
                    u"применить_формулу: as_values в Постобработка_xml игнорируется",
                )
            _direct_ui_notify_source(target_path, sheet_name)
            if _is_ods(target_path):
                ok, note = _direct_pp_apply_formula_ods(
                    target_path, sheet_name, col_name, formula, as_values
                )
            else:
                ok, note = _direct_pp_apply_formula_xlsx(
                    target_path, sheet_name, col_name, formula, as_values
                )
            status = u"ok" if ok else u"ошибка"
            _direct_journal(
                u"",
                sheet_name,
                col_name,
                u"",
                status,
                u"применить_формулу: %s" % note,
            )
            if ok:
                _direct_xml_pp_formula_done.add(
                    (
                        unicode(sheet_name or u"").strip().casefold(),
                        unicode(col_name or u"").strip().casefold(),
                    )
                )
                # as_values для формул в XML запрещён — в pending не ставим.
    if _direct_xml_pp_formula_done:
        _log(
            "xml_postprocess formulas in file: %s"
            % (u", ".join(u"%s→%s" % (s, c) for s, c in sorted(_direct_xml_pp_formula_done)))
        )
    _direct_trace(u"Постобработка_xml", u"done steps=%d" % len(pipeline))


# --- ВПР (xml_excel_ods): в файле, до reopen UNO ---


def _direct_vlookup_resolve_sheet_name(path, sheet_ref):
    """Имя или 1-based индекс листа → имя листа в файле (xlsx/ods)."""
    ref = unicode(sheet_ref or u"").strip()
    if ref == u"":
        return None
    if ref.isdigit():
        idx = int(ref) - 1
        try:
            names = _ods_sheet_names(path) if _is_ods(path) else _xlsx_sheet_names(path)
        except Exception:
            names = []
        if 0 <= idx < len(names):
            return names[idx]
        return None
    # точное (без учёта регистра) или N_ИмяЛиста
    key = _merge_identity_key(ref)
    try:
        names = _ods_sheet_names(path) if _is_ods(path) else _xlsx_sheet_names(path)
    except Exception:
        names = []
    for nm in names:
        if _merge_identity_key(nm) == key:
            return nm
    for nm in names:
        m = re.match(r"^\d+_(.+)$", unicode(nm or u"").strip())
        if m is not None and _merge_identity_key(m.group(1)) == key:
            return nm
    return None


def _direct_vlookup_parse_col_token(token, header_row_vals, start_col0=0):
    """
    token: '1' (1-based индекс), 'A' (буквы), 'Имя' (по заголовку/шаблону).
    Возвращает 0-based индекс столбца на листе.
    """
    t = unicode(token or u"").strip()
    if t == u"":
        return None
    try:
        from libre_macros_param_codec import unwrap_column_name_token
        bare, exact = unwrap_column_name_token(t)
        if exact:
            t = unicode(bare or u"").strip()
            if t == u"":
                return None
        else:
            t = unicode(bare or t).strip()
    except Exception:
        pass
    if t.isdigit():
        return int(t) - 1
    if _is_col_letters(t):
        return _col_letters_to_index(t)
    # по заголовку (с поддержкой * в начале/конце)
    key = _merge_identity_key(t)
    ci = 0
    while ci < len(header_row_vals):
        hv = unicode(header_row_vals[ci] if header_row_vals[ci] is not None else u"").strip()
        if hv != u"" and (_merge_identity_key(hv) == key or _text_match(hv, t)):
            return int(start_col0) + int(ci)
        ci += 1
    return None


def _direct_vlookup_parse_extract_token(token, header_row_vals, start_col0=0):
    """
    Токен extract → (col_idx0, out_title).
    Поддержка «Кол -> Новое» / «Кол = Новое»; без alias — заголовок справа.
    """
    raw = unicode(token or u"").strip()
    src = raw
    as_name = None
    try:
        from libre_macros_param_codec import parse_vlookup_extract_alias
        src, as_name = parse_vlookup_extract_alias(raw)
    except Exception:
        pass
    idx = _direct_vlookup_parse_col_token(src, header_row_vals, start_col0=start_col0)
    if idx is None:
        return (None, None)
    if as_name is not None and unicode(as_name).strip() != u"":
        return (idx, unicode(as_name).strip())
    # заголовок с правого листа
    try:
        local = int(idx) - int(start_col0)
        if 0 <= local < len(header_row_vals):
            hv = header_row_vals[local]
            if hv is not None and unicode(hv).strip() != u"":
                return (idx, unicode(hv).strip())
    except Exception:
        pass
    return (idx, unicode(src))


def _direct_vlookup_key_part(value, trim=False):
    if value is None:
        return u"__ПУСТО__"
    if isinstance(value, float):
        try:
            if value == int(value):
                return unicode(int(value))
        except Exception:
            pass
    if isinstance(value, int):
        return unicode(value)
    s = unicode(value)
    if trim:
        s2 = s.strip()
        return u"__ПУСТО__" if s2 == u"" else s2.casefold()
    return u"__ПУСТО__" if s == u"" else s.casefold()


def _direct_vlookup_composite_key(values, trim=False):
    return u"|".join(_direct_vlookup_key_part(v, trim=trim) for v in values)


def _direct_vlookup_not_found_fill(spec):
    nf = spec.get("not_found_fill")
    if nf is None:
        return u"#Н/Д"
    return unicode(nf)


def _direct_vlookup_highlight_color(spec):
    """
    Цвет подсветки (как в параметрах ВПР) для direct XLSX.
    Поддержка: "#RRGGBB" или "Метка: #RRGGBB". "Нет"/пусто → None.
    """
    raw = unicode((spec or {}).get("highlight_color") or u"").strip()
    if raw == u"":
        return None
    if u":" in raw and not raw.lstrip().startswith(u"#"):
        _lbl, _sep, raw = raw.partition(u":")
        raw = unicode(raw).strip()
    if raw == u"":
        return None
    if raw.strip().lower() in (u"нет", u"no", u"false", u"0"):
        return None
    if raw.startswith(u"#") and len(raw) == 7:
        return raw[1:].upper()
    # если формат не распознан — не подсвечиваем
    return None


def _direct_vlookup_derive_from_range_a1(text):
    """A1 → 1-based start_row/start_col/header_row (+ end_*)."""
    import re

    s = unicode(text or u"").strip().replace(u"$", u"").replace(u" ", u"")
    if s == u"":
        return {}
    try:
        import libre_macros_lib as lm

        col_letters_to_index = lm.col_letters_to_index
    except Exception:
        return {}
    out = {}
    m = re.match(r"^([A-Za-z]+)(\d+):([A-Za-z]+)(\d+)$", s)
    if m is not None:
        sc = col_letters_to_index(m.group(1))
        sr = int(m.group(2))
        ec = col_letters_to_index(m.group(3))
        er = int(m.group(4))
        if sc < 0 or ec < 0 or sr < 1 or er < 1:
            return {}
        if ec < sc:
            sc, ec = ec, sc
        if er < sr:
            sr, er = er, sr
        out["start_row"] = sr
        out["start_col"] = sc + 1
        out["header_row"] = sr
        out["end_row"] = er
        out["end_col"] = ec + 1
        return out
    m = re.match(r"^([A-Za-z]+)(\d+)$", s)
    if m is not None:
        sc = col_letters_to_index(m.group(1))
        sr = int(m.group(2))
        if sc < 0 or sr < 1:
            return {}
        out["start_row"] = sr
        out["start_col"] = sc + 1
        out["header_row"] = sr
        return out
    m = re.match(r"^([A-Za-z]+):([A-Za-z]+)$", s)
    if m is not None:
        sc = col_letters_to_index(m.group(1))
        ec = col_letters_to_index(m.group(2))
        if sc < 0 or ec < 0:
            return {}
        if ec < sc:
            sc, ec = ec, sc
        out["start_col"] = sc + 1
        out["end_col"] = ec + 1
        return out
    m = re.match(r"^(\d+):(\d+)$", s)
    if m is not None:
        sr = int(m.group(1))
        er = int(m.group(2))
        if sr < 1 or er < 1:
            return {}
        if er < sr:
            sr, er = er, sr
        out["start_row"] = sr
        out["header_row"] = sr
        out["end_row"] = er
        return out
    return {}


def _direct_vlookup_apply_range_to_spec(spec):
    """Если задан left/right.range — заполнить start_*/header_row."""
    out = dict(spec or {})
    for prefix in (u"left", u"right"):
        range_txt = unicode(out.get(prefix + u"_range") or u"").strip()
        if range_txt == u"":
            continue
        derived = _direct_vlookup_derive_from_range_a1(range_txt)
        for key in (u"start_row", u"start_col", u"header_row", u"end_row", u"end_col"):
            if key in derived and derived[key] is not None:
                out[prefix + u"_" + key] = derived[key]
    return out


def _direct_vlookup_normalize_join_type(raw, default=u"left"):
    s = unicode(raw or u"").strip().casefold()
    if s in (u"", u"left", u"левое", u"left_join", u"left outer", u"left_outer"):
        return u"left"
    if s in (
        u"inner",
        u"внутреннее",
        u"inner_join",
        u"inner join",
        u"внутреннее_соединение",
        u"внутреннее соединение",
    ):
        return u"inner"
    if s in (
        u"full",
        u"полное",
        u"full_outer",
        u"full outer",
        u"полное_соединение",
        u"полное соединение",
    ):
        return u"full"
    return default


def _direct_vlookup_normalize_multi_match(raw, default=u"all"):
    s = unicode(raw or u"").strip().casefold()
    if s in (u"", u"all", u"все", u"*", u"all_rows", u"все_совпадения"):
        return u"all"
    if s in (u"first", u"первый", u"первая", u"первое", u"first_row", u"1"):
        return u"first"
    if s in (u"last", u"последний", u"последняя", u"последнее", u"last_row", u"-1"):
        return u"last"
    return default


def _direct_vlookup_select_matches(matches, mode):
    if not matches:
        return matches
    m = _direct_vlookup_normalize_multi_match(mode, default=u"all")
    if m == u"first":
        return [matches[0]]
    if m == u"last":
        return [matches[len(matches) - 1]]
    return matches


def _direct_vlookup_mirror_spec(spec):
    out = dict(spec)
    out["left_sheet"] = spec.get("right_sheet")
    out["right_sheet"] = spec.get("left_sheet")
    out["left_range"] = spec.get("right_range")
    out["right_range"] = spec.get("left_range")
    for fld in (u"start_row", u"start_col", u"header_row", u"end_row", u"end_col"):
        out["left_" + fld] = spec.get("right_" + fld)
        out["right_" + fld] = spec.get("left_" + fld)
    out["join_keys"] = list(spec.get("join_keys_right") or spec.get("join_keys") or [])
    out["join_keys_left"] = list(spec.get("join_keys_right") or spec.get("join_keys") or [])
    out["join_keys_right"] = list(spec.get("join_keys_left") or spec.get("join_keys") or [])
    out["extract_cols"] = list(spec.get("extract_cols_right") or [])
    out["join_type"] = u"left"
    out["_vlookup_skip_full_dispatch"] = True
    return out


def _direct_vlookup_append_right_only_xlsx(path, spec, run_suffix=u""):
    """Дописать на левый лист строки правой без ключа слева (full join, xlsx)."""
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    spec = _direct_vlookup_apply_range_to_spec(spec)
    left_name = _direct_vlookup_resolve_sheet_name(path, spec.get("left_sheet"))
    right_name = _direct_vlookup_resolve_sheet_name(path, spec.get("right_sheet"))
    if left_name is None or right_name is None:
        return {"rows_added": 0, "rows_processed": 0}
    wb = load_workbook(path)
    try:
        left_ws = wb[left_name]
        right_ws = wb[right_name]
        l_header = int(spec.get("left_header_row") or spec.get("left_start_row") or 1) - 1
        r_header = int(spec.get("right_header_row") or spec.get("right_start_row") or 1) - 1
        l_start_col = int(spec.get("left_start_col") or 1) - 1
        r_start_col = int(spec.get("right_start_col") or 1) - 1
        l_hdr_vals = []
        r_hdr_vals = []
        max_lc = int(left_ws.max_column or 1)
        max_rc = int(right_ws.max_column or 1)
        for ci in range(l_start_col + 1, max_lc + 1):
            l_hdr_vals.append(left_ws.cell(row=l_header + 1, column=ci).value)
        for ci in range(r_start_col + 1, max_rc + 1):
            r_hdr_vals.append(right_ws.cell(row=r_header + 1, column=ci).value)
        left_key_tokens = spec.get("join_keys_left") or spec.get("join_keys") or []
        right_key_tokens = spec.get("join_keys_right") or spec.get("join_keys") or []
        extract_tokens = spec.get("extract_cols") or []
        trim_keys = bool(spec.get("trim_keys", False))
        l_key_cols = []
        for tok in left_key_tokens:
            idx = _direct_vlookup_parse_col_token(tok, l_hdr_vals, start_col0=l_start_col)
            if idx is None:
                return {"rows_added": 0, "rows_processed": 0}
            l_key_cols.append(idx)
        r_key_cols = []
        for tok in right_key_tokens:
            idx = _direct_vlookup_parse_col_token(tok, r_hdr_vals, start_col0=r_start_col)
            if idx is None:
                return {"rows_added": 0, "rows_processed": 0}
            r_key_cols.append(idx)
        r_extract_cols = []
        r_extract_titles = []
        for tok in extract_tokens:
            idx, title = _direct_vlookup_parse_extract_token(
                tok, r_hdr_vals, start_col0=r_start_col
            )
            if idx is None:
                continue
            r_extract_cols.append(idx)
            if title is None or unicode(title).strip() == u"":
                hv = right_ws.cell(row=r_header + 1, column=idx + 1).value
                title = unicode(hv if hv is not None else tok)
            r_extract_titles.append(unicode(title) + unicode(run_suffix or u""))
        left_keys = set()
        l_data_start = l_header + 2
        l_max_row = int(left_ws.max_row or l_data_start)
        rr = l_data_start
        while rr <= l_max_row:
            parts = []
            for ci in l_key_cols:
                parts.append(left_ws.cell(row=rr, column=ci + 1).value)
            key = _direct_vlookup_composite_key(parts, trim=trim_keys)
            if key not in (None, u""):
                left_keys.add(key)
            rr += 1
        # Столбцы extract на левом
        out_cols = []
        for title in r_extract_titles:
            found = None
            for ci in range(1, max_lc + 1):
                hv = left_ws.cell(row=l_header + 1, column=ci).value
                if unicode(hv or u"").strip() == unicode(title).strip():
                    found = ci - 1
                    break
            if found is None:
                max_lc += 1
                left_ws.cell(row=l_header + 1, column=max_lc, value=title)
                found = max_lc - 1
            out_cols.append(found)
        rows_added = 0
        rows_processed = 0
        r_data_start = r_header + 2
        r_max_row = int(right_ws.max_row or r_data_start)
        append_at = l_max_row + 1
        rr = r_data_start
        while rr <= r_max_row:
            rows_processed += 1
            parts = []
            for ci in r_key_cols:
                parts.append(right_ws.cell(row=rr, column=ci + 1).value)
            key = _direct_vlookup_composite_key(parts, trim=trim_keys)
            if key in (None, u"") or key in left_keys:
                rr += 1
                continue
            for ki, ci in enumerate(l_key_cols):
                val = parts[ki] if ki < len(parts) else None
                left_ws.cell(row=append_at, column=ci + 1, value=val)
            for ei, src_ci in enumerate(r_extract_cols):
                left_ws.cell(
                    row=append_at,
                    column=out_cols[ei] + 1,
                    value=right_ws.cell(row=rr, column=src_ci + 1).value,
                )
            left_keys.add(key)
            append_at += 1
            rows_added += 1
            rr += 1
        return {"rows_added": rows_added, "rows_processed": rows_processed}
    finally:
        try:
            wb.save(path)
        except Exception:
            pass
        try:
            wb.close()
        except Exception:
            pass


def _direct_vlookup_run_for_xlsx(path, spec, run_suffix=u""):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    spec = _direct_vlookup_apply_range_to_spec(spec)
    if not spec.get("_vlookup_skip_full_dispatch"):
        if _direct_vlookup_normalize_join_type(spec.get("join_type")) == u"full":
            left_spec = dict(spec)
            left_spec["join_type"] = u"left"
            left_spec["_vlookup_skip_full_dispatch"] = True
            r1 = _direct_vlookup_run_for_xlsx(path, left_spec, run_suffix=run_suffix)
            mirror = _direct_vlookup_mirror_spec(spec)
            r2 = _direct_vlookup_run_for_xlsx(path, mirror, run_suffix=run_suffix)
            try:
                r3 = _direct_vlookup_append_right_only_xlsx(path, spec, run_suffix=run_suffix)
            except Exception:
                r3 = {"rows_added": 0, "rows_processed": 0}
            return {
                "rows_processed": int(r1.get("rows_processed", 0))
                + int(r2.get("rows_processed", 0))
                + int(r3.get("rows_processed", 0)),
                "rows_added": int(r1.get("rows_added", 0))
                + int(r2.get("rows_added", 0))
                + int(r3.get("rows_added", 0)),
                "columns_added": int(r1.get("columns_added", 0))
                + int(r2.get("columns_added", 0)),
                "keymap_size": int(r1.get("keymap_size", 0))
                + int(r2.get("keymap_size", 0)),
                "join_type": u"full",
            }

    left_name = _direct_vlookup_resolve_sheet_name(path, spec.get("left_sheet"))
    right_name = _direct_vlookup_resolve_sheet_name(path, spec.get("right_sheet"))
    if left_name is None or right_name is None:
        raise ValueError(u"ВПР: не найдены листы (left/right)")
    if _merge_identity_key(left_name) == _merge_identity_key(right_name):
        raise ValueError(u"ВПР: левый и правый листы совпадают")

    wb = load_workbook(path)
    try:
        left_ws = wb[left_name]
        right_ws = wb[right_name]

        # Параметры: 1-based → 0-based
        l_header = int(spec.get("left_header_row") or spec.get("left_start_row") or 1) - 1
        r_header = int(spec.get("right_header_row") or spec.get("right_start_row") or 1) - 1
        l_start_col = int(spec.get("left_start_col") or 1) - 1
        r_start_col = int(spec.get("right_start_col") or 1) - 1

        # Заголовки
        l_hdr_vals = []
        r_hdr_vals = []
        max_lc = int(left_ws.max_column or 1)
        max_rc = int(right_ws.max_column or 1)
        for ci in range(l_start_col + 1, max_lc + 1):
            v = left_ws.cell(row=l_header + 1, column=ci).value
            if v in (None, u""):
                # не обрываем: заголовки могут быть с пропусками
                pass
            l_hdr_vals.append(v)
        for ci in range(r_start_col + 1, max_rc + 1):
            v = right_ws.cell(row=r_header + 1, column=ci).value
            r_hdr_vals.append(v)

        left_key_tokens = spec.get("join_keys_left") or spec.get("join_keys") or []
        right_key_tokens = spec.get("join_keys_right") or spec.get("join_keys") or []
        extract_tokens = spec.get("extract_cols") or []

        l_key_cols = []
        for tok in left_key_tokens:
            idx = _direct_vlookup_parse_col_token(tok, l_hdr_vals, start_col0=l_start_col)
            if idx is None:
                raise ValueError(u"ВПР: не найден ключ на левом листе: %s" % tok)
            l_key_cols.append(idx)
        r_key_cols = []
        for tok in right_key_tokens:
            idx = _direct_vlookup_parse_col_token(tok, r_hdr_vals, start_col0=r_start_col)
            if idx is None:
                raise ValueError(u"ВПР: не найден ключ на правом листе: %s" % tok)
            r_key_cols.append(idx)

        r_extract_cols = []
        r_extract_titles = []
        for tok in extract_tokens:
            idx, title = _direct_vlookup_parse_extract_token(
                tok, r_hdr_vals, start_col0=r_start_col
            )
            if idx is None:
                raise ValueError(u"ВПР: не найден столбец извлечения: %s" % tok)
            r_extract_cols.append(idx)
            # Название выходного столбца: alias или заголовок правой таблицы.
            if title is None or unicode(title).strip() == u"":
                try:
                    hv = right_ws.cell(row=r_header + 1, column=idx + 1).value
                except Exception:
                    hv = None
                if hv is None or unicode(hv).strip() == u"":
                    hv = tok
                title = unicode(hv)
            r_extract_titles.append(unicode(title))

        trim_keys = bool(spec.get("trim_keys", False))
        fill_duplicates = bool(spec.get("fill_duplicates", True))
        highlight_hex = _direct_vlookup_highlight_color(spec)
        keymap = {}
        rows_added = 0
        r_data_start = r_header + 2
        r_max_row = int(right_ws.max_row or (r_header + 1))
        for row in right_ws.iter_rows(
            min_row=r_data_start,
            max_row=r_max_row,
            min_col=1,
            max_col=max_rc,
            values_only=True,
        ):
            key_parts = []
            for ci in r_key_cols:
                key_parts.append(row[ci] if ci < len(row) else None)
            key = _direct_vlookup_composite_key(key_parts, trim=trim_keys)
            if key not in keymap:
                keymap[key] = []
            row_obj = {}
            for ei in range(len(r_extract_cols)):
                ci = r_extract_cols[ei]
                row_obj[r_extract_titles[ei]] = row[ci] if ci < len(row) else None
            keymap[key].append(row_obj)
        # выходные колонки на левом листе
        out_cols = []
        col_out_index_by_title = {}
        header_row_1 = l_header + 1
        max_lc = int(left_ws.max_column or 1)
        used_id_keys = set()
        ci = 1
        while ci <= max_lc:
            hv = left_ws.cell(row=header_row_1, column=ci).value
            used_id_keys.add(_merge_identity_key(unicode(hv if hv is not None else u"")))
            ci = ci + 1

        for title in r_extract_titles:
            out_title = unicode(title)
            if bool(spec.get("column_suffix", True)) and run_suffix:
                out_title = out_title + unicode(run_suffix)

            # Внутри одного ВПР: не дублировать одинаковые extract_cols.
            if out_title in col_out_index_by_title:
                out_cols.append(col_out_index_by_title[out_title])
                continue

            desired_key = _merge_identity_key(out_title)

            # 1) Коллизия имён: если out_title уже есть среди заголовков,
            #    то новый столбец должен получить уникальное имя out_title_2, out_title_3, ...
            header_to_write = out_title
            if desired_key in used_id_keys:
                nn = 2
                cand = unicode(out_title) + u"_%d" % nn
                while _merge_identity_key(cand) in used_id_keys:
                    nn = nn + 1
                    cand = unicode(out_title) + u"_%d" % nn
                header_to_write = cand

            # 2) Добавляем новый столбец справа
            found = max_lc  # 0-based индекс в out_cols
            left_ws.cell(row=header_row_1, column=found + 1, value=header_to_write)
            out_cols.append(found)
            col_out_index_by_title[out_title] = found

            used_id_keys.add(_merge_identity_key(header_to_write))
            max_lc = max_lc + 1

            # Оформление заголовка: копируем стиль слева (столбец found)
            try:
                from copy import copy as _copy

                if found >= 1:
                    src_cell = left_ws.cell(row=header_row_1, column=found)
                    dst_cell = left_ws.cell(row=header_row_1, column=found + 1)
                    if getattr(src_cell, "has_style", False):
                        dst_cell._style = _copy(src_cell._style)
                    try:
                        dst_cell.font = _copy(src_cell.font)
                        dst_cell.fill = _copy(src_cell.fill)
                        dst_cell.border = _copy(src_cell.border)
                        dst_cell.alignment = _copy(src_cell.alignment)
                        dst_cell.number_format = src_cell.number_format
                        dst_cell.protection = _copy(src_cell.protection)
                        dst_cell.comment = src_cell.comment
                    except Exception:
                        pass
            except Exception:
                pass

            # Автоширина только по заголовку (полный скан столбца на тысячах строк тормозит ВПР)
            try:
                col_1 = found + 1
                v = left_ws.cell(row=header_row_1, column=col_1).value
                max_len = len(unicode(v if v is not None else u""))
                width = max(8, min(60, int(max_len * 1.1) + 2))
                try:
                    col_letter = _col_index_to_letters(found)
                    left_ws.column_dimensions[col_letter].width = float(width)
                except Exception:
                    pass
            except Exception:
                pass
        nf = _direct_vlookup_not_found_fill(spec)
        multi_one = bool(spec.get("multi_match_one_cell", False))
        multi_match_mode = _direct_vlookup_normalize_multi_match(
            spec.get("multi_match"), default=u"all"
        )
        is_inner = _direct_vlookup_normalize_join_type(spec.get("join_type"), default=u"left") == u"inner"
        hl_fill = None
        if highlight_hex is not None:
            try:
                from openpyxl.styles import PatternFill

                hl_fill = PatternFill(patternType="solid", fgColor=highlight_hex)
            except Exception:
                hl_fill = None

        data_start = l_header + 2
        max_row = int(left_ws.max_row or (l_header + 1))
        input_rows = []
        for row in left_ws.iter_rows(
            min_row=data_start,
            max_row=max_row,
            min_col=1,
            max_col=max_lc,
            values_only=True,
        ):
            input_rows.append(list(row))

        # (row_vals, single_match|None, multi_matches|None)
        output_specs = []
        processed = 0
        dbg_logged = 0
        for row_vals in input_rows:
            processed += 1
            if _direct_debug_enabled():
                if dbg_logged < 20 or (
                    processed % 50 == 0 and dbg_logged < 200
                ):
                    try:
                        _direct_journal(
                            u"",
                            unicode(left_name or u""),
                            u"",
                            u"",
                            u"инфо",
                            u"ВПР_XML debug rr=%d/%d"
                            % (int(processed), int(len(input_rows))),
                        )
                        dbg_logged += 1
                    except Exception:
                        pass
            key_parts = []
            for ci in l_key_cols:
                key_parts.append(row_vals[ci] if ci < len(row_vals) else None)
            key = _direct_vlookup_composite_key(key_parts, trim=trim_keys)
            matches = keymap.get(key) or []
            matches = _direct_vlookup_select_matches(matches, multi_match_mode)
            if not matches:
                if is_inner:
                    continue
                output_specs.append((row_vals, None, None))
            elif multi_one and len(matches) > 1:
                output_specs.append((row_vals, None, matches))
            elif fill_duplicates and len(matches) > 1:
                mi = 0
                while mi < len(matches):
                    output_specs.append((row_vals, matches[mi], None))
                    mi += 1
                rows_added += len(matches) - 1
            else:
                output_specs.append((row_vals, matches[0], None))

        def _write_vlookup_out_cells(rr, single, multi):
            if single is None and multi is None:
                for oc in out_cols:
                    cell = left_ws.cell(row=rr, column=oc + 1, value=nf)
                    if hl_fill is not None:
                        try:
                            cell.fill = hl_fill
                        except Exception:
                            pass
                return
            if multi is not None:
                for ei in range(len(out_cols)):
                    base = r_extract_titles[ei]
                    parts = []
                    for m in multi:
                        parts.append(
                            unicode(m.get(base) if m.get(base) is not None else u"")
                        )
                    cell = left_ws.cell(
                        row=rr,
                        column=out_cols[ei] + 1,
                        value=u"\n".join(parts),
                    )
                    if hl_fill is not None:
                        try:
                            cell.fill = hl_fill
                        except Exception:
                            pass
                return
            for ei in range(len(out_cols)):
                base = r_extract_titles[ei]
                cell = left_ws.cell(
                    row=rr,
                    column=out_cols[ei] + 1,
                    value=single.get(base),
                )
                if hl_fill is not None:
                    try:
                        cell.fill = hl_fill
                    except Exception:
                        pass

        old_n = len(input_rows)
        new_n = len(output_specs)
        if rows_added > 0 or new_n != old_n:
            # Одна пересборка блока данных вместо insert_rows на каждой строке.
            if old_n > 0:
                left_ws.delete_rows(data_start, old_n)
            if new_n > 0:
                left_ws.insert_rows(data_start, new_n)
            i = 0
            while i < new_n:
                row_vals, single, multi = output_specs[i]
                rr = data_start + i
                cc = 0
                while cc < len(row_vals):
                    left_ws.cell(row=rr, column=cc + 1, value=row_vals[cc])
                    cc += 1
                _write_vlookup_out_cells(rr, single, multi)
                i += 1
        else:
            i = 0
            while i < new_n:
                row_vals, single, multi = output_specs[i]
                rr = data_start + i
                _write_vlookup_out_cells(rr, single, multi)
                i += 1

        if _direct_debug_enabled():
            try:
                _direct_journal(
                    u"",
                    unicode(left_name or u""),
                    u"",
                    u"",
                    u"инфо",
                    u"ВПР_XML: строки обработаны (%d), сохранение xlsx…"
                    % int(processed),
                )
            except Exception:
                pass
        return {
            "rows_processed": processed,
            "rows_added": rows_added,
            "columns_added": len(out_cols),
            "keymap_size": len(keymap),
            "join_type": u"inner" if is_inner else u"left",
        }
    finally:
        try:
            wb.save(path)
        except Exception:
            pass
        try:
            wb.close()
        except Exception:
            pass


def direct_run_vlookup(target_path, settings, one_sheet):
    """
    ВПР для xml_excel_ods: выполняется в файле результата до reopen UNO.
    Берёт спецификации из settings['processing_pipeline'] (kind='vlookup').
    """
    pipeline = (settings or {}).get("processing_pipeline") or []
    vsteps = [s for s in pipeline if (s or {}).get("kind") == "vlookup"]
    if not vsteps:
        return
    run_n = 0
    for step in vsteps:
        spec = (step or {}).get("spec") or {}
        run_n += 1
        run_suffix = ""
        if bool(spec.get("column_suffix", True)):
            run_suffix = u"_R%d" % int(run_n)
        label = u"%s → %s" % (spec.get("left_sheet", u"?"), spec.get("right_sheet", u"?"))
        _direct_journal(u"", unicode(spec.get("left_sheet") or u""), u"", u"", u"ВПР", u"старт: %s" % label)
        if _is_xlsx_target(target_path):
            res = _direct_vlookup_run_for_xlsx(target_path, spec, run_suffix=run_suffix)
        else:
            res = _direct_vlookup_run_for_ods(target_path, spec, run_suffix=run_suffix)
        _direct_journal(
            u"",
            unicode(spec.get("left_sheet") or u""),
            u"",
            unicode(res.get("rows_processed", 0)),
            u"ok",
            u"ВПР: добавлено=%d, столбцов=%d, ключей=%d"
            % (
                int(res.get("rows_added", 0)),
                int(res.get("columns_added", 0)),
                int(res.get("keymap_size", 0)),
            ),
        )


def _direct_run_vlookup_step(target_path, spec, run_n):
    """Один шаг ВПР_XML в direct-пайплайне."""
    run_suffix = u""
    if bool(spec.get("column_suffix", True)):
        run_suffix = u"_R%d" % int(run_n)
    label = u"%s → %s" % (spec.get("left_sheet", u"?"), spec.get("right_sheet", u"?"))
    _direct_journal(
        u"",
        unicode(spec.get("left_sheet") or u""),
        u"",
        u"",
        u"ВПР_XML",
        u"старт: %s" % label,
    )
    if _is_xlsx_target(target_path):
        res = _direct_vlookup_run_for_xlsx(target_path, spec, run_suffix=run_suffix)
    else:
        res = _direct_vlookup_run_for_ods(target_path, spec, run_suffix=run_suffix)
    _direct_journal(
        u"",
        unicode(spec.get("left_sheet") or u""),
        u"",
        unicode(res.get("rows_processed", 0)),
        u"ok",
        u"ВПР_XML: добавлено=%d, столбцов=%d, ключей=%d"
        % (
            int(res.get("rows_added", 0)),
            int(res.get("columns_added", 0)),
            int(res.get("keymap_size", 0)),
        ),
    )


def _direct_vlookup_run_for_ods(path, spec, run_suffix=u""):
    """
    Direct ВПР для ODS (odfpy). Поддерживает опции:
    - trim_keys (и регистронезависимость через casefold в _direct_vlookup_key_part)
    - column_suffix
    - not_found_fill (включая '_ПУСТО_' → пустая ячейка)
    - multi_match_one_cell
    - fill_duplicates (вставка строк)
    Подсветка highlight_color для ODS пока не реализуется (безопасный пропуск).
    """
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell

    spec = _direct_vlookup_apply_range_to_spec(spec)
    if not spec.get("_vlookup_skip_full_dispatch"):
        if _direct_vlookup_normalize_join_type(spec.get("join_type")) == u"full":
            left_spec = dict(spec)
            left_spec["join_type"] = u"left"
            left_spec["_vlookup_skip_full_dispatch"] = True
            r1 = _direct_vlookup_run_for_ods(path, left_spec, run_suffix=run_suffix)
            mirror = _direct_vlookup_mirror_spec(spec)
            r2 = _direct_vlookup_run_for_ods(path, mirror, run_suffix=run_suffix)
            return {
                "rows_processed": int(r1.get("rows_processed", 0))
                + int(r2.get("rows_processed", 0)),
                "rows_added": int(r1.get("rows_added", 0)) + int(r2.get("rows_added", 0)),
                "columns_added": int(r1.get("columns_added", 0))
                + int(r2.get("columns_added", 0)),
                "keymap_size": int(r1.get("keymap_size", 0))
                + int(r2.get("keymap_size", 0)),
                "join_type": u"full",
            }

    left_name = _direct_vlookup_resolve_sheet_name(path, spec.get("left_sheet"))
    right_name = _direct_vlookup_resolve_sheet_name(path, spec.get("right_sheet"))
    if left_name is None or right_name is None:
        raise ValueError(u"ВПР: не найдены листы (left/right)")
    if _merge_identity_key(left_name) == _merge_identity_key(right_name):
        raise ValueError(u"ВПР: левый и правый листы совпадают")

    doc = load(path)
    try:
        left_tbl = None
        right_tbl = None
        for t in doc.spreadsheet.getElementsByType(Table):
            nm = unicode(t.getAttribute("name") or u"")
            if nm == unicode(left_name):
                left_tbl = t
            elif nm == unicode(right_name):
                right_tbl = t
        if left_tbl is None or right_tbl is None:
            raise ValueError(u"ВПР: лист не найден в ODS")

        left_rows = left_tbl.getElementsByType(TableRow)
        right_rows = right_tbl.getElementsByType(TableRow)

        l_header = int(spec.get("left_header_row") or spec.get("left_start_row") or 1) - 1
        r_header = int(spec.get("right_header_row") or spec.get("right_start_row") or 1) - 1
        l_start_col = int(spec.get("left_start_col") or 1) - 1
        r_start_col = int(spec.get("right_start_col") or 1) - 1

        # header values
        l_hdr_vals = []
        r_hdr_vals = []
        r_hdr_by_col = {}
        if l_header < len(left_rows):
            for col_idx, val in _ods_row_col_values(left_rows[l_header]):
                if col_idx >= l_start_col:
                    l_hdr_vals.append(val)
        if r_header < len(right_rows):
            for col_idx, val in _ods_row_col_values(right_rows[r_header]):
                if col_idx >= r_start_col:
                    r_hdr_vals.append(val)
                    r_hdr_by_col[int(col_idx)] = val

        left_key_tokens = spec.get("join_keys_left") or spec.get("join_keys") or []
        right_key_tokens = spec.get("join_keys_right") or spec.get("join_keys") or []
        extract_tokens = spec.get("extract_cols") or []

        l_key_cols = []
        for tok in left_key_tokens:
            idx = _direct_vlookup_parse_col_token(tok, l_hdr_vals, start_col0=l_start_col)
            if idx is None:
                raise ValueError(u"ВПР: не найден ключ на левом листе: %s" % tok)
            l_key_cols.append(idx)
        r_key_cols = []
        for tok in right_key_tokens:
            idx = _direct_vlookup_parse_col_token(tok, r_hdr_vals, start_col0=r_start_col)
            if idx is None:
                raise ValueError(u"ВПР: не найден ключ на правом листе: %s" % tok)
            r_key_cols.append(idx)

        r_extract_cols = []
        r_extract_titles = []
        for tok in extract_tokens:
            idx, title = _direct_vlookup_parse_extract_token(
                tok, r_hdr_vals, start_col0=r_start_col
            )
            if idx is None:
                raise ValueError(u"ВПР: не найден столбец извлечения: %s" % tok)
            r_extract_cols.append(idx)
            if title is None or unicode(title).strip() == u"":
                hv = r_hdr_by_col.get(int(idx))
                if hv is None or unicode(hv).strip() == u"":
                    hv = tok
                title = unicode(hv)
            r_extract_titles.append(unicode(title))

        trim_keys = bool(spec.get("trim_keys", False))
        fill_duplicates = bool(spec.get("fill_duplicates", True))
        nf = _direct_vlookup_not_found_fill(spec)
        nf_is_empty = _merge_identity_key(nf) in (_merge_identity_key(u"_ПУСТО_"), _merge_identity_key(u"__ПУСТО__"))
        multi_one = bool(spec.get("multi_match_one_cell", False))
        multi_match_mode = _direct_vlookup_normalize_multi_match(
            spec.get("multi_match"), default=u"all"
        )
        is_inner = _direct_vlookup_normalize_join_type(spec.get("join_type"), default=u"left") == u"inner"

        # keymap for right
        keymap = {}
        for rr in range(r_header + 1, len(right_rows)):
            row_elem = right_rows[rr]
            # build dict for extract
            row_obj = {}
            # read key parts + extract parts via _ods_row_col_values
            by_col = {}
            for col_idx, val in _ods_row_col_values(row_elem):
                by_col[int(col_idx)] = val
            key_parts = [by_col.get(ci) for ci in r_key_cols]
            key = _direct_vlookup_composite_key(key_parts, trim=trim_keys)
            for ei in range(len(r_extract_cols)):
                row_obj[r_extract_titles[ei]] = by_col.get(r_extract_cols[ei])
            if key not in keymap:
                keymap[key] = []
            keymap[key].append(row_obj)

        # output columns: append at end, ensure unique header names
        used_keys = set()
        if l_header < len(left_rows):
            for col_idx, val in _ods_row_col_values(left_rows[l_header]):
                used_keys.add(_merge_identity_key(unicode(val if val is not None else u"")))
        out_cols = []
        col_out_by_title = {}
        # find current max col index in header row
        max_col0 = -1
        if l_header < len(left_rows):
            for col_idx, _val in _ods_row_col_values(left_rows[l_header]):
                max_col0 = max(max_col0, int(col_idx))
        if max_col0 < 0:
            max_col0 = l_start_col - 1

        for title in r_extract_titles:
            out_title = unicode(title)
            if bool(spec.get("column_suffix", True)) and run_suffix:
                out_title = out_title + unicode(run_suffix)
            if out_title in col_out_by_title:
                out_cols.append(col_out_by_title[out_title])
                continue
            header_to_write = out_title
            if _merge_identity_key(header_to_write) in used_keys:
                nn = 2
                cand = unicode(out_title) + u"_%d" % nn
                while _merge_identity_key(cand) in used_keys:
                    nn += 1
                    cand = unicode(out_title) + u"_%d" % nn
                header_to_write = cand
            max_col0 += 1
            target_col = max_col0
            # ensure header row exists
            while l_header >= len(left_rows):
                new_r = TableRow()
                left_tbl.addElement(new_r)
                left_rows.append(new_r)
            _direct_ods_set_cell_at_col(left_rows[l_header], target_col, _ods_make_cell(header_to_write))
            used_keys.add(_merge_identity_key(header_to_write))
            out_cols.append(target_col)
            col_out_by_title[out_title] = target_col

        # apply to left rows (план в памяти, как в XLSX)
        rows_added = 0
        processed = 0
        dbg_logged = 0
        input_rows = []
        rr = l_header + 1
        while rr < len(left_rows):
            by_col = {}
            for col_idx, val in _ods_row_col_values(left_rows[rr]):
                by_col[int(col_idx)] = val
            input_rows.append(by_col)
            rr += 1

        output_specs = []
        for by_col in input_rows:
            processed += 1
            if _direct_debug_enabled():
                if dbg_logged < 20 or (
                    processed % 50 == 0 and dbg_logged < 200
                ):
                    try:
                        _direct_journal(
                            u"",
                            unicode(left_name or u""),
                            u"",
                            u"",
                            u"инфо",
                            u"ВПР_XML debug rr=%d/%d"
                            % (int(processed), int(len(input_rows))),
                        )
                        dbg_logged += 1
                    except Exception:
                        pass
            key_parts = [by_col.get(ci) for ci in l_key_cols]
            key = _direct_vlookup_composite_key(key_parts, trim=trim_keys)
            matches = keymap.get(key) or []
            matches = _direct_vlookup_select_matches(matches, multi_match_mode)
            if not matches:
                if is_inner:
                    continue
                output_specs.append((by_col, None, None))
            elif multi_one and len(matches) > 1:
                output_specs.append((by_col, None, matches))
            elif fill_duplicates and len(matches) > 1:
                mi = 0
                while mi < len(matches):
                    output_specs.append((by_col, matches[mi], None))
                    mi += 1
                rows_added += len(matches) - 1
            else:
                output_specs.append((by_col, matches[0], None))

        base_max_col0 = max_col0

        def _vlookup_out_values(single, multi):
            if single is None and multi is None:
                out = {}
                for oc in out_cols:
                    out[int(oc)] = u"" if nf_is_empty else nf
                return out
            if multi is not None:
                out = {}
                for ei in range(len(out_cols)):
                    base = r_extract_titles[ei]
                    parts = []
                    for m in multi:
                        parts.append(
                            unicode(m.get(base) if m.get(base) is not None else u"")
                        )
                    out[int(out_cols[ei])] = u"\n".join(parts)
                return out
            out = {}
            for ei in range(len(out_cols)):
                base = r_extract_titles[ei]
                out[int(out_cols[ei])] = single.get(base)
            return out

        def _write_vlookup_out_row(row_elem, single, multi):
            out_vals = _vlookup_out_values(single, multi)
            for oc, val in out_vals.items():
                _direct_ods_set_cell_at_col(
                    row_elem,
                    int(oc),
                    _ods_make_cell(val if val is not None else u""),
                )

        def _build_vlookup_data_row(by_col, single, multi):
            from odf.table import TableRow

            out_vals = _vlookup_out_values(single, multi)
            max_c = int(base_max_col0)
            if out_cols:
                max_c = max(max_c, max(int(oc) for oc in out_cols))
            nr = TableRow()
            c = 0
            while c <= max_c:
                if c in out_vals:
                    nr.addElement(
                        _ods_make_cell(
                            out_vals[c] if out_vals[c] is not None else u""
                        )
                    )
                else:
                    val = by_col.get(c)
                    nr.addElement(_ods_make_cell(val if val is not None else u""))
                c += 1
            return nr

        if rows_added > 0 or len(output_specs) != len(input_rows):
            while len(left_rows) > l_header + 1:
                left_tbl.removeChild(left_rows[-1])
                left_rows.pop()
            i = 0
            while i < len(output_specs):
                by_col, single, multi = output_specs[i]
                nr = _build_vlookup_data_row(by_col, single, multi)
                left_tbl.addElement(nr)
                left_rows.append(nr)
                i += 1
        else:
            i = 0
            while i < len(output_specs):
                by_col, single, multi = output_specs[i]
                row_elem = left_rows[l_header + 1 + i]
                _write_vlookup_out_row(row_elem, single, multi)
                i += 1

        if _direct_debug_enabled():
            try:
                _direct_journal(
                    u"",
                    unicode(left_name or u""),
                    u"",
                    u"",
                    u"инфо",
                    u"ВПР_XML: строки обработаны (%d), сохранение ods…"
                    % int(processed),
                )
            except Exception:
                pass
        return {
            "rows_processed": processed,
            "rows_added": rows_added,
            "columns_added": len(out_cols),
            "keymap_size": len(keymap),
            "join_type": u"inner" if is_inner else u"left",
        }
    finally:
        try:
            doc.save(path)
        except Exception:
            pass


def _direct_build_read_plan(settings, one_sheet):
    files = settings.get("files") or []
    is_current_book_flags = settings.get("is_current_book_flags") or []
    per_file_blocks = settings.get("per_file_blocks") or []
    per_file_sheet_specs = settings.get("per_file_sheet_specs") or []
    per_file_col_specs = settings.get("per_file_col_specs") or []
    per_file_skip_row_specs = settings.get("per_file_skip_row_specs") or []
    per_file_source_extra = settings.get("per_file_source_extra") or []
    overrides = (settings.get("direct_source_overrides") or {}) if isinstance(settings, dict) else {}
    cell_plan = settings.get("cell_to_column_plan") or []
    ctc_per_file = _direct_cell_specs_per_file(cell_plan, len(files))
    target_path = unicode(settings.get("direct_target_path") or u"").strip()
    path_display = (settings.get("path_display") or "none").strip().lower()

    if not per_file_sheet_specs:
        raise ValueError("xml_excel_ods: пустые per_file_sheet_specs")

    plan = []
    fi = 0
    while fi < len(files):
        is_curr = is_current_book_flags[fi] if fi < len(is_current_book_flags) else False
        blk = (
            per_file_blocks[fi]
            if fi < len(per_file_blocks)
            else (per_file_blocks[0] if per_file_blocks else None)
        )
        sheet_specs = (
            per_file_sheet_specs[fi]
            if fi < len(per_file_sheet_specs)
            else per_file_sheet_specs[0]
        )
        col_specs = per_file_col_specs[fi] if fi < len(per_file_col_specs) else []
        file_note = _direct_collect_file_note(blk, sheet_specs, col_specs)

        if is_curr:
            if target_path == u"" or not os.path.isfile(target_path):
                _direct_journal(
                    u"(текущая книга)",
                    u"",
                    u"",
                    u"",
                    u"проблема",
                    u"«Эта_книга»: книга результата не сохранена на диск",
                )
                fi += 1
                continue
            src_path = os.path.abspath(target_path)
            read_path = src_path
            _direct_journal(
                u"(текущая книга)",
                u"",
                u"",
                u"",
                u"файл",
                file_note,
            )
            _log("источник «Эта_книга»: %s" % read_path)
        else:
            src_path = unicode(files[fi] or u"").strip()
            _src_is_remote_url = False
            try:
                from libre_macros_bundle_paths import is_remote_url

                _src_is_remote_url = bool(is_remote_url(src_path))
            except Exception:
                _src_is_remote_url = u"://" in src_path and not src_path.lower().startswith(u"file://")
            if _src_is_remote_url:
                _direct_journal(
                    _format_source_path(src_path, path_display if path_display != "none" else "short"),
                    u"",
                    u"",
                    u"",
                    u"проблема",
                    u"URL-источник: xml_excel_ods — только локальные файлы; smb:// — режим UNO",
                )
                fi += 1
                continue
            src_path = os.path.abspath(src_path)
            read_path = os.path.abspath(overrides.get(src_path, src_path))
            try:
                xml_temps = settings.setdefault("_direct_xml_temp_paths", [])
                src_extra = per_file_source_extra[fi] if fi < len(per_file_source_extra) else None
                read_path = _direct_resolve_read_path(
                    read_path, xml_temps, sheet_specs=sheet_specs, source_extra=src_extra
                )
            except ValueError as err:
                _direct_journal(
                    _format_source_path(read_path, path_display if path_display != "none" else "short"),
                    u"",
                    u"",
                    u"",
                    u"проблема",
                    unicode(err),
                )
                fi += 1
                continue
            try:
                from libre_macros_source_extra_lib import source_extra_last_row_column

                last_row_col_spec = source_extra_last_row_column(src_extra)
            except Exception:
                last_row_col_spec = u""
            _direct_journal(
                _format_source_path(read_path, path_display if path_display != "none" else "short"),
                u"",
                u"",
                u"",
                u"файл",
                file_note,
            )

        hdr_row0 = int(blk.get("hdr_row", 0) if blk else 0)
        start_col0 = int(blk.get("start_col", 0) if blk else 0)
        row_count = blk.get("row_count") if blk else None
        col_count = blk.get("col_count") if blk else None

        col_kind = _classify_column_specs(col_specs)
        if col_kind == "MIXED":
            raise ValueError("xml_excel_ods: MIXED в фильтре «Столбцы» не поддержан")

        sheet_kind = _classify_sheet_specs(sheet_specs)
        if sheet_kind == "MIXED":
            raise ValueError("xml_excel_ods: MIXED в фильтре «Листы» не поддержан")

        if _is_xlsx(read_path) or _is_xlsm(read_path):
            sheet_names = _xlsx_sheet_names(read_path)
            src_kind = "xlsx"
        elif _is_ods(read_path):
            sheet_names = _ods_sheet_names(read_path)
            src_kind = "ods"
        else:
            msg = u"неподдерживаемый источник: %s" % read_path
            _log(msg)
            _direct_journal(
                _format_source_path(read_path, "short"),
                u"",
                u"",
                u"",
                u"проблема",
                msg,
            )
            fi += 1
            continue

        skip_text = per_file_skip_row_specs[fi] if fi < len(per_file_skip_row_specs) else u""

        for slot in _iter_sheet_slots_file(sheet_names, sheet_specs, sheet_kind):
            if is_curr and _direct_is_service_sheet_name(slot.get("source_name")):
                continue
            target_key = _DIRECT_MERGE_ONE_SHEET if one_sheet else slot["target_key"]
            src_index = slot["source_index"]
            if sheet_kind == "NAME":
                try:
                    src_index = sheet_names.index(slot["source_name"])
                except Exception:
                    src_index = 0
            if src_index < 0 or src_index >= len(sheet_names):
                continue
            src_sheet_name = unicode(sheet_names[src_index] or u"")
            if isinstance(settings, dict):
                skipped_map = settings.setdefault("result_skipped_top_rows", {})
                try:
                    prev_sk = int(skipped_map.get(target_key) or 0)
                except Exception:
                    prev_sk = 0
                if int(hdr_row0) > prev_sk:
                    skipped_map[unicode(target_key)] = int(hdr_row0)

            # sheet в JSON — имя листа источника; пересчёт на каждый лист
            skip_parsed = _direct_parse_skip_spec(skip_text, src_sheet_name)
            if skip_parsed.get("mode") == "formula":
                _log(
                    "Пропуск_строк_источника: режим формулы не поддерживается в xml_excel_ods — %s"
                    % skip_text
                )
                _direct_journal(
                    _direct_journal_source_label(
                        {"is_current_book": is_curr, "src_path": src_path},
                        path_display,
                    ),
                    src_sheet_name,
                    u"",
                    u"",
                    u"инфо",
                    u"Пропуск_строк_источника: режим формулы не поддерживается — %s"
                    % skip_text,
                )
                skip_parsed = {"mode": "none"}

            import time
            t_bounds = time.time()
            used_end_col, used_end_row = _direct_resolve_sheet_used_bounds(
                read_path, src_kind, src_index, hdr_row0, start_col0,
                last_row_col_spec=last_row_col_spec,
            )
            bounds_elapsed_ms = (time.time() - t_bounds) * 1000.0
            last_col, last_data_row, first_data_row = _direct_effective_limits(
                used_end_col, used_end_row, hdr_row0, start_col0, row_count, col_count
            )
            if unicode(last_row_col_spec or u"").strip() != u"":
                rows_to_copy = 0
                if last_data_row >= first_data_row:
                    rows_to_copy = int(last_data_row) - int(first_data_row) + 1
                _direct_journal(
                    _direct_journal_source_label(
                        {
                            "is_current_book": is_curr,
                            "src_path": src_path,
                            "src_sheet_name": src_sheet_name,
                        },
                        path_display,
                    ),
                    src_sheet_name,
                    u"",
                    rows_to_copy,
                    u"last_row_column",
                    u"last_row_column=%s; строк_к_копированию=%d; диапазон=%d..%d"
                    % (
                        unicode(last_row_col_spec),
                        rows_to_copy,
                        int(first_data_row) + 1,
                        int(last_data_row) + 1,
                    ) + u"; время=%.1f ms" % bounds_elapsed_ms if rows_to_copy > 0 else
                    u"last_row_column=%s; строк_к_копированию=0; время=%.1f ms"
                    % (unicode(last_row_col_spec), bounds_elapsed_ms),
                )
            if last_data_row < first_data_row or last_col < start_col0:
                _direct_journal(
                    _direct_journal_source_label(
                        {
                            "is_current_book": is_curr,
                            "src_path": src_path,
                            "src_sheet_name": src_sheet_name,
                        },
                        path_display,
                    ),
                    src_sheet_name,
                    u"",
                    0,
                    u"пропуск",
                    u"нет данных в блоке",
                )
                _log(
                    "пропуск: нет данных в блоке — %s / «%s»"
                    % (src_path, src_sheet_name)
                )
                # «На разные листы» + MERGE_DO_EMPTY_SHEETS: лист с заголовком и сообщением.
                if (not one_sheet) and _direct_do_empty_sheets(settings):
                    empty_columns = []
                    try:
                        if last_col >= start_col0:
                            header_pairs = _direct_read_header_pairs(
                                read_path, src_kind, src_index, hdr_row0, start_col0, last_col
                            )
                            spec = (blk or {}).get("header_spec")
                            if spec:
                                try:
                                    from libre_macros_header_lib import flatten_headers_path
                                    header_pairs = flatten_headers_path(
                                        read_path, src_kind, src_index, spec,
                                        start_col0, (blk or {}).get("col_count"),
                                    )
                                except Exception:
                                    pass
                            empty_columns = _direct_build_column_list(
                                header_pairs, col_specs, col_kind, start_col0, col_count
                            )
                    except Exception:
                        empty_columns = []
                    plan.append(
                        {
                            "fi": fi,
                            "is_current_book": is_curr,
                            "src_path": src_path,
                            "read_path": read_path,
                            "src_kind": src_kind,
                            "target_key": target_key,
                            "src_index": src_index,
                            "src_sheet_name": src_sheet_name,
                            "hdr_row0": hdr_row0,
                            "start_col0": start_col0,
                            "row_count": row_count,
                            "col_count": col_count,
                            "first_data_row": int(first_data_row),
                            "last_data_row": int(first_data_row) - 1,
                            "last_col": last_col,
                            "columns": empty_columns,
                            "skip_parsed": {"mode": "none"},
                            "cell_values": {},
                            "empty_placeholder": True,
                        }
                    )
                continue

            header_pairs = _direct_read_header_pairs(
                read_path, src_kind, src_index, hdr_row0, start_col0, last_col
            )
            spec = (blk or {}).get("header_spec")
            if spec:
                try:
                    from libre_macros_header_lib import flatten_headers_path
                    header_pairs = flatten_headers_path(
                        read_path, src_kind, src_index, spec,
                        start_col0, (blk or {}).get("col_count"),
                    )
                except Exception:
                    pass
            columns = _direct_build_column_list(
                header_pairs, col_specs, col_kind, start_col0, col_count
            )
            if len(columns) == 0:
                _direct_journal(
                    _direct_journal_source_label(
                        {
                            "is_current_book": is_curr,
                            "src_path": src_path,
                            "src_sheet_name": src_sheet_name,
                        },
                        path_display,
                    ),
                    src_sheet_name,
                    u"",
                    0,
                    u"пропуск",
                    u"нет столбцов",
                )
                _log(
                    "пропуск: нет столбцов — %s / «%s»"
                    % (src_path, src_sheet_name)
                )
                continue

            if skip_parsed.get("mode") == "columns":
                skip_parsed = dict(skip_parsed)
                skip_parsed["col_indices"] = _direct_resolve_skip_indices(
                    header_pairs, skip_parsed.get("parts") or []
                )
                if not skip_parsed["col_indices"]:
                    _direct_journal(
                        _direct_journal_source_label(
                            {
                                "is_current_book": is_curr,
                                "src_path": src_path,
                            },
                            path_display,
                        ),
                        src_sheet_name,
                        u"",
                        u"",
                        u"инфо",
                        u"Пропуск_строк_источника: столбцы не найдены — %s"
                        % skip_text,
                    )
                    _log(
                        "Пропуск_строк_источника: столбцы не найдены — %s"
                        % skip_text
                    )
                    skip_parsed["mode"] = "none"

            cell_specs = ctc_per_file[fi] if fi < len(ctc_per_file) else []
            if cell_specs:
                cell_specs = [
                    s
                    for s in cell_specs
                    if _direct_cell_spec_applies_to_target_sheet(s, target_key)
                    and _direct_cell_spec_applies_to_source_sheet(s, src_sheet_name)
                ]
            cell_values = _direct_prefetch_cell_values(read_path, src_kind, src_index, cell_specs)
            _direct_prefetch_add_to_variables_map(
                fi,
                src_path if not is_curr else u"",
                is_curr,
                path_display,
                src_sheet_name,
                read_path,
                src_kind,
                src_index,
            )

            plan.append(
                {
                    "fi": fi,
                    "is_current_book": is_curr,
                    "src_path": src_path,
                    "read_path": read_path,
                    "src_kind": src_kind,
                    "target_key": target_key,
                    "src_index": src_index,
                    "src_sheet_name": src_sheet_name,
                    "hdr_row0": hdr_row0,
                    "start_col0": start_col0,
                    "row_count": row_count,
                    "col_count": col_count,
                    "first_data_row": first_data_row,
                    "last_data_row": last_data_row,
                    "last_col": last_col,
                    "columns": columns,
                    "skip_parsed": skip_parsed,
                    "cell_values": cell_values,
                }
            )
        fi += 1
    return plan


def _direct_run_merge(target_path, settings, one_sheet):
    path_mode = (settings.get("path_display") or "none").strip().lower()
    datetime_mode = (settings.get("datetime_display") or "none").strip().lower()
    service_cols = []
    if path_mode != "none":
        service_cols.append(u"#Путь")
    if datetime_mode != "none":
        service_cols.append(u"#ДатаВремяФайла")

    try:
        return _direct_run_merge_core(target_path, settings, one_sheet, service_cols)
    finally:
        _direct_cleanup_xml_temp_paths(settings)


def _direct_run_merge_core(target_path, settings, one_sheet, service_cols):
    path_mode = (settings.get("path_display") or "none").strip().lower()
    datetime_mode = (settings.get("datetime_display") or "none").strip().lower()
    plan = _direct_build_read_plan(settings, one_sheet)
    if not plan:
        _log("план чтения пуст")
        _direct_journal(u"", u"", u"", u"", u"инфо", u"план чтения пуст — нет данных для переноса")
        return

    agg = {}

    def _get_bucket(target_key):
        if target_key not in agg:
            headers = list(service_cols)
            agg[target_key] = {"headers": headers, "h2i": {h: i for i, h in enumerate(headers)}}
        return agg[target_key]

    def _ensure_header(bucket, h):
        hh = unicode(h or u"").strip()
        if hh == u"":
            return -1
        if hh in bucket["h2i"]:
            return bucket["h2i"][hh]
        bucket["h2i"][hh] = len(bucket["headers"])
        bucket["headers"].append(hh)
        return bucket["h2i"][hh]

    for item in plan:
        bucket = _get_bucket(item["target_key"])
        for _c, h in item["columns"]:
            _ensure_header(bucket, h)
        for col_name in (item.get("cell_values") or {}).keys():
            _ensure_header(bucket, col_name)

    if _is_ods(target_path):
        import odf_bundled  # noqa: F401
        from odf.opendocument import load

        doc = load(target_path)
        style_cache = (
            _direct_ods_style_cache_new(doc)
            if _direct_wants_convert_to_numbers()
            else None
        )
        tables = {}
        for target_key, bucket in agg.items():
            table = _ods_get_or_create_table(doc, target_key)
            _ods_clear_table_rows(table)
            _ods_add_row(table, bucket["headers"])
            tables[target_key] = table
        written = {}
        for item in plan:
            target_key = item["target_key"]
            table = tables.get(target_key)
            bucket = agg.get(target_key) or {}
            headers_out = bucket.get("headers") or []
            h2i = bucket.get("h2i") or {}
            if table is None or not headers_out:
                continue
            _log("source: %s -> %s" % (item["src_path"], target_key))
            _direct_ui_notify_source(item["src_path"], item["src_sheet_name"])
            _direct_ui_yield(force=True)
            col_headers = [h for _c, h in item["columns"]]
            gen = _direct_iter_filtered_rows(
                item["read_path"],
                item["src_kind"],
                item["src_index"],
                item["first_data_row"],
                item["last_data_row"],
                item["columns"],
                item["skip_parsed"],
            )
            item_rows = 0
            for row in gen:
                out = [u""] * len(headers_out)
                si = 0
                if path_mode != "none":
                    out[si] = _direct_service_path_value(item, path_mode)
                    si += 1
                if datetime_mode != "none":
                    out[si] = _direct_service_datetime_value(item, datetime_mode)
                    si += 1
                for ci in range(min(len(col_headers), len(row))):
                    h = col_headers[ci]
                    idx = h2i.get(h)
                    if idx is None:
                        continue
                    out[idx] = row[ci]
                for col_name, val in (item.get("cell_values") or {}).items():
                    idx = h2i.get(col_name)
                    if idx is not None:
                        out[idx] = val
                _direct_ods_add_row_coerced(table, out, style_cache)
                item_rows += 1
                written[target_key] = written.get(target_key, 0) + 1
                if written[target_key] % 500 == 0:
                    _log("written %s: %d" % (target_key, written[target_key]))
                    _direct_ui_yield()
            _direct_journal(
                _direct_journal_source_label(item, path_mode),
                item.get("src_sheet_name") or u"",
                target_key,
                item_rows,
                u"скопировано",
                u"→ «%s»" % target_key,
            )
        if _direct_do_empty_sheets(settings):
            msg = _direct_empty_sheet_message()
            for target_key, bucket in agg.items():
                if int(written.get(target_key) or 0) > 0:
                    continue
                table = tables.get(target_key)
                if table is None:
                    continue
                headers_out = bucket.get("headers") or list(service_cols)
                row = [u""] * max(1, len(headers_out))
                row[0] = msg
                _ods_add_row(table, row)
                written[target_key] = 1
                _direct_journal(u"", u"", target_key, 0, u"пусто", msg)
        else:
            # Не оставлять пустые листы без данных.
            for target_key in list(agg.keys()):
                if int(written.get(target_key) or 0) > 0:
                    continue
                try:
                    # table уже создан с одной строкой заголовка — удалить лист из ODS сложно;
                    # оставляем только заголовок без сообщения (как раньше для empty plan skip).
                    pass
                except Exception:
                    pass
        doc.save(target_path)
        _log("saved(ods): %s sheets=%d" % (target_path, len(agg)))
        return

    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(target_path)
    try:
        ws_by_key = {}
        written = {}
        for target_key, bucket in agg.items():
            if target_key in wb.sheetnames:
                ws = wb[target_key]
                ws.delete_rows(1, ws.max_row or 1)
            else:
                ws = wb.create_sheet(title=target_key)
            ws.append(bucket["headers"])
            ws_by_key[target_key] = ws
            written[target_key] = 0

        for item in plan:
            target_key = item["target_key"]
            ws = ws_by_key.get(target_key)
            bucket = agg.get(target_key) or {}
            headers_out = bucket.get("headers") or []
            h2i = bucket.get("h2i") or {}
            if ws is None or not headers_out:
                continue
            _log("source: %s -> %s" % (item["src_path"], target_key))
            _direct_ui_notify_source(item["src_path"], item["src_sheet_name"])
            _direct_ui_yield(force=True)
            col_headers = [h for _c, h in item["columns"]]
            gen = _direct_iter_filtered_rows(
                item["read_path"],
                item["src_kind"],
                item["src_index"],
                item["first_data_row"],
                item["last_data_row"],
                item["columns"],
                item["skip_parsed"],
            )
            item_rows = 0
            for row in gen:
                out = [None] * len(headers_out)
                si = 0
                if path_mode != "none":
                    out[si] = _direct_service_path_value(item, path_mode)
                    si += 1
                if datetime_mode != "none":
                    out[si] = _direct_service_datetime_value(item, datetime_mode)
                    si += 1
                for ci in range(min(len(col_headers), len(row))):
                    h = col_headers[ci]
                    idx = h2i.get(h)
                    if idx is None:
                        continue
                    out[idx] = row[ci]
                for col_name, val in (item.get("cell_values") or {}).items():
                    idx = h2i.get(col_name)
                    if idx is not None:
                        out[idx] = val
                _direct_xlsx_write_row_coerced(ws, out)
                item_rows += 1
                written[target_key] = written.get(target_key, 0) + 1
                if written[target_key] % 500 == 0:
                    _log("written %s: %d" % (target_key, written[target_key]))
                    _direct_ui_yield()
            _direct_journal(
                _direct_journal_source_label(item, path_mode),
                item.get("src_sheet_name") or u"",
                target_key,
                item_rows,
                u"скопировано",
                u"→ «%s»" % target_key,
            )
        if _direct_do_empty_sheets(settings):
            msg = _direct_empty_sheet_message()
            for target_key, bucket in agg.items():
                if int(written.get(target_key) or 0) > 0:
                    continue
                ws = ws_by_key.get(target_key)
                if ws is None:
                    continue
                headers_out = bucket.get("headers") or list(service_cols)
                row = [u""] * max(1, len(headers_out) if headers_out else 1)
                row[0] = msg
                _direct_xlsx_write_row_coerced(ws, row)
                written[target_key] = 1
                _direct_journal(u"", u"", target_key, 0, u"пусто", msg)
        else:
            for target_key in list(agg.keys()):
                if int(written.get(target_key) or 0) > 0:
                    continue
                ws = ws_by_key.get(target_key)
                if ws is None:
                    continue
                try:
                    # Удалить лист без данных (только заголовок).
                    if target_key in wb.sheetnames and len(wb.sheetnames) > 1:
                        del wb[target_key]
                except Exception:
                    pass
        wb.save(target_path)
        _log("saved(xlsx): %s sheets=%d" % (target_path, len(agg)))
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _xlsx_read_sheet_rows(path, sheet_index, hdr_row0, start_col0, row_count, col_count):
    """
    Прочитать строки с первого листа XLSX.

    Возвращает (headers, data_rows), где:
      headers: list[str]
      data_rows: list[list[Any]]
    """
    import openpyxl_bundled  # noqa: F401 — регистрирует import hook
    from openpyxl import load_workbook

    if _direct_t0 is None:
        _direct_since_start()
    _log("xlsx: open %s" % path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        try:
            _log('xlsx: sheet=%s idx=%d' % (unicode(getattr(ws, "title", "")), int(sheet_index)))
        except Exception:
            pass
        # --- Границы блока (Excel 1-based) ---
        header_row_1 = int(hdr_row0) + 1
        data_r0 = header_row_1 + 1
        c0 = int(start_col0) + 1
        c1_cap = c0 + (int(col_count) - 1) if col_count is not None else (c0 + 500 - 1)

        # Верхняя граница строк для чтения (если row_count задан — это число строк данных)
        if row_count is not None:
            data_r1 = data_r0 + int(row_count) - 1
        else:
            data_r1 = ws.max_row or data_r0

        # --- Заголовки (пакетно, без ws.cell по одной) ---
        headers = []
        try:
            header_vals = next(
                ws.iter_rows(
                    min_row=header_row_1,
                    max_row=header_row_1,
                    min_col=c0,
                    max_col=c1_cap,
                    values_only=True,
                ),
                (),
            )
        except Exception:
            header_vals = ()
        ci = 0
        for v in header_vals:
            t = _direct_clean_header_text(v) if v is not None else u""
            if col_count is None and t == u"":
                break
            headers.append(t if t != u"" else (u"COL_%d" % (ci + 1)))
            ci += 1

        # --- Данные (пакетно) ---
        data_rows = []
        if len(headers) == 0:
            return headers, data_rows
        max_col_used = c0 + len(headers) - 1

        rr = data_r0
        for row_vals in ws.iter_rows(
            min_row=data_r0,
            max_row=data_r1,
            min_col=c0,
            max_col=max_col_used,
            values_only=True,
        ):
            if rr % 200 == 0:
                _direct_ui_yield()
                try:
                    done = rr - data_r0 + 1
                    total = (data_r1 - data_r0 + 1) if data_r1 >= data_r0 else 0
                    _log(
                        "xlsx: read rows %d/%d (t=%.1fs)"
                        % (int(done), int(total), float(_direct_since_start()))
                    )
                except Exception:
                    pass

            # row_vals уже tuple значений; ограничиваем по headers
            row_vals = row_vals[: len(headers)]
            if col_count is None:
                empty = True
                for v in row_vals:
                    if v not in (None, ""):
                        empty = False
                        break
                if empty:
                    break
            data_rows.append(list(row_vals))
            rr += 1
        return headers, data_rows
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _xlsx_read_sheet_headers(path, sheet_index, hdr_row0, start_col0, col_count):
    """Быстро прочитать только заголовки (строка hdr_row0)."""
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        header_row_1 = int(hdr_row0) + 1
        c0 = int(start_col0) + 1
        c1_cap = c0 + (int(col_count) - 1) if col_count is not None else (c0 + 500 - 1)
        try:
            header_vals = next(
                ws.iter_rows(
                    min_row=header_row_1,
                    max_row=header_row_1,
                    min_col=c0,
                    max_col=c1_cap,
                    values_only=True,
                ),
                (),
            )
        except Exception:
            header_vals = ()
        headers = []
        ci = 0
        for v in header_vals:
            t = _direct_clean_header_text(v) if v is not None else u""
            if col_count is None and t == u"":
                break
            headers.append(t if t != u"" else (u"COL_%d" % (ci + 1)))
            ci += 1
        return headers
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _xlsx_iter_sheet_data_rows(path, sheet_index, hdr_row0, start_col0, row_count, headers, col_count):
    """
    Стриминг строк данных: yield list значений (только колонки headers).
    """
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    if _direct_t0 is None:
        _direct_since_start()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        header_row_1 = int(hdr_row0) + 1
        data_r0 = header_row_1 + 1
        c0 = int(start_col0) + 1
        if row_count is not None:
            data_r1 = data_r0 + int(row_count) - 1
        else:
            data_r1 = ws.max_row or data_r0
        if not headers:
            return
        max_col_used = c0 + len(headers) - 1
        rr = data_r0
        for row_vals in ws.iter_rows(
            min_row=data_r0,
            max_row=data_r1,
            min_col=c0,
            max_col=max_col_used,
            values_only=True,
        ):
            if rr % 200 == 0:
                _direct_ui_yield()
                try:
                    done = rr - data_r0 + 1
                    total = (data_r1 - data_r0 + 1) if data_r1 >= data_r0 else 0
                    _log(
                        "xlsx: read rows %d/%d (t=%.1fs)"
                        % (int(done), int(total), float(_direct_since_start()))
                    )
                except Exception:
                    pass
            row_vals = row_vals[: len(headers)]
            if col_count is None:
                empty = True
                for v in row_vals:
                    if v not in (None, ""):
                        empty = False
                        break
                if empty:
                    break
            yield list(row_vals)
            rr += 1
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _ods_cell_text(cell):
    try:
        from odf.text import P

        parts = cell.getElementsByType(P)
        if not parts:
            return u""
        node = parts[0].firstChild
        return u"" if node is None else unicode(node.data)
    except Exception:
        return u""


def _ods_cell_value(cell):
    try:
        vt = unicode(cell.getAttribute("valuetype") or u"").strip().lower()
    except Exception:
        vt = u""
    try:
        if vt in ("float", "currency", "percentage"):
            v = cell.getAttribute("value")
            try:
                return float(v)
            except Exception:
                return _ods_cell_text(cell)
        if vt == "boolean":
            b = unicode(cell.getAttribute("booleanvalue") or u"").strip().lower()
            if b in ("true", "1"):
                return True
            if b in ("false", "0"):
                return False
            return _ods_cell_text(cell)
        if vt == "date":
            dv = unicode(cell.getAttribute("datevalue") or u"").strip()
            if dv:
                try:
                    if "T" in dv:
                        return datetime.datetime.fromisoformat(dv)
                    return datetime.date.fromisoformat(dv)
                except Exception:
                    return dv
        if vt == "time":
            tv = unicode(cell.getAttribute("timevalue") or u"").strip()
            return tv if tv else _ods_cell_text(cell)
    except Exception:
        pass
    return _ods_cell_text(cell)


_RE_INT = re.compile(r"^[+-]?\d+$")
_RE_FLOAT = re.compile(r"^[+-]?\d+(?:[.,]\d+)?$")
_RE_DATE_DMY = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$")
_RE_DATE_DMY_WD = re.compile(
    r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})\s+(?:[А-Яа-яЁё]{2,12}|[A-Za-z]{2,9})\.?$",
    re.UNICODE,
)
_RE_DATE_YMD = re.compile(r"^(\d{2,4})-(\d{1,2})-(\d{1,2})$")
_RE_DATE_SLASH = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$")
_RE_DATE_DMY_MMM = re.compile(
    r"^(\d{1,2})\.([A-Za-zА-Яа-яЁё\.]+)\.(\d{2,4})$", re.UNICODE
)
_RE_DATE_DMY_TEXT = re.compile(
    r"^(\d{1,2})\s+([A-Za-zА-Яа-яЁё\.]+)\s+(\d{2,4})(?:\s*г\.?)?$",
    re.UNICODE | re.IGNORECASE,
)
_RE_DT_DMY = re.compile(
    r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?$"
)
_RE_DT_DMY_T = re.compile(
    r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})T(\d{1,2}):(\d{2})(?::(\d{2}))?$"
)
_RE_DT_YMD = re.compile(
    r"^(\d{2,4})-(\d{1,2})-(\d{1,2})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?$"
)
_RE_DT_YMD_T = re.compile(
    r"^(\d{2,4})-(\d{1,2})-(\d{1,2})T(\d{1,2}):(\d{2})(?::(\d{2}))?$"
)

_RU_MONTH_SHORT = {
    u"янв": 1,
    u"фев": 2,
    u"мар": 3,
    u"апр": 4,
    u"май": 5,
    u"мая": 5,
    u"июн": 6,
    u"июл": 7,
    u"авг": 8,
    u"сен": 9,
    u"окт": 10,
    u"ноя": 11,
    u"дек": 12,
}
_RU_MONTH_PREFIX = (
    (u"январ", 1),
    (u"феврал", 2),
    (u"март", 3),
    (u"апрел", 4),
    (u"мая", 5),
    (u"май", 5),
    (u"июн", 6),
    (u"июл", 7),
    (u"август", 8),
    (u"сентябр", 9),
    (u"октябр", 10),
    (u"ноябр", 11),
    (u"декабр", 12),
)
_EN_MONTH_SHORT = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}
_EN_MONTH_PREFIX = (
    ("january", 1),
    ("february", 2),
    ("march", 3),
    ("april", 4),
    ("may", 5),
    ("june", 6),
    ("july", 7),
    ("august", 8),
    ("september", 9),
    ("october", 10),
    ("november", 11),
    ("december", 12),
)


def _direct_parse_month_token(token):
    t = unicode(token or u"").strip().lower().rstrip(u".")
    if t == u"":
        return None
    if t.isdigit():
        mo = int(t)
        if 1 <= mo <= 12:
            return mo
        return None
    if t in _RU_MONTH_SHORT:
        return _RU_MONTH_SHORT[t]
    if t in _EN_MONTH_SHORT:
        return _EN_MONTH_SHORT[t]
    for prefix, num in _RU_MONTH_PREFIX:
        if t.startswith(prefix) or prefix.startswith(t):
            return num
    for prefix, num in _EN_MONTH_PREFIX:
        if t.startswith(prefix) or prefix.startswith(t):
            return num
    return None


def _direct_normalize_year(y):
    y = int(y)
    if y < 100:
        y += 2000
    return y


def _direct_is_month_name_only(text):
    """Целая ячейка — только имя/аббревиатура месяца (не цифра, не полная дата)."""
    s = unicode(text or u"").strip()
    if s == u"":
        return False
    # «5» / «12» не считаем month-only — иначе ломаем числа
    core = s.lower().rstrip(u".")
    if core.isdigit():
        return False
    if re.search(r"\d", s):
        return False
    return _direct_parse_month_token(s) is not None


def _direct_parse_month_construct_template(token):
    """
    Шаблон month-only → (canonical, day_spec, year_spec) или None.

    day_spec: 'first' | 'last' | int(1..31)
    year_spec: 'year' | 'next_year' | int(YYYY)
    canonical: нормализованная строка вида FIRST_DAY.MMMM.YEAR
    """
    raw = unicode(token or u"").strip()
    if raw == u"":
        return None
    parts = re.split(r"[./]", raw)
    if len(parts) != 3:
        return None
    day_p = unicode(parts[0] or u"").strip()
    mon_p = unicode(parts[1] or u"").strip()
    year_p = unicode(parts[2] or u"").strip()
    if day_p == u"" or mon_p == u"" or year_p == u"":
        return None
    mon_u = mon_p.upper()
    if mon_u not in _MONTH_NAME_FMT_TOKENS and mon_p not in _MONTH_NAME_FMT_TOKENS:
        return None
    # канон месяца: латиница MMMM/MMM, кириллица как в токене (верхний регистр)
    if mon_u in ("MMMM", "MMM"):
        mon_canon = mon_u
    else:
        mon_canon = mon_u if mon_u in _MONTH_NAME_FMT_TOKENS else mon_p

    day_u = day_p.upper()
    if day_u == "FIRST_DAY":
        day_spec = "first"
        day_canon = "FIRST_DAY"
    elif day_u == "LAST_DAY":
        day_spec = "last"
        day_canon = "LAST_DAY"
    elif day_p.isdigit():
        d = int(day_p)
        if d < 1 or d > 31:
            return None
        day_spec = d
        day_canon = "%02d" % d
    else:
        return None

    year_u = year_p.upper()
    if year_u == "YEAR":
        year_spec = "year"
        year_canon = "YEAR"
    elif year_u == "NEXT_YEAR":
        year_spec = "next_year"
        year_canon = "NEXT_YEAR"
    elif year_p.isdigit():
        year_spec = _direct_normalize_year(year_p)
        year_canon = str(year_spec)
    else:
        return None

    canonical = u"%s.%s.%s" % (day_canon, mon_canon, year_canon)
    return (canonical, day_spec, year_spec)


def _direct_build_date_from_month_construct(month, day_spec, year_spec):
    if year_spec == "year":
        y = datetime.date.today().year
    elif year_spec == "next_year":
        y = datetime.date.today().year + 1
    else:
        y = int(year_spec)
    if day_spec == "first":
        d = 1
    elif day_spec == "last":
        d = calendar.monthrange(y, month)[1]
    else:
        d = int(day_spec)
        last = calendar.monthrange(y, month)[1]
        if d > last:
            return None
    try:
        return datetime.date(y, month, d)
    except Exception:
        return None


def _direct_month_construct_specs(allowed_dt, formats_raw=None):
    """
    Список (canonical, day_spec, year_spec).

    Порядок — как в formats_raw (`,`/`;`); если raw пуст — стабильный sort по ключам allowed_dt.
    """
    out = []
    seen = set()
    raw = unicode(formats_raw if formats_raw is not None else u"").strip()
    if raw == u"" and formats_raw is None:
        raw = unicode(
            str((_direct_options or {}).get("datetime_formats", "") or "")
        ).strip()
    if raw != u"":
        for part in re.split(r"[,;]", raw):
            key = _direct_canonicalize_datetime_format_token(part)
            if key is None:
                continue
            if allowed_dt is not None and key not in allowed_dt:
                continue
            parsed = _direct_parse_month_construct_template(key)
            if parsed is None:
                continue
            canon = parsed[0]
            if canon in seen:
                continue
            seen.add(canon)
            out.append(parsed)
        if len(out) > 0:
            return out
    for key in sorted(unicode(k) for k in (allowed_dt or ())):
        parsed = _direct_parse_month_construct_template(key)
        if parsed is None:
            continue
        canon = parsed[0]
        if canon in seen:
            continue
        seen.add(canon)
        out.append(parsed)
    return out


def _direct_try_parse_month_only_date(s, allowed_dt, formats_raw=None):
    """
    Ячейка «Январь» / «янв.» / «January» + шаблон FIRST_DAY.MMMM.YEAR → date.
    """
    specs = _direct_month_construct_specs(allowed_dt, formats_raw=formats_raw)
    if len(specs) == 0:
        return None
    if not _direct_is_month_name_only(s):
        return None
    mo = _direct_parse_month_token(s)
    if mo is None:
        return None
    _canon, day_spec, year_spec = specs[0]
    return _direct_build_date_from_month_construct(mo, day_spec, year_spec)


def _direct_try_parse_slash_date(s, allowed_dt):
    if not (
        "DD/MM/YYYY" in allowed_dt
        or "DD/MM/YY" in allowed_dt
        or "MM/DD/YYYY" in allowed_dt
        or "MM/DD/YY" in allowed_dt
    ):
        return None
    m = _RE_DATE_SLASH.match(s)
    if not m:
        return None
    a, b, y = int(m.group(1)), int(m.group(2)), _direct_normalize_year(m.group(3))
    dmy_ok = "DD/MM/YYYY" in allowed_dt or "DD/MM/YY" in allowed_dt
    mdy_ok = "MM/DD/YYYY" in allowed_dt or "MM/DD/YY" in allowed_dt
    lang, country = _direct_options_locale()
    prefer_mdy = lang.startswith("en") and country in ("US", "USA")
    if dmy_ok and not mdy_ok:
        d, mo = a, b
    elif mdy_ok and not dmy_ok:
        mo, d = a, b
    elif prefer_mdy:
        mo, d = a, b
    else:
        d, mo = a, b
    try:
        return datetime.date(y, mo, d)
    except Exception:
        return None


def _direct_try_parse_mmm_date(s, allowed_dt):
    if not (
        "DD.MMM.YYYY" in allowed_dt
        or "DD.MMM.YY" in allowed_dt
        or "D MMMM YYYY" in allowed_dt
        or "DD MMM YYYY" in allowed_dt
    ):
        return None
    m = _RE_DATE_DMY_MMM.match(s)
    if m and ("DD.MMM.YYYY" in allowed_dt or "DD.MMM.YY" in allowed_dt):
        d = int(m.group(1))
        mo = _direct_parse_month_token(m.group(2))
        if mo is not None:
            y = _direct_normalize_year(m.group(3))
            try:
                return datetime.date(y, mo, d)
            except Exception:
                pass
    m = _RE_DATE_DMY_TEXT.match(s)
    if m and ("D MMMM YYYY" in allowed_dt or "DD MMM YYYY" in allowed_dt):
        d = int(m.group(1))
        mo = _direct_parse_month_token(m.group(2))
        if mo is not None:
            y = _direct_normalize_year(m.group(3))
            try:
                return datetime.date(y, mo, d)
            except Exception:
                pass
    return None


def _direct_try_parse_datetime_text(text, allowed_dt=None, formats_raw=None):
    """
    Разбор текста в date/datetime по MERGE_XML_DATETIME_FORMATS / базовому набору.

    Ведущий апостроф снимается. Не число — только дата/время; иначе None.
    formats_raw — исходная строка форматов (порядок month-only шаблонов).
    """
    s = _direct_strip_leading_apostrophe(text)
    if s == u"":
        return None
    if allowed_dt is None:
        allowed_dt = _direct_allowed_datetime_formats()
        if formats_raw is None:
            formats_raw = str(
                (_direct_options or {}).get("datetime_formats", "") or ""
            )
    if (
        "DD.MM.YYYY HH:MM" in allowed_dt
        or "DD.MM.YYYY HH:MM:SS" in allowed_dt
        or "DD.MM.YYYYTHH:MM" in allowed_dt
        or "DD.MM.YYYYTHH:MM:SS" in allowed_dt
        or "DD.MM.YY HH:MM" in allowed_dt
        or "DD.MM.YY HH:MM:SS" in allowed_dt
        or "DD.MM.YYTHH:MM" in allowed_dt
        or "DD.MM.YYTHH:MM:SS" in allowed_dt
    ):
        m = _RE_DT_DMY.match(s)
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            hh, mm = int(m.group(4)), int(m.group(5))
            ss = int(m.group(6)) if m.group(6) is not None else 0
            if y < 100:
                y += 2000
            try:
                if m.group(6) is None and (
                    "DD.MM.YYYY HH:MM" not in allowed_dt
                    and "DD.MM.YY HH:MM" not in allowed_dt
                ):
                    return None
                if m.group(6) is not None and (
                    "DD.MM.YYYY HH:MM:SS" not in allowed_dt
                    and "DD.MM.YY HH:MM:SS" not in allowed_dt
                ):
                    return None
                return datetime.datetime(y, mo, d, hh, mm, ss)
            except Exception:
                return None
        m = _RE_DT_DMY_T.match(s)
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            hh, mm = int(m.group(4)), int(m.group(5))
            ss = int(m.group(6)) if m.group(6) is not None else 0
            if y < 100:
                y += 2000
            try:
                if m.group(6) is None and (
                    "DD.MM.YYYYTHH:MM" not in allowed_dt
                    and "DD.MM.YYTHH:MM" not in allowed_dt
                ):
                    return None
                if m.group(6) is not None and (
                    "DD.MM.YYYYTHH:MM:SS" not in allowed_dt
                    and "DD.MM.YYTHH:MM:SS" not in allowed_dt
                ):
                    return None
                return datetime.datetime(y, mo, d, hh, mm, ss)
            except Exception:
                return None
    m = _RE_DATE_DMY.match(s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), _direct_normalize_year(m.group(3))
        try:
            return datetime.date(y, mo, d)
        except Exception:
            return None
    m = _RE_DATE_DMY_WD.match(s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), _direct_normalize_year(m.group(3))
        try:
            return datetime.date(y, mo, d)
        except Exception:
            return None
    parsed = _direct_try_parse_slash_date(s, allowed_dt)
    if parsed is not None:
        return parsed
    parsed = _direct_try_parse_mmm_date(s, allowed_dt)
    if parsed is not None:
        return parsed
    parsed = _direct_try_parse_month_only_date(s, allowed_dt, formats_raw=formats_raw)
    if parsed is not None:
        return parsed
    m = _RE_DATE_YMD.match(s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mo, d)
        except Exception:
            return None
    if (
        "YYYY-MM-DD HH:MM" in allowed_dt
        or "YYYY-MM-DD HH:MM:SS" in allowed_dt
        or "YYYY-MM-DDTHH:MM" in allowed_dt
        or "YYYY-MM-DDTHH:MM:SS" in allowed_dt
        or "YY-MM-DD HH:MM" in allowed_dt
        or "YY-MM-DD HH:MM:SS" in allowed_dt
        or "YY-MM-DDTHH:MM" in allowed_dt
        or "YY-MM-DDTHH:MM:SS" in allowed_dt
    ):
        m = _RE_DT_YMD.match(s)
        if m:
            y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            hh, mm = int(m.group(4)), int(m.group(5))
            ss = int(m.group(6)) if m.group(6) is not None else 0
            if y < 100:
                y += 2000
            try:
                if m.group(6) is None and (
                    "YYYY-MM-DD HH:MM" not in allowed_dt
                    and "YY-MM-DD HH:MM" not in allowed_dt
                ):
                    return None
                if m.group(6) is not None and (
                    "YYYY-MM-DD HH:MM:SS" not in allowed_dt
                    and "YY-MM-DD HH:MM:SS" not in allowed_dt
                ):
                    return None
                return datetime.datetime(y, mo, d, hh, mm, ss)
            except Exception:
                return None
        m = _RE_DT_YMD_T.match(s)
        if m:
            y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
            hh, mm = int(m.group(4)), int(m.group(5))
            ss = int(m.group(6)) if m.group(6) is not None else 0
            if y < 100:
                y += 2000
            try:
                if m.group(6) is None and (
                    "YYYY-MM-DDTHH:MM" not in allowed_dt
                    and "YY-MM-DDTHH:MM" not in allowed_dt
                ):
                    return None
                if m.group(6) is not None and (
                    "YYYY-MM-DDTHH:MM:SS" not in allowed_dt
                    and "YY-MM-DDTHH:MM:SS" not in allowed_dt
                ):
                    return None
                return datetime.datetime(y, mo, d, hh, mm, ss)
            except Exception:
                return None
    return None


def _direct_has_grouped_number_separators(s):
    """Пробел/апостроф между цифрами или несколько ,/. — признак группировки разрядов."""
    if re.search(r"[\s\u00a0\u202f\u2009]", s):
        return True
    if re.search(r"(?<=\d)['\u2019](?=\d)", s):
        return True
    if s.count(u",") > 1:
        return True
    if s.count(u".") > 1:
        return True
    return False


def _direct_try_parse_number_text(text):
    """
    Текст → int/float или None.

    Поддерживает: 700; 1,5; 1 000 000,00; 1 000 000.00; 1'000'000,00;
    ведущий апостроф (текстовый префикс Calc). Строки-даты не перехватывает.
    """
    s = _direct_strip_leading_apostrophe(unicode(text or u"")).strip()
    if s == u"":
        return None
    sign = 1
    if s[0] in (u"+", u"-"):
        sign = -1 if s[0] == u"-" else 1
        s = s[1:].strip()
    if s == u"":
        return None
    if _RE_INT.match(s):
        try:
            return sign * int(s)
        except Exception:
            return None
    if _RE_FLOAT.match(s):
        try:
            return sign * float(s.replace(u",", u"."))
        except Exception:
            return None
    if not _direct_has_grouped_number_separators(s):
        return None
    if _direct_try_parse_datetime_text(s) is not None:
        return None
    s2 = re.sub(r"[\s\u00a0\u202f\u2009]+", u"", s)
    s2 = re.sub(r"(?<=\d)['\u2019](?=\d)", u"", s2)
    if u"," in s2 and u"." in s2:
        if s2.rfind(u",") > s2.rfind(u"."):
            s2 = s2.replace(u".", u"")
            s2 = s2.replace(u",", u".")
        else:
            s2 = s2.replace(u",", u"")
    elif u"," in s2:
        parts = s2.split(u",")
        if len(parts) == 2 and 1 <= len(parts[1]) <= 3:
            s2 = parts[0] + u"." + parts[1]
        else:
            s2 = s2.replace(u",", u"")
    elif s2.count(u".") > 1:
        s2 = s2.replace(u".", u"")
    if not _RE_GROUPED_NUM_BODY.match(s2):
        return None
    try:
        if u"." in s2:
            return sign * float(s2)
        return sign * int(s2)
    except Exception:
        return None


def _coerce_text_value(text):
    s = unicode(text or u"").strip()
    if s == u"":
        return u""
    if _RE_INT.match(s):
        try:
            return int(s)
        except Exception:
            pass
    if _RE_FLOAT.match(s):
        try:
            return float(s.replace(",", "."))
        except Exception:
            pass
    parsed = _direct_try_parse_datetime_text(s)
    if parsed is not None:
        return parsed
    num = _direct_try_parse_number_text(s)
    if num is not None:
        return num
    return s


def _ods_read_sheet_rows(path, sheet_index, hdr_row0, start_col0, row_count, col_count):
    """
    Прочитать строки с первого листа ODS.
    Возвращает (headers, data_rows) как в _xlsx_read_first_sheet_rows.
    """
    import odf_bundled  # noqa: F401 — регистрирует import hook
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell

    if _direct_t0 is None:
        _direct_since_start()
    _log("ods: open %s" % path)
    doc = load(path)
    sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
    rows = sheet.getElementsByType(TableRow)
    hdr_r = int(hdr_row0)
    if hdr_r < 0 or hdr_r >= len(rows):
        return [], []
    cells = rows[hdr_r].getElementsByType(TableCell)
    c0 = int(start_col0)
    headers = []
    max_cols = int(col_count) if col_count is not None else 500
    for i in range(max_cols):
        idx = c0 + i
        t = u""
        if idx < len(cells):
            t = unicode(_ods_cell_text(cells[idx])).strip()
        if col_count is None and t == u"":
            break
        headers.append(t if t != u"" else (u"COL_%d" % (i + 1)))
    data_rows = []
    data_r0 = hdr_r + 1
    data_r1 = data_r0 + int(row_count) - 1 if row_count is not None else (len(rows) - 1)
    if data_r1 >= len(rows):
        data_r1 = len(rows) - 1
    for rr in range(data_r0, data_r1 + 1):
        if rr % 200 == 0:
            _direct_ui_yield()
        if rr % 200 == 0:
            try:
                done = rr - data_r0 + 1
                total = (data_r1 - data_r0 + 1) if data_r1 >= data_r0 else 0
                _log(
                    "ods: read rows %d/%d (t=%.1fs)"
                    % (int(done), int(total), float(_direct_since_start()))
                )
            except Exception:
                pass
        rcells = rows[rr].getElementsByType(TableCell)
        row_vals = []
        empty = True
        for ci in range(len(headers)):
            idx = c0 + ci
            v = u""
            if idx < len(rcells):
                v = _ods_cell_value(rcells[idx])
            if unicode(v).strip() != u"":
                empty = False
            row_vals.append(_direct_coerce_cell_value(v))
        if col_count is None and empty:
            break
        data_rows.append(row_vals)
    return headers, data_rows


def _ods_read_sheet_headers(path, sheet_index, hdr_row0, start_col0, col_count):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell

    doc = load(path)
    sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
    rows = sheet.getElementsByType(TableRow)
    hdr_r = int(hdr_row0)
    if hdr_r < 0 or hdr_r >= len(rows):
        return []
    cells = rows[hdr_r].getElementsByType(TableCell)
    c0 = int(start_col0)
    max_cols = int(col_count) if col_count is not None else 500
    headers = []
    for i in range(max_cols):
        idx = c0 + i
        t = u""
        if idx < len(cells):
            t = unicode(_ods_cell_text(cells[idx])).strip()
        if col_count is None and t == u"":
            break
        headers.append(t if t != u"" else (u"COL_%d" % (i + 1)))
    return headers


def _ods_iter_sheet_data_rows(path, sheet_index, hdr_row0, start_col0, row_count, headers, col_count):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell

    if _direct_t0 is None:
        _direct_since_start()
    doc = load(path)
    sheet = doc.spreadsheet.getElementsByType(Table)[int(sheet_index)]
    rows = sheet.getElementsByType(TableRow)
    hdr_r = int(hdr_row0)
    data_r0 = hdr_r + 1
    if row_count is not None:
        data_r1 = data_r0 + int(row_count) - 1
    else:
        data_r1 = len(rows) - 1
    if data_r1 >= len(rows):
        data_r1 = len(rows) - 1
    c0 = int(start_col0)
    w = len(headers or [])
    rr = data_r0
    while rr <= data_r1:
        if rr % 200 == 0:
            _direct_ui_yield()
            try:
                done = rr - data_r0 + 1
                total = (data_r1 - data_r0 + 1) if data_r1 >= data_r0 else 0
                _log(
                    "ods: read rows %d/%d (t=%.1fs)"
                    % (int(done), int(total), float(_direct_since_start()))
                )
            except Exception:
                pass
        rcells = rows[rr].getElementsByType(TableCell)
        row_vals = []
        empty = True
        ci = 0
        while ci < w:
            idx = c0 + ci
            v = u""
            if idx < len(rcells):
                v = _ods_cell_value(rcells[idx])
            if unicode(v).strip() != u"":
                empty = False
            row_vals.append(_direct_coerce_cell_value(v))
            ci += 1
        if col_count is None and empty:
            break
        yield row_vals
        rr += 1


def _ods_get_or_create_table(doc, name):
    from odf.table import Table

    for t in doc.spreadsheet.getElementsByType(Table):
        if unicode(t.getAttribute("name") or u"") == unicode(name):
            return t
    t = Table(name=unicode(name))
    doc.spreadsheet.addElement(t)
    return t


def _ods_clear_table_rows(table):
    from odf.table import TableRow

    # Удаляем только TableRow, оставляя свойства таблицы.
    to_remove = []
    for ch in list(getattr(table, "childNodes", []) or []):
        try:
            if isinstance(ch, TableRow):
                to_remove.append(ch)
        except Exception:
            pass
    for ch in to_remove:
        try:
            table.removeChild(ch)
        except Exception:
            try:
                table.childNodes.remove(ch)
            except Exception:
                pass


def _ods_row_cell_by_col(row_elem):
    """col_index → TableCell (первая ячейка группы repeat)."""
    from odf.table import TableCell

    out = {}
    col = 0
    for cell in row_elem.getElementsByType(TableCell):
        try:
            repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
        except Exception:
            repeat = 1
        if repeat < 1:
            repeat = 1
        out[col] = cell
        col += repeat
    return out


def _direct_pp_ods_clone_cell_light(src_cell):
    """
    Быстрое копирование ячейки ODS без deepcopy.

    style-name / valuetype / value / date-value / formula + отображаемый текст.
    (attrs.items() с tuple-ключами через setAttribute не работает в odfpy.)
    """
    from odf.table import TableCell
    from odf.text import P

    dst = TableCell()
    for name in (
        "stylename",
        "valuetype",
        "value",
        "datevalue",
        "timevalue",
        "booleanvalue",
        "formula",
        "stringvalue",
        "currencysymbol",
    ):
        try:
            val = src_cell.getAttribute(name)
            if val is not None:
                dst.setAttribute(name, val)
        except Exception:
            pass
    try:
        dst.addElement(P(text=unicode(_ods_cell_text(src_cell))))
    except Exception:
        pass
    return dst


def _ods_make_cell(value):
    from odf.table import TableCell
    from odf.text import P

    if value is None:
        cell = TableCell()
        cell.addElement(P(text=u""))
        return cell
    if isinstance(value, bool):
        cell = TableCell(
            valuetype="boolean",
            booleanvalue="true" if value else "false",
        )
        cell.addElement(P(text=u"TRUE" if value else u"FALSE"))
        return cell
    if isinstance(value, int):
        cell = TableCell(valuetype="float", value=str(float(value)))
        cell.addElement(P(text=unicode(value)))
        return cell
    if isinstance(value, float):
        cell = TableCell(valuetype="float", value=str(value))
        cell.addElement(P(text=unicode(value)))
        return cell
    if isinstance(value, datetime.datetime):
        iso = value.replace(microsecond=0).isoformat()
        cell = TableCell(valuetype="date", datevalue=iso)
        cell.addElement(P(text=iso.replace("T", " ")))
        return cell
    if isinstance(value, datetime.date):
        iso = value.isoformat()
        cell = TableCell(valuetype="date", datevalue=iso)
        cell.addElement(P(text=iso))
        return cell
    if isinstance(value, datetime.time):
        sec = value.hour * 3600 + value.minute * 60 + value.second
        tv = "PT%dS" % int(sec)
        cell = TableCell(valuetype="time", timevalue=tv)
        cell.addElement(P(text=value.strftime("%H:%M:%S")))
        return cell
    cell = TableCell()
    cell.addElement(P(text=unicode(value)))
    return cell


def _ods_add_row(table, values):
    from odf.table import TableRow

    row = TableRow()
    for v in values:
        row.addElement(_ods_make_cell(v))
    table.addElement(row)


def _format_source_path(path, path_display):
    if path_display == "none":
        path_display = "short"
    abspath = os.path.abspath(path)
    if path_display == "short":
        base = os.path.basename(abspath)
        stem, _ext = os.path.splitext(base)
        if stem != "":
            return stem
        return base
    return abspath


def _format_source_label(path, path_display, sheet_name=None):
    label = _format_source_path(path, path_display)
    if sheet_name is not None and unicode(sheet_name).strip() != u"":
        return u"%s#%s" % (label, unicode(sheet_name).strip())
    return label


def _format_file_mtime(path, datetime_mode="long"):
    if datetime_mode == "none":
        return u""
    try:
        mtime = os.path.getmtime(os.path.abspath(path))
        if datetime_mode == "short":
            return datetime.datetime.fromtimestamp(mtime).strftime("%Y-%m-%d")
        return datetime.datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return u""


def _service_path_value(src_path, path_mode, sheet_name=None):
    if path_mode == "none":
        return u""
    return _format_source_label(src_path, path_mode, sheet_name)


def _service_datetime_value(src_path, datetime_mode):
    if datetime_mode == "none":
        return u""
    dt = _file_mtime_datetime(src_path)
    if dt is None:
        return u""
    if datetime_mode == "short":
        return DirectServiceDatetime(dt.year, dt.month, dt.day)
    return DirectServiceDatetime(
        dt.year, dt.month, dt.day, dt.hour, dt.minute, dt.second
    )


def _source_sheet_name(read_path, sheet_index, src_kind):
    try:
        if src_kind == "xlsx":
            names = _xlsx_sheet_names(read_path)
        elif src_kind == "ods":
            names = _ods_sheet_names(read_path)
        else:
            names = _xlsx_sheet_names(read_path) if _is_xlsx(read_path) or _is_xlsm(read_path) else _ods_sheet_names(read_path)
        if sheet_index >= 0 and sheet_index < len(names):
            return unicode(names[sheet_index] or u"")
    except Exception:
        pass
    return u""


def _file_mtime_datetime(path):
    try:
        ts = os.path.getmtime(path)
        return datetime.datetime.fromtimestamp(ts).replace(microsecond=0)
    except Exception:
        return None


def _direct_copy_row_trim_parts(header_row, start_row, header_last=None):
    """1-based: top до первой строки шапки; after_header после последней строки шапки."""
    try:
        hr = max(1, int(header_row))
    except (TypeError, ValueError):
        hr = 1
    try:
        sr = max(1, int(start_row))
    except (TypeError, ValueError):
        sr = hr
    last = header_last if header_last not in (None, u"") else hr
    try:
        last = int(last)
    except (TypeError, ValueError):
        last = hr
    last = max(hr, last)
    top = max(0, hr - 1)
    after = max(0, sr - last - 1) if sr > last + 1 else 0
    return top, after


def _direct_block_header_first_last_1(blk):
    if not blk:
        return (1, 1)
    if blk.get("no_header"):
        return (0, 0)
    try:
        first_1 = int(blk.get("header_row_first_1") or 0)
        last_1 = int(blk.get("header_row_last_1") or 0)
    except (TypeError, ValueError):
        first_1 = 0
        last_1 = 0
    if first_1 > 0 and last_1 > 0:
        if last_1 < first_1:
            last_1 = first_1
        return (first_1, last_1)
    try:
        hr = int(str(blk.get("header_row_disp") or blk.get("start_row_disp") or "1"))
    except (TypeError, ValueError):
        hr = 1
    if hr <= 0:
        return (0, 0)
    return (hr, hr)


def _direct_copy_sheets_name_taken(candidate, used_names):
    if candidate in used_names:
        return True
    return False


def _direct_copy_sheets_resolve_name(source_sheet_name, source_index, used_names):
    """Имя вкладки N_имя (как merge_on_copy_sheets)."""
    base = unicode(source_sheet_name or u"").strip()
    if base == u"":
        base = u"Лист"
    prefix = u"%d_" % int(source_index)
    max_base = 31 - len(prefix)
    if max_base < 1:
        max_base = 1
    candidate = prefix + base[:max_base]
    suffix = 2
    while _direct_copy_sheets_name_taken(candidate, used_names):
        tail = u"_%d" % suffix
        head_len = 31 - len(prefix) - len(tail)
        if head_len < 1:
            head_len = 1
        candidate = prefix + base[:head_len] + tail
        suffix += 1
        if suffix > 999:
            break
    used_names.add(candidate)
    return candidate


def _direct_copy_sheets_seed_names(target_doc):
    """used_names и next_prefix_by_sheet_key из существующих листов ODS."""
    from odf.table import Table

    used_names = set()
    for svc in _DIRECT_SERVICE_SHEET_NAMES:
        used_names.add(unicode(svc))
    max_prefix_by_sheet_key = {}
    for table in target_doc.spreadsheet.getElementsByType(Table):
        existing_name = unicode(table.getAttribute("name") or u"").strip()
        if existing_name != u"":
            used_names.add(existing_name)
        m = re.match(r"^(\d+)_(.+)$", existing_name)
        if m is not None:
            try:
                n = int(m.group(1))
                key = _merge_identity_key(m.group(2))
                prev = max_prefix_by_sheet_key.get(key, 0)
                if n > prev:
                    max_prefix_by_sheet_key[key] = n
            except Exception:
                pass
    next_prefix_by_sheet_key = {
        k: int(v) + 1 for k, v in max_prefix_by_sheet_key.items()
    }
    return used_names, next_prefix_by_sheet_key


def _direct_copy_has_explicit_delete_top(sheet_name, file_index, settings):
    mapping = (settings or {}).get("delete_top_rows_map") or {}
    per_file = (settings or {}).get("delete_top_rows_per_file") or []
    n = _direct_resolve_delete_top_count(sheet_name, mapping)
    if n is not None:
        try:
            if int(n) > 0:
                return True
        except (TypeError, ValueError):
            pass
    if file_index is not None and file_index < len(per_file):
        try:
            if int(per_file[file_index] or 0) > 0:
                return True
        except (TypeError, ValueError):
            pass
    return False


def _direct_copy_resolve_delete_top_rows(sheet_name, file_index, settings):
    mapping = (settings or {}).get("delete_top_rows_map") or {}
    n = _direct_resolve_delete_top_count(sheet_name, mapping)
    if n is not None:
        try:
            if int(n) > 0:
                return int(n)
        except (TypeError, ValueError):
            pass
    per_file = (settings or {}).get("delete_top_rows_per_file") or []
    if file_index is not None and file_index < len(per_file):
        try:
            pf_n = int(per_file[file_index] or 0)
        except (TypeError, ValueError):
            pf_n = 0
        if pf_n > 0 and _direct_resolve_delete_top_count(sheet_name, mapping) is None:
            return pf_n
    return 0


def _direct_copy_filter_rows(rows, blk, target_name, file_index, settings):
    """Автообрезка и удаление верхних строк (как UNO + trim + delete_top)."""
    rows = list(rows or [])
    if not rows:
        return rows
    flags = (settings or {}).get("merge_flags") or {}
    auto_trim = True
    if "MERGE_COPY_AUTO_TRIM" in flags:
        auto_trim = bool(flags.get("MERGE_COPY_AUTO_TRIM"))
    if auto_trim and not _direct_copy_has_explicit_delete_top(
        target_name, file_index, settings
    ):
        hr = int(
            str((blk or {}).get("header_row_disp") or (blk or {}).get("start_row_disp") or "1")
        )
        sr = int(str((blk or {}).get("start_row_disp") or "1"))
        _hr, last = _direct_block_header_first_last_1(blk)
        if _hr > 0:
            hr = _hr
        top, after = _direct_copy_row_trim_parts(hr, sr, last)
        if top > 0:
            rows = rows[top:]
        if after > 0 and len(rows) > 1:
            rows = [rows[0]] + rows[1 + after :]
    delete_n = _direct_copy_resolve_delete_top_rows(target_name, file_index, settings)
    if delete_n > 0:
        rows = rows[delete_n:]
    return rows


def _direct_ods_find_table(doc, sheet_name, sheet_index=-1):
    from odf.table import Table

    want = unicode(sheet_name or u"").strip()
    tables = doc.spreadsheet.getElementsByType(Table)
    if want != u"":
        for table in tables:
            if unicode(table.getAttribute("name") or u"") == want:
                return table
    if sheet_index >= 0 and sheet_index < len(tables):
        return tables[int(sheet_index)]
    return None


def _direct_ods_read_table_rows(table):
    from odf.table import TableRow

    if table is None:
        return []
    out = []
    for row in table.getElementsByType(TableRow):
        row_map = dict(_ods_row_col_values(row))
        if not row_map:
            out.append([])
            continue
        w = max(row_map.keys()) + 1
        out.append([row_map.get(c, u"") for c in range(w)])
    return out


def _direct_xlsx_read_all_rows(path, sheet_index):
    import openpyxl_bundled  # noqa: F401
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[int(sheet_index)]
        dim_bounds = _direct_xlsx_bounds_from_dimension(ws)
        if dim_bounds is not None:
            end_col, end_row = dim_bounds
        else:
            end_col, end_row = _direct_scan_bounds_xlsx(ws, 0, 0)
        if end_row < 0 or end_col < 0:
            return []
        rows = []
        rr = 0
        for row_vals in ws.iter_rows(
            min_row=1,
            max_row=int(end_row) + 1,
            min_col=1,
            max_col=int(end_col) + 1,
            values_only=True,
        ):
            if rr % 200 == 0:
                _direct_ui_yield()
            rows.append(list(row_vals))
            rr += 1
        return rows
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _direct_ods_read_all_rows(path, sheet_index):
    import odf_bundled  # noqa: F401
    from odf.opendocument import load
    from odf.table import Table

    doc = load(path)
    tables = doc.spreadsheet.getElementsByType(Table)
    if int(sheet_index) < 0 or int(sheet_index) >= len(tables):
        return []
    return _direct_ods_read_table_rows(tables[int(sheet_index)])


def _direct_copy_read_source_rows( read_path, src_kind, sheet_index, sheet_name, source_doc=None ):
    if source_doc is not None:
        table = _direct_ods_find_table(source_doc, sheet_name, sheet_index)
        return _direct_ods_read_table_rows(table)
    if src_kind == "xlsx":
        return _direct_xlsx_read_all_rows(read_path, sheet_index)
    return _direct_ods_read_all_rows(read_path, sheet_index)


def _direct_ods_add_row_coerced(table, values, style_cache):
    from odf.table import TableRow

    row = TableRow()
    for v in values or []:
        coerced = _direct_coerce_cell_value(v)
        cell = _ods_make_cell(coerced)
        if style_cache is not None and _direct_wants_convert_to_numbers():
            sn = _direct_ods_prebuilt_cell_style(style_cache, coerced)
            if sn:
                try:
                    cell.setAttribute(u"stylename", sn)
                except Exception:
                    pass
            else:
                _direct_ods_apply_converted_cell_style(style_cache, cell, coerced)
        row.addElement(cell)
    table.addElement(row)


def _direct_run_copy_sheets_to_ods(target_path, settings):
    """
    «Копирование листов» в ODS без UNO: чтение openpyxl/odfpy, запись с coerce.
    """
    import odf_bundled  # noqa: F401
    from odf.opendocument import load

    if not _is_ods(target_path):
        raise ValueError("copy_sheets(ods): цель должна быть .ods")
    files = (settings or {}).get("files") or []
    if not files:
        raise ValueError("copy_sheets(ods): нет файлов-источников")
    is_current_book_flags = (settings or {}).get("is_current_book_flags") or []
    per_file_sheet_specs = (settings or {}).get("per_file_sheet_specs") or []
    per_file_blocks = (settings or {}).get("per_file_blocks") or []
    overrides = (settings or {}).get("direct_source_overrides") or {}
    target_book = unicode((settings or {}).get("direct_target_path") or target_path).strip()
    path_display = ((settings or {}).get("path_display") or "none").strip().lower()
    ctc_per_file = _direct_cell_specs_per_file(
        (settings or {}).get("cell_to_column_plan") or [], len(files)
    )

    doc = load(target_path)
    want_convert = _direct_wants_convert_to_numbers()
    style_cache = _direct_ods_style_cache_new(doc) if want_convert else None
    used_names, next_prefix_by_sheet_key = _direct_copy_sheets_seed_names(doc)

    copied_names = []
    block_by_name = {}
    file_index_by_name = {}
    any_sheet = False
    fi = 0
    while fi < len(files):
        is_curr = is_current_book_flags[fi] if fi < len(is_current_book_flags) else False
        blk = (
            per_file_blocks[fi]
            if fi < len(per_file_blocks)
            else (per_file_blocks[0] if per_file_blocks else {})
        )
        sheet_specs = (
            per_file_sheet_specs[fi]
            if fi < len(per_file_sheet_specs)
            else (per_file_sheet_specs[0] if per_file_sheet_specs else [])
        )
        sheet_kind = _classify_sheet_specs(sheet_specs)
        if sheet_kind == "MIXED":
            raise ValueError("copy_sheets(ods): MIXED в фильтре «Листы» не поддержан")

        if is_curr:
            if target_book == u"" or not os.path.isfile(target_book):
                _direct_journal(
                    u"(текущая книга)",
                    u"",
                    u"",
                    u"",
                    u"проблема",
                    u"«Эта_книга»: книга результата не сохранена на диск",
                )
                fi += 1
                continue
            read_path = os.path.abspath(target_book)
            src_label = u"(текущая книга)"
            source_doc = doc
            sheet_names = _ods_sheet_names(read_path)
            src_kind = "ods"
        else:
            src_path = os.path.abspath(files[fi])
            read_path = os.path.abspath(overrides.get(src_path, src_path))
            source_doc = None
            src_label = _format_source_path(
                read_path, path_display if path_display != "none" else "short"
            )
            if _is_xlsx(read_path) or _is_xlsm(read_path):
                sheet_names = _xlsx_sheet_names(read_path)
                src_kind = "xlsx"
            elif _is_ods(read_path):
                sheet_names = _ods_sheet_names(read_path)
                src_kind = "ods"
            else:
                _direct_journal(
                    src_label, u"", u"", u"", u"проблема", u"неподдерживаемый источник"
                )
                fi += 1
                continue

        file_note = _direct_collect_file_note(blk, sheet_specs, [])
        _direct_journal(src_label, u"", u"", u"", u"файл", file_note)
        _direct_ui_notify_source(read_path)
        _direct_ui_yield(force=True)

        for slot in _iter_sheet_slots_file(sheet_names, sheet_specs, sheet_kind):
            if is_curr and _direct_is_service_sheet_name(slot.get("source_name")):
                continue
            tname = unicode(slot.get("source_name") or u"").strip()
            if tname == u"":
                continue
            src_index = int(slot.get("source_index", -1))
            if sheet_kind == "NAME" or src_index < 0:
                try:
                    src_index = sheet_names.index(tname)
                except Exception:
                    continue
            if src_index < 0 or src_index >= len(sheet_names):
                continue

            sheet_key = _merge_identity_key(tname)
            source_index = int(next_prefix_by_sheet_key.get(sheet_key, 1))
            target_name = _direct_copy_sheets_resolve_name(
                tname, source_index, used_names
            )
            next_prefix_by_sheet_key[sheet_key] = source_index + 1

            _direct_ui_notify_source(read_path, tname)
            rows = _direct_copy_read_source_rows(
                read_path, src_kind, src_index, tname, source_doc=source_doc
            )
            cell_specs = ctc_per_file[fi] if fi < len(ctc_per_file) else []
            if cell_specs:
                cell_specs = [
                    s
                    for s in cell_specs
                    if s is not None
                    and _direct_cell_spec_applies_to_source_sheet(s, tname)
                    and _direct_cell_spec_applies_to_target_sheet(s, target_name)
                ]
            cell_values = _direct_copy_prefetch_cell_values_from_rows(rows, cell_specs)
            rows = _direct_copy_filter_rows(rows, blk, target_name, fi, settings)
            rename_items = []
            per_file_rename = (settings or {}).get("per_file_col_rename_items") or []
            if fi < len(per_file_rename):
                rename_items = per_file_rename[fi] or []
            if rename_items:
                rows = _direct_copy_apply_col_rename_to_rows(rows, blk, rename_items)
            if cell_values:
                rows = _direct_copy_inject_cell_to_column(rows, cell_values)
            if not rows:
                _direct_journal(
                    src_label,
                    tname,
                    u"",
                    0,
                    u"пропуск",
                    u"нет данных после обрезки",
                )
                continue

            table = _ods_get_or_create_table(doc, target_name)
            _ods_clear_table_rows(table)
            ri = 0
            while ri < len(rows):
                if ri % 200 == 0:
                    _direct_ui_yield()
                _direct_ods_add_row_coerced(table, rows[ri], style_cache)
                ri += 1

            any_sheet = True
            copied_names.append(target_name)
            block_by_name[target_name] = blk
            file_index_by_name[target_name] = fi
            _direct_journal(
                src_label,
                target_name,
                tname,
                len(rows),
                u"скопирован лист",
                file_note,
            )
            _log(
                u"copy_sheets(ods): «%s» ← «%s» (%d строк)"
                % (target_name, tname, len(rows))
            )

        fi += 1

    if not any_sheet:
        raise ValueError("copy_sheets(ods): нет скопированных листов по заданным фильтрам")

    doc.save(target_path)
    if isinstance(settings, dict):
        settings["copy_sheets_block_by_name"] = block_by_name
        settings["copy_sheets_file_index_by_name"] = file_index_by_name
    return copied_names


def direct_copy_sheets_to_ods(target_ods_path, settings):
    """
    Прямое «Копирование листов» в ODS: openpyxl/odfpy → ODS с coerce при записи.

    settings — dict из collect_workbooks.parse_collect_settings (+ direct_target_path).
    Возвращает список имён добавленных листов результата.
    """
    if not _is_ods(target_ods_path):
        raise ValueError("xml_excel_ods: цель должна быть .ods (сейчас %s)" % target_ods_path)
    if not os.path.isfile(target_ods_path):
        raise ValueError("xml_excel_ods: не найден файл результата: %s" % target_ods_path)
    _log("target(ods,copy): %s" % target_ods_path)
    return _direct_run_copy_sheets_to_ods(target_ods_path, settings)


def direct_merge_on_one_sheet_to_ods(target_ods_path, settings):
    """
    Прямой перенос в лист «Merge» внутри ODS-файла.

    settings — dict из collect_workbooks.parse_collect_settings.
    """
    if not _is_ods(target_ods_path):
        raise ValueError("xml_excel_ods: цель должна быть .ods (сейчас %s)" % target_ods_path)
    if not os.path.isfile(target_ods_path):
        raise ValueError("xml_excel_ods: не найден файл результата: %s" % target_ods_path)
    _log("target(ods,one): %s" % target_ods_path)
    _direct_run_merge(target_ods_path, settings, one_sheet=True)


def direct_merge_on_one_sheet_to_xlsx(target_xlsx_path, settings):
    """Прямой перенос в лист «Merge» внутри XLSX-файла."""
    if not _is_xlsx_target(target_xlsx_path):
        raise ValueError(
            "xml_excel_ods: цель должна быть .xlsx или .ods (сейчас %s)"
            % target_xlsx_path
        )
    if not os.path.isfile(target_xlsx_path):
        raise ValueError("xml_excel_ods: не найден файл результата: %s" % target_xlsx_path)
    _log("target(xlsx,one): %s" % target_xlsx_path)
    _direct_run_merge(target_xlsx_path, settings, one_sheet=True)


def direct_merge_on_many_sheets_to_ods(target_ods_path, settings):
    """Прямой перенос в режиме «На разные листы» (ODS)."""
    if not _is_ods(target_ods_path):
        raise ValueError(
            "xml_excel_ods: цель должна быть .ods (сейчас %s)" % target_ods_path
        )
    if not os.path.isfile(target_ods_path):
        raise ValueError("xml_excel_ods: не найден файл результата: %s" % target_ods_path)
    _log("target(ods,many): %s" % target_ods_path)
    _direct_run_merge(target_ods_path, settings, one_sheet=False)


def direct_merge_on_many_sheets_to_xlsx(target_xlsx_path, settings):
    """Прямой перенос в режиме «На разные листы» (XLSX)."""
    if not _is_xlsx_target(target_xlsx_path):
        raise ValueError(
            "xml_excel_ods: цель должна быть .xlsx или .ods (сейчас %s)"
            % target_xlsx_path
        )
    if not os.path.isfile(target_xlsx_path):
        raise ValueError("xml_excel_ods: не найден файл результата: %s" % target_xlsx_path)
    _log("target(xlsx,many): %s" % target_xlsx_path)
    _direct_run_merge(target_xlsx_path, settings, one_sheet=False)

