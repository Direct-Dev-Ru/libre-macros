# -*- coding: utf-8 -*-
MACRO_VERSION = "3.10.499"
"""
Проверка bundled openpyxl + smoke цикла xml_excel_ods (без полного сбора).

1) test_openpyxl_bundle — запись/чтение .xlsx через openpyxl.
2) test_xml_excel_ods_file_cycle — как deferred UNO у collect_workbooks:
   save активной книги → копия в temp → правка openpyxl → UNO open/правка/save/close
   → close исходной → replace на место → reopen → MessageBox.
   Книга должна быть уже сохранена на диск (.xlsx/.xlsm/.ods).
"""
import os
import shutil
import tempfile
import time

import uno
from com.sun.star.beans import PropertyValue

import openpyxl_bundled  # noqa: F401 — регистрирует import hook
from openpyxl import Workbook, load_workbook, __version__ as OPENPYXL_VERSION


def _log(msg):
    try:
        print("[test_openpyxl] %s" % msg)
    except Exception:
        pass


def _show_info(title, text, doc=None):
    ctx = uno.getComponentContext()
    smgr = ctx.ServiceManager
    toolkit = smgr.createInstanceWithContext("com.sun.star.awt.Toolkit", ctx)
    if doc is None:
        desktop = smgr.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        doc = desktop.getCurrentComponent()
    parent = getattr(
        getattr(getattr(doc, "CurrentController", None), "Frame", None),
        "ContainerWindow",
        None,
    )
    if parent is None:
        parent = toolkit.getDesktopWindow()
    from com.sun.star.awt.MessageBoxButtons import BUTTONS_OK
    from com.sun.star.awt.MessageBoxType import INFOBOX

    box = toolkit.createMessageBox(parent, INFOBOX, BUTTONS_OK, title, text)
    _log("MessageBox.execute()…")
    box.execute()
    _log("MessageBox.execute() returned OK")


def _desktop():
    ctx = uno.getComponentContext()
    smgr = ctx.ServiceManager
    return smgr.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)


def _current_doc():
    try:
        return XSCRIPTCONTEXT.getDocument()
    except Exception:
        pass
    return _desktop().getCurrentComponent()


def _doc_system_path(doc):
    if doc is None:
        return ""
    try:
        url = doc.getURL()
        if not url:
            return ""
        return uno.fileUrlToSystemPath(url)
    except Exception:
        return ""


def _load_prop(name, value):
    p = PropertyValue()
    p.Name = name
    p.Value = value
    return p


def _dispatch(doc, cmd):
    try:
        frame = doc.getCurrentController().getFrame()
        ctx = uno.getComponentContext()
        smgr = ctx.ServiceManager
        helper = smgr.createInstanceWithContext(
            "com.sun.star.frame.DispatchHelper", ctx
        )
        helper.executeDispatch(frame, cmd, "", 0, ())
    except Exception as err:
        _log("dispatch %s failed: %s" % (cmd, err))


def _save_doc(doc):
    if doc is None:
        return
    try:
        if getattr(doc, "isModified", None) is not None and (not doc.isModified()):
            return
    except Exception:
        pass
    _dispatch(doc, ".uno:Save")


def _close_doc(doc, save=False):
    if doc is None:
        return
    try:
        doc.close(bool(save))
    except Exception:
        try:
            doc.dispose()
        except Exception:
            pass


def _open_path(path, hidden=False):
    path = os.path.abspath(str(path or "").strip())
    if path == "" or not os.path.isfile(path):
        raise ValueError("файл не найден: %s" % path)
    url = uno.systemPathToFileUrl(path)
    props = (_load_prop("Hidden", bool(hidden)), _load_prop("ReadOnly", False))
    opened = _desktop().loadComponentFromURL(url, "_blank", 0, props)
    if opened is None:
        raise ValueError("loadComponentFromURL вернул None: %s" % path)
    return opened


def _make_work_copy(target_path):
    ext = os.path.splitext(target_path)[1] or ".xlsx"
    fd, tmp_path = tempfile.mkstemp(prefix="libre_macros_xml_work_", suffix=ext)
    try:
        os.close(fd)
    except Exception:
        pass
    try:
        os.unlink(tmp_path)
    except Exception:
        pass
    shutil.copy2(target_path, tmp_path)
    return tmp_path


def _replace_path(src_path, dst_path):
    src_path = os.path.abspath(str(src_path or "").strip())
    dst_path = os.path.abspath(str(dst_path or "").strip())
    try:
        os.replace(src_path, dst_path)
        return dst_path
    except Exception:
        shutil.copy2(src_path, dst_path)
        try:
            os.unlink(src_path)
        except Exception:
            pass
        return dst_path


def test_openpyxl_bundle(*args):
    _log("старт bundle, openpyxl %s" % OPENPYXL_VERSION)

    fd, xlsx_path = tempfile.mkstemp(prefix="libre_macros_openpyxl_", suffix=".xlsx")
    os.close(fd)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Проверка"
        ws["A1"] = "Ключ"
        ws["B1"] = "Значение"
        ws["A2"] = "alpha"
        ws["B2"] = 123
        ws["A3"] = "beta"
        ws["B3"] = 456
        wb.save(xlsx_path)
        _log("записано: %s" % xlsx_path)

        wb2 = load_workbook(xlsx_path, read_only=True, data_only=True)
        sheet = wb2.active
        rows = []
        for row in sheet.iter_rows(min_row=1, values_only=True):
            rows.append(row)
        wb2.close()
        _log("прочитано строк: %d" % len(rows))

        doc = _current_doc()
        lo_sheet = doc.getCurrentController().getActiveSheet()

        out = (
            ("openpyxl", OPENPYXL_VERSION),
            ("Файл", xlsx_path),
            ("Строк прочитано", len(rows)),
            ("A2", rows[1][0] if len(rows) > 1 else None),
            ("B2", rows[1][1] if len(rows) > 1 else None),
            ("A3", rows[2][0] if len(rows) > 2 else None),
            ("B3", rows[2][1] if len(rows) > 2 else None),
        )
        lo_sheet.getCellRangeByPosition(0, 0, 1, len(out) - 1).setDataArray(out)

        _show_info(
            "openpyxl bundled",
            "openpyxl %s работает.\n\n"
            "Временный файл:\n%s\n\n"
            "На лист записана сводка (колонки A:B, с первой строки)."
            % (OPENPYXL_VERSION, xlsx_path),
            doc,
        )
    finally:
        pass

    _log("конец bundle")


def test_xml_excel_ods_file_cycle(*args):
    """
    Smoke: save → temp copy → openpyxl edit → UNO open/edit/save/close
    → close original → OS replace → reopen → MessageBox.
    """
    t0 = time.time()

    def _t(step):
        _log("t=%.2fs | %s" % (time.time() - t0, step))

    doc = _current_doc()
    if doc is None or not hasattr(doc, "getSheets"):
        _show_info("xml cycle", "Нужна открытая книга Calc.", None)
        return

    path = _doc_system_path(doc)
    if path == "":
        _show_info(
            "xml cycle",
            "Книга не сохранена на диск.\n"
            "Сохраните .xlsx/.xlsm/.ods и запустите снова.",
            doc,
        )
        return

    _t("1 save active: %s" % path)
    _save_doc(doc)

    work = _make_work_copy(path)
    _t("2 work_copy: %s" % work)

    stamp = "XML_%s" % int(time.time())
    wb = load_workbook(work)
    ws = wb.active
    ws["Z1"] = stamp
    wb.save(work)
    wb.close()
    _t("3 openpyxl edit Z1=%s" % stamp)

    temp_doc = _open_path(work, hidden=False)
    _t("4 UNO open temp")
    try:
        sheet = temp_doc.getSheets().getByIndex(0)
        cell = sheet.getCellByPosition(25, 1)
        cell.String = "UNO_%s" % stamp
        try:
            cell.CellBackColor = 0xF2F2F2
        except Exception:
            pass
        _t("5 UNO edit Z2 + CellBackColor")
        _save_doc(temp_doc)
        _t("6 UNO save temp")
    finally:
        _close_doc(temp_doc, save=False)
        _t("7 UNO close temp")

    _t("8 close original without save")
    _close_doc(doc, save=False)

    _t("9 replace temp -> original")
    _replace_path(work, path)

    _t("10 reopen final")
    final_doc = _open_path(path, hidden=False)
    try:
        final_doc.getCurrentController().getFrame().getContainerWindow().setVisible(True)
    except Exception:
        pass

    z1 = ""
    z2 = ""
    try:
        sh = final_doc.getSheets().getByIndex(0)
        z1 = str(sh.getCellByPosition(25, 0).String or "")
        z2 = str(sh.getCellByPosition(25, 1).String or "")
    except Exception as err:
        _t("readback failed: %s" % err)

    msg = (
        "Цикл xml_excel_ods (smoke) завершён.\n\n"
        "Файл:\n%s\n\n"
        "Z1 (openpyxl)=%s\n"
        "Z2 (UNO)=%s\n\n"
        "Если AO упадёт после OK — краш на MessageBox/после reopen."
        % (path, z1, z2)
    )
    _t("11 MessageBox (жмите OK)")
    _show_info("xml_excel_ods file cycle", msg, final_doc)
    _t("12 after MessageBox — если видите это, OK не уронил процесс")


g_exportedScripts = (test_openpyxl_bundle, test_xml_excel_ods_file_cycle)
