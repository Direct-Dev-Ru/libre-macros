# -*- coding: utf-8 -*-
"""Подмена Файлы-Источники / Папка_результата / прогресс / Способ_переноса в копии шаблона."""
from __future__ import print_function, unicode_literals

import os
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET

try:
    unicode
except NameError:
    unicode = str

NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}
for _prefix, _uri in NS.items():
    ET.register_namespace(_prefix, _uri)

P_FILES = u"Файлы-Источники"
P_RESULT_DIR = u"Папка_результата"
P_PROGRESS = u"Отображение_прогресса"
P_MODE = u"Режим"
P_TRANSFER = u"Способ_переноса"
# Headless: буфер обмена пуст — принудительно По_API; xml_excel_ods не трогаем.
TRANSFER_API = u"По_API"
_XML_EXCEL_ODS_TOKENS = frozenset(
    (
        u"xml_excel_ods",
        u"excel_ods",
        u"xlsx_ods",
        u"direct",
        u"direct_files",
        u"по_файлам",
        u"прямой",
        u"прямой_перенос",
    )
)
PARAM_PREFIXES = (
    u"collect_params",
    u"Параметры_Объединения",
    u"параметры_объединения",
)
TABLE_NS = NS["table"]
TEXT_NS = NS["text"]
OFFICE_NS = NS["office"]


def _norm_name(raw):
    return u" ".join(_as_text(raw).replace(u"_", u" ").split()).casefold()


def _as_text(raw):
    if raw is None:
        return u""
    try:
        return unicode(raw)
    except Exception:
        return str(raw)


def _is_param_sheet(name):
    s = _as_text(name).strip()
    for prefix in PARAM_PREFIXES:
        if s == prefix or s.startswith(prefix + u"_") or s.startswith(prefix):
            return True
    return False


def _looks_like_path(raw):
    s = _as_text(raw).strip()
    if s == u"":
        return False
    low = s.casefold()
    if any(
        tok in low
        for tok in (
            u"source_",
            u"/sources/",
            u"\\sources\\",
            u".xlsx",
            u".xls",
            u".ods",
            u".csv",
            u".xlsm",
            u".xml",
        )
    ):
        return True
    if os.path.sep in s or (len(s) >= 2 and s[1] == u":"):
        return True
    return False


def remap_source_path(raw, sources_dir):
    s = _as_text(raw).strip()
    if s == u"":
        return s
    sources_dir = os.path.abspath(sources_dir)
    unified = s.replace(u"\\", u"/")
    marker = u"/sources/"
    idx = unified.casefold().rfind(marker)
    if idx >= 0:
        tail = unified[idx + len(marker) :]
        cand = os.path.join(sources_dir, tail.replace(u"/", os.sep))
        return cand
    if os.path.exists(s):
        return os.path.abspath(s)
    base = os.path.basename(unified)
    if base and base not in (u".", u"..", u"*"):
        direct = os.path.join(sources_dir, base)
        if os.path.exists(direct):
            return direct
        for root, _dirs, files in os.walk(sources_dir):
            if base in files:
                return os.path.join(root, base)
            if any(name.casefold() == base.casefold() for name in files):
                for name in files:
                    if name.casefold() == base.casefold():
                        return os.path.join(root, name)
    return s


def patch_workbook_copy(src_path, dest_path, sources_dir, result_dir=None):
    """Скопировать шаблон и поправить пути. Возвращает dest_path."""
    src_path = os.path.abspath(src_path)
    dest_path = os.path.abspath(dest_path)
    parent = os.path.dirname(dest_path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    shutil.copy2(src_path, dest_path)
    ext = os.path.splitext(dest_path)[1].lower()
    if ext in (".xlsx", ".xltx", ".xlsm"):
        _patch_xlsx(dest_path, sources_dir, result_dir)
    elif ext in (".ods", ".ots"):
        _patch_ods(dest_path, sources_dir, result_dir)
    return dest_path


def _transfer_token(raw):
    return (
        _as_text(raw)
        .strip()
        .casefold()
        .replace(u" ", u"_")
        .replace(u"-", u"_")
    )


def _force_transfer_api_if_needed(raw):
    """
    Для копии шаблона в headless:
    - пусто / Буфер* / clipboard* → По_API;
    - xml_excel_ods (и алиасы) — без изменений;
    - уже По_API / прочее — без изменений.
    """
    text = _as_text(raw).strip()
    tok = _transfer_token(text)
    if tok in _XML_EXCEL_ODS_TOKENS:
        return text
    if tok == u"" or tok.startswith(u"буфер") or tok.startswith(u"clipboard"):
        return TRANSFER_API
    return text


def _ensure_cell(cells, index):
    while len(cells) <= index:
        cells.append(u"")


def _apply_row_cells(cells, sources_dir, result_dir):
    """cells: list[str]; col0 = имя параметра. Правит inplace."""
    if not cells:
        return
    key = _norm_name(cells[0])
    if key == _norm_name(P_FILES):
        i = 1
        while i < len(cells):
            if _looks_like_path(cells[i]) or _as_text(cells[i]).strip():
                cells[i] = remap_source_path(cells[i], sources_dir)
            i += 1
    elif key == _norm_name(P_RESULT_DIR) and result_dir:
        _ensure_cell(cells, 1)
        cells[1] = os.path.abspath(result_dir)
    elif key == _norm_name(P_PROGRESS):
        _ensure_cell(cells, 1)
        cells[1] = u"status"
    elif key == _norm_name(P_TRANSFER):
        # Отдельная строка «Способ_переноса» — значение в B.
        _ensure_cell(cells, 1)
        cells[1] = _force_transfer_api_if_needed(cells[1])
    elif key == _norm_name(P_MODE):
        # Inline способ переноса — колонка C у «Режим».
        _ensure_cell(cells, 2)
        cells[2] = _force_transfer_api_if_needed(cells[2])


def _patch_xlsx(path, sources_dir, result_dir):
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Для .xlsx/.xltx нужен openpyxl (pip install openpyxl)")
    wb = openpyxl.load_workbook(path)
    try:
        for ws in wb.worksheets:
            if not _is_param_sheet(ws.title):
                continue
            max_row = ws.max_row or 1
            max_col = ws.max_column or 1
            r = 1
            while r <= max_row:
                cells = []
                c = 1
                while c <= max_col:
                    val = ws.cell(r, c).value
                    cells.append(u"" if val is None else _as_text(val))
                    c += 1
                _apply_row_cells(cells, sources_dir, result_dir)
                c = 1
                while c <= len(cells):
                    old = ws.cell(r, c).value
                    new = cells[c - 1]
                    if _as_text(old or u"") != new:
                        ws.cell(r, c).value = new if new != u"" else None
                    c += 1
                r += 1
        wb.save(path)
    finally:
        wb.close()


def _odf_cell_text(cell):
    parts = []
    for p in cell.findall(u".//{%s}p" % TEXT_NS):
        parts.append(u"".join(p.itertext()))
    if parts:
        return u"\n".join(parts)
    val = cell.get(u"{%s}value" % OFFICE_NS)
    if val:
        return _as_text(val)
    return u""


def _set_odf_cell_text(cell, text):
    text = _as_text(text)
    cell.set(u"{%s}value-type" % OFFICE_NS, u"string")
    for attr in (
        u"{%s}value" % OFFICE_NS,
        u"{%s}date-value" % OFFICE_NS,
        u"{%s}boolean-value" % OFFICE_NS,
        u"{%s}formula" % TABLE_NS,
    ):
        if attr in cell.attrib:
            del cell.attrib[attr]
    for child in list(cell):
        cell.remove(child)
    p = ET.SubElement(cell, u"{%s}p" % TEXT_NS)
    p.text = text


def _expand_ods_row(row):
    cells = []
    for child in list(row):
        if child.tag != u"{%s}table-cell" % TABLE_NS:
            continue
        repeat = int(child.get(u"{%s}number-columns-repeated" % TABLE_NS) or u"1")
        if repeat > 64:
            repeat = 64
        i = 0
        while i < repeat:
            cells.append(child)
            i += 1
    return cells


def _ods_ensure_row_width(row, odf_cells, need_len):
    """Добавить пустые table-cell, если в texts нужно больше колонок."""
    while len(odf_cells) < need_len:
        cell = ET.SubElement(row, u"{%s}table-cell" % TABLE_NS)
        odf_cells.append(cell)
    return odf_cells


def _patch_ods(path, sources_dir, result_dir):
    tmp_dir = tempfile.mkdtemp(prefix="lm_qa_ods_")
    try:
        with zipfile.ZipFile(path, "r") as zf:
            zf.extractall(tmp_dir)
        content_path = os.path.join(tmp_dir, "content.xml")
        if not os.path.isfile(content_path):
            return
        tree = ET.parse(content_path)
        root = tree.getroot()
        for table in root.iter(u"{%s}table" % TABLE_NS):
            name = table.get(u"{%s}name" % TABLE_NS) or u""
            if not _is_param_sheet(name):
                continue
            for row in table.findall(u"{%s}table-row" % TABLE_NS):
                odf_cells = _expand_ods_row(row)
                if not odf_cells:
                    continue
                texts = [_odf_cell_text(c) for c in odf_cells]
                _apply_row_cells(texts, sources_dir, result_dir)
                odf_cells = _ods_ensure_row_width(row, odf_cells, len(texts))
                i = 0
                while i < len(odf_cells) and i < len(texts):
                    if _odf_cell_text(odf_cells[i]) != texts[i]:
                        _set_odf_cell_text(odf_cells[i], texts[i])
                    i += 1
        tree.write(content_path, encoding="UTF-8", xml_declaration=True)
        tmp_zip = path + u".tmpzip"
        with zipfile.ZipFile(tmp_zip, "w") as zf:
            mime = os.path.join(tmp_dir, "mimetype")
            if os.path.isfile(mime):
                zf.write(mime, "mimetype", compress_type=zipfile.ZIP_STORED)
            for folder, _dirs, files in os.walk(tmp_dir):
                for name in files:
                    if name == "mimetype" and os.path.abspath(folder) == os.path.abspath(tmp_dir):
                        continue
                    full = os.path.join(folder, name)
                    rel = os.path.relpath(full, tmp_dir).replace(os.sep, "/")
                    zf.write(full, rel, compress_type=zipfile.ZIP_DEFLATED)
        os.replace(tmp_zip, path)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
