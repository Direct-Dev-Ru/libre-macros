# -*- coding: utf-8 -*-
"""Проверки сохранённого результата (без UNO)."""
from __future__ import print_function, unicode_literals

import os
import zipfile
from xml.etree import ElementTree as ET

try:
    unicode
except NameError:
    unicode = str

TABLE_NS = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
OFFICE_NS = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"

LOG_PREFIXES = (u"Сбор_книг_лог",)
SKIP_SHEET_PREFIXES = (
    u"collect_params",
    u"Параметры_Объединения",
    u"Сбор_книг_лог",
    u"Справка_макроса",
)


def _as_text(raw):
    if raw is None:
        return u""
    try:
        return unicode(raw)
    except Exception:
        return str(raw)


def _cf(raw):
    return _as_text(raw).strip().casefold()


def col_letter_to_index(cell_a1):
    s = _as_text(cell_a1).strip().upper()
    col = 0
    i = 0
    while i < len(s) and s[i].isalpha():
        col = col * 26 + (ord(s[i]) - 64)
        i += 1
    row = int(s[i:] or u"1")
    return col - 1, row - 1


def load_workbook_grid(path):
    """{sheet_name: list[list[str]]}  (прямоугольник, пустые = '')."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx"):
        return _load_xlsx(path)
    if ext in (".ods", ".ots"):
        return _load_ods(path)
    raise RuntimeError("Неизвестный формат результата: %s" % path)


def _load_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([u"" if c is None else _as_text(c) for c in row])
            out[ws.title] = rows
    finally:
        wb.close()
    return out


def _odf_text(cell):
    parts = []
    for p in cell.findall(".//{%s}p" % TEXT_NS):
        parts.append(u"".join(p.itertext()))
    if parts:
        return u"\n".join(parts)
    val = cell.get("{%s}value" % OFFICE_NS)
    if val:
        return _as_text(val)
    return u""


def _load_ods(path):
    with zipfile.ZipFile(path, "r") as zf:
        xml = zf.read("content.xml")
    root = ET.fromstring(xml)
    out = {}
    for table in root.iter("{%s}table" % TABLE_NS):
        name = table.get("{%s}name" % TABLE_NS) or u""
        grid = []
        for row_el in table.findall("{%s}table-row" % TABLE_NS):
            row_rep = int(row_el.get("{%s}number-rows-repeated" % TABLE_NS) or u"1")
            if row_rep > 5000:
                row_rep = 1
            cells = []
            for child in list(row_el):
                if child.tag != "{%s}table-cell" % TABLE_NS:
                    continue
                crep = int(child.get("{%s}number-columns-repeated" % TABLE_NS) or u"1")
                if crep > 256:
                    crep = 1
                txt = _odf_text(child)
                ci = 0
                while ci < crep:
                    cells.append(txt)
                    ci += 1
            ri = 0
            while ri < row_rep:
                grid.append(list(cells))
                ri += 1
        out[name] = grid
    return out


def _find_sheet(grids, name):
    want = _as_text(name).strip()
    if want in grids:
        return want, grids[want]
    for key in grids:
        if _cf(key) == _cf(want):
            return key, grids[key]
    return None, None


def _is_skip_sheet(name):
    s = _as_text(name)
    for prefix in SKIP_SHEET_PREFIXES:
        if s == prefix or s.startswith(prefix + u"_") or s.startswith(prefix):
            return True
    return False


def _log_sheet_name(grids):
    for name in grids:
        for prefix in LOG_PREFIXES:
            if _as_text(name).startswith(prefix):
                return name
    return None


def _log_blob(grids):
    name = _log_sheet_name(grids)
    if not name:
        return u""
    parts = []
    for row in grids.get(name) or []:
        parts.append(u"\t".join(_as_text(c) for c in row))
    return u"\n".join(parts)


def _status_col_values(grids):
    name = _log_sheet_name(grids)
    if not name:
        return []
    rows = grids.get(name) or []
    out = []
    i = 1
    while i < len(rows):
        row = rows[i]
        if len(row) > 5:
            out.append(_as_text(row[5]))
        i += 1
    return out


def _data_row_count(grid):
    if not grid:
        return 0
    n = 0
    i = 1
    while i < len(grid):
        if any(_as_text(c).strip() for c in grid[i]):
            n += 1
        i += 1
    return n


def _header_row(grid):
    if not grid:
        return []
    return [_as_text(c) for c in grid[0]]


def _cell(grid, a1):
    col, row = col_letter_to_index(a1)
    if row < 0 or col < 0 or row >= len(grid):
        return u""
    line = grid[row]
    if col >= len(line):
        return u""
    return _as_text(line[col])


def _values_close(got, expect, abs_tol):
    if abs_tol is None:
        return _as_text(got).strip() == _as_text(expect).strip()
    try:
        return abs(float(_as_text(got).replace(u",", u".").replace(u" ", u"")) - float(expect)) <= float(
            abs_tol
        )
    except Exception:
        return _as_text(got).strip() == _as_text(expect).strip()


def run_checks(path, checks, version_hint=u"", sources_dir=None):
    """Вернуть list[{id,type,status,detail}] status=pass|fail|skip."""
    del sources_dir
    results = []
    if not checks:
        return results
    if not os.path.isfile(path):
        return [
            {
                "id": 0,
                "type": "file",
                "status": "fail",
                "detail": "нет файла результата: %s" % path,
            }
        ]
    grids = load_workbook_grid(path)
    i = 0
    while i < len(checks):
        spec = checks[i] or {}
        i += 1
        ctype = _as_text(spec.get("type") or u"").strip()
        item = {"id": i, "type": ctype, "status": "pass", "detail": u""}
        if spec.get("gui"):
            item["status"] = "skip"
            item["detail"] = "gui-check в headless пропущен"
            results.append(item)
            continue
        try:
            _eval_check(item, spec, grids, version_hint)
        except Exception as err:
            item["status"] = "fail"
            item["detail"] = _as_text(err)
        results.append(item)
    return results


def _eval_check(item, spec, grids, version_hint):
    ctype = _cf(spec.get("type"))
    if ctype == u"log_exists":
        name = _log_sheet_name(grids)
        if not name:
            item["status"] = "fail"
            item["detail"] = "нет листа журнала"
        else:
            item["detail"] = name
        return
    if ctype == u"log_macro_version":
        blob = _log_blob(grids)
        ver = _as_text(spec.get("version") or version_hint).strip()
        if ver and ver not in blob:
            item["status"] = "fail"
            item["detail"] = "в журнале нет версии %s" % ver
        elif not blob.strip():
            item["status"] = "fail"
            item["detail"] = "журнал пуст"
        else:
            item["detail"] = "версия найдена" if ver else "журнал не пуст"
        return
    if ctype == u"log_no_status":
        needles = spec.get("values") or [u"ошибка", u"error"]
        if isinstance(needles, (str, bytes)):
            needles = [needles]
        statuses = _status_col_values(grids)
        hit = []
        for st in statuses:
            low = _cf(st)
            for n in needles:
                if _cf(n) and _cf(n) in low:
                    hit.append(st)
                    break
        if hit:
            item["status"] = "fail"
            item["detail"] = u"статусы: " + u", ".join(hit[:8])
        else:
            item["detail"] = "нет запрещённых статусов"
        return
    if ctype == u"log_contains":
        text = _as_text(spec.get("text") or u"")
        blob = _log_blob(grids)
        if text and text not in blob and _cf(text) not in _cf(blob):
            item["status"] = "fail"
            item["detail"] = "нет подстроки «%s»" % text
        return
    if ctype == u"sheet_exists":
        name, grid = _find_sheet(grids, spec.get("name"))
        if grid is None:
            item["status"] = "fail"
            item["detail"] = "нет листа «%s»" % spec.get("name")
        else:
            item["detail"] = name
        return
    if ctype == u"sheet_absent":
        name, grid = _find_sheet(grids, spec.get("name"))
        if grid is not None:
            item["status"] = "fail"
            item["detail"] = "лист «%s» есть" % name
        return
    if ctype in (u"min_data_rows", u"max_data_rows"):
        sheet = spec.get("sheet") or spec.get("name")
        name, grid = _find_sheet(grids, sheet) if sheet else (None, None)
        if grid is None and not sheet:
            for key, g in grids.items():
                if not _is_skip_sheet(key):
                    name, grid = key, g
                    break
        if grid is None:
            item["status"] = "fail"
            item["detail"] = "лист для подсчёта строк не найден"
            return
        n = _data_row_count(grid)
        need = int(spec.get("n") or 0)
        if ctype == u"min_data_rows" and n < need:
            item["status"] = "fail"
        if ctype == u"max_data_rows" and n > need:
            item["status"] = "fail"
        item["detail"] = "%s: %d строк" % (name, n)
        return
    if ctype == u"header_contains":
        name, grid = _find_sheet(grids, spec.get("sheet") or spec.get("name"))
        if grid is None:
            item["status"] = "fail"
            item["detail"] = "лист не найден"
            return
        headers = _header_row(grid)
        blob = u" | ".join(headers)
        names = spec.get("names") or []
        missing = []
        for tok in names:
            ok = False
            for h in headers:
                if _as_text(tok) in _as_text(h) or _cf(tok) in _cf(h):
                    ok = True
                    break
            if not ok:
                missing.append(_as_text(tok))
        if missing:
            item["status"] = "fail"
            item["detail"] = u"нет заголовков %s; есть: %s" % (missing, blob)
        else:
            item["detail"] = blob[:200]
        return
    if ctype == u"cell_equals":
        name, grid = _find_sheet(grids, spec.get("sheet"))
        if grid is None:
            item["status"] = "fail"
            item["detail"] = "лист не найден"
            return
        got = _cell(grid, spec.get("cell") or u"A1")
        expect = spec.get("value")
        if not _values_close(got, expect, spec.get("abs_tol")):
            item["status"] = "fail"
            item["detail"] = "ожидали %r, получили %r" % (expect, got)
        else:
            item["detail"] = got
        return
    if ctype == u"cell_not_empty":
        name, grid = _find_sheet(grids, spec.get("sheet"))
        if grid is None:
            item["status"] = "fail"
            item["detail"] = "лист не найден"
            return
        got = _cell(grid, spec.get("cell") or u"A1")
        if not got.strip():
            item["status"] = "fail"
            item["detail"] = "пусто"
        else:
            item["detail"] = got[:80]
        return
    if ctype == u"golden_sheet":
        gpath = spec.get("path")
        if not gpath or not os.path.isfile(gpath):
            item["status"] = "fail"
            item["detail"] = "нет эталона %s" % gpath
            return
        gold = load_workbook_grid(gpath)
        sname = spec.get("sheet")
        _, g1 = _find_sheet(grids, sname)
        _, g2 = _find_sheet(gold, sname)
        if g1 is None or g2 is None:
            item["status"] = "fail"
            item["detail"] = "лист «%s» нет в результате или эталоне" % sname
            return
        rng = _as_text(spec.get("range") or u"").strip()
        if rng and u":" in rng:
            a, b = rng.split(u":", 1)
            c1, r1 = col_letter_to_index(a)
            c2, r2 = col_letter_to_index(b)
        else:
            r1 = c1 = 0
            r2 = max(len(g1), len(g2)) - 1
            c2 = max((len(r) for r in g1 + g2), default=1) - 1
        diffs = 0
        rr = r1
        while rr <= r2:
            cc = c1
            while cc <= c2:
                v1 = u""
                v2 = u""
                if rr < len(g1) and cc < len(g1[rr]):
                    v1 = _as_text(g1[rr][cc]).strip()
                if rr < len(g2) and cc < len(g2[rr]):
                    v2 = _as_text(g2[rr][cc]).strip()
                if v1 != v2:
                    diffs += 1
                cc += 1
            rr += 1
        if diffs:
            item["status"] = "fail"
            item["detail"] = "расхождений: %d" % diffs
        else:
            item["detail"] = "совпало"
        return
    item["status"] = "fail"
    item["detail"] = "неизвестный type: %s" % spec.get("type")


def default_checks():
    return [
        {"type": "log_exists"},
        {"type": "log_macro_version"},
        {"type": "log_no_status", "values": ["ошибка", "error", "critical"]},
        {"type": "min_data_rows", "n": 1},
    ]
