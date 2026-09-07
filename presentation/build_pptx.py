#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка презентации libre-macros (python-pptx).

Раскадровку задают параметром --storyboard (файл или id: full / short).
Выходной файл — --out.

  pip install -r requirements.txt
  python build_pptx.py --storyboard STORYBOARD.md --out full.pptx
  python build_pptx.py --storyboard short --out short.pptx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Callable

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
except ImportError:
    print(
        "Нужен python-pptx: pip install -r presentation/requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(1)

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"

# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

C_NAVY = RGBColor(0x0B, 0x3D, 0x5C)
C_TEAL = RGBColor(0x1A, 0x9B, 0x8E)
C_LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
C_MUTED = RGBColor(0x5A, 0x6A, 0x78)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_DARK = RGBColor(0x1A, 0x1A, 0x1A)
C_ACCENT_SOFT = RGBColor(0xD6, 0xEE, 0xEA)
C_WARN = RGBColor(0xC0, 0x4A, 0x3A)


def _set_run(run, *, size=18, bold=False, color=C_DARK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_textbox(slide, left, top, width, height, text, *, size=18, bold=False,
                 color=C_DARK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def _add_bullets(slide, left, top, width, height, items, *, size=18, color=C_DARK,
                 bullet="•"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"{bullet} {item}"
        _set_run(run, size=size, color=color)
    return box


def _add_rect(slide, left, top, width, height, fill: RGBColor, *, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def _add_round_rect(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def _find_asset(basename: str) -> Path | None:
    stem = Path(basename).stem
    for ext in (".png", ".jpg", ".jpeg", ".webp", ".PNG", ".JPG"):
        p = ASSETS / f"{stem}{ext}"
        if p.is_file():
            return p
    return None


def _find_asset_any(*basenames: str) -> tuple[Path | None, str]:
    for name in basenames:
        path = _find_asset(name)
        if path is not None:
            return path, name
    return None, basenames[0] if basenames else ""


def _picture_or_placeholder(slide, basename: str, left, top, width, height, *,
                            hint: str = "", fallbacks: tuple[str, ...] = ()):
    names = (basename,) + fallbacks
    path, shown = _find_asset_any(*names)
    if path is not None:
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        return True
    _add_round_rect(slide, left, top, width, height, C_ACCENT_SOFT)
    label = hint or f"Положите картинку:\nassets/{shown}"
    _add_textbox(
        slide,
        left + Inches(0.2),
        top + height / 2 - Inches(0.4),
        width - Inches(0.4),
        Inches(0.9),
        label,
        size=12,
        color=C_MUTED,
        align=PP_ALIGN.CENTER,
    )
    return False


def _footer(slide, page: int, total: int, org: str):
    _add_textbox(
        slide,
        Inches(0.5),
        Inches(7.05),
        Inches(9),
        Inches(0.35),
        f"libre-macros · {org}",
        size=10,
        color=C_MUTED,
    )
    _add_textbox(
        slide,
        Inches(11.5),
        Inches(7.05),
        Inches(1.3),
        Inches(0.35),
        f"{page} / {total}",
        size=10,
        color=C_MUTED,
        align=PP_ALIGN.RIGHT,
    )


def _accent_bar(slide):
    _add_rect(slide, Inches(0), Inches(0), Inches(0.12), SLIDE_H, C_TEAL)


def _blank_slide(prs):
    blank = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(blank)


def build_full(org: str, year: str, out: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 14

    # --- 01 title ---
    s = _blank_slide(prs)
    _add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, C_NAVY)
    _picture_or_placeholder(
        s, "01_title_hero.png",
        Inches(6.8), Inches(0), Inches(6.533), SLIDE_H,
        hint="assets/01_title_hero.png\n(промпт в STORYBOARD)",
    )
    # dim overlay on right if image present — keep navy left panel
    _add_rect(s, Inches(0), Inches(0), Inches(7.2), SLIDE_H, C_NAVY)
    _add_textbox(
        s, Inches(0.7), Inches(1.8), Inches(6), Inches(1.8),
        "От Excel к AlterOffice:\nавтоматизация отчётов\nбез потери процессов",
        size=28, bold=True, color=C_WHITE,
    )
    _add_textbox(
        s, Inches(0.7), Inches(4.0), Inches(6), Inches(1.0),
        "Пакет макросов libre-macros\nимпортозамещение · Calc",
        size=16, color=C_ACCENT_SOFT,
    )
    _add_textbox(
        s, Inches(0.7), Inches(6.4), Inches(6), Inches(0.4),
        f"Для руководства · {org} · {year}",
        size=12, color=C_MUTED,
    )

    # --- 02 agenda ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "О чём поговорим", size=28, bold=True, color=C_NAVY)
    items = [
        "Зачем уходить с Excel",
        "Что остаётся болевой точкой после установки AlterOffice",
        "Решение: AlterOffice + libre-macros",
        "Возможности и типовые сценарии",
        "Внедрение, риски, пилот",
        "Решение к принятию",
    ]
    for i, t in enumerate(items):
        y = Inches(1.2) + Inches(0.85) * i
        _add_round_rect(s, Inches(0.6), y, Inches(0.7), Inches(0.7), C_TEAL)
        _add_textbox(
            s, Inches(0.6), y + Inches(0.12), Inches(0.7), Inches(0.5),
            f"{i + 1:02d}", size=18, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        _add_textbox(s, Inches(1.55), y + Inches(0.15), Inches(10), Inches(0.5),
                     t, size=20, color=C_DARK)
    _footer(s, 2, total, org)

    # --- 03 context ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(7), Inches(0.6),
                 "Импортозамещение офисного контура", size=26, bold=True, color=C_NAVY)
    _add_textbox(
        s, Inches(0.6), Inches(1.1), Inches(7), Inches(1.0),
        "Переход на AlterOffice закрывает лицензии и суверенитет ПО, "
        "но не закрывает привычные Excel-процессы сам по себе.",
        size=16, color=C_MUTED,
    )
    _add_bullets(
        s, Inches(0.6), Inches(2.3), Inches(7), Inches(3.5),
        [
            "Лицензии Microsoft 365 / Office",
            "Power Query, VBA, «сводные как привыкли»",
            "Сотни книг-источников у подразделений",
            "Риск: «поставили Calc — отчёты собираем руками»",
        ],
        size=18,
    )
    _picture_or_placeholder(
        s, "03_risk_gap.png",
        Inches(8.0), Inches(1.3), Inches(4.8), Inches(5.0),
        hint="assets/03_risk_gap.png",
    )
    _footer(s, 3, total, org)

    # --- 04 pain ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Как сейчас (типичный Excel-контур)", size=26, bold=True, color=C_NAVY)
    _add_bullets(
        s, Inches(0.6), Inches(1.3), Inches(7), Inches(4.5),
        [
            "Сбор из папки / почты / сетевого диска вручную",
            "Power Query или «скопировать → вставить»",
            "VBA / макросы, привязанные к Excel",
            "Разные версии книг → разные ошибки",
            "Ключевые знания у 1–2 сотрудников",
        ],
        size=18,
    )
    _picture_or_placeholder(
        s, "04_excel_chaos.png",
        Inches(8.0), Inches(1.3), Inches(4.8), Inches(5.0),
        hint="assets/04_excel_chaos.png",
    )
    _footer(s, 4, total, org)

    # --- 05 solution ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Решение", size=26, bold=True, color=C_NAVY)
    _add_round_rect(s, Inches(0.6), Inches(1.1), Inches(12.1), Inches(1.6), C_LIGHT)
    _add_textbox(
        s, Inches(0.9), Inches(1.3), Inches(11.5), Inches(1.3),
        "AlterOffice Calc + libre-macros = воспроизводимый конвейер\n"
        "«собрать → очистить → свести → оформить» без Microsoft Excel и без Power Query.",
        size=18, bold=True, color=C_NAVY,
    )
    pillars = [
        ("1. Сбор", "Книги, листы, CSV/XML,\nмаски, «эта книга»"),
        ("2. Постобработка", "Очистка, форма таблицы,\nформулы, оформление"),
        ("3. Финал", "Автофильтр, сводная,\nпечать, журнал"),
    ]
    for i, (title, body) in enumerate(pillars):
        x = Inches(0.6) + Inches(4.1) * i
        _add_round_rect(s, x, Inches(3.0), Inches(3.8), Inches(2.0), C_ACCENT_SOFT)
        _add_textbox(s, x + Inches(0.2), Inches(3.15), Inches(3.4), Inches(0.5),
                     title, size=16, bold=True, color=C_TEAL)
        _add_textbox(s, x + Inches(0.2), Inches(3.7), Inches(3.4), Inches(1.1),
                     body, size=14, color=C_DARK)
    _picture_or_placeholder(
        s, "05_pipeline.png",
        Inches(0.6), Inches(5.2), Inches(12.1), Inches(1.5),
        hint="assets/05_pipeline.png — горизонтальный pipeline",
    )
    _footer(s, 5, total, org)

    # --- 06 what is ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "libre-macros — что это", size=26, bold=True, color=C_NAVY)
    _add_bullets(
        s, Inches(0.6), Inches(1.2), Inches(12), Inches(4.0),
        [
            "Пакет Python-макросов для LibreOffice / AlterOffice Calc",
            "Управляется листом параметров — не «чёрный ящик»",
            "Визард параметров на русском языке",
            "Совместимость с фильтром скриптов AlterOffice 2026",
            "Расширяемость: свои шаги в functions_pp.py / functions_final.py",
        ],
        size=18,
    )
    _add_round_rect(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(1.0), C_LIGHT)
    _add_textbox(
        s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(0.7),
        "Не замена всего Excel — слой автоматизации отчётности поверх отечественного офиса.",
        size=16, bold=True, color=C_NAVY,
    )
    _footer(s, 6, total, org)

    # --- 07 architecture ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Как устроен прогон", size=26, bold=True, color=C_NAVY)
    flow = [
        "Параметры_\nОбъединения",
        "Источники",
        "Сбор",
        "Постобработка\n(JSON)",
        "Финал",
        "Журнал",
    ]
    for i, label in enumerate(flow):
        x = Inches(0.4) + Inches(2.15) * i
        _add_round_rect(s, x, Inches(1.2), Inches(1.95), Inches(1.35), C_NAVY)
        _add_textbox(
            s, x + Inches(0.05), Inches(1.4), Inches(1.85), Inches(1.0),
            label, size=12, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        if i < len(flow) - 1:
            _add_textbox(
                s, x + Inches(1.85), Inches(1.55), Inches(0.35), Inches(0.5),
                "→", size=18, bold=True, color=C_TEAL, align=PP_ALIGN.CENTER,
            )
    _add_bullets(
        s, Inches(0.6), Inches(2.9), Inches(6.5), Inches(3.0),
        [
            "Один конфиг или оркестратор нескольких (collect_pack)",
            "Повторяемый сценарий: сохранили книгу параметров — повторили отчёт",
            "Журнал прогона для аудита и разбора сбоев",
        ],
        size=16,
    )
    _picture_or_placeholder(
        s, "07_architecture.png",
        Inches(7.5), Inches(2.9), Inches(5.3), Inches(3.5),
        hint="assets/07_architecture.png",
    )
    _footer(s, 7, total, org)

    # --- 08 capabilities ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Возможности вместо Excel / Power Query", size=24, bold=True, color=C_NAVY)
    cols = [
        ("Сбор", [
            "Много файлов и маски",
            "CSV / XML",
            "«Эта книга»",
            "Вовлечение листов",
        ]),
        ("Очистка и форма", [
            "Дубликаты, замена",
            "Текст, столбцы",
            "Формулы, условный столбец",
            "Разделение столбцов",
        ]),
        ("Аналитика", [
            "ВПР / lookup",
            "Сводная таблица",
            "Количество значений",
            "Группировка",
        ]),
        ("Оформление", [
            "Сетка, зебра, градиент",
            "Автофильтр",
            "Печать, закрепление",
            "Форматы денег/дат",
        ]),
    ]
    for i, (title, bullets) in enumerate(cols):
        x = Inches(0.4) + Inches(3.2) * i
        _add_round_rect(s, x, Inches(1.15), Inches(3.05), Inches(5.2), C_LIGHT)
        _add_rect(s, x, Inches(1.15), Inches(3.05), Inches(0.55), C_TEAL)
        _add_textbox(
            s, x + Inches(0.1), Inches(1.22), Inches(2.85), Inches(0.45),
            title, size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        _add_bullets(
            s, x + Inches(0.15), Inches(1.9), Inches(2.75), Inches(4.0),
            bullets, size=13,
        )
    _add_textbox(
        s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.4),
        "Не обещаем «100% Power Query» — закрываем типовые сценарии отчётности; остальное — roadmap.",
        size=12, color=C_MUTED,
    )
    _footer(s, 8, total, org)

    # --- 09 scenario ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Типовой сценарий предприятия", size=26, bold=True, color=C_NAVY)
    steps = [
        "Подразделения кладут книги в папку",
        "Ответственный открывает книгу параметров",
        "Запускает сбор (или collect_pack)",
        "Макросы чистят, обогащают, сводят",
        "Готовый лист + автофильтр / сводная",
        "При необходимости — QR, меню, AD-списки",
    ]
    for i, t in enumerate(steps):
        y = Inches(1.1) + Inches(0.7) * i
        _add_round_rect(s, Inches(0.6), y, Inches(0.55), Inches(0.55), C_TEAL)
        _add_textbox(
            s, Inches(0.6), y + Inches(0.08), Inches(0.55), Inches(0.4),
            str(i + 1), size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        _add_textbox(s, Inches(1.4), y + Inches(0.1), Inches(6.5), Inches(0.45),
                     t, size=16, color=C_DARK)
    _picture_or_placeholder(
        s, "09_scenario.png",
        Inches(8.2), Inches(1.2), Inches(4.6), Inches(5.2),
        hint="assets/09_scenario.png",
    )
    _footer(s, 9, total, org)

    # --- 10 UX ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(7), Inches(0.6),
                 "Удобство для пользователей", size=26, bold=True, color=C_NAVY)
    _add_bullets(
        s, Inches(0.6), Inches(1.2), Inches(7), Inches(4.5),
        [
            "Визард параметров — без ручного JSON",
            "Подписи на русском, выбор столбцов из заголовков",
            "Пресеты типовых сценариев",
            "Меню и контекстное меню в Calc",
            "Версия пакета видна (show_version)",
        ],
        size=18,
    )
    _picture_or_placeholder(
        s, "10_wizard_ui.png",
        Inches(8.0), Inches(1.2), Inches(4.8), Inches(5.2),
        hint="Лучше реальный скрин\nassets/10_wizard_ui.png",
    )
    _footer(s, 10, total, org)

    # --- 11 deploy ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(7), Inches(0.6),
                 "Внедрение", size=26, bold=True, color=C_NAVY)
    _add_bullets(
        s, Inches(0.6), Inches(1.2), Inches(7), Inches(4.5),
        [
            "Установка скриптов пользователя (Linux / Windows)",
            "LibreOffice и AlterOffice",
            "Проверки совместимости AO 2026 в поставке",
            "Параметры и сценарии — в книге, версионируются с процессом",
            "Свои Python-шаги без переписывания ядра",
        ],
        size=18,
    )
    _picture_or_placeholder(
        s, "11_deploy.png",
        Inches(8.0), Inches(1.2), Inches(4.8), Inches(5.2),
        hint="assets/11_deploy.png",
    )
    _footer(s, 11, total, org)

    # --- 12 effects ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Эффект и риски", size=26, bold=True, color=C_NAVY)
    _add_round_rect(s, Inches(0.5), Inches(1.15), Inches(6.0), Inches(5.2), C_LIGHT)
    _add_textbox(s, Inches(0.8), Inches(1.35), Inches(5.5), Inches(0.5),
                 "Эффект", size=18, bold=True, color=C_TEAL)
    _add_bullets(
        s, Inches(0.8), Inches(2.0), Inches(5.4), Inches(4.0),
        [
            "Снижение зависимости от Excel",
            "Повторяемые отчёты",
            "Меньше ручного копирования",
            "Контроль через параметры и журнал",
            "Развитие под свои регламенты",
        ],
        size=15,
    )
    _add_round_rect(s, Inches(6.8), Inches(1.15), Inches(6.0), Inches(5.2), C_LIGHT)
    _add_textbox(s, Inches(7.1), Inches(1.35), Inches(5.5), Inches(0.5),
                 "Риски → митигация", size=18, bold=True, color=C_WARN)
    _add_bullets(
        s, Inches(7.1), Inches(2.0), Inches(5.4), Inches(4.0),
        [
            "Не всё из PQ сразу → roadmap + пилот 1–2 отчёта",
            "Обучение → визард + инструкция",
            "Качество источников → шаги очистки в pipeline",
        ],
        size=15,
    )
    _footer(s, 12, total, org)

    # --- 13 proposal ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.6),
                 "Предложение", size=26, bold=True, color=C_NAVY)
    proposals = [
        ("01", "Утвердить стек: AlterOffice + libre-macros для регламентной отчётности"),
        ("02", "Пилот: 1–2 критичных отчёта (срок N недель)"),
        ("03", "Назначить владельца книги параметров"),
        ("04", "По итогам пилота — план тиража и обучение"),
    ]
    for i, (num, text) in enumerate(proposals):
        y = Inches(1.3) + Inches(1.15) * i
        _add_round_rect(s, Inches(0.6), y, Inches(12.1), Inches(1.0), C_LIGHT)
        _add_round_rect(s, Inches(0.85), y + Inches(0.2), Inches(0.7), Inches(0.6), C_TEAL)
        _add_textbox(
            s, Inches(0.85), y + Inches(0.28), Inches(0.7), Inches(0.45),
            num, size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        _add_textbox(
            s, Inches(1.8), y + Inches(0.28), Inches(10.5), Inches(0.5),
            text, size=18, color=C_DARK,
        )
    _footer(s, 13, total, org)

    # --- 14 closing ---
    s = _blank_slide(prs)
    _add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, C_NAVY)
    _picture_or_placeholder(
        s, "14_closing.png",
        Inches(0), Inches(0), SLIDE_W, SLIDE_H,
        hint="",
    )
    # ensure readable overlay
    _add_rect(s, Inches(0), Inches(2.2), SLIDE_W, Inches(3.0), C_NAVY)
    _add_textbox(
        s, Inches(0.8), Inches(2.5), Inches(11.5), Inches(0.8),
        "Вопросы", size=36, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
    )
    _add_textbox(
        s, Inches(0.8), Inches(3.5), Inches(11.5), Inches(1.2),
        "Демо стенда · документация docs/ · версия пакета в macro-lib/version.txt\n"
        f"{org} · {year}",
        size=16, color=C_ACCENT_SOFT, align=PP_ALIGN.CENTER,
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


FULL_ASSETS = (
    "01_title_hero.png",
    "03_risk_gap.png",
    "04_excel_chaos.png",
    "05_pipeline.png",
    "07_architecture.png",
    "09_scenario.png",
    "10_wizard_ui.png",
    "11_deploy.png",
    "14_closing.png",
)

SHORT_ASSETS = (
    "s01_title.png",
    "s02_vba_wall.png",
    "s03_pipeline.png",
    "s04_practice.png",
    "s05_ask.png",
)


def build_short(org: str, year: str, out: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 5

    # --- 01 title ---
    s = _blank_slide(prs)
    _add_rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, C_NAVY)
    _picture_or_placeholder(
        s, "s01_title.png",
        Inches(6.8), Inches(0), Inches(6.533), SLIDE_H,
        hint="assets/s01_title.png\n(или 01_title_hero.png)",
        fallbacks=("01_title_hero.png",),
    )
    _add_rect(s, Inches(0), Inches(0), Inches(7.2), SLIDE_H, C_NAVY)
    _add_textbox(
        s, Inches(0.7), Inches(1.7), Inches(6), Inches(2.0),
        "Excel → AlterOffice:\nотчёты без переписывания\nвсего VBA",
        size=28, bold=True, color=C_WHITE,
    )
    _add_textbox(
        s, Inches(0.7), Inches(4.1), Inches(6), Inches(1.0),
        "libre-macros · импортозамещение · Calc",
        size=16, color=C_ACCENT_SOFT,
    )
    _add_textbox(
        s, Inches(0.7), Inches(6.4), Inches(6), Inches(0.4),
        f"Кратко для руководства · {org} · {year}",
        size=12, color=C_MUTED,
    )

    # --- 02 problem ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(7), Inches(0.55),
                 "Две задачи сразу", size=26, bold=True, color=C_NAVY)
    _add_textbox(
        s, Inches(0.6), Inches(1.05), Inches(7), Inches(0.9),
        "Поставить AlterOffice — необходимо.\n"
        "Сохранить автоматизацию отчётов — отдельная задача.",
        size=16, color=C_MUTED,
    )
    _add_bullets(
        s, Inches(0.6), Inches(2.2), Inches(7), Inches(4.0),
        [
            "Лицензии и суверенитет офисного ПО",
            "В Excel накоплен парк макросов VBA (сбор, чистка, сводные…)",
            "При смене платформы их нужно переписывать — дорого и рискованно",
            "«Голый» Calc этот слой сам по себе не заменяет",
        ],
        size=16,
    )
    _picture_or_placeholder(
        s, "s02_vba_wall.png",
        Inches(8.0), Inches(1.2), Inches(4.8), Inches(5.2),
        hint="assets/s02_vba_wall.png",
        fallbacks=("03_risk_gap.png",),
    )
    _footer(s, 2, total, org)

    # --- 03 killer ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.55),
                 "Killer feature: конвейер вместо точечного VBA", size=24,
                 bold=True, color=C_NAVY)
    _add_round_rect(s, Inches(0.6), Inches(1.0), Inches(12.1), Inches(1.45), C_LIGHT)
    _add_textbox(
        s, Inches(0.85), Inches(1.15), Inches(11.6), Inches(1.2),
        "libre-macros закрывает типовой класс задач, ради которых писали VBA:\n"
        "собрать книги → очистить → свести → оформить — параметрами и визардом,\n"
        "без отдельного макроса на каждый отчёт.",
        size=15, bold=True, color=C_NAVY,
    )
    _add_textbox(s, Inches(0.6), Inches(2.7), Inches(6.5), Inches(0.4),
                 "Что закрывает (часть парка VBA)", size=14, bold=True, color=C_TEAL)
    _add_bullets(
        s, Inches(0.6), Inches(3.15), Inches(6.5), Inches(2.4),
        [
            "Сбор из папки / многих файлов",
            "Очистка, дубликаты, замена, столбцы",
            "ВПР / lookup, сводная, оформление, автофильтр",
            "Повторяемый сценарий на листе параметров",
        ],
        size=15,
    )
    _add_round_rect(s, Inches(7.4), Inches(2.7), Inches(5.3), Inches(3.3), C_ACCENT_SOFT)
    _add_textbox(s, Inches(7.65), Inches(2.9), Inches(4.9), Inches(0.4),
                 "Честно", size=14, bold=True, color=C_WARN)
    _add_textbox(
        s, Inches(7.65), Inches(3.4), Inches(4.9), Inches(2.3),
        "Не 100% всех VBA.\n"
        "Кастомную логику — пилотом\n"
        "или своими шагами Python.\n\n"
        "Массовый «сбор → свести →\n"
        "оформить» не пишем с нуля.",
        size=14, color=C_DARK,
    )
    _picture_or_placeholder(
        s, "s03_pipeline.png",
        Inches(0.6), Inches(5.7), Inches(12.1), Inches(1.0),
        hint="assets/s03_pipeline.png",
        fallbacks=("05_pipeline.png",),
    )
    _footer(s, 3, total, org)

    # --- 04 practice ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5),
                 "Как выглядит работа", size=26, bold=True, color=C_NAVY)
    _add_textbox(
        s, Inches(0.6), Inches(0.9), Inches(7), Inches(0.7),
        "Один сохранённый сценарий на листе параметров —\n"
        "повторяемый отчёт без «ручного Excel».",
        size=15, color=C_MUTED,
    )
    steps = [
        "Книги подразделений в папке",
        "Открыть книгу параметров",
        "Запуск сбора / pack",
        "Очистка и обогащение в pipeline",
        "Сводная / автофильтр / оформление",
        "Журнал прогона",
    ]
    for i, t in enumerate(steps):
        y = Inches(1.75) + Inches(0.7) * i
        _add_round_rect(s, Inches(0.6), y, Inches(0.5), Inches(0.5), C_TEAL)
        _add_textbox(
            s, Inches(0.6), y + Inches(0.05), Inches(0.5), Inches(0.4),
            str(i + 1), size=13, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        _add_textbox(s, Inches(1.3), y + Inches(0.08), Inches(6.2), Inches(0.4),
                     t, size=15, color=C_DARK)
    _add_round_rect(s, Inches(7.8), Inches(1.75), Inches(4.9), Inches(3.6), C_LIGHT)
    _add_textbox(s, Inches(8.05), Inches(1.95), Inches(4.5), Inches(0.4),
                 "Выгоды", size=14, bold=True, color=C_TEAL)
    _add_bullets(
        s, Inches(8.05), Inches(2.5), Inches(4.5), Inches(2.6),
        [
            "Сценарий версионируется вместе с книгой",
            "Визард на русском — без программиста на каждый отчёт",
            "Совместимость AlterOffice 2026",
        ],
        size=13,
    )
    _picture_or_placeholder(
        s, "s04_practice.png",
        Inches(7.8), Inches(5.5), Inches(4.9), Inches(1.15),
        hint="assets/s04_practice.png",
        fallbacks=("09_scenario.png",),
    )
    _footer(s, 4, total, org)

    # --- 05 ask ---
    s = _blank_slide(prs)
    _accent_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.55),
                 "Что предлагаем", size=26, bold=True, color=C_NAVY)
    proposals = [
        ("01", "Стек: AlterOffice + libre-macros"),
        ("02", "Инвентаризация VBA: отчёты «сбор/очистка/свод» → пилот на libre-macros"),
        ("03", "Остальное — roadmap / точечная разработка"),
        ("04", "Решение: пилот 1–2 отчёта, владелец книги параметров"),
    ]
    for i, (num, text) in enumerate(proposals):
        y = Inches(1.15) + Inches(1.05) * i
        _add_round_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.9), C_LIGHT)
        _add_round_rect(s, Inches(0.85), y + Inches(0.15), Inches(0.7), Inches(0.6), C_TEAL)
        _add_textbox(
            s, Inches(0.85), y + Inches(0.25), Inches(0.7), Inches(0.45),
            num, size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER,
        )
        _add_textbox(
            s, Inches(1.8), y + Inches(0.25), Inches(10.5), Inches(0.5),
            text, size=16, color=C_DARK,
        )
    _add_textbox(
        s, Inches(0.6), Inches(5.6), Inches(12), Inches(0.8),
        f"Вопросы · демо · docs/\n{org} · {year}",
        size=14, color=C_MUTED, align=PP_ALIGN.CENTER,
    )
    _footer(s, 5, total, org)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


# id / имя файла раскадровки → (builder, default out name, asset checklist, label)
STORYBOARDS: dict[str, tuple[Callable[..., Path], str, tuple[str, ...], str]] = {
    "full": (build_full, "libre_macros_import_substitution.pptx", FULL_ASSETS, "полная (14)"),
    "short": (build_short, "libre_macros_short.pptx", SHORT_ASSETS, "краткая (5)"),
    "STORYBOARD.md": (build_full, "libre_macros_import_substitution.pptx", FULL_ASSETS, "полная (14)"),
    "STORYBOARD_short.md": (build_short, "libre_macros_short.pptx", SHORT_ASSETS, "краткая (5)"),
}


def resolve_storyboard(raw: str) -> tuple[Callable[..., Path], Path, tuple[str, ...], str, str]:
    """Возвращает builder, default_out, assets, label, resolved_key."""
    key = raw.strip()
    path = Path(key)
    # абсолютный / относительный путь к md
    if path.suffix.lower() == ".md" or "/" in key or "\\" in key:
        name = path.name
        if name in STORYBOARDS:
            builder, default_name, assets, label = STORYBOARDS[name]
            return builder, HERE / default_name, assets, label, name
        # по содержимому имени
        low = name.lower()
        if "short" in low:
            builder, default_name, assets, label = STORYBOARDS["short"]
            return builder, HERE / default_name, assets, label, name
        builder, default_name, assets, label = STORYBOARDS["full"]
        return builder, HERE / default_name, assets, label, name
    # id
    low = key.lower()
    if low in STORYBOARDS:
        builder, default_name, assets, label = STORYBOARDS[low]
        return builder, HERE / default_name, assets, label, low
    known = ", ".join(sorted({k for k in ("full", "short")}))
    raise SystemExit(f"Неизвестная раскадровка {raw!r}. Доступно: {known} или путь к .md")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        description="Сборка презентации libre-macros (PPTX). "
        "Раскадровка — --storyboard, выход — --out.",
    )
    p.add_argument(
        "--storyboard",
        default="full",
        help="Раскадровка: full | short | путь к STORYBOARD*.md (по умолчанию full)",
    )
    p.add_argument(
        "--out",
        type=Path,
        default=None,
        help="Имя/путь выходного .pptx (по умолчанию из раскадровки)",
    )
    p.add_argument("--org", default="Предприятие", help="Подпись организации в футере")
    p.add_argument("--year", default="2026", help="Год на титуле")
    args = p.parse_args(argv)

    builder, default_out, assets, label, key = resolve_storyboard(args.storyboard)
    if args.out is None:
        out = default_out
    else:
        out = Path(args.out)
        if not out.is_absolute():
            out = (HERE / out) if out.parent == Path(".") else (Path.cwd() / out).resolve()

    print(f"Раскадровка: {key} ({label})")
    out = builder(args.org, args.year, out)
    print(f"OK: {out}")
    missing = [name for name in assets if _find_asset(name) is None]
    if missing:
        print("Нет картинок (заглушки / fallback):")
        for m in missing:
            print(f"  - assets/{m}")
        print("Промпты: STORYBOARD.md / STORYBOARD_short.md")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
