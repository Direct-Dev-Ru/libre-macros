#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Источник для ручной проверки убрать_объединение (Word→Calc merges).

sources/unmerge/source_unmerge.xlsx
  Таблица — шапка 3 строки + пустая + 2 блока данных (как на скрине ТЗ)

Сценарий: 18_unmerge_collapse.xltx
"""
from __future__ import print_function, unicode_literals

import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TEST_DIR = os.path.join(ROOT, "test", "collect_workbooks")
OUT_DIR = os.path.join(TEST_DIR, "sources", "unmerge")
OUT_XLSX = os.path.join(OUT_DIR, "source_unmerge.xlsx")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side
except ImportError:
    sys.path.insert(0, os.path.join(ROOT, "macro-lib", "pythonpath"))
    import openpyxl_bundled  # noqa: F401
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side


def _thin_border():
    thin = Side(style="thin", color="FF808080")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _style_range(ws, r0, r1, c0, c1):
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = _thin_border()
    r = r0
    while r <= r1:
        c = c0
        while c <= c1:
            cell = ws.cell(r, c)
            cell.alignment = align
            cell.border = border
            c += 1
        r += 1


def generate_unmerge_sources(out_path=None):
    out_path = out_path or OUT_XLSX
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(u"Таблица")

    # --- Шапка (строки 1–3), как на скрине ---
    ws["A1"] = u"Заголовок знач1"
    ws["A2"] = u"Заголовк знач2"
    # A3 пусто

    ws["B1"] = u"Заголовок объед знач"
    ws.merge_cells("B1:B3")

    ws["C1"] = u"заголовок знач2.1"
    ws["C2"] = u"заголовок знач2.2."
    ws["C3"] = u"заголовок знач2.3"

    # Горизонтальный merge в шапке (проверка: снять, столбцы не удалять)
    ws["D1"] = u"Горз объед"
    ws.merge_cells("D1:E1")
    ws["D2"] = u"D2"
    ws["E2"] = u"E2"
    ws["D3"] = u"D3"
    ws["E3"] = u"E3"

    # Строка 4 — пустая прослойка (не входит в header_rows=3)

    # --- Блок данных 1 (строки 5–6) ---
    ws["A5"] = u"данные строка 1"
    ws.merge_cells("A5:A6")
    ws["B5"] = u"данные строка 1 кол2 — 1"
    ws["B6"] = u"данные строка 1 кол2 — 2"
    ws["C5"] = u"данные строка 1 кол3 — 1"
    ws["C6"] = u"данные строка 1 кол3 — 2"
    ws["D5"] = u"д1 D"
    ws["E5"] = u"д1 E"
    ws["D6"] = u"д1 D2"
    ws["E6"] = u"д1 E2"

    # --- Блок данных 2 (строки 7–8) ---
    ws["A7"] = u"данные строка 2"
    ws.merge_cells("A7:A8")
    ws["B7"] = u"данные строка 2 кол2 — 1"
    ws["B8"] = u"данные строка 2 кол2 — 2"
    ws["C7"] = u"данные строка 2 кол3 — 1"
    ws.merge_cells("C7:C8")
    ws["D7"] = u"д2 DE"
    ws.merge_cells("D7:E8")

    _style_range(ws, 1, 8, 1, 5)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    for r in range(1, 9):
        ws.row_dimensions[r].height = 22

    # Второй лист — только данные с постоянной высотой блока (для block_rows=2)
    ws2 = wb.create_sheet(u"Блоки_фиксир")
    ws2["A1"] = u"Код"
    ws2["B1"] = u"Значение"
    ws2.merge_cells("A1:A2")
    ws2["A1"] = u"Код / подпись"
    ws2["B1"] = u"Знач верх"
    ws2["B2"] = u"Знач низ"
    ws2["A3"] = u"row2"
    ws2.merge_cells("A3:A4")
    ws2["B3"] = u"v2a"
    ws2["B4"] = u"v2b"
    _style_range(ws2, 1, 4, 1, 2)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 14

    wb.save(out_path)
    readme = os.path.join(os.path.dirname(out_path), "README.md")
    with open(readme, "w", encoding="utf-8") as f:
        f.write(
            u"# Источник убрать_объединение\n\n"
            u"`source_unmerge.xlsx`:\n\n"
            u"- **Таблица** — шапка 3 строки + пустая + 2 блока (маркер A), "
            u"плюс гориз. merge D:E\n"
            u"- **Блоки_фиксир** — равномерные блоки по 2 строки (`block_rows`)\n\n"
            u"Сценарий: `18_unmerge_collapse.xltx`\n\n"
            u"Ожидание после `header_rows=3`, `marker=A`, `sep=\" \"` на «Таблица»:\n"
            u"- 1 строка шапки (значения через пробел)\n"
            u"- 2 строки данных без merge\n"
            u"- столбцы D/E на месте\n"
        )
    print("  %s" % out_path)
    return out_path


def main():
    print("Источники unmerge:")
    generate_unmerge_sources()


if __name__ == "__main__":
    main()
