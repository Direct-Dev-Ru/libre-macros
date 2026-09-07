#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Источники для ручной проверки анкета_в_таблицу / таблица_в_анкету.

Файл: sources/forms/source_forms.xlsx

Листы:
  Анкета_ФИО   — пары Q/A, нарезка by_repeat_key (ключ ФИО)
  Анкета_пусто — блоки через пустую строку
  Анкета_шапка — шапка листа + шапка блока + тело
  Wide_анкеты  — широкая таблица для таблица_в_анкету

Сценарии: 14_form_to_table.xltx, 15_table_to_form.xltx
  (генерируются generate_manual_workbooks.py / generate.sh)

Использование:
  python test/collect_workbooks/generate_form_sources.py
  ./test/collect_workbooks/generate.sh
"""
from __future__ import print_function, unicode_literals

import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TEST_DIR = os.path.join(ROOT, "test", "collect_workbooks")
FORMS_DIR = os.path.join(TEST_DIR, "sources", "forms")
FORMS_XLSX = os.path.join(FORMS_DIR, "source_forms.xlsx")

try:
    from openpyxl import Workbook
except ImportError:
    _MACRO_PYTHONPATH = os.path.join(ROOT, "macro-lib", "pythonpath")
    if _MACRO_PYTHONPATH not in sys.path:
        sys.path.insert(0, _MACRO_PYTHONPATH)
    try:
        import openpyxl_bundled  # noqa: F401
        from openpyxl import Workbook
    except ImportError:
        print("Требуется openpyxl: pip install openpyxl", file=sys.stderr)
        sys.exit(1)


def _write_qa_sheet(wb, title, rows, header=("Вопрос", "Ответ")):
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for q, a in rows:
        ws.append([q, a])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 22
    return ws


def _qa_fio_blocks():
    """Три анкеты с повторяющимся ключом ФИО."""
    people = (
        ("Иванов", "30", "Казань", "инженер"),
        ("Петров", "25", "Уфа", "менеджер"),
        ("Сидорова", "41", "Самара", "аналитик"),
    )
    rows = []
    for fio, age, city, job in people:
        rows.extend(
            [
                ("ФИО", fio),
                ("Возраст", age),
                ("Город", city),
                ("Должность", job),
            ]
        )
    return rows


def _qa_blank_blocks():
    """Две анкеты, разделённые пустой строкой (без повтора ключа)."""
    rows = [
        ("Товар", "Хлеб"),
        ("Кол-во", "2"),
        ("Цена", "45"),
        (None, None),
        ("Товар", "Молоко"),
        ("Кол-во", "1"),
        ("Цена", "89"),
    ]
    return rows


def _qa_preamble_stream():
    """
    Шапка листа (строки 1–3 абсолютные на листе после заголовка Вопрос/Ответ —
    здесь данные с строки 2 файла; в сценарии forms_start_row учитывает копирование).

    В файле:
      row1: Вопрос | Ответ
      row2–4: шапка листа (Название / Период / Орг)
      row5: пусто
      row6+: Бланк+Дата (шапка блока) + ФИО… (тело), повтор
    """
    rows = [
        ("Название опроса", "Опрос сотрудников 2026"),
        ("Период", "январь–март"),
        ("Организация", "ООО Тест"),
        (None, None),
        ("Бланк №", "1001"),
        ("Дата", "01.02.2026"),
        ("ФИО", "Иванов"),
        ("Возраст", "30"),
        ("Город", "Казань"),
        ("Бланк №", "1002"),
        ("Дата", "02.02.2026"),
        ("ФИО", "Петров"),
        ("Возраст", "25"),
        ("Город", "Уфа"),
    ]
    return rows


def _wide_rows():
    header = ["#Блок", "Бланк №", "Дата", "ФИО", "Возраст", "Город", "Должность"]
    body = [
        [1, "1001", "01.02.2026", "Иванов", 30, "Казань", "инженер"],
        [2, "1002", "02.02.2026", "Петров", 25, "Уфа", "менеджер"],
        [3, "1003", "03.02.2026", "Сидорова", 41, "Самара", "аналитик"],
    ]
    return header, body


def generate_form_sources(out_path=None):
    out_path = out_path or FORMS_XLSX
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    _write_qa_sheet(wb, "Анкета_ФИО", _qa_fio_blocks())

    ws_blank = _write_qa_sheet(wb, "Анкета_пусто", [])
    # rewrite with possible blank rows
    ws_blank.delete_rows(1, ws_blank.max_row or 1)
    ws_blank.append(["Вопрос", "Ответ"])
    for q, a in _qa_blank_blocks():
        if q is None and a is None:
            ws_blank.append([None, None])
        else:
            ws_blank.append([q, a])
    ws_blank.column_dimensions["A"].width = 14
    ws_blank.column_dimensions["B"].width = 14

    _write_qa_sheet(wb, "Анкета_шапка", _qa_preamble_stream())

    header, body = _wide_rows()
    ws_wide = wb.create_sheet("Wide_анкеты")
    ws_wide.append(header)
    for row in body:
        ws_wide.append(row)
    for col in ("A", "B", "C", "D", "E", "F", "G"):
        ws_wide.column_dimensions[col].width = 14

    wb.save(out_path)

    readme = os.path.join(os.path.dirname(out_path), "README.md")
    with open(readme, "w", encoding="utf-8") as f:
        f.write(
            "# Источники анкета ↔ таблица\n\n"
            "`source_forms.xlsx` — листы:\n\n"
            "| Лист | Назначение |\n"
            "|------|------------|\n"
            "| **Анкета_ФИО** | Q/A, `by_repeat_key` (ключ ФИО) → `анкета_в_таблицу` |\n"
            "| **Анкета_пусто** | Q/A через пустую строку → `by_blank_row` |\n"
            "| **Анкета_шапка** | шапка листа + шапка блока + тело |\n"
            "| **Wide_анкеты** | wide → `таблица_в_анкету` |\n\n"
            "Сценарии: `14_form_to_table.xltx`, `15_table_to_form.xltx` "
            "(режим «Копирование листов»).\n\n"
            "Генерация:\n\n"
            "```bash\n"
            "python test/collect_workbooks/generate_form_sources.py\n"
            "./test/collect_workbooks/generate.sh\n"
            "```\n"
        )
    print("  %s" % out_path)
    return out_path


def main():
    print("Источники анкет:")
    generate_form_sources()


if __name__ == "__main__":
    main()
