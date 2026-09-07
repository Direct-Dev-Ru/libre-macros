#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Источник для ручной проверки транспонировать_таблицу.

sources/transpose/source_transpose.xlsx
  Матрица      — обычная таблица 3×4
  Метки_колонка — пример headers_from_column (Имя | Q1|Q2|Q3)

Сценарий: 16_transpose_table.xltx
"""
from __future__ import print_function, unicode_literals

import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TEST_DIR = os.path.join(ROOT, "test", "collect_workbooks")
OUT_DIR = os.path.join(TEST_DIR, "sources", "transpose")
OUT_XLSX = os.path.join(OUT_DIR, "source_transpose.xlsx")

try:
    from openpyxl import Workbook
except ImportError:
    sys.path.insert(0, os.path.join(ROOT, "macro-lib", "pythonpath"))
    import openpyxl_bundled  # noqa: F401
    from openpyxl import Workbook


def generate_transpose_sources(out_path=None):
    out_path = out_path or OUT_XLSX
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(u"Матрица")
    ws.append([u"A", u"B", u"C", u"D"])
    ws.append([1, 2, 3, 4])
    ws.append([5, 6, 7, 8])

    ws2 = wb.create_sheet(u"Метки_колонка")
    ws2.append([u"Имя", u"Q1", u"Q2", u"Q3"])
    ws2.append([u"Иван", 5, 3, 8])
    ws2.append([u"Пётр", 2, 7, 1])

    wb.save(out_path)
    readme = os.path.join(os.path.dirname(out_path), "README.md")
    with open(readme, "w", encoding="utf-8") as f:
        f.write(
            u"# Источник транспонирования\n\n"
            u"`source_transpose.xlsx`:\n\n"
            u"- **Матрица** — обычный transpose / result_headers\n"
            u"- **Метки_колонка** — `headers_from_column`\n\n"
            u"Сценарий: `16_transpose_table.xltx`\n"
        )
    print("  %s" % out_path)
    return out_path


def main():
    print("Источники transpose:")
    generate_transpose_sources()


if __name__ == "__main__":
    main()
