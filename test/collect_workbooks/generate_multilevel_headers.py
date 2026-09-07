#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Примеры листов с многоэтажными заголовками для ручной проверки collect_workbooks.

Создаёт:
  sources/multilevel_headers/multilevel_vertical.xlsx
  sources/multilevel_headers/multilevel_mixed.xlsx

Запуск:
  python3 test/collect_workbooks/generate_multilevel_headers.py
"""
from __future__ import print_function, unicode_literals

import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
except ImportError:
    print("Требуется openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sources", "multilevel_headers")


def _save(wb, name):
    os.makedirs(BASE, exist_ok=True)
    path = os.path.join(BASE, name)
    wb.save(path)
    print("written:", path)


def build_vertical():
    """Вертикальные объединения на разных уровнях (строки 2–5)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Vertical"
    ws["A1"] = "Отчёт: многоэтажная шапка (вертикаль)"
    ws.merge_cells("A2:A5")
    ws["A2"] = "Регион"
    ws.merge_cells("B2:B4")
    ws["B2"] = "Отдел"
    ws.merge_cells("C2:D2")
    ws["C2"] = "2024"
    ws["C3"] = "Q1"
    ws["D3"] = "Q2"
    ws["C4"] = "Янв"
    ws["D4"] = "Фев"
    ws["C5"] = "Сумма"
    ws["D5"] = "Кол-во"
    for r in range(6, 16):
        ws.cell(row=r, column=1, value="Центр")
        ws.cell(row=r, column=2, value="Продажи")
        ws.cell(row=r, column=3, value=r * 10)
        ws.cell(row=r, column=4, value=r)
    for cell in ("A2", "B2", "C2", "D2", "C5", "D5"):
        ws[cell].font = Font(bold=True)
        ws[cell].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _save(wb, "multilevel_vertical.xlsx")


def build_mixed():
    """Вертикальные и горизонтальные объединения (строки 2–4)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mixed"
    ws["A1"] = "Отчёт: смешанные объединения"
    ws.merge_cells("A2:A4")
    ws["A2"] = "Блок"
    ws.merge_cells("B2:C2")
    ws["B2"] = "Показатели"
    ws["B3"] = "План"
    ws["C3"] = "Факт"
    ws.merge_cells("B4:C4")
    ws["B4"] = "Итого"
    ws["D2"] = "Комментарий"
    ws["D3"] = "Комментарий"
    ws["D4"] = "Комментарий"
    for r in range(5, 12):
        ws.cell(row=r, column=1, value="A")
        ws.cell(row=r, column=2, value=100 + r)
        ws.cell(row=r, column=3, value=90 + r)
        ws.cell(row=r, column=4, value="ok")
    for ref in ("A2", "B2", "B3", "C3", "B4", "D2"):
        ws[ref].font = Font(bold=True)
        ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    _save(wb, "multilevel_mixed.xlsx")


if __name__ == "__main__":
    build_vertical()
    build_mixed()
