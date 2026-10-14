#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Источник для ручной проверки добавить_объединение.

sources/add_merge/source_add_merge.xlsx
  По_маркеру — столбец A уже с merge; B/C без merge (склеить по высоте A)
  По_блокам — плоская таблица; block_rows=2 для столбцов A,B

Сценарий: 19_add_merge.xltx
"""
from __future__ import print_function, unicode_literals

import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
TEST_DIR = os.path.join(ROOT, "test", "collect_workbooks")
OUT_DIR = os.path.join(TEST_DIR, "sources", "add_merge")
OUT_XLSX = os.path.join(OUT_DIR, "source_add_merge.xlsx")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side
except ImportError:
    sys.path.insert(0, os.path.join(ROOT, "macro-lib", "pythonpath"))
    import openpyxl_bundled  # noqa: F401
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side


def _thin_border():
    thin = Side(style="thin", color="FF808080")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _style_range(ws, r0, r1, c0, c1):
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = _thin_border()
    r = r0
    while r <= r1:
        c = c0
        while c <= c1:
            cell = ws.cell(r, c)
            cell.alignment = align
            cell.border = border
            c += 1
        r += 1


def generate_add_merge_sources(out_path=None):
    out_path = out_path or OUT_XLSX
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    # Лист 1: опорный A уже объединён; B/C — отдельные значения в блоке
    ws = wb.create_sheet(u"По_маркеру")
    ws["A1"] = u"Код"
    ws["B1"] = u"Кол2"
    ws["C1"] = u"Кол3"
    ws["A2"] = u"строка 1"
    ws.merge_cells("A2:A3")
    ws["B2"] = u"b1-1"
    ws["B3"] = u"b1-2"
    ws["C2"] = u"c1-1"
    ws["C3"] = u"c1-2"
    ws["A4"] = u"строка 2"
    ws.merge_cells("A4:A5")
    ws["B4"] = u"b2-1"
    ws["B5"] = u"b2-2"
    ws["C4"] = u"c2-1"
    ws["C5"] = u"c2-2"
    _style_range(ws, 1, 5, 1, 3)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12

    # Лист 2: без merge — block_rows=2
    ws2 = wb.create_sheet(u"По_блокам")
    ws2["A1"] = u"Группа"
    ws2["B1"] = u"Значение"
    ws2["A2"] = u"G1a"
    ws2["B2"] = u"v1a"
    ws2["A3"] = u"G1b"
    ws2["B3"] = u"v1b"
    ws2["A4"] = u"G2a"
    ws2["B4"] = u"v2a"
    ws2["A5"] = u"G2b"
    ws2["B5"] = u"v2b"
    _style_range(ws2, 1, 5, 1, 2)
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 12

    wb.save(out_path)
    readme = os.path.join(os.path.dirname(out_path), "README.md")
    with open(readme, "w", encoding="utf-8") as f:
        f.write(
            u"# Источник добавить_объединение\n\n"
            u"`source_add_merge.xlsx`:\n\n"
            u"- **По_маркеру** — A уже merge; B/C склеить по высоте A (`marker=A`)\n"
            u"- **По_блокам** — плоские пары строк (`block_rows=2` на A,B)\n\n"
            u"Сценарий: `19_add_merge.xltx`\n"
        )
    print("  %s" % out_path)
    return out_path


def main():
    print("Источники add_merge:")
    generate_add_merge_sources()


if __name__ == "__main__":
    main()
